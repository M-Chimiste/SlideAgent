#!/usr/bin/env python
"""Dev tool: Introspect a PPTX file and output a shape map for schema authoring.

Usage:
    python scripts/analyze_template.py <path-to-template.pptx> [--pretty]
"""
import argparse
import json
import sys
from pathlib import Path

# Add backend to path so we can import models
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "backend"))

from pptx import Presentation  # noqa: E402

from app.models.dev_tools import ShapeInfo, ShapeMap  # noqa: E402

EMU_PER_INCH = 914400


def analyze_template(pptx_path: str) -> ShapeMap:
    prs = Presentation(pptx_path)
    shapes: list[ShapeInfo] = []

    for slide_idx, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            info = ShapeInfo(
                shape_id=shape.shape_id,
                shape_name=shape.name,
                slide_index=slide_idx,
                shape_type=str(shape.shape_type),
                has_text=shape.has_text_frame,
                current_text=(
                    shape.text_frame.text[:100] if shape.has_text_frame else None
                ),
                position={
                    "x": round(shape.left / EMU_PER_INCH, 3) if shape.left else 0,
                    "y": round(shape.top / EMU_PER_INCH, 3) if shape.top else 0,
                    "w": round(shape.width / EMU_PER_INCH, 3) if shape.width else 0,
                    "h": round(shape.height / EMU_PER_INCH, 3) if shape.height else 0,
                },
                placeholder_type=(
                    str(shape.placeholder_format.type)
                    if shape.is_placeholder
                    else None
                ),
            )
            shapes.append(info)

    return ShapeMap(
        template_path=pptx_path,
        slide_count=len(prs.slides),
        shapes=shapes,
    )


def main():
    parser = argparse.ArgumentParser(
        description="Introspect a PPTX file and output a shape map"
    )
    parser.add_argument("pptx_path", help="Path to the .pptx file")
    parser.add_argument(
        "--pretty", action="store_true", help="Pretty-print the JSON output"
    )
    args = parser.parse_args()

    if not Path(args.pptx_path).exists():
        print(f"Error: File not found: {args.pptx_path}", file=sys.stderr)
        sys.exit(1)

    shape_map = analyze_template(args.pptx_path)
    indent = 2 if args.pretty else None
    print(json.dumps(shape_map.model_dump(), indent=indent))


if __name__ == "__main__":
    main()
