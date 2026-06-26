from app.services.visual_qa_agent import VisualQAAgent
from app.models.outline import SlideOutline
from app.services.rendered_slide_audit import RenderedSlideAudit
from app.services.rendering import RenderingError

import zipfile
from io import BytesIO

from PIL import Image
from pptx.dml.color import RGBColor
from pptx import Presentation
from pptx.util import Inches, Pt


def test_extract_json_from_fenced_block() -> None:
    agent = VisualQAAgent()
    report = """Some explanation
```json
{"issues":[{"severity":"CRITICAL","message":"Overlap"}]}
```
"""
    payload = agent._extract_json(report)
    assert payload is not None
    assert payload["issues"][0]["message"] == "Overlap"


def test_parse_report_schema_fallback_warning() -> None:
    agent = VisualQAAgent()
    issues = agent._parse_report("not-json", 2)
    assert len(issues) == 1
    assert issues[0].severity == "WARNING"
    assert issues[0].slide_index == 2


def test_parse_report_uses_actual_image_index_for_single_slide_inspection() -> None:
    agent = VisualQAAgent()
    report = (
        '{"issues":[{"severity":"CRITICAL","message":"Overlap",'
        '"category":"overlap","slide_index":0}]}'
    )

    issues = agent._parse_report(report, 4)

    assert len(issues) == 1
    assert issues[0].slide_index == 4


def test_rule_checks_ignore_stale_source_placeholder_when_refs_are_valid() -> None:
    outline = SlideOutline(
        id="outline-source-stale",
        job_id="job-source",
        slide_index=0,
        mode="flexible",
        label="Context rot makes long-running work unreliable",
        content_json={
            "sources": ["[source needed]", "Beyond Vibe Coding > Context Rot"],
            "source_refs": ["doc-1:section:context-rot"],
        },
        layout_json={"layout": "chart", "visual_elements": ["charts"]},
        created_at="2026-01-01T00:00:00Z",
    )

    issues = VisualQAAgent()._rule_based_checks([outline])

    assert "source_placeholder" not in {issue.category for issue in issues}


def test_pptx_structure_checks_flag_empty_inherited_placeholders(tmp_path) -> None:
    pptx_path = tmp_path / "empty-placeholder.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[1])
    prs.save(pptx_path.as_posix())

    issues = VisualQAAgent()._pptx_structure_checks(pptx_path, [])

    assert "unfilled_placeholder" in {issue.category for issue in issues}
    assert any("fill or delete" in issue.message for issue in issues)


def test_rendered_slide_audit_flags_bad_copy_and_generic_diagram(tmp_path) -> None:
    pptx_path = tmp_path / "output.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(1))
    box.text = "Convert the framework consists of five layers into an owned action."
    placeholder = slide.shapes.add_textbox(Inches(0.5), Inches(1.6), Inches(8), Inches(1))
    placeholder.text = "The proliferation of AI models has outpaced the development of (Owner / Next)"
    raw = slide.shapes.add_textbox(Inches(0.5), Inches(2.7), Inches(8), Inches(1))
    raw.text = "Artifact | Purpose | Update trigger"
    filler = slide.shapes.add_textbox(Inches(0.5), Inches(3.8), Inches(8), Inches(1))
    filler.text = "Review evidence for what goes in? When conditions change. When conditions change. When conditions change."
    dangling = slide.shapes.add_textbox(Inches(0.5), Inches(4.9), Inches(8), Inches(1))
    dangling.text = "The proliferation of AI models has outpaced the development of"
    meta = slide.shapes.add_textbox(Inches(0.5), Inches(5.8), Inches(3), Inches(0.4))
    meta.text = "Cover Slide"
    generic = slide.shapes.add_textbox(Inches(3.8), Inches(5.8), Inches(5), Inches(0.4))
    generic.text = "Current readout. Source claim. Operating implication."
    bullet = slide.shapes.add_textbox(Inches(0.5), Inches(6.25), Inches(5), Inches(0.4))
    bullet.text = "\u2022 A unicode bullet slipped into rendered text"
    renderer_filler = slide.shapes.add_textbox(Inches(5.6), Inches(6.25), Inches(5), Inches(0.4))
    renderer_filler.text = "Tie the claim to a source-backed evaluation artifact."
    soft_renderer_filler = slide.shapes.add_textbox(Inches(5.6), Inches(6.65), Inches(5), Inches(0.4))
    soft_renderer_filler.text = "Connect the source evidence to the decision before scaling."
    prs.save(pptx_path.as_posix())
    diagram_dir = tmp_path / "output-diagrams"
    diagram_dir.mkdir()
    (diagram_dir / "01-framework_cycle.svg").write_text(
        """
        <svg xmlns="http://www.w3.org/2000/svg">
          <text>Frame</text><text>Ground</text><text>Build</text><text>Review</text>
        </svg>
        """,
        encoding="utf-8",
    )

    payload, issues = RenderedSlideAudit().inspect(pptx_path, [], tmp_path / "qa")

    categories = {issue.category for issue in issues}
    severities = {issue.category: issue.severity for issue in issues}
    assert payload["passed"] is False
    assert payload["issues"]
    assert "content_quality" in categories
    assert "placeholder_text" in categories
    assert "raw_artifact" in categories
    assert "meta_copy" in categories
    assert "generic_comparison_copy" in categories
    assert "diagram_semantic_fit" in categories
    assert "repeated_placeholder" in categories
    assert "unicode_bullets" in categories
    assert "renderer_filler_copy" in categories
    assert severities["raw_artifact"] == "CRITICAL"
    assert severities["incomplete_content"] == "CRITICAL"
    assert (tmp_path / "qa" / "rendered-slide-audit.json").exists()


def test_rendered_slide_audit_flags_wrapped_renderer_filler(tmp_path) -> None:
    pptx_path = tmp_path / "wrapped-filler.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    lead = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(4), Inches(0.4))
    lead.text = "Tie the claim"
    detail = slide.shapes.add_textbox(Inches(0.5), Inches(1.0), Inches(5), Inches(0.4))
    detail.text = "source-backed evaluation artifact."
    prs.save(pptx_path.as_posix())

    _payload, issues = RenderedSlideAudit().inspect(pptx_path, [], tmp_path / "qa")

    assert any(
        issue.category == "renderer_filler_copy" and issue.severity == "CRITICAL"
        for issue in issues
    )


def test_rendered_slide_audit_flags_dangling_adjective_endings(tmp_path) -> None:
    pptx_path = tmp_path / "dangling-adjective.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bad = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(5), Inches(0.4))
    bad.text = "Bootstrapped benchmarks carry varying"
    source_slice = slide.shapes.add_textbox(Inches(0.5), Inches(1.1), Inches(7), Inches(0.4))
    source_slice.text = "The proliferation of AI models has outpaced the development of evaluation"
    prs.save(pptx_path.as_posix())

    payload, issues = RenderedSlideAudit().inspect(pptx_path, [], tmp_path / "qa")

    assert payload["passed"] is False
    assert any(
        issue.category == "incomplete_content" and issue.severity == "CRITICAL"
        for issue in issues
    )


def test_rendered_slide_audit_flags_modal_heading_fragments(tmp_path) -> None:
    pptx_path = tmp_path / "modal-heading-fragment.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bad = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(5), Inches(0.4))
    bad.text = "Synthetic benchmarks can"
    prs.save(pptx_path.as_posix())

    payload, issues = RenderedSlideAudit().inspect(pptx_path, [], tmp_path / "qa")

    assert payload["passed"] is False
    assert any(
        issue.category == "incomplete_content" and issue.severity == "CRITICAL"
        for issue in issues
    )


def test_rendered_slide_audit_flags_nonsensical_title_fragments(tmp_path) -> None:
    pptx_path = tmp_path / "nonsense-fragments.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bad = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(5), Inches(0.4))
    bad.text = "Synthetic benchmarks cannot is"
    meta_one = slide.shapes.add_textbox(Inches(0.5), Inches(1.1), Inches(4), Inches(0.4))
    meta_one.text = "Executive Summary"
    meta_two = slide.shapes.add_textbox(Inches(0.5), Inches(1.7), Inches(4), Inches(0.4))
    meta_two.text = "Executive Summary"
    prs.save(pptx_path.as_posix())

    payload, issues = RenderedSlideAudit().inspect(pptx_path, [], tmp_path / "qa")

    categories = {issue.category for issue in issues}
    assert payload["passed"] is False
    assert "nonsensical_copy" in categories
    assert "meta_copy" in categories


def test_rendered_slide_audit_flags_layout_overflow_and_overlap(tmp_path) -> None:
    pptx_path = tmp_path / "layout-risk.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tight = slide.shapes.add_textbox(Inches(0.7), Inches(0.7), Inches(1.1), Inches(0.28))
    tight.text = (
        "This text is intentionally too long for a tiny text box and should "
        "be treated as a cut-off risk."
    )
    left = slide.shapes.add_textbox(Inches(2.0), Inches(1.5), Inches(3.0), Inches(1.0))
    left.text = "First visible text region"
    right = slide.shapes.add_textbox(Inches(2.4), Inches(1.65), Inches(3.0), Inches(1.0))
    right.text = "Second visible text region"
    title = slide.shapes.add_textbox(Inches(0.7), Inches(3.1), Inches(3.5), Inches(0.7))
    title.text = "This title is partially covered by the following card"
    blocker = slide.shapes.add_shape(1, Inches(3.25), Inches(3.1), Inches(2.2), Inches(0.9))
    blocker.fill.solid()
    blocker.fill.fore_color.rgb = RGBColor(60, 80, 96)
    prs.save(pptx_path.as_posix())

    payload, issues = RenderedSlideAudit().inspect(pptx_path, [], tmp_path / "qa")

    categories = {issue.category for issue in issues}
    assert "cut-off-text" in categories
    assert "overlap" in categories
    assert "occluded_text" in categories
    diagnostics = payload["slides"][0]["layout_diagnostics"]
    assert diagnostics["overflow_risk_count"] == 1
    assert diagnostics["overlap_pair_count"] >= 1
    assert diagnostics["occlusion_pair_count"] >= 1


def test_rendered_slide_audit_flags_small_body_text(tmp_path) -> None:
    pptx_path = tmp_path / "small-text-risk.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1.0), Inches(1.0), Inches(5.0), Inches(0.5))
    frame = box.text_frame
    frame.clear()
    run = frame.paragraphs[0].add_run()
    run.text = "This support statement is intentionally too small to read comfortably."
    run.font.size = Pt(8)
    prs.save(pptx_path.as_posix())

    payload, issues = RenderedSlideAudit().inspect(pptx_path, [], tmp_path / "qa")

    categories = {issue.category for issue in issues}
    assert "small_text" in categories
    diagnostics = payload["slides"][0]["layout_diagnostics"]
    assert diagnostics["small_text_risk_count"] == 1
    assert diagnostics["small_text_risks"][0]["font_size"] == 8.0


class DummyVisionClient:
    def __init__(self) -> None:
        self.calls = 0

    def complete_vision(self, **kwargs) -> str:
        self.calls += 1
        return '{"issues":[{"severity":"INFO","message":"Looks clean","category":"visual","slide_index":0}]}'


def test_openai_vision_client_is_used_for_image_inspection(tmp_path) -> None:
    image_path = tmp_path / "slide-1.jpg"
    image_path.write_bytes(b"fake-image")
    client = DummyVisionClient()
    agent = VisualQAAgent(openai_client=client)

    report = agent._inspect_image(image_path)

    assert client.calls == 1
    assert "Looks clean" in report


def test_vision_qa_is_bounded_for_large_decks(tmp_path, monkeypatch) -> None:
    images = []
    for index in range(10):
        image_path = tmp_path / f"slide-{index + 1}.jpg"
        image_path.write_bytes(b"fake-image")
        images.append(image_path)
    pptx_path = tmp_path / "deck.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(pptx_path.as_posix())
    monkeypatch.setattr(
        "app.services.visual_qa_agent.render_pptx_to_images",
        lambda *args, **kwargs: images,
    )
    client = DummyVisionClient()

    result, rendered = VisualQAAgent(
        openai_client=client,
        vision_max_slides=4,
    ).inspect_deck(pptx_path, tmp_path / "preview", [])

    assert rendered == images
    assert client.calls == 4
    assert any(issue.category == "vision_sampling" for issue in result.issues)


class CapturingVisionClient:
    def __init__(self) -> None:
        self.image_bytes = b""

    def complete_vision(self, **kwargs) -> str:
        self.image_bytes = kwargs["image_bytes"]
        return '{"issues":[]}'


def test_openai_vision_client_receives_downscaled_slide_image(tmp_path) -> None:
    image_path = tmp_path / "large-slide.jpg"
    Image.new("RGB", (2400, 1350), "#ffffff").save(image_path)
    client = CapturingVisionClient()

    VisualQAAgent(openai_client=client)._inspect_image(image_path)

    with Image.open(BytesIO(client.image_bytes)) as inspected:
        assert max(inspected.size) <= 1280


class CriticalVisionClient:
    def complete_vision(self, **kwargs) -> str:
        return '{"issues":[{"severity":"CRITICAL","message":"Possible cutoff","category":"cut-off-text","slide_index":0}]}'


class FailingVisionClient:
    def complete_vision(self, **kwargs) -> str:
        raise TimeoutError("vision timed out")


def test_vision_failure_warns_without_crashing(tmp_path, monkeypatch) -> None:
    image_path = tmp_path / "slide-1.jpg"
    image_path.write_bytes(b"fake-image")
    pptx_path = tmp_path / "deck.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(pptx_path.as_posix())
    monkeypatch.setattr(
        "app.services.visual_qa_agent.render_pptx_to_images",
        lambda *args, **kwargs: [image_path],
    )

    result, images = VisualQAAgent(openai_client=FailingVisionClient()).inspect_deck(
        pptx_path, tmp_path / "preview", []
    )

    assert images == [image_path]
    assert result.passed is True
    assert any(issue.category == "vision_unavailable" for issue in result.issues)


def test_render_unavailable_blocks_final_qa(tmp_path) -> None:
    outline = SlideOutline(
        id="outline-1",
        job_id="job-1",
        slide_index=0,
        mode="flexible",
        label="Slide",
        content_json={"title": "Action title"},
        layout_json={"layout": "two_column", "visual_elements": ["structured_text"]},
        created_at="2026-01-01T00:00:00Z",
    )
    agent = VisualQAAgent()

    result, images = agent.inspect_deck(
        tmp_path / "missing.pptx", tmp_path / "preview", [outline]
    )

    assert images == []
    assert result.passed is False
    assert result.issues[0].category == "render_unavailable"
    assert any(issue.category == "office_compatibility" for issue in result.issues)


def test_render_unavailable_uses_pptx_preview_fallback(tmp_path, monkeypatch) -> None:
    def fail_render(*args, **kwargs):
        raise RenderingError("LibreOffice is unavailable")

    pptx_path = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text = "A clear action title"
    prs.save(pptx_path.as_posix())
    outline = SlideOutline(
        id="outline-1",
        job_id="job-1",
        slide_index=0,
        mode="flexible",
        label="Slide",
        content_json={"title": "Action title", "sources": ["Uploaded source"]},
        layout_json={"layout": "two_column", "visual_elements": ["structured_text"]},
        created_at="2026-01-01T00:00:00Z",
    )
    monkeypatch.setattr("app.services.visual_qa_agent.render_pptx_to_images", fail_render)

    result, images = VisualQAAgent().inspect_deck(
        pptx_path, tmp_path / "preview", [outline]
    )

    assert images
    assert images[0].exists()
    assert result.passed is False
    assert {issue.category for issue in result.issues} >= {
        "render_unavailable",
        "render_fallback",
    }


def test_fallback_preview_vision_criticals_are_non_blocking(tmp_path, monkeypatch) -> None:
    def fail_render(*args, **kwargs):
        raise RenderingError("LibreOffice is unavailable")

    pptx_path = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text = "A clear action title"
    prs.save(pptx_path.as_posix())
    outline = SlideOutline(
        id="outline-1",
        job_id="job-1",
        slide_index=0,
        mode="flexible",
        label="Slide",
        content_json={"title": "Action title", "sources": ["Uploaded source"]},
        layout_json={"layout": "two_column", "visual_elements": ["structured_text"]},
        created_at="2026-01-01T00:00:00Z",
    )
    monkeypatch.setattr("app.services.visual_qa_agent.render_pptx_to_images", fail_render)

    result, images = VisualQAAgent(openai_client=CriticalVisionClient()).inspect_deck(
        pptx_path, tmp_path / "preview", [outline]
    )

    assert images
    assert result.passed is False
    assert any(issue.severity == "CRITICAL" and issue.category == "render_unavailable" for issue in result.issues)
    assert any(issue.category == "cut-off-text" for issue in result.issues)
    assert any(
        issue.message.startswith("Approximate preview finding:")
        for issue in result.issues
        if issue.category == "cut-off-text"
    )


def test_rule_based_checks_detect_source_and_layout_risks() -> None:
    outlines = [
        SlideOutline(
            id=f"outline-{idx}",
            job_id="job-1",
            slide_index=idx,
            mode="flexible",
            label=f"Slide {idx}",
            content_json={"title": f"Slide {idx}", "sources": []},
            layout_json={"layout": "two_column", "visual_elements": ["structured_text"]},
            created_at="2026-01-01T00:00:00Z",
        )
        for idx in range(4)
    ]

    issues = VisualQAAgent()._rule_based_checks(outlines)
    categories = {issue.category for issue in issues}

    assert "source_coverage" in categories
    assert "layout_variety" in categories
    assert "layout_repetition" in categories


def test_rule_based_checks_detect_repeated_heavy_comparison_treatment() -> None:
    outlines = []
    for idx, layout in enumerate(
        ["comparison_table", "callouts", "comparison_table", "icon_rows"]
    ):
        outlines.append(
            SlideOutline(
                id=f"outline-heavy-repeat-{idx}",
                job_id="job-heavy-repeat",
                slide_index=idx,
                mode="flexible",
                label=f"Slide {idx}",
                content_json={
                    "title": f"Slide {idx}",
                    "sources": ["Uploaded source"],
                    "exhibit_spec": {
                        "type": layout if layout != "icon_rows" else "callouts",
                        "columns": ["Dimension", "Current state", "Target state"],
                        "rows": [{"label": "Process", "values": ["Ad hoc", "Managed"]}],
                    },
                },
                layout_json={"layout": layout, "visual_elements": ["structured_text"]},
                created_at="2026-01-01T00:00:00Z",
            )
        )

    issues = VisualQAAgent()._rule_based_checks(outlines)

    assert "visual_repetition" in {issue.category for issue in issues}


def test_pptx_structure_checks_detect_text_density(tmp_path) -> None:
    pptx_path = tmp_path / "dense.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(0.4))
    box.text = " ".join(["dense"] * 120)
    prs.save(pptx_path.as_posix())

    issues = VisualQAAgent()._pptx_structure_checks(pptx_path, [])

    assert any(issue.category == "overflow_risk" for issue in issues)


def test_pptx_structure_checks_allow_short_structured_exhibit_labels(tmp_path) -> None:
    pptx_path = tmp_path / "structured-labels.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    labels = [
        "FILE",
        "ROLE",
        "UPDATE TRIGGER",
        "projectbrief.md",
        "Purpose and scope",
        "Scope changes",
        "productContext.md",
        "User goals",
        "Insight changes",
        "systemPatterns.md",
        "Architecture rules",
        "Design changes",
        "activeContext.md",
        "Current focus",
        "Each session",
    ]
    for idx, label in enumerate(labels):
        row = idx // 3
        col = idx % 3
        box = slide.shapes.add_textbox(
            Inches(1 + col * 3.2),
            Inches(1.2 + row * 0.52),
            Inches(2.6),
            Inches(0.3),
        )
        box.text = label
    prs.save(pptx_path.as_posix())
    outline = SlideOutline(
        id="outline-1",
        job_id="job-1",
        slide_index=0,
        mode="flexible",
        label="Reference map",
        content_json={"archetype": "table_reference", "sources": ["Uploaded source"]},
        layout_json={"layout": "table_reference"},
        created_at="2026-01-01T00:00:00Z",
    )

    issues = VisualQAAgent()._pptx_structure_checks(pptx_path, [outline])

    assert not any(issue.category == "scanability" for issue in issues)


def test_pptx_structure_checks_still_flags_many_paragraph_objects(tmp_path) -> None:
    pptx_path = tmp_path / "many-paragraphs.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    paragraph = "This paragraph-style evidence box forces the reader to parse another sentence."
    for idx in range(12):
        row = idx // 3
        col = idx % 3
        box = slide.shapes.add_textbox(
            Inches(0.7 + col * 4.0),
            Inches(1.1 + row * 0.76),
            Inches(3.2),
            Inches(0.42),
        )
        box.text = paragraph
    prs.save(pptx_path.as_posix())
    outline = SlideOutline(
        id="outline-1",
        job_id="job-1",
        slide_index=0,
        mode="flexible",
        label="Dense slide",
        content_json={"archetype": "two_column", "sources": ["Uploaded source"]},
        layout_json={"layout": "two_column"},
        created_at="2026-01-01T00:00:00Z",
    )

    issues = VisualQAAgent()._pptx_structure_checks(pptx_path, [outline])

    assert any(issue.category == "scanability" for issue in issues)


def test_pptx_structure_checks_allow_valid_package(tmp_path) -> None:
    pptx_path = tmp_path / "valid.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text = "A clear action title"
    prs.save(pptx_path.as_posix())

    issues = VisualQAAgent()._pptx_structure_checks(pptx_path, [])

    assert not any(issue.category == "office_compatibility" for issue in issues)


def test_pptx_structure_checks_detect_missing_relationship_target(tmp_path) -> None:
    pptx_path = tmp_path / "broken-rel.pptx"
    _write_basic_pptx(pptx_path)
    rel_name = "ppt/slides/_rels/slide1.xml.rels"
    with zipfile.ZipFile(pptx_path, "r") as source_zip:
        entries = {name: source_zip.read(name) for name in source_zip.namelist()}
    rel_xml = entries[rel_name].replace(b"../slideLayouts/slideLayout7.xml", b"../slideLayouts/missing.xml")
    entries[rel_name] = rel_xml
    _write_zip_entries(pptx_path, entries)

    issues = VisualQAAgent()._pptx_structure_checks(pptx_path, [])

    assert any(
        issue.category == "office_compatibility"
        and "Relationship target is missing" in issue.message
        for issue in issues
    )


def test_pptx_structure_checks_detect_missing_content_type(tmp_path) -> None:
    pptx_path = tmp_path / "missing-content-type.pptx"
    _write_basic_pptx(pptx_path)
    with zipfile.ZipFile(pptx_path, "r") as source_zip:
        entries = {name: source_zip.read(name) for name in source_zip.namelist()}
    entries["ppt/customData/custom.dat"] = b"custom"
    _write_zip_entries(pptx_path, entries)

    issues = VisualQAAgent()._pptx_structure_checks(pptx_path, [])

    assert any(
        issue.category == "office_compatibility"
        and "no content type declaration" in issue.message
        for issue in issues
    )


def _write_basic_pptx(path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text = "A clear action title"
    prs.save(path.as_posix())


def _write_zip_entries(path, entries: dict[str, bytes]) -> None:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as output_zip:
        for name, data in entries.items():
            output_zip.writestr(name, data)
