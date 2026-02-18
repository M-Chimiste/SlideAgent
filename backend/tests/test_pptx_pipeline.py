"""Integration tests for PPTXPipeline using the real sample template."""
from pathlib import Path

import pytest

from app.models.schemas import InjectionTarget, PipelineResult
from app.services.pptx_pipeline import PPTXPipeline
from app.services.xml_injector import XMLInjector
from app.storage.local import LocalStorage

TEMPLATES_DIR = Path(__file__).resolve().parent.parent.parent / "templates"
SAMPLE_TEMPLATE = TEMPLATES_DIR / "novartis-status-weekly" / "template.pptx"


@pytest.fixture
def storage(tmp_path):
    # Copy template into storage
    store = LocalStorage(str(tmp_path / "store"))
    template_bytes = SAMPLE_TEMPLATE.read_bytes()
    import asyncio
    asyncio.get_event_loop().run_until_complete(
        store.write_bytes("templates/novartis-status-weekly/template.pptx", template_bytes)
    )
    return store


@pytest.fixture
def pipeline(storage, tmp_path):
    injector = XMLInjector()
    staging = str(tmp_path / "staging")
    return PPTXPipeline(storage=storage, injector=injector, staging_base=staging)


@pytest.mark.asyncio
async def test_pipeline_happy_path(pipeline, storage):
    """Full pipeline: inject into real template, produce valid PPTX."""
    targets = [
        InjectionTarget(slide_index=1, shape_id=2, field_name="deck_title", value="Weekly Status Report"),
        InjectionTarget(slide_index=1, shape_id=3, field_name="report_date", value="17 Feb 2026"),
        InjectionTarget(slide_index=2, shape_id=2, field_name="project_name", value="Ariadne"),
        InjectionTarget(slide_index=2, shape_id=3, field_name="summary", value="On track for Q1 delivery. All milestones met."),
        InjectionTarget(slide_index=2, shape_id=4, field_name="status", value="Green"),
        InjectionTarget(slide_index=2, shape_id=5, field_name="owner", value="Christian"),
    ]

    result = await pipeline.execute(
        template_path="templates/novartis-status-weekly/template.pptx",
        targets=targets,
        job_id="test-job-001",
    )

    assert result.success is True
    assert result.error is None
    assert result.output_path == "outputs/test-job-001/output.pptx"
    assert len(result.warnings) == 0

    # Verify output exists in storage
    assert await storage.exists(result.output_path)

    # Verify output is a valid PPTX that python-pptx can open
    output_bytes = await storage.read_bytes(result.output_path)
    from pptx import Presentation
    import io
    prs = Presentation(io.BytesIO(output_bytes))
    assert len(prs.slides) == 2

    # Verify injected text appears
    slide1 = prs.slides[0]
    texts_1 = [shape.text_frame.text for shape in slide1.shapes if shape.has_text_frame]
    assert "Weekly Status Report" in texts_1
    assert "17 Feb 2026" in texts_1

    slide2 = prs.slides[1]
    texts_2 = [shape.text_frame.text for shape in slide2.shapes if shape.has_text_frame]
    assert "Ariadne" in texts_2
    assert "Green" in texts_2
    assert "Christian" in texts_2


@pytest.mark.asyncio
async def test_pipeline_staging_cleanup(pipeline, tmp_path):
    """Staging directory should be cleaned up on success."""
    targets = [
        InjectionTarget(slide_index=1, shape_id=2, field_name="deck_title", value="Test"),
    ]
    result = await pipeline.execute(
        template_path="templates/novartis-status-weekly/template.pptx",
        targets=targets,
        job_id="cleanup-test",
    )
    assert result.success
    staging = tmp_path / "staging" / "cleanup-test"
    assert not staging.exists()


@pytest.mark.asyncio
async def test_pipeline_missing_template(pipeline):
    """Pipeline should fail gracefully with missing template."""
    targets = [
        InjectionTarget(slide_index=1, shape_id=2, field_name="title", value="Test"),
    ]
    result = await pipeline.execute(
        template_path="templates/nonexistent/template.pptx",
        targets=targets,
        job_id="missing-template",
    )
    assert not result.success
    assert result.error is not None
    assert "not found" in result.error.lower()


@pytest.mark.asyncio
async def test_pipeline_missing_shape_warns(pipeline):
    """Missing shape ID should produce a warning but not fail."""
    targets = [
        InjectionTarget(slide_index=1, shape_id=2, field_name="deck_title", value="Title"),
        InjectionTarget(slide_index=1, shape_id=999, field_name="nonexistent", value="X"),
    ]
    result = await pipeline.execute(
        template_path="templates/novartis-status-weekly/template.pptx",
        targets=targets,
        job_id="missing-shape",
    )
    assert result.success
    assert len(result.warnings) == 1
    assert "999" in result.warnings[0]
