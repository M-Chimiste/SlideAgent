from pathlib import Path
import shutil

import pytest
from pptx import Presentation

from app.models.brand import BrandDNA
from app.models.generation import GeneratedSlideSpec
from app.models.outline import SlideOutline
from app.services.concept_diagram_renderer import (
    ConceptDiagramRenderer,
    DiagramRenderError,
)
from app.services.pptx_renderer import DeterministicPptxRenderer


def _outline(
    layout: str = "dependency_map",
    exhibit_spec: dict | None = None,
    diagram_spec: dict | None = None,
    slide_index: int = 0,
) -> SlideOutline:
    content_json = {
        "action_title": "Make context dependencies explicit before scaling",
        "subheading": "Evidence from uploaded source",
        "bullets": [
            "[Diagram Description: arrows] Product context persists across sessions.",
            "Rules constrain the next run.",
            "Review gates catch unsupported claims.",
        ],
        "sources": ["Uploaded source"],
        "archetype": layout,
    }
    if exhibit_spec is not None:
        content_json["exhibit_spec"] = exhibit_spec
    if diagram_spec is not None:
        content_json["diagram_spec"] = diagram_spec
    return SlideOutline(
        id=f"outline-{slide_index}",
        job_id="job-1",
        slide_index=slide_index,
        mode="flexible",
        label="Make context dependencies explicit before scaling",
        content_json=content_json,
        layout_json={"layout": layout, "visual_elements": ["diagram"]},
        created_at="2026-01-01T00:00:00Z",
    )


def _has_node_sharp() -> bool:
    backend_dir = Path(__file__).resolve().parents[1]
    return shutil.which("node") is not None and (backend_dir / "node_modules" / "sharp").exists()


def test_generated_slide_spec_accepts_diagram_spec() -> None:
    slide = GeneratedSlideSpec(
        slide_number=1,
        slide_type="framework",
        action_title="Show the operating loop leaders can manage",
        diagram_spec={
            "kind": "cycle",
            "center_label": "Agent loop",
            "steps": [{"label": "Frame"}, {"label": "Review"}],
        },
    )

    assert slide.diagram_spec
    assert slide.diagram_spec["kind"] == "cycle"


def test_dependency_exhibit_derives_diagram_spec() -> None:
    renderer = ConceptDiagramRenderer()
    spec = renderer.diagram_spec(
        _outline(
            exhibit_spec={
                "type": "dependency_map",
                "left_node": "Source packet",
                "middle_nodes": ["Memory bank", "Rules files", "Review gate"],
                "right_outcome": "Reliable output",
                "connector_labels": ["grounds", "checks"],
            }
        )
    )

    assert spec["kind"] == "dependency_flow"
    assert spec["left_node"] == "Source packet"
    assert spec["middle_nodes"] == ["Memory bank", "Rules files", "Review gate"]
    assert spec["right_outcome"] == "Reliable output"


def test_cycle_exhibit_derives_diagram_spec() -> None:
    renderer = ConceptDiagramRenderer()
    spec = renderer.diagram_spec(
        _outline(
            layout="framework_cycle",
            exhibit_spec={
                "type": "cycle",
                "center_label": "Agent loop",
                "steps": [
                    {"label": "Frame"},
                    {"label": "Prime"},
                    {"label": "Generate"},
                    {"label": "Review"},
                ],
            },
        )
    )

    assert spec["kind"] == "cycle"
    assert spec["center_label"] == "Agent loop"
    assert [step["label"] for step in spec["steps"]] == [
        "Frame",
        "Prime",
        "Generate",
        "Review",
    ]


def test_label_fit_prefers_complete_clause_before_truncating() -> None:
    renderer = ConceptDiagramRenderer()

    label = renderer._fit_label(
        "New sessions start from zero; previous architectural decisions are lost",
        54,
    )

    assert label == "New sessions start from zero"


def test_svg_validator_rejects_missing_text_class() -> None:
    svg = """
    <svg viewBox="0 0 960 520" xmlns="http://www.w3.org/2000/svg">
      <defs><marker id="arrow"><path d="M0 0L1 1" /></marker></defs>
      <text x="10" y="10" dominant-baseline="central">Bad</text>
    </svg>
    """

    with pytest.raises(DiagramRenderError, match="text node"):
        ConceptDiagramRenderer().validate_svg(svg)


def test_svg_validator_rejects_filled_connectors() -> None:
    svg = """
    <svg viewBox="0 0 960 520" xmlns="http://www.w3.org/2000/svg">
      <defs><marker id="arrow"><path d="M0 0L1 1" /></marker></defs>
      <path class="arr" d="M0 0L10 10" fill="black" />
      <text class="th" x="10" y="10" dominant-baseline="central">Bad</text>
    </svg>
    """

    with pytest.raises(DiagramRenderError, match="fill='none'"):
        ConceptDiagramRenderer().validate_svg(svg)


def test_svg_validator_rejects_out_of_bounds_text() -> None:
    svg = """
    <svg viewBox="0 0 960 520" xmlns="http://www.w3.org/2000/svg">
      <defs><marker id="arrow"><path d="M0 0L1 1" /></marker></defs>
      <text class="th" x="9999" y="10" dominant-baseline="central">Bad</text>
    </svg>
    """

    with pytest.raises(DiagramRenderError, match="viewBox"):
        ConceptDiagramRenderer().validate_svg(svg)


@pytest.mark.skipif(not _has_node_sharp(), reason="Node/sharp worker is not installed")
def test_node_worker_rasterizes_sample_svg(tmp_path: Path) -> None:
    renderer = ConceptDiagramRenderer()
    svg = renderer.svg_for_spec(
        {
            "kind": "cycle",
            "center_label": "Agent loop",
            "steps": [{"label": "Frame"}, {"label": "Prime"}, {"label": "Review"}, {"label": "Reset"}],
        },
        BrandDNA(),
    )
    svg_path = tmp_path / "sample.svg"
    png_path = tmp_path / "sample.png"
    svg_path.write_text(svg, encoding="utf-8")

    renderer.rasterize_svg(svg_path, png_path)

    assert png_path.exists()
    assert png_path.stat().st_size > 1000


@pytest.mark.skipif(not _has_node_sharp(), reason="Node/sharp worker is not installed")
def test_renderer_inserts_diagram_png_and_writes_debug_artifacts(tmp_path: Path) -> None:
    outline = _outline(
        exhibit_spec={
            "type": "dependency_map",
            "left_node": "Source packet",
            "middle_nodes": ["Memory bank", "Rules files", "Review gate"],
            "right_outcome": "Reliable output",
        }
    )
    output_path = tmp_path / "diagram-deck.pptx"

    warnings = DeterministicPptxRenderer().render([outline], BrandDNA(), output_path)

    rendered = Presentation(output_path.as_posix())
    assert warnings == []
    assert any(shape.shape_type == 13 for shape in rendered.slides[0].shapes)
    diagram_dir = tmp_path / "diagram-deck-diagrams"
    assert len(list(diagram_dir.glob("*.svg"))) == 1
    assert len(list(diagram_dir.glob("*.html"))) == 1
    assert len(list(diagram_dir.glob("*.png"))) == 1
    assert "Product context persists" not in "\n".join(
        shape.text
        for shape in rendered.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
    )
    assert "Memory bank" in next(diagram_dir.glob("*.svg")).read_text(encoding="utf-8")


def test_renderer_falls_back_to_native_shapes_with_warning(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def fail_render(*args, **kwargs):
        raise DiagramRenderError("worker unavailable")

    monkeypatch.setattr(ConceptDiagramRenderer, "render", fail_render)
    output_path = tmp_path / "fallback-deck.pptx"

    warnings = DeterministicPptxRenderer().render([_outline()], BrandDNA(), output_path)

    rendered = Presentation(output_path.as_posix())
    text = "\n".join(
        shape.text
        for shape in rendered.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
    )
    assert output_path.exists()
    assert any(warning["field"] == "diagram_render" for warning in warnings)
    assert "Product context persists across sessions." in text
    assert "Diagram Description" not in text


@pytest.mark.skipif(not _has_node_sharp(), reason="Node/sharp worker is not installed")
def test_renderer_can_disable_diagram_pipeline_for_strict_replacements(tmp_path: Path) -> None:
    output_path = tmp_path / "strict-replacement.pptx"

    warnings = DeterministicPptxRenderer().render(
        [_outline()], BrandDNA(), output_path, enable_diagrams=False
    )

    assert warnings == []
    assert output_path.exists()
    assert not (tmp_path / "strict-replacement-diagrams").exists()
