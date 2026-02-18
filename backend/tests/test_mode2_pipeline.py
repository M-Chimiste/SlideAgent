"""Tests for Mode 2 pipeline: slide creation from layouts and content injection."""

import io
from pathlib import Path

import pytest
from pptx import Presentation

from app.models.schemas import DeckOutline, InjectionTarget, SlideOutlineEntry
from app.models.templates import ContentType
from app.services.pptx_pipeline import PPTXPipeline
from app.services.template_registry import TemplateRegistry
from app.services.xml_injector import XMLInjector
from app.storage.local import LocalStorage

TEMPLATES_DIR = Path(__file__).resolve().parent.parent.parent / "templates"


@pytest.fixture
async def registry():
    reg = TemplateRegistry(templates_dir=str(TEMPLATES_DIR))
    await reg.load_all()
    return reg


@pytest.fixture
def pipeline(tmp_path):
    storage_base = str(TEMPLATES_DIR.parent)
    storage = LocalStorage(base_path=storage_base)
    injector = XMLInjector()
    return PPTXPipeline(
        storage=storage,
        injector=injector,
        staging_base=str(tmp_path / "staging"),
    )


@pytest.fixture
def sample_outline():
    return DeckOutline(
        deck_title="Test Deck",
        audience="Testers",
        narrative_arc="Intro, body, conclusion",
        slides=[
            SlideOutlineEntry(
                slide_number=1,
                layout_name="title_slide",
                title="Test Deck Title",
                content_summary="Opening slide",
                content_type=ContentType.title,
            ),
            SlideOutlineEntry(
                slide_number=2,
                layout_name="content_bullets",
                title="Key Points",
                content_summary="Main bullet points",
                content_type=ContentType.bullets,
            ),
            SlideOutlineEntry(
                slide_number=3,
                layout_name="two_column",
                title="Comparison",
                content_summary="Side by side analysis",
                content_type=ContentType.two_column,
            ),
            SlideOutlineEntry(
                slide_number=4,
                layout_name="closing",
                title="Thank You",
                content_summary="Closing slide",
                content_type=ContentType.closing,
            ),
        ],
        total_slides=4,
    )


@pytest.mark.asyncio
async def test_mode2_pipeline_creates_slides(pipeline, registry, sample_outline):
    """Mode 2 pipeline should create slides from layouts and inject content."""
    layouts = registry.get_layouts_for_template("corporate-deck")
    assert layouts is not None

    targets = [
        InjectionTarget(slide_index=1, shape_id=2, field_name="deck_title", value="My Test Deck"),
        InjectionTarget(slide_index=1, shape_id=3, field_name="subtitle", value="A test presentation"),
        InjectionTarget(slide_index=2, shape_id=2, field_name="slide_title", value="Key Points"),
        InjectionTarget(slide_index=2, shape_id=3, field_name="body", value="Point 1\nPoint 2\nPoint 3"),
        InjectionTarget(slide_index=3, shape_id=2, field_name="slide_title", value="Comparison"),
        InjectionTarget(slide_index=3, shape_id=3, field_name="left_content", value="Pro: Fast"),
        InjectionTarget(slide_index=3, shape_id=4, field_name="right_content", value="Con: Complex"),
        InjectionTarget(slide_index=4, shape_id=2, field_name="closing_title", value="Thank You!"),
    ]

    template_path = "templates/corporate-deck/template.pptx"

    result = await pipeline.execute_mode2(
        template_path=template_path,
        outline=sample_outline,
        layouts=layouts,
        targets=targets,
        job_id="test-mode2-001",
    )

    assert result.success, f"Pipeline failed: {result.error}"
    assert result.output_path == "outputs/test-mode2-001/output.pptx"

    # Read and verify the output PPTX
    output_bytes = await pipeline._storage.read_bytes(result.output_path)
    prs = Presentation(io.BytesIO(output_bytes))

    # Should have 4 slides
    assert len(prs.slides) == 4

    # Verify slide 1 (title slide)
    slide1 = prs.slides[0]
    assert slide1.slide_layout.name == "Title Slide"
    texts = [s.text_frame.text for s in slide1.shapes if s.has_text_frame]
    assert "My Test Deck" in texts

    # Verify slide 2 (content bullets)
    slide2 = prs.slides[1]
    assert slide2.slide_layout.name == "Title and Content"

    # Verify slide 3 (two column)
    slide3 = prs.slides[2]
    assert slide3.slide_layout.name == "Two Content"

    # Verify slide 4 (closing)
    slide4 = prs.slides[3]
    assert slide4.slide_layout.name == "Title Only"


@pytest.mark.asyncio
async def test_mode2_pipeline_empty_outline(pipeline, registry):
    """Pipeline should handle an empty outline."""
    layouts = registry.get_layouts_for_template("corporate-deck")
    outline = DeckOutline(
        deck_title="Empty",
        audience="Nobody",
        narrative_arc="None",
        slides=[],
        total_slides=0,
    )

    result = await pipeline.execute_mode2(
        template_path="templates/corporate-deck/template.pptx",
        outline=outline,
        layouts=layouts,
        targets=[],
        job_id="test-mode2-empty",
    )

    assert result.success
    output_bytes = await pipeline._storage.read_bytes(result.output_path)
    prs = Presentation(io.BytesIO(output_bytes))
    assert len(prs.slides) == 0


@pytest.mark.asyncio
async def test_mode2_pipeline_unknown_layout_warns(pipeline, registry, sample_outline):
    """Pipeline should warn about unknown layouts but not crash."""
    layouts = registry.get_layouts_for_template("corporate-deck")

    # Add a slide with unknown layout
    sample_outline.slides.append(
        SlideOutlineEntry(
            slide_number=5,
            layout_name="nonexistent_layout",
            title="Ghost Slide",
            content_summary="This layout doesn't exist",
            content_type=ContentType.bullets,
        )
    )
    sample_outline.total_slides = 5

    result = await pipeline.execute_mode2(
        template_path="templates/corporate-deck/template.pptx",
        outline=sample_outline,
        layouts=layouts,
        targets=[],
        job_id="test-mode2-unknown",
    )

    assert result.success
    # Should still create 4 valid slides (the 5th is skipped)
    output_bytes = await pipeline._storage.read_bytes(result.output_path)
    prs = Presentation(io.BytesIO(output_bytes))
    assert len(prs.slides) == 4
