from pathlib import Path

from app.models.outline import SlideOutline
from app.models.qa import QAIssue, QAResult
from app.services.design_agent import DesignAgent
from app.tools.all_mode_smoke import (
    _actionable_issue_signature,
    _dedupe_warning_dicts,
    _generic_title_frames,
    _has_planner_fallback,
    _is_initial_consulting_warning,
    _model_slug,
    _normalize_modes,
    _parse_modes,
    _planner_fallback_reason,
    deck_report,
    evaluate_deck_quality,
    serialize_qa_issues,
    summarize_qa,
)


def test_summarize_qa_counts_severities_and_categories() -> None:
    summary = summarize_qa(
        [
            QAIssue(severity="CRITICAL", message="Broken", category="office"),
            QAIssue(severity="WARNING", message="Needs source", category="source"),
            QAIssue(severity="INFO", message="Rendered", category="render"),
            QAIssue(severity="WARNING", message="Again", category="source"),
        ]
    )

    assert summary == {
        "count": 4,
        "critical": 1,
        "warning": 2,
        "info": 1,
        "categories": {"office": 1, "source": 2, "render": 1},
    }


def test_deck_report_shape_is_machine_readable(tmp_path: Path) -> None:
    output = tmp_path / "deck.pptx"
    preview = tmp_path / "slide-1.png"
    result = deck_report(
        mode="freeform",
        output_path=output,
        slide_count=1,
        titles=["Improve delivery discipline"],
        layouts=["two_column"],
        planning_warnings=[],
        build_warnings=[],
        qa_result=QAResult(
            issues=[
                QAIssue(
                    severity="WARNING",
                    message="Spacing is tight.",
                    category="spacing",
                    slide_index=0,
                )
            ],
            passed=True,
        ),
        preview_images=[preview],
    )

    assert result["mode"] == "freeform"
    assert result["output"] == output.as_posix()
    assert result["qa_passed"] is True
    assert result["qa_rounds"] == 0
    assert result["qa_history"] == [result["qa"]]
    assert result["deck_quality"] == {
        "available": False,
        "passed": True,
        "issues": [],
        "metrics": {},
    }
    assert result["planner_fallback"] is False
    assert result["consulting_repair_history"] == []
    assert result["qa_issues"] == [
        {
            "severity": "WARNING",
            "category": "spacing",
            "message": "Spacing is tight.",
            "slide_index": 0,
        }
    ]
    assert result["preview_images"] == [preview.as_posix()]
    assert result["contact_sheet"] is None


def test_evaluate_deck_quality_flags_generic_layouts(tmp_path: Path) -> None:
    output = tmp_path / "generic.pptx"
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    blank = prs.slide_layouts[6]
    for idx in range(8):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(0.6)).text = (
            f"Use the same frame for slide {idx}"
        )
    prs.save(output)

    quality = evaluate_deck_quality(
        mode="freeform",
        output_path=output,
        titles=["Use the same frame" for _ in range(8)],
        layouts=["two_column" for _ in range(8)],
        outlines=[],
    )

    assert quality["available"] is True
    assert quality["passed"] is False
    assert any("distinct layouts" in issue for issue in quality["issues"])
    assert any("same layout" in issue for issue in quality["issues"])


def test_evaluate_deck_quality_flags_meta_action_titles(tmp_path: Path) -> None:
    output = tmp_path / "meta-title.pptx"
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    prs = Presentation()
    blank = prs.slide_layouts[6]
    for idx in range(8):
        slide = prs.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(7), Inches(0.5))
        box.text = f"Slide {idx}"
        slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.5),
            Inches(1.4),
            Inches(2),
            Inches(1),
        )
    prs.save(output)
    titles = [
        "Current benchmarks need harnesses that discover operational truth",
        "Translate executive summary of the benchmarking into a distinct operating decision",
        "Use source-grounded harnesses instead of synthetic benchmark shortcuts",
        "Harness-centric design turns existing workflows into evaluation evidence",
        "Implicit ground truth discovery makes benchmark creation scalable",
        "Five harness layers connect contracts, data, execution, and review",
        "Enterprise data catalogs make benchmark discovery scalable",
        "Commit to the recommendation with named ownership",
        "Translate business case for shifting from manual into a distinct operating decision",
        "Translate evidence for why generic benchmarks into a distinct operating decision",
    ]

    quality = evaluate_deck_quality(
        mode="freeform",
        output_path=output,
        titles=titles,
        layouts=[
            "cover",
            "executive_summary",
            "anti_patterns",
            "table_reference",
            "code_panel",
            "dependency_map",
            "callouts",
            "closing_recommendation",
        ],
        outlines=[],
    )

    assert _generic_title_frames(titles) == [
        "Translate executive summary of the benchmarking into a distinct operating decision",
        "Commit to the recommendation with named ownership",
        "Translate business case for shifting from manual into a distinct operating decision",
        "Translate evidence for why generic benchmarks into a distinct operating decision",
    ]
    assert quality["passed"] is False
    assert quality["metrics"]["generic_title_count"] == 4
    assert any("generic/meta action titles" in issue for issue in quality["issues"])


def test_evaluate_deck_quality_flags_placeholder_visual_text(tmp_path: Path) -> None:
    output = tmp_path / "placeholder-text.pptx"
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    prs = Presentation()
    blank = prs.slide_layouts[6]
    layouts = [
        "cover",
        "executive_summary",
        "anti_patterns",
        "table_reference",
        "code_panel",
        "dependency_map",
        "callouts",
        "closing_recommendation",
    ]
    for idx in range(8):
        slide = prs.slides.add_slide(blank)
        text = "Predictions" if idx == 4 else f"Slide {idx}"
        slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(7), Inches(0.5)).text = text
        slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.5),
            Inches(1.4),
            Inches(2),
            Inches(1),
        )
    prs.save(output)

    quality = evaluate_deck_quality(
        mode="freeform",
        output_path=output,
        titles=[f"Improve benchmark harness {idx}" for idx in range(8)],
        layouts=layouts,
        outlines=[],
    )

    assert quality["passed"] is False
    assert quality["metrics"]["placeholder_text_count"] == 1
    assert any("placeholder visual text" in issue for issue in quality["issues"])


def test_evaluate_deck_quality_flags_sparse_card_slides(tmp_path: Path) -> None:
    output = tmp_path / "sparse-card.pptx"
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    prs = Presentation()
    blank = prs.slide_layouts[6]
    colors = [
        "0F2A38",
        "1B7A8C",
        "D89B2C",
        "E7F0F2",
        "F4F6F8",
        "334155",
        "7C3AED",
        "0F766E",
    ]
    for idx, color in enumerate(colors):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_textbox(
            Inches(0.5),
            Inches(0.45),
            Inches(7.4),
            Inches(0.5),
        ).text = f"Benchmark harness slide {idx + 1}"
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.7),
            Inches(1.3),
            Inches(2.2),
            Inches(1.0),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color)
        shape.line.color.rgb = RGBColor.from_string(color)
    prs.save(output)

    layouts = [
        "cover",
        "callouts",
        "anti_patterns",
        "table_reference",
        "code_panel",
        "dependency_map",
        "framework_cycle",
        "checklist",
    ]
    outlines = [
        _smoke_outline(index, layout, sparse=(index == 1))
        for index, layout in enumerate(layouts)
    ]

    quality = evaluate_deck_quality(
        mode="freeform",
        output_path=output,
        titles=[f"Improve benchmark harness evidence path {idx}" for idx in range(8)],
        layouts=layouts,
        outlines=outlines,
    )

    assert quality["passed"] is False
    assert quality["metrics"]["sparse_outline_slide_count"] == 1
    assert quality["metrics"]["sparse_outline_slides"] == [2]
    assert any("sparse source-backed slides" in issue for issue in quality["issues"])


def test_deck_report_keeps_consulting_repair_history_separate(tmp_path: Path) -> None:
    history = [
        {
            "slide_index": 1,
            "field": "consulting_qa",
            "message": "Round 0: Slide body did not support the title.",
        }
    ]

    result = deck_report(
        mode="freeform",
        output_path=tmp_path / "deck.pptx",
        slide_count=1,
        titles=["Improve delivery discipline"],
        layouts=["two_column"],
        planning_warnings=[],
        build_warnings=[],
        qa_result=QAResult(issues=[], passed=True),
        preview_images=[],
        consulting_repair_history=history,
    )

    assert result["planning_warnings"] == []
    assert result["consulting_repair_history"] == history


def test_initial_horizontal_flow_warning_is_consulting_repair_history() -> None:
    assert _is_initial_consulting_warning(
        {
            "slide_index": None,
            "field": "horizontal_flow",
            "message": "Repeated action titles weaken the deck storyline.",
        }
    )
    assert not _is_initial_consulting_warning(
        {
            "slide_index": None,
            "field": "llm_planning",
            "message": "LLM planning was unavailable or malformed.",
        }
    )


def test_deck_report_dedupes_repeated_planning_warnings(tmp_path: Path) -> None:
    repeated = {
        "slide_index": 3,
        "field": "consulting_qa",
        "message": "Round final: Slide body does not clearly support the action title.",
    }

    result = deck_report(
        mode="freeform",
        output_path=tmp_path / "deck.pptx",
        slide_count=1,
        titles=["Improve delivery discipline"],
        layouts=["two_column"],
        planning_warnings=[repeated, dict(repeated)],
        build_warnings=[],
        qa_result=QAResult(issues=[], passed=True),
        preview_images=[],
    )

    assert result["planning_warnings"] == [repeated]
    assert _dedupe_warning_dicts([repeated, dict(repeated)]) == [repeated]


def test_deck_report_fails_quality_when_consulting_warnings_remain(
    tmp_path: Path,
) -> None:
    output = tmp_path / "deck.pptx"
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    blank = prs.slide_layouts[6]
    for idx in range(6):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(7), Inches(0.5)).text = (
            f"Slide {idx}"
        )
    prs.save(output)

    result = deck_report(
        mode="freeform",
        output_path=output,
        slide_count=6,
        titles=[f"Improve benchmark harness {idx}" for idx in range(6)],
        layouts=["cover", "executive_summary", "framework_cycle", "table_reference", "callouts", "checklist"],
        planning_warnings=[
            {
                "slide_index": 2,
                "field": "consulting_qa",
                "message": "Round final: Action title should state a conclusion with a verb.",
            }
        ],
        build_warnings=[],
        qa_result=QAResult(issues=[], passed=True),
        preview_images=[],
    )

    assert result["deck_quality"]["passed"] is False
    assert any("unresolved consulting" in issue for issue in result["deck_quality"]["issues"])


def test_deck_report_fails_quality_when_source_or_spec_warnings_remain(
    tmp_path: Path,
) -> None:
    output = tmp_path / "deck.pptx"
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    blank = prs.slide_layouts[6]
    for idx in range(6):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(7), Inches(0.5)).text = (
            f"Slide {idx}"
        )
    prs.save(output)

    result = deck_report(
        mode="freeform",
        output_path=output,
        slide_count=6,
        titles=[f"Improve benchmark harness {idx}" for idx in range(6)],
        layouts=["cover", "executive_summary", "framework_cycle", "table_reference", "callouts", "checklist"],
        planning_warnings=[
            {
                "slide_index": 2,
                "field": "spec_gate",
                "message": "Slide still contains [source needed] after numeric grounding.",
            }
        ],
        build_warnings=[],
        qa_result=QAResult(issues=[], passed=True),
        preview_images=[],
    )

    assert result["deck_quality"]["passed"] is False
    assert result["deck_quality"]["metrics"]["unresolved_source_spec_warnings"] == 1
    assert any("source/spec" in issue for issue in result["deck_quality"]["issues"])


def test_deck_report_marks_planner_fallback(tmp_path: Path) -> None:
    result = deck_report(
        mode="freeform",
        output_path=tmp_path / "deck.pptx",
        slide_count=1,
        titles=["Improve delivery discipline"],
        layouts=["two_column"],
        planning_warnings=[
            {
                "slide_index": None,
                "field": "llm_planning",
                "message": "LLM planning was unavailable or malformed.",
            }
        ],
        build_warnings=[],
        qa_result=QAResult(issues=[], passed=True),
        preview_images=[],
    )

    assert result["planner_fallback"] is True
    assert _has_planner_fallback(result["planning_warnings"]) is True
    assert (
        _planner_fallback_reason(result["planning_warnings"])
        == "LLM planning was unavailable or malformed."
    )


def test_serialize_qa_issues_preserves_issue_details() -> None:
    issue = QAIssue(
        severity="CRITICAL",
        category="overlap",
        message="Title overlaps subheading.",
        slide_index=2,
    )

    assert serialize_qa_issues([issue]) == [
        {
            "severity": "CRITICAL",
            "category": "overlap",
            "message": "Title overlaps subheading.",
            "slide_index": 2,
        }
    ]


def test_model_slug_is_filesystem_friendly() -> None:
    assert _model_slug("minimax-m2.7") == "minimax-m2-7"
    assert _model_slug("qwen3.6-35b-a3b-mtp") == "qwen3-6-35b-a3b-mtp"


def test_parse_modes_deduplicates_and_normalizes() -> None:
    assert _parse_modes("freeform, brand,freeform") == ["freeform", "brand"]


def test_normalize_modes_defaults_to_all_modes() -> None:
    assert _normalize_modes([]) == ["freeform", "brand", "strict"]


def test_normalize_modes_rejects_unknown_mode() -> None:
    try:
        _normalize_modes(["freeform", "unknown"])
    except ValueError as exc:
        assert "unknown" in str(exc)
    else:
        raise AssertionError("Expected unsupported mode to raise ValueError")


def test_actionable_issue_signature_sorts_deck_and_slide_level_issues() -> None:
    result = QAResult(
        passed=False,
        issues=[
            QAIssue(
                severity="WARNING",
                category="layout",
                message="Deck-level layout variety issue.",
                slide_index=None,
            ),
            QAIssue(
                severity="CRITICAL",
                category="overlap",
                message="Slide-level overlap.",
                slide_index=1,
            ),
        ],
    )

    signature = _actionable_issue_signature(DesignAgent(), result)

    assert signature[0][0] == -1
    assert signature[1][0] == 1


def _smoke_outline(index: int, layout: str, sparse: bool = False) -> SlideOutline:
    if sparse:
        bullets = ["Validated truth exists."]
        exhibit = {"type": "callouts", "points": ["Validated truth exists."]}
    else:
        bullets = [
            "Operational traces reveal expected outcomes before labeling begins.",
            "Reusable contracts turn those traces into repeatable benchmark cases.",
        ]
        exhibit = {
            "type": layout,
            "points": bullets,
        }
    return SlideOutline(
        id=f"outline-{index}",
        job_id="smoke-quality",
        slide_index=index,
        mode="flexible",
        label=f"Benchmark harness evidence path {index}",
        content_json={
            "action_title": f"Improve benchmark harness evidence path {index}",
            "subheading": "Source-backed review should show more than one thin card.",
            "archetype": layout,
            "bullets": bullets,
            "exhibit_spec": exhibit,
            "source_refs": ["source-doc:section:1:harness"],
        },
        layout_json={"layout": layout, "visual_elements": ["structured_exhibit"]},
        created_at="2026-01-01T00:00:00Z",
    )
