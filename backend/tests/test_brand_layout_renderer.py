from pathlib import Path

from pptx import Presentation

from app.models.outline import SlideOutline
from app.models.template import TemplateProfile
from app.services.brand_layout_renderer import (
    BrandLayoutInstantiationRenderer,
    role_for_layout,
)
from app.services.pptx_builder import PptxBuilder
from app.services.template_analyzer import TemplateAnalyzer


def _brand_template(tmp_path: Path) -> TemplateProfile:
    src = tmp_path / "brand.pptx"
    Presentation().save(src.as_posix())  # default template ships the standard layout library
    profile, _ = TemplateAnalyzer().analyze(src, "Brand", "brand")
    return profile


def _outline(index: int, role: str, family: str, title: str, bullets: list[str]) -> SlideOutline:
    return SlideOutline(
        id=f"s{index}",
        job_id="brand-layout-job",
        created_at="t",
        slide_index=index,
        mode="flexible",
        label=title,
        content_json={
            "action_title": title,
            "narrative_role": role,
            "composition_family": family,
            "bullets": bullets,
            "subheading": "context",
        },
        layout_json={},
    )


def _sample_outlines() -> list[SlideOutline]:
    return [
        _outline(0, "cover", "editorial_cover", "Welcome to the deck", []),
        _outline(1, "problem", "reframe_comparison", "Old way versus new way", ["Old is slow", "New is fast"]),
        _outline(2, "evidence", "why_it_matters_cards", "Three reasons it matters", ["One", "Two", "Three"]),
        _outline(3, "section", "section_divider", "Part two", []),
        _outline(4, "closing", "path_forward_close", "Commit to the plan", ["Step one", "Step two"]),
    ]


def test_extract_layout_library_captures_roles(tmp_path: Path) -> None:
    profile = _brand_template(tmp_path)
    assert len(profile.layout_library) >= 8
    roles = {layout.role for layout in profile.layout_library}
    # the standard template exposes at least these distinct roles
    assert {"cover", "section", "comparison", "content"} <= roles


def test_role_for_layout_classifies_standard_layouts(tmp_path: Path) -> None:
    src = tmp_path / "t.pptx"
    Presentation().save(src.as_posix())
    prs = Presentation(src.as_posix())
    roles_by_name = {layout.name: role_for_layout(layout) for layout in prs.slide_layouts}
    assert roles_by_name.get("Title Slide") == "cover"
    assert roles_by_name.get("Section Header") == "section"
    assert roles_by_name.get("Comparison") == "comparison"


def test_instantiation_produces_varied_layouts(tmp_path: Path) -> None:
    profile = _brand_template(tmp_path)
    outlines = _sample_outlines()
    out = tmp_path / "out.pptx"

    ok, warnings = BrandLayoutInstantiationRenderer().render(profile, outlines, out, tmp_path)

    assert ok is True
    res = Presentation(out.as_posix())
    slides = list(res.slides)
    assert len(slides) == len(outlines)
    # distinct, role-appropriate layouts — not "the same 4 slides"
    layout_names = {s.slide_layout.name for s in slides}
    assert len(layout_names) >= 3
    assert slides[0].slide_layout.name == "Title Slide"
    assert slides[1].slide_layout.name == "Comparison"
    # titles are filled
    first_texts = [sh.text_frame.text for sh in slides[0].shapes if sh.has_text_frame]
    assert any("Welcome to the deck" in t for t in first_texts)


def test_instantiation_missing_source_returns_false(tmp_path: Path) -> None:
    profile = TemplateProfile(
        id="b",
        name="B",
        type="brand",
        brand=_brand_template(tmp_path).brand,
        slides=[],
        source_file=str(tmp_path / "does-not-exist.pptx"),
        created_at="t",
        updated_at="t",
    )
    ok, warnings = BrandLayoutInstantiationRenderer().render(profile, _sample_outlines(), tmp_path / "o.pptx", tmp_path)
    assert ok is False
    assert warnings


def test_build_deck_routes_to_instantiation_when_flag_on(tmp_path: Path) -> None:
    profile = _brand_template(tmp_path)
    builder = PptxBuilder(node_runner=object(), brand_layout_instantiation=True)
    out = tmp_path / "deck.pptx"

    builder.build_deck(profile, _sample_outlines(), out, tmp_path)

    res = Presentation(out.as_posix())
    layout_names = {s.slide_layout.name for s in res.slides}
    # instantiation reuses the template's named layouts (clone/generated would not)
    assert "Title Slide" in layout_names
    assert len(list(res.slides)) == len(_sample_outlines())
