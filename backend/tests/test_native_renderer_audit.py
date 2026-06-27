"""Geometry/layout regression guard for the native renderer.

Builds a deck that exercises every native primitive with realistic
action-title + subheading + multi-item content, renders it, then runs the
rendered-slide audit and asserts the deck is free of the layout defects the
metis e2e surfaced: header title/subhead overlap, off-canvas decorative shapes,
opaque-shape text occlusion, unicode bullet glyphs, and sub-floor body text.
Deterministic — no LLM.
"""

import uuid

from pptx import Presentation

from app.models.brand import BrandDNA
from app.models.outline import SlideOutline
from app.services.pptx_native.renderer import NativePptxRenderer
from app.services.rendered_slide_audit import EMU_PER_INCH, RenderedSlideAudit

_TITLE = "Memory Bank architecture provides persistent external context for agents"
_SUB = "Six core files in a dependency hierarchy replace ephemeral chat context across sessions"
_ITEMS = [
    {"title": "Project brief", "body": "Defines scope, goals, and non-negotiable constraints up front"},
    {"title": "Active context", "body": "Tracks the current task focus and recent operating decisions"},
    {"title": "System patterns", "body": "Records architecture choices and the rationale behind them"},
    {"title": "Progress log", "body": "Captures what works, what is left, and known open issues"},
    {"title": "Tech context", "body": "Lists the stack, tooling, and environment setup details"},
    {"title": "Decision record", "body": "Explains why key tradeoffs were made for future readers"},
]


def _outline(i: int, primitive: str, slide_type: str, **extra) -> SlideOutline:
    content = {
        "action_title": _TITLE,
        "title": _TITLE,
        "subheading": _SUB,
        "summary": _SUB,
        "pinned_primitive": primitive,
        "slide_type": slide_type,
        "narrative_role": extra.pop("role", "evidence"),
        "bullets": [it["title"] + ". " + it["body"] for it in _ITEMS],
        "exhibit_spec": {"type": "callouts", "points": _ITEMS},
        "metrics": [
            {"value": "92%", "label": "Context retained", "description": "across long sessions"},
            {"value": "6", "label": "Core files", "description": "in the memory bank"},
            {"value": "3x", "label": "Fewer resets", "description": "per project"},
        ],
        "source_labels": ["Uploaded source"],
        **extra,
    }
    return SlideOutline(
        id=str(uuid.uuid4()),
        job_id="t",
        slide_index=i,
        mode="flexible",
        label=_TITLE,
        content_json=content,
        layout_json={"archetype": primitive},
        created_at="2026-01-01T00:00:00Z",
    )


_PRIMITIVES = [
    ("cover", "cover", {"role": "cover", "deck_title": "Beyond Vibe Coding: An Operating Model"}),
    ("statement", "statement", {"role": "problem"}),
    ("cards", "content", {}),
    ("callout_list", "callouts", {}),
    ("rows", "reference", {}),
    ("steps", "checklist", {}),
    ("columns", "content", {}),
    ("layers", "framework", {}),
    ("timeline", "process", {}),
    (
        "matrix",
        "matrix",
        {"exhibit_spec": {"type": "matrix_2x2", "points": _ITEMS[:4],
                          "quadrant_labels": ["High / Low", "High / High", "Low / Low", "Low / High"]}},
    ),
    ("metrics", "chart", {}),
    ("big_stat", "stat", {}),
    ("quote", "quote", {"role": "decision"}),
    ("split", "deep_dive", {}),
    (
        "table",
        "comparison",
        {"exhibit_spec": {"type": "comparison_table", "columns": ["Dimension", "Vibe coding", "Agentic coding"],
                          "rows": [{"label": "Context", "values": ["Ephemeral", "Persistent"]},
                                   {"label": "Scale", "values": ["Fails", "Holds"]},
                                   {"label": "Review", "values": ["Skipped", "Mandatory"]}]}},
    ),
    (
        "closing",
        "closing",
        {"role": "closing", "exhibit_spec": {"type": "recommendation",
                                             "next_steps": [it["title"] + ". " + it["body"] for it in _ITEMS[:4]],
                                             "decision_ask": "Approve the first governed pilot this quarter"}},
    ),
]

# Layout defects that must never appear in a rendered native deck.
_FORBIDDEN = {"overlap", "occluded_text", "shape_bounds", "unicode_bullets", "small_text", "cut-off-text"}


def test_native_renderer_has_no_layout_defects(tmp_path):
    outlines = [_outline(i, prim, stype, **extra) for i, (prim, stype, extra) in enumerate(_PRIMITIVES)]
    out = tmp_path / "all-primitives.pptx"
    NativePptxRenderer().render(outlines, BrandDNA(), out)

    audit = RenderedSlideAudit()
    payload, _ = audit.inspect(out, outlines)

    offenders = [
        f"slide {issue['slide_index']} [{_PRIMITIVES[issue['slide_index']][0]}]: "
        f"{issue['category']} — {issue['message']}"
        for issue in payload["issues"]
        if issue.get("category") in _FORBIDDEN
    ]
    assert not offenders, "native renderer produced layout defects:\n" + "\n".join(offenders)


def test_native_renderer_does_not_truncate_untitled_points_mid_fragment(tmp_path):
    """Long, colon-less exhibit points (no clean lead) fall back to a truncated
    title display; that fallback must cut on a word boundary and never leave a
    dangling fragment like '...experience goals, a' (the metis slide-3 bug)."""
    points = [
        "Projectbrief.md Defines purpose, scope, and success criteria",
        "ProductContext.md Captures user problem, experience goals, and project scope",
        "SystemPatterns.md Records architecture, patterns, and constraints",
        "ActiveContext.md Tracks current focus, decisions, and next actions",
    ]
    outline = _outline(0, "rows", "reference", exhibit_spec={"type": "callouts", "points": points})
    out = tmp_path / "untitled-points.pptx"
    NativePptxRenderer().render([outline], BrandDNA(), out)

    audit = RenderedSlideAudit()
    payload, _ = audit.inspect(out, [outline])
    offenders = [i for i in payload["issues"] if i.get("category") in _FORBIDDEN]
    assert not offenders, [i["message"] for i in offenders]


def test_native_renderer_keeps_all_shapes_in_bounds(tmp_path):
    """Decorative motifs and content alike must stay within the 13.333x7.5in canvas."""
    outlines = [_outline(i, prim, stype, **extra) for i, (prim, stype, extra) in enumerate(_PRIMITIVES)]
    out = tmp_path / "bounds.pptx"
    NativePptxRenderer().render(outlines, BrandDNA(), out)

    prs = Presentation(out.as_posix())
    w = prs.slide_width / EMU_PER_INCH
    h = prs.slide_height / EMU_PER_INCH
    oob: list[str] = []
    for i, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            left = shape.left / EMU_PER_INCH
            top = shape.top / EMU_PER_INCH
            right = left + (shape.width or 0) / EMU_PER_INCH
            bottom = top + (shape.height or 0) / EMU_PER_INCH
            if left < -0.02 or top < -0.02 or right > w + 0.05 or bottom > h + 0.05:
                oob.append(f"slide {i} [{_PRIMITIVES[i][0]}] '{getattr(shape, 'name', '')}' "
                           f"L{left:.2f} T{top:.2f} R{right:.2f} B{bottom:.2f}")
    assert not oob, "shapes outside slide bounds:\n" + "\n".join(oob)
