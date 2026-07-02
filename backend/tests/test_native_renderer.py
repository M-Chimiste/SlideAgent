from pptx import Presentation

from app.models.brand import BrandDNA
from app.models.outline import SlideOutline
from app.services.pptx_native import geometry as geo
from app.services.pptx_native import primitives as prim
from app.services.pptx_native import theme as th
from app.services.pptx_native.renderer import NativePptxRenderer


def _brand(dl: str = "editorial_serif") -> BrandDNA:
    return BrandDNA(
        colors={"primary": "14213D", "secondary": "1F7A8C", "accent": "C8893B", "background_light": "EAEEF5"},
        fonts={"heading": "Aptos Display", "body": "Aptos"},
        layout_profile={"design_language": dl},
    )


def _outline(i: int, primitive: str, title: str, bullets: list[str]) -> SlideOutline:
    return SlideOutline(
        id=f"s{i}", job_id="j", created_at="t", slide_index=i, mode="flexible", label=title,
        content_json={
            "action_title": title, "pinned_primitive": primitive, "slide_type": primitive,
            "narrative_role": "evidence", "subheading": "context here", "bullets": bullets,
            "metrics": [{"value": "85%", "label": "adoption", "description": "of teams"}] if primitive in ("metrics", "big_stat") else [],
            "exhibit_spec": {"columns": ["A", "B"], "rows": [{"label": "x", "values": ["1"]}]} if primitive == "table" else {},
        },
        layout_json={},
    )


def test_px_to_inch_is_96_per_inch():
    assert geo.px(96) == 1.0
    assert geo.px(48) == 0.5


def test_parse_padding_shorthand():
    assert geo.parse_padding_px("70px 84px 60px") == (70, 84, 60, 84)
    assert geo.parse_padding_px("26px 28px") == (26, 28, 26, 28)
    assert geo.parse_padding_px("40px") == (40, 40, 40, 40)


def test_content_area_insets_by_padding():
    area = geo.content_area("70px 84px 60px")
    assert round(area.x, 3) == round(geo.px(84), 3)
    assert round(area.y, 3) == round(geo.px(70), 3)
    assert round(area.right, 3) == round(geo.SLIDE_W_IN - geo.px(84), 3)
    assert round(area.bottom, 3) == round(geo.SLIDE_H_IN - geo.px(60), 3)


def test_grid_reflows_by_count():
    area = geo.Rect(0, 0, 12, 6)
    one_row = geo.grid(3, 3, area, gap=0.2)
    assert len(one_row) == 3
    # equal widths, single row (same y)
    assert len({round(c.w, 4) for c in one_row}) == 1
    assert len({round(c.y, 4) for c in one_row}) == 1
    two_by_two = geo.grid(4, 2, area, gap=0.2)
    assert len(two_by_two) == 4
    assert len({round(c.y, 4) for c in two_by_two}) == 2  # two rows


def test_column_split_respects_ratios():
    area = geo.Rect(0, 0, 10, 6)
    cols = geo.column_split(area, [0.82, 1.18], gap=0.0)
    assert round(cols[1].w / cols[0].w, 2) == round(1.18 / 0.82, 2)
    assert round(cols[0].w + cols[1].w, 4) == 10.0


def test_stack_rows():
    area = geo.Rect(0, 0, 10, 6)
    rows = geo.stack(area, 4, gap=0.1)
    assert len(rows) == 4
    assert all(round(r.w, 4) == 10.0 for r in rows)
    assert rows[1].y > rows[0].y


def test_theme_pt_scales_display_by_heading_scale():
    assert th.pt(60) == 45.0           # 60px * 0.75
    assert th.pt(60, 1.14) == 51.3     # bold_minimal scale


def test_theme_helpers_resolve_from_design_language():
    # a freeform-style brand carrying a design language
    brand = BrandDNA(
        colors={"primary": "14213D", "secondary": "1F7A8C", "accent": "C8893B", "background_light": "EAEEF5"},
        fonts={"heading": "Newsreader", "body": "Inter"},
        layout_profile={"design_language": "bold_minimal"},
    )
    theme = th.resolve_theme(brand)
    assert theme.card_radius == 4  # bold_minimal
    area = th.slide_content_area(theme)
    assert area.w < geo.SLIDE_W_IN  # inset by slide padding
    assert th.card_gap_in(theme) > 0


def test_native_render_is_editable_no_images(tmp_path):
    outs = [
        _outline(0, "cover", "Beyond Vibe Coding", []),
        _outline(1, "cards", "Three failure modes", ["Hallucinations spike with ambiguity", "Silent integration breaks", "Developer becomes a PM"]),
        _outline(2, "statement", "Structured context is the antidote", ["Persistent memory beats history", "A rules file removes ambiguity"]),
        _outline(3, "rows", "The Memory Bank architecture", ["Files form a dependency graph", "Brief anchors scope", "Active context captures state"]),
        _outline(4, "quote", "It forces you to clarify your thinking", ["It forces you to clarify your own thinking before the model writes a line"]),
        _outline(5, "closing", "Commit to persistent context", ["Adopt a rules file", "Stand up a Memory Bank"]),
    ]
    out = tmp_path / "deck.pptx"
    warnings = NativePptxRenderer().render(outs, _brand(), out)
    assert warnings == []
    res = Presentation(out.as_posix())
    slides = list(res.slides)
    assert len(slides) == len(outs)
    from pptx.util import Emu

    for s in slides:
        # editable: real text frames; the only pictures allowed are small icon
        # chips / logos (like the reference decks), never rasterized content
        for sh in s.shapes:
            if sh.shape_type == 13:
                assert Emu(sh.width).inches <= 2.4, "picture larger than an icon/logo"
                assert Emu(sh.height).inches <= 1.2
        assert any(sh.has_text_frame and sh.text_frame.text.strip() for sh in s.shapes)


def test_native_render_covers_every_primitive(tmp_path):
    outs = [_outline(i, p, f"{p} slide title", ["Point one is a finished thought", "Point two is also complete", "Point three rounds it out"])
            for i, p in enumerate(prim.BUILDERS)]
    out = tmp_path / "all.pptx"
    warnings = NativePptxRenderer().render(outs, _brand(), out)
    assert warnings == []  # every primitive builds without falling back
    assert len(list(Presentation(out.as_posix()).slides)) == len(prim.BUILDERS)


def test_design_language_changes_native_geometry(tmp_path):
    # bold_minimal (radius 4, scale 1.14) vs editorial_serif (radius 16) -> different theme tokens
    bold = th.resolve_theme(_brand("bold_minimal"))
    editorial = th.resolve_theme(_brand("editorial_serif"))
    assert bold.card_radius != editorial.card_radius
    assert th.pt(60, bold.heading_scale) > th.pt(60, editorial.heading_scale)


def test_unknown_renderer_engine_resolves_to_native():
    """A stale RENDERER_ENGINE value (e.g. the removed image-based `html`
    engine) must resolve to the fully editable native renderer."""
    from app.services.pptx_builder import PptxBuilder

    builder = PptxBuilder(node_runner=object(), renderer_engine="html")
    assert builder._renderer() is builder.native_renderer
    builder = PptxBuilder(node_runner=object(), renderer_engine="native")
    assert builder._renderer() is builder.native_renderer
    builder = PptxBuilder(node_runner=object(), renderer_engine="authored")
    assert builder._renderer() is builder.authored_renderer
    builder = PptxBuilder(node_runner=object(), renderer_engine="legacy")
    assert builder._renderer() is builder.legacy_renderer


def test_editability_audit_flags_picture_dominated_slide(tmp_path):
    """The post-build audit warns when a generated slide is mostly a picture
    (rasterized-deck regression guard) and stays silent for native decks."""
    from PIL import Image
    from pptx.util import Inches

    from app.services.pptx_builder import PptxBuilder

    builder = PptxBuilder(node_runner=object(), renderer_engine="native")

    # Native deck: zero pictures -> no editability warnings.
    outs = [_outline(0, "cards", "A fully editable native slide title", [
        "Point one is a finished thought", "Point two is also complete"])]
    native_path = tmp_path / "native.pptx"
    NativePptxRenderer().render(outs, _brand(), native_path)
    assert builder._editability_audit(native_path) == []

    # Deck with a near-full-bleed picture -> flagged.
    img_path = tmp_path / "img.png"
    Image.new("RGB", (32, 32), "navy").save(img_path)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(img_path.as_posix(), Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    flat_path = tmp_path / "flat.pptx"
    prs.save(flat_path.as_posix())
    warnings = builder._editability_audit(flat_path)
    assert len(warnings) == 1
    assert warnings[0]["field"] == "editability"


def test_untitled_items_render_full_body_not_chopped(tmp_path):
    """An untitled item's full sentence must reach the slide — never a 50-60
    char pseudo-title with the remainder discarded."""
    long_body = (
        "The recommended shift is from manual benchmark creation to governed "
        "evidence discovery inside validated workflows"
    )
    for primitive in ("rows", "callout_list", "layers", "timeline", "columns", "matrix", "cards"):
        outs = [_outline(0, primitive, "A grammatical action title states the point", [
            {"title": "Lead point", "body": "A titled point keeps its lead and body."},
            {"title": "", "body": long_body},
        ])]
        out = tmp_path / f"{primitive}.pptx"
        NativePptxRenderer().render(outs, _brand(), out)
        texts = " ".join(
            sh.text_frame.text for s in Presentation(out.as_posix()).slides
            for sh in s.shapes if sh.has_text_frame
        )
        assert "governed evidence discovery inside validated workflows" in texts, primitive


def test_icon_circles_carry_glyph_or_monogram(tmp_path):
    from app.services.pptx_native.primitives import _item_glyph

    assert _item_glyph({"title": "Validated proof points", "body": ""}) == "✓"
    assert _item_glyph({"title": "Execution risk", "body": ""}) == "!"
    assert _item_glyph({"icon": "arrow-shift", "title": "", "body": "x"}) == "→"
    # monogram fallback: first letter of the lead
    assert _item_glyph({"title": "Data extraction", "body": ""}) == "D"
    assert _item_glyph({"title": "", "body": "memory files feed the loop"}) == "M"


def test_icon_resolver_maps_hints_keywords_and_rotates():
    from app.services.pptx_native import icons

    # explicit react-icons hint wins
    assert icons.resolve_icon("FiZap", "anything", "", 0) == "FiZap"
    # keyword stems over the lead
    assert icons.resolve_icon("", "Benchmark saturation", "", 0) == "FiTrendingUp"
    assert icons.resolve_icon("risk", "", "", 0) == "FiAlertTriangle"
    assert icons.resolve_icon("", "Prohibitive curation cost", "", 0) == "FiDollarSign"
    # no signal -> rotating defaults, varied across items
    a = icons.resolve_icon("", "zzz", "zzz", 0)
    b = icons.resolve_icon("", "zzz", "zzz", 1)
    assert a != b


def test_native_cards_carry_icon_chips_when_node_available(tmp_path):
    import shutil

    import pytest

    from app.services.pptx_native import icons

    if shutil.which("node") is None or icons.icon_file("FiTarget", "FFFFFF") is None:
        pytest.skip("node icon worker unavailable")
    outs = [_outline(0, "cards", "Three failure modes erode benchmark value", [
        {"title": "Benchmark saturation", "body": "Frontier training pollutes public benchmarks."},
        {"title": "Weak correlation", "body": "Scores rarely predict performance on actual use cases."},
        {"title": "Curation cost", "body": "Expert-labeled ground truth is expensive to maintain."},
    ])]
    out = tmp_path / "icons.pptx"
    NativePptxRenderer().render(outs, _brand(), out)
    from pptx.util import Emu

    pics = [sh for s in Presentation(out.as_posix()).slides for sh in s.shapes if sh.shape_type == 13]
    assert len(pics) == 3  # one icon chip per card
    assert all(Emu(p.width).inches <= 0.6 for p in pics)


def test_native_renderer_places_brand_logo_on_cover(tmp_path):
    from PIL import Image

    logo = tmp_path / "logo.png"
    Image.new("RGB", (200, 80), "navy").save(logo)
    from app.models.brand import BrandLogo

    brand = _brand()
    brand.logo = BrandLogo(path=logo.as_posix())
    outs = [
        _outline(0, "cover", "Brand Deck", []),
        _outline(1, "cards", "A content slide keeps the canvas clean", [
            {"title": "Point", "body": "Body sentence for the card."}]),
    ]
    out = tmp_path / "logo.pptx"
    NativePptxRenderer().render(outs, brand, out)
    prs = Presentation(out.as_posix())
    cover_pics = [sh for sh in prs.slides[0].shapes if sh.shape_type == 13]
    assert cover_pics, "cover carries the brand logo"


def test_brand_mode_renders_through_native_layouts_by_default(tmp_path):
    """brand_render_mode='native' (default): brand decks use the same layout
    system as freeform, themed by the template's BrandDNA — no clone path."""
    from app.models.template import TemplateProfile
    from app.services.pptx_builder import PptxBuilder

    template = TemplateProfile(
        id="b", name="b", type="brand", brand=_brand(), slides=[],
        source_file="", created_at="t", updated_at="t",
    )
    builder = PptxBuilder(node_runner=object(), renderer_engine="native")
    assert builder.brand_render_mode == "native"
    outs = [
        _outline(0, "cover", "Brand Deck Title", []),
        _outline(1, "cards", "Brand slides share the freeform layout system", [
            {"title": "Consistency", "body": "Both modes produce the same design language."}]),
    ]
    out = tmp_path / "brand.pptx"
    warnings = builder.build_deck(template, outs, out, tmp_path)
    assert out.exists()
    assert not [w for w in warnings if w.get("field") == "editability"]
    prs = Presentation(out.as_posix())
    assert len(prs.slides) == 2
    assert any(sh.has_text_frame and "layout system" in sh.text_frame.text
               for sh in prs.slides[1].shapes)
