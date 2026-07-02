from datetime import UTC, datetime
import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.config import Settings
from app.infra.local_storage import LocalStorage
from app.infra.sqlite_store import SQLiteStore
from app.models.brand import BrandDNA
from app.models.document import DocumentBundle, DocumentMetadata, DocumentSection
from app.models.job import FREEFORM_TEMPLATE_ID, JobRecord
from app.models.outline import SlideOutline
from app.models.qa import QAIssue, QAResult
from app.models.template import SlideField, SlideSchema, SlideSpec, TemplateProfile
from app.services.content_planner import ContentPlanner
from app.services.design_agent import DesignAgent
from app.services.orchestrator import JobOrchestrator
from app.services.pptx_builder import PptxBuilder
from app.services.visual_qa_agent import VisualQAAgent


class StaticIngester:
    def ingest_documents(self, job_id: str, file_paths: list[Path]) -> tuple[list, DocumentBundle]:
        return [], DocumentBundle(
            job_id=job_id,
            sections=[
                DocumentSection(
                    title="Vibe Coding Risk",
                    level=1,
                    content=(
                        "Vibe coding accelerates prototypes but fails in production "
                        "when teams lose context, skip review, and cannot reproduce AI decisions."
                    ),
                    source_doc_id="doc-1",
                ),
                DocumentSection(
                    title="External Brain",
                    level=1,
                    content=(
                        "A memory bank and structured specification workflow preserve "
                        "context so AI agents can deliver reliable software across sessions."
                    ),
                    source_doc_id="doc-1",
                ),
            ],
            tables=[],
            metrics=[],
            metadata=DocumentMetadata(title="Beyond Vibe Coding"),
            content_inventory=[],
        )


class WarningThenCleanQAAgent:
    def __init__(self) -> None:
        self.calls = 0

    def inspect_deck(self, output_path, output_dir, outlines):
        self.calls += 1
        if self.calls == 1:
            return (
                QAResult(
                    passed=True,
                    issues=[
                        QAIssue(
                            severity="WARNING",
                            category="text_wall",
                            message="Slide reads as a text wall.",
                            slide_index=0,
                        )
                    ],
                ),
                [],
            )
        return QAResult(passed=True, issues=[]), []

    def export_pdf(self, output_path, working_dir):
        return None


class CleanQAAgent:
    def inspect_deck(self, output_path, output_dir, outlines):
        return QAResult(passed=True, issues=[]), []

    def export_pdf(self, output_path, working_dir):
        return None


class PersistentCriticalQAAgent:
    def __init__(self) -> None:
        self.calls = 0

    def inspect_deck(self, output_path, output_dir, outlines):
        self.calls += 1
        return (
            QAResult(
                passed=False,
                issues=[
                    QAIssue(
                        severity="CRITICAL",
                        category="content_quality",
                        message="Rendered slide still contains incomplete text.",
                        slide_index=0,
                    )
                ],
            ),
            [],
        )

    def export_pdf(self, output_path, working_dir):
        return None


class RecordingSQLiteStore(SQLiteStore):
    def __init__(self, settings: Settings) -> None:
        super().__init__(settings)
        self.status_updates: list[str] = []

    async def update_job(self, job_id: str, **fields) -> None:
        if "status" in fields:
            self.status_updates.append(str(fields["status"]))
        await super().update_job(job_id, **fields)


class NonActionableWarningQAAgent:
    def inspect_deck(self, output_path, output_dir, outlines):
        return (
            QAResult(
                passed=True,
                issues=[
                    QAIssue(
                        severity="INFO",
                        category="render_fallback",
                        message="Used approximate preview rendering.",
                        slide_index=None,
                    )
                ],
            ),
            [],
        )

    def export_pdf(self, output_path, working_dir):
        return None


class WarningEditingContract:
    def build(
        self,
        template: TemplateProfile,
        outlines: list[SlideOutline],
        phase: str = "planned",
    ) -> dict:
        return {
            "artifact": "editing-contract",
            "standard": "claude-pptx-editing-v1",
            "phase": phase,
            "status": "warning",
            "issue_count": 1,
            "requirements": [
                {
                    "id": "varied_composition_families",
                    "label": "Vary visible composition families",
                    "status": "warning",
                    "message": "Only one visible composition family used.",
                }
            ],
            "warnings": ["Only one visible composition family used."],
        }


class CountingIngester(StaticIngester):
    def __init__(self) -> None:
        self.calls = 0

    def ingest_documents(self, job_id: str, file_paths: list[Path]) -> tuple[list, DocumentBundle]:
        self.calls += 1
        return super().ingest_documents(job_id, file_paths)


class CountingPlanner(ContentPlanner):
    def __init__(self) -> None:
        super().__init__()
        self.calls = 0

    def plan(
        self,
        template: TemplateProfile,
        bundle: DocumentBundle,
        instructions: str = "",
        generation_mode: str | None = None,
        quality_profile: str = "balanced",
        length_strategy: str = "auto",
        presentation_style: str = "consulting",
        design_language: str = "editorial_serif",
    ):
        self.calls += 1
        return super().plan(
            template,
            bundle,
            instructions=instructions,
            generation_mode=generation_mode,
            quality_profile=quality_profile,
            length_strategy=length_strategy,
            presentation_style=presentation_style,
            design_language=design_language,
        )


class DuplicateOutlinePlanner(ContentPlanner):
    def plan(
        self,
        template: TemplateProfile,
        bundle: DocumentBundle,
        instructions: str = "",
        generation_mode: str | None = None,
        quality_profile: str = "balanced",
        length_strategy: str = "auto",
        presentation_style: str = "consulting",
        design_language: str = "editorial_serif",
    ):
        created_at = "2026-01-01T00:00:00Z"
        outlines = []
        for index in range(2):
            outlines.append(
                SlideOutline(
                    id=f"weak-outline-{index}",
                    job_id=bundle.job_id,
                    slide_index=index,
                    mode="flexible",
                    label="Overview",
                    content_json={
                        "action_title": "Overview",
                        "title": "Overview",
                        "narrative_role": "evidence",
                        "archetype": "comparison_table",
                        "bullets": ["Preserve context before work begins"],
                        "content_blocks": [
                            {
                                "type": "bullets",
                                "body": ["Preserve context before work begins"],
                            }
                        ],
                        "sources": ["Uploaded source"],
                    },
                    layout_json={
                        "layout": "comparison_table",
                        "archetype": "comparison_table",
                    },
                    created_at=created_at,
                )
            )
        return outlines, []


class ArtifactFailingPlanner(ContentPlanner):
    def plan(
        self,
        template: TemplateProfile,
        bundle: DocumentBundle,
        instructions: str = "",
        generation_mode: str | None = None,
        quality_profile: str = "balanced",
        length_strategy: str = "auto",
        presentation_style: str = "consulting",
        design_language: str = "editorial_serif",
    ):
        self.last_planning_artifacts = {
            "story-map": {
                "status": "unavailable",
                "fallback_reason": "model timed out",
            }
        }
        raise RuntimeError("planner unavailable")


def _timestamp() -> str:
    return datetime.now(UTC).isoformat().replace("+00:00", "Z")


def _settings(tmp_path: Path) -> Settings:
    return Settings(
        data_dir=tmp_path / "data",
        sqlite_path=tmp_path / "data" / "jobs.sqlite",
        templates_dir=tmp_path / "data" / "templates",
        jobs_dir=tmp_path / "data" / "jobs",
        LLM_PROVIDER="none",
        BEDROCK_VALIDATE=False,
    )


def test_orchestrator_selects_deep_planner_profile(tmp_path: Path) -> None:
    settings = Settings(
        data_dir=tmp_path / "data",
        sqlite_path=tmp_path / "data" / "jobs.sqlite",
        templates_dir=tmp_path / "data" / "templates",
        jobs_dir=tmp_path / "data" / "jobs",
        LLM_PROVIDER="openai_compatible",
        DEEP_PLANNER_BASE_URL="http://athena.local:1240/v1",
        DEEP_PLANNER_MODEL="minimax-m2.7",
        BEDROCK_VALIDATE=False,
    )
    store = SQLiteStore(settings)
    storage = LocalStorage(settings)
    orchestrator = _orchestrator(settings, store, storage)
    job = _job("deep-job", FREEFORM_TEMPLATE_ID, "freeform")
    job.config_json = {"generation_mode": "freeform", "planner_profile": "deep"}

    planner = orchestrator._planner_for_job(job)

    assert planner is orchestrator.deep_planner
    assert planner.llm_client is not None
    assert planner.llm_client.base_url == "http://athena.local:1240/v1"
    assert planner.llm_client.model == "minimax-m2.7"
    # Deep profile authors slides one-by-one; fast profile batches.
    assert planner.slide_generation_strategy == "per_slide"
    assert orchestrator._planner_for_job(
        _job("fast-job", FREEFORM_TEMPLATE_ID, "freeform")
    ).slide_generation_strategy == "batched"


def _job(job_id: str, template_id: str, mode: str) -> JobRecord:
    return JobRecord(
        id=job_id,
        template_id=template_id,
        instructions="Create an executive deck about moving beyond vibe coding.",
        config_json={"generation_mode": mode},
        status="queued",
        progress=0,
        qa_rounds=0,
        warnings=[],
        created_at=_timestamp(),
    )


def _orchestrator(settings: Settings, store: SQLiteStore, storage: LocalStorage) -> JobOrchestrator:
    return JobOrchestrator(
        settings=settings,
        store=store,
        storage=storage,
        ingester=StaticIngester(),
        planner=ContentPlanner(),
        designer=DesignAgent(),
        builder=PptxBuilder(node_runner=object()),
        qa_agent=VisualQAAgent(),
    )


def _orchestrator_with_qa(
    settings: Settings,
    store: SQLiteStore,
    storage: LocalStorage,
    qa_agent,
) -> JobOrchestrator:
    return JobOrchestrator(
        settings=settings,
        store=store,
        storage=storage,
        ingester=StaticIngester(),
        planner=ContentPlanner(),
        designer=DesignAgent(),
        builder=PptxBuilder(node_runner=object()),
        qa_agent=qa_agent,
    )


@pytest.mark.asyncio
async def test_orchestrator_runs_freeform_job_without_template(tmp_path: Path) -> None:
    settings = _settings(tmp_path)
    store = SQLiteStore(settings)
    storage = LocalStorage(settings)
    await store.init()
    await store.create_job(_job("freeform-job", FREEFORM_TEMPLATE_ID, "freeform"))

    await _orchestrator_with_qa(settings, store, storage, CleanQAAgent()).run_job(
        "freeform-job"
    )

    job = await store.get_job("freeform-job")
    outlines = await store.list_slide_outlines("freeform-job")

    assert job is not None
    assert job.status == "done"
    assert job.result_file
    assert Path(job.result_file).exists()
    assert outlines
    assert {outline.layout_json["generation_mode"] for outline in outlines} == {"freeform"}
    assert len(Presentation(job.result_file).slides) == len(outlines)
    planning_dir = storage.job_dir("freeform-job") / "planning"
    assert (planning_dir / "source-compression.json").exists()
    assert (planning_dir / "story-map.json").exists()
    assert (planning_dir / "spec-gate.json").exists()


@pytest.mark.asyncio
async def test_orchestrator_persists_planning_artifacts_when_planner_fails(
    tmp_path: Path,
) -> None:
    settings = _settings(tmp_path)
    store = SQLiteStore(settings)
    storage = LocalStorage(settings)
    await store.init()
    await store.create_job(_job("planner-error-job", FREEFORM_TEMPLATE_ID, "freeform"))

    orchestrator = JobOrchestrator(
        settings=settings,
        store=store,
        storage=storage,
        ingester=StaticIngester(),
        planner=ArtifactFailingPlanner(),
        designer=DesignAgent(),
        builder=PptxBuilder(node_runner=object()),
        qa_agent=CleanQAAgent(),
    )
    await orchestrator.run_job("planner-error-job")

    job = await store.get_job("planner-error-job")
    assert job is not None
    assert job.status == "error"
    assert job.error_message == "planner unavailable"
    assert storage.planning_artifact_path("planner-error-job", "story-map").exists()


@pytest.mark.asyncio
async def test_orchestrator_plan_only_persists_plan_without_rendering(tmp_path: Path) -> None:
    settings = _settings(tmp_path)
    store = SQLiteStore(settings)
    storage = LocalStorage(settings)
    await store.init()
    job = _job("plan-only-job", FREEFORM_TEMPLATE_ID, "freeform")
    job.config_json = {"generation_mode": "freeform", "plan_only": True}
    await store.create_job(job)

    await _orchestrator_with_qa(settings, store, storage, CleanQAAgent()).run_job("plan-only-job")

    planned = await store.get_job("plan-only-job")
    outlines = await store.list_slide_outlines("plan-only-job")
    planning_dir = storage.job_dir("plan-only-job") / "planning"

    assert planned is not None
    assert planned.status == "planned"
    assert planned.progress == 0.5
    assert planned.result_file is None
    assert planned.preview_dir is None
    assert outlines
    assert all(outline.layout_json.get("composition_family") for outline in outlines)
    assert (planning_dir / "document-bundle.json").exists()
    assert (planning_dir / "source-compression.json").exists()
    assert (planning_dir / "editing-contract.json").exists()
    assert not (storage.job_dir("plan-only-job") / "output.pptx").exists()
    assert not list((storage.job_dir("plan-only-job") / "preview").glob("slide-*.jpg"))


@pytest.mark.asyncio
async def test_orchestrator_renders_from_persisted_plan_without_replanning(tmp_path: Path) -> None:
    settings = _settings(tmp_path)
    store = SQLiteStore(settings)
    storage = LocalStorage(settings)
    ingester = CountingIngester()
    planner = CountingPlanner()
    await store.init()
    job = _job("resume-plan-job", FREEFORM_TEMPLATE_ID, "freeform")
    job.config_json = {"generation_mode": "freeform", "plan_only": True}
    await store.create_job(job)
    orchestrator = JobOrchestrator(
        settings=settings,
        store=store,
        storage=storage,
        ingester=ingester,
        planner=planner,
        designer=DesignAgent(),
        builder=PptxBuilder(node_runner=object()),
        qa_agent=CleanQAAgent(),
    )

    await orchestrator.run_job("resume-plan-job")
    planned = await store.get_job("resume-plan-job")
    assert planned is not None
    assert planned.status == "planned"
    persisted_bundle = storage.load_document_bundle("resume-plan-job")
    assert persisted_bundle is not None

    config = dict(planned.config_json or {})
    config["plan_only"] = False
    config["render_from_plan"] = True
    await store.update_job(
        "resume-plan-job",
        status="queued",
        progress=0.5,
        config_json=config,
        result_file=None,
        preview_dir=None,
        completed_at=None,
    )

    await orchestrator.run_job("resume-plan-job")

    rendered = await store.get_job("resume-plan-job")
    assert rendered is not None
    assert rendered.status == "done"
    assert rendered.result_file
    assert Path(rendered.result_file).exists()
    assert ingester.calls == 1
    assert planner.calls == 1
    assert storage.load_document_bundle("resume-plan-job") == persisted_bundle


@pytest.mark.asyncio
async def test_orchestrator_repairs_actionable_warning_and_persists_outline(
    tmp_path: Path,
) -> None:
    settings = _settings(tmp_path)
    store = RecordingSQLiteStore(settings)
    storage = LocalStorage(settings)
    qa_agent = WarningThenCleanQAAgent()
    await store.init()
    await store.create_job(_job("repair-job", FREEFORM_TEMPLATE_ID, "freeform"))

    await _orchestrator_with_qa(settings, store, storage, qa_agent).run_job("repair-job")

    job = await store.get_job("repair-job")
    outlines = await store.list_slide_outlines("repair-job")

    assert job is not None
    assert job.status == "done"
    assert job.qa_rounds == 1
    assert qa_agent.calls == 2
    assert outlines[0].layout_json["qa_repair"]["applied"] is True
    assert outlines[0].qa_status == "pass"
    assert "repairing" in store.status_updates
    assert store.status_updates.count("qa") >= 2
    assert (storage.job_dir("repair-job") / "qa" / "round-0.json").exists()
    assert (storage.job_dir("repair-job") / "qa" / "round-1.json").exists()
    round_0 = json.loads(
        (storage.job_dir("repair-job") / "qa" / "round-0.json").read_text(
            encoding="utf-8"
        )
    )
    round_1 = json.loads(
        (storage.job_dir("repair-job") / "qa" / "round-1.json").read_text(
            encoding="utf-8"
        )
    )
    assert round_0["repair_applied"] is True
    assert round_0["actionable_issue_count"] == 1
    assert round_0["stop_reason"] == "repair_applied"
    assert round_1["repair_applied"] is False
    assert round_1["actionable_issue_count"] == 0
    assert round_1["stop_reason"] == "no_actionable_issues"


@pytest.mark.asyncio
async def test_orchestrator_logs_non_actionable_qa_without_repair(
    tmp_path: Path,
) -> None:
    settings = _settings(tmp_path)
    store = SQLiteStore(settings)
    storage = LocalStorage(settings)
    await store.init()
    await store.create_job(_job("non-actionable-qa-job", FREEFORM_TEMPLATE_ID, "freeform"))

    await _orchestrator_with_qa(
        settings,
        store,
        storage,
        NonActionableWarningQAAgent(),
    ).run_job("non-actionable-qa-job")

    job = await store.get_job("non-actionable-qa-job")
    round_0 = json.loads(
        (storage.job_dir("non-actionable-qa-job") / "qa" / "round-0.json").read_text(
            encoding="utf-8"
        )
    )

    assert job is not None
    assert job.status == "done"
    assert job.qa_rounds == 0
    assert round_0["repair_applied"] is False
    assert round_0["actionable_issue_count"] == 0
    assert round_0["stop_reason"] == "no_actionable_issues"


@pytest.mark.asyncio
async def test_orchestrator_marks_review_failed_when_final_editing_contract_warns(
    tmp_path: Path,
) -> None:
    settings = _settings(tmp_path)
    store = SQLiteStore(settings)
    storage = LocalStorage(settings)
    await store.init()
    await store.create_job(_job("contract-review-failed-job", FREEFORM_TEMPLATE_ID, "freeform"))
    orchestrator = _orchestrator_with_qa(settings, store, storage, CleanQAAgent())
    orchestrator.editing_contract = WarningEditingContract()

    await orchestrator.run_job("contract-review-failed-job")

    job = await store.get_job("contract-review-failed-job")
    assert job is not None
    assert job.status == "review_failed"
    assert job.result_file
    assert Path(job.result_file).exists()
    assert job.error_message
    assert "editing contract" in job.error_message
    assert any(warning["field"] == "editing_contract" for warning in job.warnings)
    final_contract = json.loads(
        (storage.job_dir("contract-review-failed-job") / "planning" / "editing-contract.json").read_text(
            encoding="utf-8"
        )
    )
    assert final_contract["phase"] == "final"


@pytest.mark.asyncio
async def test_orchestrator_marks_review_failed_when_final_qa_still_fails(
    tmp_path: Path,
) -> None:
    settings = _settings(tmp_path)
    store = SQLiteStore(settings)
    storage = LocalStorage(settings)
    qa_agent = PersistentCriticalQAAgent()
    await store.init()
    await store.create_job(_job("final-review-failed-job", FREEFORM_TEMPLATE_ID, "freeform"))

    await _orchestrator_with_qa(settings, store, storage, qa_agent).run_job(
        "final-review-failed-job"
    )

    job = await store.get_job("final-review-failed-job")
    assert job is not None
    assert job.status == "review_failed"
    assert job.result_file
    assert Path(job.result_file).exists()
    assert job.error_message and "failed final review" in job.error_message
    assert list((storage.job_dir("final-review-failed-job") / "qa").glob("round-*.json"))
    assert (
        storage.job_dir("final-review-failed-job") / "outline" / "render-input.json"
    ).exists()
    assert (
        storage.job_dir("final-review-failed-job") / "outline" / "repair-round-1.json"
    ).exists()


@pytest.mark.asyncio
async def test_orchestrator_repairs_consulting_issues_before_build(
    tmp_path: Path,
) -> None:
    settings = _settings(tmp_path)
    store = SQLiteStore(settings)
    storage = LocalStorage(settings)
    await store.init()
    await store.create_job(_job("consulting-repair-job", FREEFORM_TEMPLATE_ID, "freeform"))
    orchestrator = JobOrchestrator(
        settings=settings,
        store=store,
        storage=storage,
        ingester=StaticIngester(),
        planner=DuplicateOutlinePlanner(),
        designer=DesignAgent(),
        builder=PptxBuilder(node_runner=object()),
        qa_agent=CleanQAAgent(),
    )

    await orchestrator.run_job("consulting-repair-job")

    job = await store.get_job("consulting-repair-job")
    outlines = await store.list_slide_outlines("consulting-repair-job")

    assert job is not None
    assert job.status == "done"
    assert job.result_file
    assert Path(job.result_file).exists()
    assert len({outline.label for outline in outlines}) == len(outlines)
    assert all(outline.label != "Overview" for outline in outlines)
    assert all(outline.content_json.get("source_refs") for outline in outlines)
    assert all(outline.content_json.get("exhibit_spec") for outline in outlines)
    assert any(warning["field"] == "consulting_qa" for warning in job.warnings)


@pytest.mark.asyncio
async def test_orchestrator_runs_brand_job_from_template_profile(tmp_path: Path) -> None:
    settings = _settings(tmp_path)
    store = SQLiteStore(settings)
    storage = LocalStorage(settings)
    await store.init()
    template = TemplateProfile(
        id="brand-template",
        name="Brand Template",
        type="brand",
        brand=BrandDNA(primary_color="#111827", secondary_color="#2563eb"),
        slides=[],
        source_file="",
        created_at=_timestamp(),
        updated_at=_timestamp(),
    )
    await store.create_template(template)
    await store.create_job(_job("brand-job", "brand-template", "brand"))

    await _orchestrator_with_qa(settings, store, storage, CleanQAAgent()).run_job(
        "brand-job"
    )

    job = await store.get_job("brand-job")
    outlines = await store.list_slide_outlines("brand-job")

    assert job is not None
    assert job.status == "done"
    assert job.result_file
    assert Path(job.result_file).exists()
    assert outlines
    assert {outline.layout_json["generation_mode"] for outline in outlines} == {"brand"}
    assert len(Presentation(job.result_file).slides) == len(outlines)


@pytest.mark.asyncio
async def test_orchestrator_runs_strict_job_with_xml_injection(tmp_path: Path) -> None:
    settings = _settings(tmp_path)
    store = SQLiteStore(settings)
    storage = LocalStorage(settings)
    await store.init()
    template_path = tmp_path / "strict-template.pptx"
    _write_strict_template(template_path)
    template = TemplateProfile(
        id="strict-template",
        name="Strict Template",
        type="strict",
        brand=BrandDNA(),
        source_file=template_path.as_posix(),
        slides=[
            SlideSpec(
                index=0,
                mode="strict",
                label="Executive Summary",
                schema=SlideSchema(
                    fields=[
                        SlideField(
                            id="project_title",
                            type="text",
                            location="shape:ProjectTitle",
                            required=True,
                            max_chars=90,
                        ),
                        SlideField(
                            id="executive_summary",
                            type="text",
                            location="shape:ExecutiveSummary",
                            required=True,
                            max_chars=260,
                        ),
                        SlideField(
                            id="key_implication",
                            type="text",
                            location="shape:KeyImplication",
                            required=True,
                            max_chars=260,
                        ),
                    ]
                ),
            )
        ],
        created_at=_timestamp(),
        updated_at=_timestamp(),
    )
    await store.create_template(template)
    await store.create_job(_job("strict-job", "strict-template", "strict"))

    await _orchestrator_with_qa(settings, store, storage, CleanQAAgent()).run_job(
        "strict-job"
    )

    job = await store.get_job("strict-job")
    outlines = await store.list_slide_outlines("strict-job")

    assert job is not None
    assert job.status == "done"
    assert job.result_file
    rendered = Presentation(job.result_file)
    texts = [shape.text for shape in rendered.slides[0].shapes if shape.has_text_frame]
    assert outlines[0].mode == "strict"
    assert "Beyond Vibe Coding" in texts
    assert any("Vibe coding" in text for text in texts)


def _write_strict_template(path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for name, text, top in [
        ("ProjectTitle", "Old project title", 0.7),
        ("ExecutiveSummary", "Old summary", 1.6),
        ("KeyImplication", "Old implication", 4.2),
    ]:
        box = slide.shapes.add_textbox(Inches(0.8), Inches(top), Inches(11.5), Inches(0.8))
        box.name = name
        box.text = text
    prs.save(path.as_posix())


@pytest.mark.asyncio
async def test_orchestrator_freeform_job_with_native_engine(tmp_path: Path) -> None:
    """Full pipeline (ingest -> plan -> design -> build -> QA) with the default
    native renderer produces a terminal, EDITABLE deck (real shapes, no images)."""
    settings = _settings(tmp_path)
    store = SQLiteStore(settings)
    storage = LocalStorage(settings)
    await store.init()
    await store.create_job(_job("native-job", FREEFORM_TEMPLATE_ID, "freeform"))

    orchestrator = JobOrchestrator(
        settings=settings,
        store=store,
        storage=storage,
        ingester=StaticIngester(),
        planner=ContentPlanner(),
        designer=DesignAgent(),
        builder=PptxBuilder(node_runner=object(), renderer_engine="native"),
        qa_agent=CleanQAAgent(),
    )
    await orchestrator.run_job("native-job")

    job = await store.get_job("native-job")
    outlines = await store.list_slide_outlines("native-job")
    assert job is not None
    assert job.status in {"done", "review_failed"}
    assert job.result_file and Path(job.result_file).exists()
    prs = Presentation(job.result_file)
    assert len(prs.slides) == len(outlines)
    # native engine emits editable shapes/text; only small icon/logo pictures
    # are allowed, never a rasterized slide
    from pptx.util import Emu

    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type == 13:
                assert Emu(shape.width).inches <= 2.4, "picture larger than an icon/logo"
        assert any(shape.has_text_frame and shape.text_frame.text.strip() for shape in slide.shapes)


class UnreachableEndpointPlanner(ContentPlanner):
    """Planner wired to a client whose endpoint preflight fails."""

    def __init__(self) -> None:
        super().__init__()

        class _DeadClient:
            def preflight(self) -> None:
                raise RuntimeError(
                    "LLM endpoint unreachable at http://metis.local:1240/v1 "
                    "(ConnectTimeout: timed out)."
                )

        self.llm_client = _DeadClient()


@pytest.mark.asyncio
async def test_orchestrator_fails_fast_when_llm_endpoint_unreachable(
    tmp_path: Path,
) -> None:
    """An unreachable model server errors the job immediately with a clear
    message instead of black-holing the worker for the generation timeout."""
    settings = _settings(tmp_path)
    store = SQLiteStore(settings)
    storage = LocalStorage(settings)
    await store.init()
    await store.create_job(_job("dead-endpoint-job", FREEFORM_TEMPLATE_ID, "freeform"))

    orchestrator = JobOrchestrator(
        settings=settings,
        store=store,
        storage=storage,
        ingester=StaticIngester(),
        planner=UnreachableEndpointPlanner(),
        designer=DesignAgent(),
        builder=PptxBuilder(node_runner=object()),
        qa_agent=CleanQAAgent(),
    )
    await orchestrator.run_job("dead-endpoint-job")

    job = await store.get_job("dead-endpoint-job")
    assert job is not None
    assert job.status == "error"
    assert "LLM endpoint unreachable" in (job.error_message or "")


@pytest.mark.asyncio
async def test_edit_slide_applies_fields_and_rebuilds(tmp_path: Path) -> None:
    """The review-cockpit edit path: title/points/layout edits persist to the
    outline and the rebuilt deck carries the new text."""
    settings = _settings(tmp_path)
    store = SQLiteStore(settings)
    storage = LocalStorage(settings)
    await store.init()
    await store.create_job(_job("edit-job", FREEFORM_TEMPLATE_ID, "freeform"))

    orchestrator = JobOrchestrator(
        settings=settings,
        store=store,
        storage=storage,
        ingester=StaticIngester(),
        planner=ContentPlanner(),
        designer=DesignAgent(),
        builder=PptxBuilder(node_runner=object(), renderer_engine="native"),
        qa_agent=CleanQAAgent(),
    )
    await orchestrator.run_job("edit-job")
    job = await store.get_job("edit-job")
    assert job.status in {"done", "review_failed"}
    outlines = await store.list_slide_outlines("edit-job")
    target = next(o for o in outlines if o.mode == "flexible" and o.slide_index > 0)

    await orchestrator.edit_slide(
        "edit-job",
        orchestrator.freeform_template(),
        target.slide_index,
        {
            "action_title": "Edited title states the user's own conclusion",
            "subheading": "Edited positioning line explains why this matters now.",
            "points": [
                {"title": "Edited lead", "body": "Edited body sentence carries the user's evidence.", "icon": "target"},
                {"title": "Second lead", "body": "Second body sentence adds a concrete consequence."},
            ],
            "layout": "icon_rows",
        },
    )

    refreshed = await store.list_slide_outlines("edit-job")
    edited = next(o for o in refreshed if o.slide_index == target.slide_index)
    assert edited.content_json["action_title"] == "Edited title states the user's own conclusion"
    assert edited.layout_json["layout"] == "icon_rows"
    assert edited.content_json["exhibit_spec"]["points"][0]["title"] == "Edited lead"

    from pptx import Presentation

    job = await store.get_job("edit-job")
    prs = Presentation(job.result_file)
    texts = " ".join(
        sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame
    )
    assert "Edited title states the user's own conclusion" in texts
    assert "Edited body sentence carries the user's evidence" in texts


@pytest.mark.asyncio
async def test_regenerate_slide_with_guidance_uses_llm_rewrite(tmp_path: Path) -> None:
    """Guided regeneration: the planner client re-authors the slide following
    the user's guidance, and the rebuilt deck carries the rewrite."""
    settings = _settings(tmp_path)
    store = SQLiteStore(settings)
    storage = LocalStorage(settings)
    await store.init()
    await store.create_job(_job("guided-job", FREEFORM_TEMPLATE_ID, "freeform"))

    class GuidedRegenLLM:
        def complete_json(self, **kwargs):
            prompt = kwargs.get("user_prompt", "")
            if "USER GUIDANCE" in prompt:
                assert "make it about risk" in prompt
                return {
                    "action_title": "Unmanaged rollout risk erodes stakeholder trust",
                    "subheading": "Risk framing repositions the slide for the steering committee.",
                    "points": [
                        {"title": "Exposure", "body": "Unreviewed changes reach production without an owner.", "icon": "risk"},
                        {"title": "Mitigation", "body": "A named reviewer gates every consequential change."},
                    ],
                    "layout": "callouts",
                }
            return None

    orchestrator = JobOrchestrator(
        settings=settings,
        store=store,
        storage=storage,
        ingester=StaticIngester(),
        planner=ContentPlanner(),
        designer=DesignAgent(),
        builder=PptxBuilder(node_runner=object(), renderer_engine="native"),
        qa_agent=CleanQAAgent(),
    )
    await orchestrator.run_job("guided-job")
    outlines = await store.list_slide_outlines("guided-job")
    target = next(o for o in outlines if o.mode == "flexible" and o.slide_index > 0)

    orchestrator.planner.llm_client = GuidedRegenLLM()
    await orchestrator.regenerate_slide(
        "guided-job",
        orchestrator.freeform_template(),
        target.slide_index,
        guidance="make it about risk",
    )

    refreshed = await store.list_slide_outlines("guided-job")
    edited = next(o for o in refreshed if o.slide_index == target.slide_index)
    assert edited.content_json["action_title"] == "Unmanaged rollout risk erodes stakeholder trust"
    assert edited.content_json["exhibit_spec"]["points"][0]["icon"] == "risk"


@pytest.mark.asyncio
async def test_regenerate_slide_applies_edits_before_llm_rework(tmp_path: Path) -> None:
    """Save & regenerate: manual edits land first, then the model reworks the
    edited slide; with a rejected rewrite the user's edits still survive."""
    settings = _settings(tmp_path)
    store = SQLiteStore(settings)
    storage = LocalStorage(settings)
    await store.init()
    await store.create_job(_job("edit-regen-job", FREEFORM_TEMPLATE_ID, "freeform"))

    class EchoEditsLLM:
        def complete_json(self, **kwargs):
            prompt = kwargs.get("user_prompt", "")
            # The PRODUCED payload must already carry the user's edit.
            assert "User edited lead" in prompt
            return {
                "action_title": "Reworked claim builds on the user's edit",
                "subheading": "The model refined the user's version of the slide.",
                "points": [
                    {"title": "User edited lead", "body": "The model kept the user's point and sharpened it."},
                ],
                "layout": "callouts",
            }

    orchestrator = JobOrchestrator(
        settings=settings,
        store=store,
        storage=storage,
        ingester=StaticIngester(),
        planner=ContentPlanner(),
        designer=DesignAgent(),
        builder=PptxBuilder(node_runner=object(), renderer_engine="native"),
        qa_agent=CleanQAAgent(),
    )
    await orchestrator.run_job("edit-regen-job")
    outlines = await store.list_slide_outlines("edit-regen-job")
    target = next(o for o in outlines if o.mode == "flexible" and o.slide_index > 0)

    orchestrator.planner.llm_client = EchoEditsLLM()
    await orchestrator.regenerate_slide(
        "edit-regen-job",
        orchestrator.freeform_template(),
        target.slide_index,
        guidance="build on my edits",
        edits={
            "action_title": None,
            "subheading": None,
            "points": [{"title": "User edited lead", "body": "The user's own sentence.", "icon": ""}],
            "layout": None,
        },
    )
    refreshed = await store.list_slide_outlines("edit-regen-job")
    edited = next(o for o in refreshed if o.slide_index == target.slide_index)
    assert edited.content_json["action_title"] == "Reworked claim builds on the user's edit"

    # Rejected rewrite (no client): the edits alone still persist.
    orchestrator.planner.llm_client = None
    await orchestrator.regenerate_slide(
        "edit-regen-job",
        orchestrator.freeform_template(),
        target.slide_index,
        guidance="",
        edits={
            "action_title": "Manual title survives a rejected rewrite",
            "subheading": None,
            "points": None,
            "layout": None,
        },
    )
    refreshed = await store.list_slide_outlines("edit-regen-job")
    edited = next(o for o in refreshed if o.slide_index == target.slide_index)
    assert edited.content_json["action_title"] == "Manual title survives a rejected rewrite"


@pytest.mark.asyncio
async def test_background_style_light_stamps_all_but_cover(tmp_path: Path) -> None:
    """background_style=light: every slide except the cover renders on the
    light background; per-slide edits can still pin a slide afterwards."""
    settings = _settings(tmp_path)
    store = SQLiteStore(settings)
    storage = LocalStorage(settings)
    await store.init()
    job = _job("bg-job", FREEFORM_TEMPLATE_ID, "freeform")
    job.config_json = {"generation_mode": "freeform", "background_style": "light"}
    await store.create_job(job)

    orchestrator = JobOrchestrator(
        settings=settings,
        store=store,
        storage=storage,
        ingester=StaticIngester(),
        planner=ContentPlanner(),
        designer=DesignAgent(),
        builder=PptxBuilder(node_runner=object(), renderer_engine="native"),
        qa_agent=CleanQAAgent(),
    )
    await orchestrator.run_job("bg-job")
    outlines = await store.list_slide_outlines("bg-job")
    assert outlines
    for outline in outlines:
        role = (outline.content_json.get("narrative_role") or outline.content_json.get("slide_type") or "").lower()
        if outline.slide_index == 0 or role == "cover":
            assert outline.content_json.get("background_mode") != "light" or role != "cover"
        else:
            assert outline.content_json.get("background_mode") == "light"

    # Per-slide override back to dark via the edit path.
    target = next(o for o in outlines if o.slide_index > 0 and o.mode == "flexible")
    await orchestrator.edit_slide(
        "bg-job", orchestrator.freeform_template(), target.slide_index,
        {"background": "dark"}, rebuild=False,
    )
    refreshed = await store.list_slide_outlines("bg-job")
    edited = next(o for o in refreshed if o.slide_index == target.slide_index)
    assert edited.content_json["background_mode"] == "dark"
