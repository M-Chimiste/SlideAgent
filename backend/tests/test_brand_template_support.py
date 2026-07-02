import json
import re
import zipfile
from io import BytesIO
from pathlib import Path

from lxml import etree
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.util import Inches

from app.models.brand import BrandDNA
from app.models.outline import SlideOutline
from app.models.template import SlideField, SlideSchema, SlideSpec, TemplateProfile
from app.services.brand_template_renderer import BrandTemplateCloneRenderer
from app.services.generation_editing_contract import GenerationEditingContract
from app.services.pptx_builder import PptxBuilder
from app.services.pptx_renderer import DeterministicPptxRenderer
from app.services.rendered_slide_audit import RenderedSlideAudit
from app.services.template_analyzer import TemplateAnalyzer


def test_brand_dna_accepts_flat_brand_aliases() -> None:
    brand = BrandDNA(
        primary_color="#111827",
        secondary_color="#2563eb",
        accent_color="#16a34a",
        font_headings="Aptos Display",
        font_body="Aptos",
    )

    assert brand.colors.primary == "#111827"
    assert brand.colors.secondary == "#2563eb"
    assert brand.colors.accent == "#16a34a"
    assert brand.fonts.heading == "Aptos Display"
    assert brand.fonts.body == "Aptos"


def test_template_analyzer_extracts_logo_and_layout_notes(tmp_path: Path) -> None:
    template_path = tmp_path / "brand-template.pptx"
    logo_path = tmp_path / "logo.png"
    Image.new("RGB", (180, 60), "#111827").save(logo_path)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(logo_path.as_posix(), Inches(11.0), Inches(0.4), Inches(1.4), Inches(0.45))
    title = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(6), Inches(0.5))
    title.text = "Brand example"
    presentation.save(template_path.as_posix())

    profile, thumbnails = TemplateAnalyzer().analyze(
        template_path,
        template_name="Brand Template",
        template_type="brand",
        template_id="template-1",
    )

    # thumbnails are produced only when render tools (soffice/pdftoppm) are present;
    # the analyzer falls back to [] otherwise, so don't couple this test to the env.
    assert isinstance(thumbnails, list)
    assert profile.brand.logo is not None
    assert Path(profile.brand.logo.path).exists()
    assert "Reusable layout examples" in (profile.brand.design_notes or "")
    assert profile.slides[0].layout_name
    frame_map = json.loads((tmp_path / "frame-map.json").read_text(encoding="utf-8"))
    assert frame_map["artifact"] == "template-frame-map"
    assert frame_map["slide_count"] == 1
    assert frame_map["slot_count"] >= 2
    assert frame_map["slides"][0]["text_inventory"] == "Brand example"
    assert frame_map["slides"][0]["media_slot_count"] == 1
    assert frame_map["slides"][0]["content_category"] in {"content", "media_story", "section_or_cover"}
    assert frame_map["slides"][0]["visual_guidance"]


def test_template_analyzer_prefers_sys_color_lastclr_hex(tmp_path: Path) -> None:
    template_path = tmp_path / "sys-color-template.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(template_path.as_posix())
    with zipfile.ZipFile(template_path, "r") as source_zip:
        entries = {name: source_zip.read(name) for name in source_zip.namelist()}
    theme_xml = entries["ppt/theme/theme1.xml"].decode("utf-8")
    theme_xml = re.sub(
        r"<a:dk1>.*?</a:dk1>",
        '<a:dk1><a:sysClr val="windowText" lastClr="111111"/></a:dk1>',
        theme_xml,
        count=1,
        flags=re.DOTALL,
    )
    entries["ppt/theme/theme1.xml"] = theme_xml.encode("utf-8")
    with zipfile.ZipFile(template_path, "w", zipfile.ZIP_DEFLATED) as target_zip:
        for name, data in entries.items():
            target_zip.writestr(name, data)

    profile, _thumbnails = TemplateAnalyzer().analyze(
        template_path,
        template_name="System Color Template",
        template_type="brand",
        template_id="sys-color-template",
    )

    assert profile.brand.colors.text_dark == "111111"
    assert DeterministicPptxRenderer()._clean_hex("windowText") == "000000"


def test_template_analyzer_adds_semantic_frame_categories_for_brand_mapping(
    tmp_path: Path,
) -> None:
    template_path = tmp_path / "brand-frames.pptx"
    presentation = Presentation()
    media_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    media_title = media_slide.shapes.add_textbox(
        Inches(0.8),
        Inches(0.8),
        Inches(4.5),
        Inches(0.5),
    )
    media_title.text = "Customer story"
    media_slide.shapes.add_shape(
        1,
        Inches(7.4),
        Inches(1.4),
        Inches(3.2),
        Inches(2.4),
    )
    table_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table_title = table_slide.shapes.add_textbox(
        Inches(0.8),
        Inches(0.8),
        Inches(4.5),
        Inches(0.5),
    )
    table_title.text = "Decision grid"
    table = table_slide.shapes.add_table(
        3,
        3,
        Inches(0.9),
        Inches(1.6),
        Inches(8.2),
        Inches(2.0),
    ).table
    table.cell(0, 0).text = "Signal"
    table.cell(0, 1).text = "Current"
    table.cell(0, 2).text = "Target"
    presentation.save(template_path.as_posix())

    profile, _thumbnails = TemplateAnalyzer().analyze(
        template_path,
        template_name="Brand Frames",
        template_type="brand",
        template_id="brand-frames",
    )

    assert profile.slides[0].content_category in {
        "content",
        "media_story",
        "section_or_cover",
    }
    assert profile.slides[1].content_category == "comparison"
    assert "table frame" in (profile.slides[1].visual_guidance or "")
    frame_map = json.loads((tmp_path / "frame-map.json").read_text(encoding="utf-8"))
    assert frame_map["slides"][1]["content_category"] == "comparison"
    assert "table frame" in frame_map["slides"][1]["visual_guidance"]


def test_template_analyzer_does_not_treat_score_language_as_chart_frame(
    tmp_path: Path,
) -> None:
    template_path = tmp_path / "score-language.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(5.8), Inches(0.6)).text = (
        "Scores only matter with source context"
    )
    for index, text in enumerate(
        [
            "Model scores drift when tests are synthetic.",
            "Source-backed evidence keeps evaluation meaningful.",
            "Reviewers need narrative proof, not a generic public test.",
        ]
    ):
        slide.shapes.add_textbox(
            Inches(0.9),
            Inches(1.5 + index * 0.8),
            Inches(6.0),
            Inches(0.5),
        ).text = text
    presentation.save(template_path.as_posix())

    profile, _thumbnails = TemplateAnalyzer().analyze(
        template_path,
        template_name="Score Language",
        template_type="brand",
        template_id="score-language",
    )

    assert profile.slides[0].content_category == "evidence_points"
    frame_map = json.loads((tmp_path / "frame-map.json").read_text(encoding="utf-8"))
    assert frame_map["slides"][0]["content_category"] == "evidence_points"


def test_editing_contract_uses_analyzed_brand_frame_categories_for_mapping(
    tmp_path: Path,
) -> None:
    template_path = tmp_path / "brand-map.pptx"
    presentation = Presentation()
    media_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    media_slide.shapes.add_textbox(
        Inches(0.8),
        Inches(0.8),
        Inches(4.5),
        Inches(0.5),
    ).text = "Customer story"
    media_slide.shapes.add_shape(
        1,
        Inches(7.4),
        Inches(1.4),
        Inches(3.2),
        Inches(2.4),
    )
    table_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table_slide.shapes.add_textbox(
        Inches(0.8),
        Inches(0.8),
        Inches(4.5),
        Inches(0.5),
    ).text = "Decision grid"
    table = table_slide.shapes.add_table(
        3,
        3,
        Inches(0.9),
        Inches(1.6),
        Inches(8.2),
        Inches(2.0),
    ).table
    table.cell(0, 0).text = "Signal"
    table.cell(0, 1).text = "Current"
    table.cell(0, 2).text = "Target"
    presentation.save(template_path.as_posix())
    profile, _thumbnails = TemplateAnalyzer().analyze(
        template_path,
        template_name="Brand Map",
        template_type="brand",
        template_id="brand-map",
    )
    outline = SlideOutline(
        id="outline-comparison",
        job_id="job-brand",
        slide_index=0,
        mode="flexible",
        label="Compare operating choices",
        content_json={
            "action_title": "Compare operating choices before committing",
            "narrative_role": "evidence",
            "exhibit_spec": {
                "type": "comparison_table",
                "rows": [["Signal", "Current", "Target"]],
            },
            "content_blocks": [{"type": "bullets", "body": ["Current vs target"]}],
        },
        layout_json={"layout": "comparison_table", "archetype": "comparison_table"},
        created_at="2026-01-01T00:00:00Z",
    )

    contract = GenerationEditingContract().build(profile, [outline])

    assert contract["slides"][0]["template_slide"]["index"] == 1
    assert contract["slides"][0]["template_slide"]["method"] == "semantic_match"
    assert contract["slides"][0]["template_slide"]["match_confidence"] == "high"
    assert contract["slides"][0]["template_slide"]["match_score"] >= 6
    assert "category:comparison" in contract["slides"][0]["template_slide"]["match_reason"]
    assert contract["slides"][0]["template_slide"]["content_category"] == "comparison"


def test_pptx_builder_attaches_brand_template_frame_metadata(tmp_path: Path) -> None:
    source_file = tmp_path / "brand-source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(
        Inches(0.8),
        Inches(0.8),
        Inches(4.0),
        Inches(0.5),
    ).text = "Comparison frame"
    presentation.save(source_file.as_posix())
    profile = TemplateAnalyzer().analyze(
        source_file,
        template_name="Brand Source",
        template_type="brand",
        template_id="brand-source",
    )[0]
    outline = SlideOutline(
        id="outline-comparison",
        job_id="job-brand",
        slide_index=0,
        mode="flexible",
        label="Compare current and target states",
        content_json={
            "action_title": "Compare current and target states",
            "narrative_role": "evidence",
            "exhibit_spec": {"type": "comparison_table"},
            "content_blocks": [{"type": "bullets", "body": ["Current vs target"]}],
        },
        layout_json={"layout": "comparison_table", "archetype": "comparison_table"},
        created_at="2026-01-01T00:00:00Z",
    )

    prepared = PptxBuilder(node_runner=object(), brand_render_mode="clone").prepare_outlines(profile, [outline])

    frame = prepared[0].layout_json["template_frame"]
    assert frame["source_file"] == source_file.as_posix()
    assert frame["source_slide"] == 1
    assert frame["reuse_mode"] == "duplicate-slide-edit"
    assert frame["match_confidence"] in {"medium", "high"}
    assert frame["match_score"] >= 3
    assert frame["slot_plan"]["action"] in {
        "inspect_template_placeholders",
        "delete_excess_template_elements",
        "fill_template_slots",
        "split_or_summarize_source_items",
    }
    assert "source_file" not in prepared[0].content_json["template_frame"]
    assert prepared[0].content_json["template_frame"]["slot_plan"] == frame["slot_plan"]


def test_pptx_builder_renders_brand_deck_by_cloning_and_editing_source_slide(
    tmp_path: Path,
) -> None:
    source_file = tmp_path / "brand-source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    rail = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(0),
        Inches(0.42),
        Inches(7.5),
    )
    rail.fill.solid()
    rail.fill.fore_color.rgb = RGBColor(0x20, 0x4B, 0x8C)
    rail.line.color.rgb = RGBColor(0x20, 0x4B, 0x8C)
    slide.shapes.add_textbox(
        Inches(0.8),
        Inches(0.7),
        Inches(6.5),
        Inches(0.7),
    ).text = "Template title placeholder"
    slide.shapes.add_textbox(
        Inches(0.9),
        Inches(1.8),
        Inches(5.8),
        Inches(0.8),
    ).text = "Template body placeholder"
    slide.shapes.add_textbox(
        Inches(0.8),
        Inches(7.0),
        Inches(2.2),
        Inches(0.3),
    ).text = "Brand footer"
    presentation.save(source_file.as_posix())
    template = TemplateProfile(
        id="brand-clone",
        name="Brand Clone",
        type="brand",
        brand=BrandDNA(),
        slides=[
            SlideSpec(
                index=0,
                mode="flexible",
                label="Evidence frame",
                layout_name="Evidence",
                content_category="evidence_points",
                visual_guidance="title and body slots with left rail",
            )
        ],
        source_file=source_file.as_posix(),
        created_at="2026-01-01T00:00:00Z",
        updated_at="2026-01-01T00:00:00Z",
    )
    outline = SlideOutline(
        id="outline-1",
        job_id="job-brand",
        slide_index=0,
        mode="flexible",
        label="Generated title replaces the template title",
        content_json={
            "action_title": "Generated title replaces the template title",
            "subheading": "Generated subheading uses the inherited body slot.",
            "content_blocks": [
                {
                    "type": "bullets",
                    "body": [
                        "Generated proof point stays editable.",
                        "Status: Ready for review.",
                    ],
                }
            ],
            "source_refs": ["doc-1:section:1"],
        },
        layout_json={"layout": "callouts", "archetype": "callouts"},
        created_at="2026-01-01T00:00:00Z",
    )
    output_path = tmp_path / "brand-output.pptx"

    warnings = PptxBuilder(node_runner=object(), brand_render_mode="clone").build_deck(
        template,
        [outline],
        output_path,
        tmp_path,
    )

    rendered = Presentation(output_path.as_posix())
    text = "\n".join(
        shape.text
        for shape in rendered.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
    )
    fills = []
    for shape in rendered.slides[0].shapes:
        try:
            fills.append(str(shape.fill.fore_color.rgb))
        except Exception:
            continue
    artifact = json.loads((tmp_path / "template-clone-edit.json").read_text(encoding="utf-8"))
    frame_map = json.loads((tmp_path / "template-frame-map.json").read_text(encoding="utf-8"))
    with zipfile.ZipFile(output_path, "r") as pptx_zip:
        slide_xml = pptx_zip.read("ppt/slides/slide1.xml")
    root = etree.fromstring(slide_xml)
    namespaces = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    bold_run_texts = [
        "".join(run.xpath(".//a:t/text()", namespaces=namespaces))
        for run in root.xpath(".//a:r[a:rPr[@b='1']]", namespaces=namespaces)
    ]
    assert warnings == []
    assert len(rendered.slides) == 1
    assert "Generated title replaces the template title" in text
    assert "Generated subheading uses the inherited body slot." in text
    assert "Status: Ready for review." in text
    assert "Template title placeholder" not in text
    assert "Template body placeholder" not in text
    assert "Brand footer" in text
    assert "204B8C" in fills
    assert artifact["artifact"] == "template-clone-edit"
    assert artifact["mappings"][0]["source_slide"] == 1
    assert frame_map["outputSlideCount"] == 1
    assert frame_map["sourceSlideCount"] == 1
    assert frame_map["omittedSourceSlideCount"] == 0
    assert frame_map["outputSlides"][0]["sourceSlide"] == 1
    assert frame_map["outputSlides"][0]["reuseMode"] == "duplicate-slide-edit"
    assert artifact["mappings"][0]["edit_target_count"] >= 2
    assert artifact["mappings"][0]["bolded_text_run_count"] >= 3
    assert "Generated title replaces the template title" in bold_run_texts
    assert "Generated subheading uses the inherited body slot." in bold_run_texts
    assert "Status:" in bold_run_texts
    actions = {target["action"] for target in artifact["mappings"][0]["editTargets"]}
    assert "rewrite" in actions
    assert artifact["mappings"][0]["editTargets"][0]["shapeId"]


def test_brand_clone_edit_blocks_weak_template_frame_mapping(tmp_path: Path) -> None:
    source_file = tmp_path / "brand-source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(6.5), Inches(0.7)).text = "Template title"
    presentation.save(source_file.as_posix())
    template = TemplateProfile(
        id="brand-template",
        name="Brand Template",
        type="brand",
        brand=BrandDNA(),
        slides=[],
        source_file=source_file.as_posix(),
        created_at="2026-01-01T00:00:00Z",
        updated_at="2026-01-01T00:00:00Z",
    )
    outline = SlideOutline(
        id="outline-weak-frame",
        job_id="job-brand",
        slide_index=0,
        mode="flexible",
        label="Generated claim",
        content_json={"action_title": "Generated claim"},
        layout_json={
            "layout": "proof_strip",
            "template_frame": {
                "index": 0,
                "label": "Fallback frame",
                "method": "low_confidence_match",
                "match_confidence": "low",
                "match_score": 0.1,
                "match_reason": "category mismatch",
                "closest_candidates": [
                    {
                        "index": 0,
                        "source_slide": 1,
                        "label": "Fallback frame",
                        "match_score": 0,
                        "match_reason": "no semantic match",
                    }
                ],
            },
        },
        created_at="2026-01-01T00:00:00Z",
    )
    output_path = tmp_path / "blocked-output.pptx"

    rendered, warnings = BrandTemplateCloneRenderer().render(
        template,
        [outline],
        output_path,
        tmp_path,
    )

    artifact = json.loads((tmp_path / "template-clone-edit.json").read_text(encoding="utf-8"))
    frame_map = json.loads((tmp_path / "template-frame-map.json").read_text(encoding="utf-8"))
    deviation_log = json.loads((tmp_path / "template-deviation-log.json").read_text(encoding="utf-8"))
    assert rendered is False
    assert not output_path.exists()
    assert warnings[0]["field"] == "template_clone_edit"
    assert "clone/edit was blocked" in warnings[0]["message"]
    assert artifact["status"] == "blocked"
    assert artifact["slide_count"] == 1
    assert artifact["mappings"][0]["clone_edit_blocked"] is True
    assert artifact["mappings"][0]["match_confidence"] == "low"
    assert artifact["mappings"][0]["block_reason"] == "method=low_confidence_match"
    assert artifact["mappings"][0]["closest_candidates"][0]["source_slide"] == 1
    assert frame_map["artifact"] == "template-frame-map"
    assert frame_map["outputSlides"][0]["outputSlide"] == 1
    assert frame_map["outputSlides"][0]["sourceSlide"] == 1
    assert frame_map["outputSlides"][0]["reuseMode"] == "blocked"
    assert frame_map["outputSlides"][0]["editTargets"] == []
    assert deviation_log["artifact"] == "template-deviation-log"
    assert deviation_log["status"] == "blocked"
    assert deviation_log["deviation_count"] >= 1
    assert deviation_log["deviations"][0]["type"] == "blocked_clone_edit"
    assert deviation_log["deviations"][0]["closest_candidates"][0]["source_slide"] == 1


def test_pptx_builder_strips_template_frames_before_clone_edit_fallback(tmp_path: Path) -> None:
    class BlockingCloneRenderer:
        def render(self, template, outlines, output_path, working_dir):
            return False, [
                {
                    "slide_index": 0,
                    "field": "template_clone_edit",
                    "message": "blocked",
                    "severity": "WARNING",
                }
            ]

    class CapturingRenderer:
        def __init__(self):
            self.rendered_outlines = []

        def author_outlines(self, outlines):
            return outlines

        def render(self, outlines, brand, output_path, enable_diagrams=True):
            self.rendered_outlines = outlines
            output_path.write_bytes(b"placeholder")
            return []

    source_file = tmp_path / "brand-source.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(source_file.as_posix())
    template = TemplateProfile(
        id="brand-template",
        name="Brand Template",
        type="brand",
        brand=BrandDNA(),
        slides=[
            SlideSpec(
                index=0,
                mode="flexible",
                label="Evidence frame",
                layout_name="Evidence",
                content_category="evidence_points",
                visual_guidance="body slots",
            )
        ],
        source_file=source_file.as_posix(),
        created_at="2026-01-01T00:00:00Z",
        updated_at="2026-01-01T00:00:00Z",
    )
    outline = SlideOutline(
        id="outline-frame-strip",
        job_id="job-brand",
        slide_index=0,
        mode="flexible",
        label="Generated evidence claim",
        content_json={
            "action_title": "Generated evidence claim",
            "narrative_role": "evidence",
            "exhibit_spec": {"type": "callouts", "points": ["Generated point."]},
        },
        layout_json={"layout": "callouts", "archetype": "callouts"},
        created_at="2026-01-01T00:00:00Z",
    )
    builder = PptxBuilder(node_runner=object(), brand_render_mode="clone")
    capture = CapturingRenderer()
    builder.brand_template_renderer = BlockingCloneRenderer()
    builder.authored_renderer = capture

    warnings = builder.build_deck(
        template,
        [outline],
        tmp_path / "fallback-output.pptx",
        tmp_path,
    )

    assert warnings[0]["message"] == "blocked"
    assert capture.rendered_outlines
    assert "template_frame" not in capture.rendered_outlines[0].layout_json
    assert "template_frame" not in capture.rendered_outlines[0].content_json


def test_brand_clone_edit_cleans_unused_template_slide_parts(tmp_path: Path) -> None:
    source_file = tmp_path / "brand-clean-source.pptx"
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[6])
    first.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(6.5), Inches(0.7)).text = (
        "Reusable title placeholder"
    )
    first.shapes.add_textbox(Inches(0.9), Inches(1.8), Inches(5.8), Inches(0.8)).text = (
        "Reusable body placeholder"
    )
    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    second.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(6.5), Inches(0.7)).text = (
        "Unused template slide should be removed"
    )
    presentation.save(source_file.as_posix())
    template = TemplateProfile(
        id="brand-clean",
        name="Brand Clean",
        type="brand",
        brand=BrandDNA(),
        slides=[
            SlideSpec(
                index=0,
                mode="flexible",
                label="Evidence frame",
                layout_name="Evidence",
                content_category="evidence_points",
                visual_guidance="title and body slots",
            ),
            SlideSpec(
                index=1,
                mode="flexible",
                label="Unused frame",
                layout_name="Unused",
                content_category="evidence_points",
                visual_guidance="unused frame",
            ),
        ],
        source_file=source_file.as_posix(),
        created_at="2026-01-01T00:00:00Z",
        updated_at="2026-01-01T00:00:00Z",
    )
    outline = SlideOutline(
        id="outline-clean",
        job_id="job-brand",
        slide_index=0,
        mode="flexible",
        label="Generated title replaces the template title",
        content_json={
            "action_title": "Generated title replaces the template title",
            "subheading": "Generated subheading uses the inherited body slot.",
            "content_blocks": [{"type": "bullets", "body": ["Generated proof point."]}],
        },
        layout_json={"layout": "callouts", "archetype": "callouts"},
        created_at="2026-01-01T00:00:00Z",
    )
    output_path = tmp_path / "brand-clean-output.pptx"

    warnings = PptxBuilder(node_runner=object(), brand_render_mode="clone").build_deck(
        template,
        [outline],
        output_path,
        tmp_path,
    )

    with zipfile.ZipFile(output_path, "r") as pptx_zip:
        names = set(pptx_zip.namelist())
        content_types = pptx_zip.read("[Content_Types].xml").decode("utf-8")
    artifact = json.loads((tmp_path / "template-clone-edit.json").read_text(encoding="utf-8"))
    cleanup = artifact["package_cleanup"]

    assert warnings == []
    assert "ppt/slides/slide1.xml" in names
    assert "ppt/slides/slide2.xml" not in names
    assert "ppt/slides/_rels/slide2.xml.rels" not in names
    assert "/ppt/slides/slide2.xml" not in content_types
    assert cleanup["deleted_unused_slide_part_count"] == 1
    assert cleanup["removed_content_type_override_count"] >= 1
    assert "ppt/slides/slide2.xml" in cleanup["deleted_parts"]


def test_brand_clone_edit_deletes_inherited_media_placeholder(tmp_path: Path) -> None:
    source_file = tmp_path / "brand-media-source.pptx"
    logo_path = tmp_path / "logo.png"
    placeholder_path = tmp_path / "placeholder.png"
    Image.new("RGB", (160, 48), "#102030").save(logo_path)
    Image.new("RGB", (640, 360), "#cccccc").save(placeholder_path)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    logo = slide.shapes.add_picture(logo_path.as_posix(), Inches(11.0), Inches(0.35), Inches(1.2), Inches(0.36))
    logo.name = "Brand logo"
    slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(6.5), Inches(0.7)).text = (
        "Template media title"
    )
    slide.shapes.add_textbox(Inches(0.9), Inches(1.6), Inches(4.1), Inches(1.0)).text = (
        "Template media body"
    )
    placeholder = slide.shapes.add_picture(
        placeholder_path.as_posix(),
        Inches(6.2),
        Inches(1.55),
        Inches(4.4),
        Inches(2.65),
    )
    placeholder.name = "Picture Placeholder 1"
    presentation.save(source_file.as_posix())

    slide_ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
    with zipfile.ZipFile(source_file, "r") as pptx_zip:
        entries = {name: pptx_zip.read(name) for name in pptx_zip.namelist()}
    slide_xml = etree.fromstring(entries["ppt/slides/slide1.xml"])
    placeholder_pic = next(
        pic
        for pic in slide_xml.xpath(".//p:pic", namespaces=slide_ns)
        if "Picture Placeholder" in (pic.find("p:nvPicPr/p:cNvPr", namespaces=slide_ns).get("name") or "")
    )
    nv_pr = placeholder_pic.find("p:nvPicPr/p:nvPr", namespaces=slide_ns)
    assert nv_pr is not None
    ph = nv_pr.find("p:ph", namespaces=slide_ns)
    if ph is None:
        ph = etree.SubElement(nv_pr, "{http://schemas.openxmlformats.org/presentationml/2006/main}ph")
    ph.set("type", "pic")
    entries["ppt/slides/slide1.xml"] = etree.tostring(
        slide_xml,
        xml_declaration=True,
        encoding="UTF-8",
        standalone=True,
    )
    with zipfile.ZipFile(source_file, "w", zipfile.ZIP_DEFLATED) as pptx_zip:
        for name, payload in entries.items():
            pptx_zip.writestr(name, payload)

    template = TemplateProfile(
        id="brand-media",
        name="Brand Media",
        type="brand",
        brand=BrandDNA(),
        slides=[
            SlideSpec(
                index=0,
                mode="flexible",
                label="Media frame",
                layout_name="Media",
                content_category="evidence_points",
                visual_guidance="title, body, and inherited media placeholder",
            )
        ],
        source_file=source_file.as_posix(),
        created_at="2026-01-01T00:00:00Z",
        updated_at="2026-01-01T00:00:00Z",
    )
    outline = SlideOutline(
        id="outline-media",
        job_id="job-brand",
        slide_index=0,
        mode="flexible",
        label="Generated evidence should not keep stale media placeholders",
        content_json={
            "action_title": "Generated evidence should not keep stale media placeholders",
            "subheading": "Generated body text uses the inherited body slot.",
        },
        layout_json={"layout": "callouts", "archetype": "callouts"},
        created_at="2026-01-01T00:00:00Z",
    )
    output_path = tmp_path / "brand-media-output.pptx"

    warnings = PptxBuilder(node_runner=object(), brand_render_mode="clone").build_deck(
        template,
        [outline],
        output_path,
        tmp_path,
    )

    rel_ns = {"rel": "http://schemas.openxmlformats.org/package/2006/relationships"}
    with zipfile.ZipFile(output_path, "r") as pptx_zip:
        output_names = set(pptx_zip.namelist())
        output_slide_xml = etree.fromstring(pptx_zip.read("ppt/slides/slide1.xml"))
        output_rels = etree.fromstring(pptx_zip.read("ppt/slides/_rels/slide1.xml.rels"))
    image_rels = [
        rel
        for rel in output_rels.findall("rel:Relationship", namespaces=rel_ns)
        if rel.get("Type", "").endswith("/image")
    ]
    artifact = json.loads((tmp_path / "template-clone-edit.json").read_text(encoding="utf-8"))
    mapping = artifact["mappings"][0]
    cleanup = artifact["package_cleanup"]

    assert warnings == []
    assert not output_slide_xml.xpath(
        ".//p:pic[contains(p:nvPicPr/p:cNvPr/@name, 'Picture Placeholder')]",
        namespaces=slide_ns,
    )
    assert output_slide_xml.xpath(".//p:pic[contains(p:nvPicPr/p:cNvPr/@name, 'Brand logo')]", namespaces=slide_ns)
    assert len(image_rels) == 1
    assert len([name for name in output_names if name.startswith("ppt/media/")]) == 1
    assert mapping["deleted_media_placeholder_count"] == 1
    assert mapping["deleted_media_relationship_count"] == 1
    assert cleanup["deleted_unreferenced_media_part_count"] == 1
    assert any(target["role"] == "media_placeholder" for target in mapping["editTargets"])


def test_brand_clone_edit_audits_unfilled_inherited_text_placeholder(tmp_path: Path) -> None:
    source_file = tmp_path / "brand-placeholder-source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(6.5), Inches(0.7)).text = (
        "Template title placeholder"
    )
    slide.shapes.add_textbox(Inches(0.9), Inches(1.7), Inches(5.8), Inches(0.8)).text = (
        "Template body placeholder"
    )
    presentation.save(source_file.as_posix())

    p_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    a_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    with zipfile.ZipFile(source_file, "r") as pptx_zip:
        entries = {name: pptx_zip.read(name) for name in pptx_zip.namelist()}
    root = etree.fromstring(entries["ppt/slides/slide1.xml"])
    sp_tree = root.find(".//p:spTree", namespaces={"p": p_ns})
    assert sp_tree is not None
    footer = etree.SubElement(sp_tree, f"{{{p_ns}}}sp")
    nv_sp_pr = etree.SubElement(footer, f"{{{p_ns}}}nvSpPr")
    etree.SubElement(nv_sp_pr, f"{{{p_ns}}}cNvPr", id="77", name="Footer Placeholder")
    etree.SubElement(nv_sp_pr, f"{{{p_ns}}}cNvSpPr")
    nv_pr = etree.SubElement(nv_sp_pr, f"{{{p_ns}}}nvPr")
    etree.SubElement(nv_pr, f"{{{p_ns}}}ph", type="ftr")
    sp_pr = etree.SubElement(footer, f"{{{p_ns}}}spPr")
    xfrm = etree.SubElement(sp_pr, f"{{{a_ns}}}xfrm")
    etree.SubElement(xfrm, f"{{{a_ns}}}off", x="731520", y="6400800")
    etree.SubElement(xfrm, f"{{{a_ns}}}ext", cx="1828800", cy="274320")
    tx_body = etree.SubElement(footer, f"{{{p_ns}}}txBody")
    etree.SubElement(tx_body, f"{{{a_ns}}}bodyPr")
    etree.SubElement(tx_body, f"{{{a_ns}}}lstStyle")
    etree.SubElement(tx_body, f"{{{a_ns}}}p")
    entries["ppt/slides/slide1.xml"] = etree.tostring(
        root,
        xml_declaration=True,
        encoding="UTF-8",
        standalone=True,
    )
    with zipfile.ZipFile(source_file, "w", zipfile.ZIP_DEFLATED) as pptx_zip:
        for name, payload in entries.items():
            pptx_zip.writestr(name, payload)

    template = TemplateProfile(
        id="brand-placeholder",
        name="Brand Placeholder",
        type="brand",
        brand=BrandDNA(),
        slides=[
            SlideSpec(
                index=0,
                mode="flexible",
                label="Evidence frame",
                layout_name="Evidence",
                content_category="evidence_points",
                visual_guidance="title and body slots with inherited footer placeholder",
            )
        ],
        source_file=source_file.as_posix(),
        created_at="2026-01-01T00:00:00Z",
        updated_at="2026-01-01T00:00:00Z",
    )
    outline = SlideOutline(
        id="outline-placeholder",
        job_id="job-brand",
        slide_index=0,
        mode="flexible",
        label="Generated title fills inherited text slots",
        content_json={
            "action_title": "Generated title fills inherited text slots",
            "subheading": "Generated body text fills the body slot.",
        },
        layout_json={"layout": "callouts", "archetype": "callouts"},
        created_at="2026-01-01T00:00:00Z",
    )
    output_path = tmp_path / "brand-placeholder-output.pptx"

    warnings = PptxBuilder(node_runner=object(), brand_render_mode="clone").build_deck(
        template,
        [outline],
        output_path,
        tmp_path,
    )

    artifact = json.loads((tmp_path / "template-clone-edit.json").read_text(encoding="utf-8"))
    deviation_log = json.loads((tmp_path / "template-deviation-log.json").read_text(encoding="utf-8"))
    mapping = artifact["mappings"][0]

    assert any("unfilled inherited placeholder" in warning["message"] for warning in warnings)
    assert mapping["unfilled_placeholder_count"] == 1
    assert mapping["unfilled_placeholders"][0]["placeholder_type"] == "ftr"
    assert any(
        deviation["type"] == "unfilled_inherited_placeholder"
        for deviation in deviation_log["deviations"]
    )


def test_brand_clone_edit_fills_inherited_table_frame(tmp_path: Path) -> None:
    source_file = tmp_path / "brand-table-source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(
        Inches(0.8),
        Inches(0.55),
        Inches(6.5),
        Inches(0.6),
    ).text = "Template table title"
    table = slide.shapes.add_table(
        4,
        3,
        Inches(0.8),
        Inches(1.45),
        Inches(8.4),
        Inches(2.4),
    ).table
    for row in range(4):
        for col in range(3):
            table.cell(row, col).text = f"Sample {row}-{col}"
    presentation.save(source_file.as_posix())
    template = TemplateProfile(
        id="brand-table",
        name="Brand Table",
        type="brand",
        brand=BrandDNA(),
        slides=[
            SlideSpec(
                index=0,
                mode="flexible",
                label="Comparison table frame",
                layout_name="Table",
                content_category="comparison",
                visual_guidance="title plus inherited table",
                schema=SlideSchema(
                    fields=[
                        SlideField(
                            id="comparison_rows",
                            type="list",
                            location="table",
                            max_items=3,
                        )
                    ]
                ),
            )
        ],
        source_file=source_file.as_posix(),
        created_at="2026-01-01T00:00:00Z",
        updated_at="2026-01-01T00:00:00Z",
    )
    outline = SlideOutline(
        id="outline-table",
        job_id="job-brand",
        slide_index=0,
        mode="flexible",
        label="Compare the operating shift in the inherited table",
        content_json={
            "action_title": "Compare the operating shift in the inherited table",
            "exhibit_spec": {
                "type": "comparison_table",
                "columns": ["Signal", "Current", "Target"],
                "rows": [
                    {"label": "Context", "values": ["Fragmented", "Persistent"]},
                    {"label": "Review", "values": ["After output", "Before merge"]},
                ],
            },
        },
        layout_json={"layout": "comparison_table", "archetype": "comparison_table"},
        created_at="2026-01-01T00:00:00Z",
    )
    output_path = tmp_path / "brand-table-output.pptx"

    warnings = PptxBuilder(node_runner=object(), brand_render_mode="clone").build_deck(
        template,
        [outline],
        output_path,
        tmp_path,
    )

    rendered = Presentation(output_path.as_posix())
    table_shape = next(shape for shape in rendered.slides[0].shapes if shape.has_table)
    rendered_rows = [
        [cell.text for cell in row.cells]
        for row in table_shape.table.rows
    ]
    all_text = "\n".join(
        shape.text
        for shape in rendered.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
    )
    artifact = json.loads((tmp_path / "template-clone-edit.json").read_text(encoding="utf-8"))
    targets = artifact["mappings"][0]["editTargets"]
    with zipfile.ZipFile(output_path, "r") as pptx_zip:
        slide_xml = pptx_zip.read("ppt/slides/slide1.xml")
    root = etree.fromstring(slide_xml)
    namespaces = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    bold_run_texts = [
        "".join(run.xpath(".//a:t/text()", namespaces=namespaces))
        for run in root.xpath(".//a:r[a:rPr[@b='1']]", namespaces=namespaces)
    ]

    assert warnings == []
    assert rendered_rows == [
        ["Signal", "Current", "Target"],
        ["Context", "Fragmented", "Persistent"],
        ["Review", "After output", "Before merge"],
    ]
    assert "Sample" not in all_text
    assert artifact["mappings"][0]["rewritten_table_cell_count"] == 9
    assert artifact["mappings"][0]["deleted_table_row_count"] == 1
    assert artifact["mappings"][0]["bolded_text_run_count"] >= 4
    assert {"Signal", "Current", "Target"}.issubset(set(bold_run_texts))
    assert artifact["mappings"][0]["slot_plan"]["action"] == "delete_excess_template_elements"
    assert artifact["mappings"][0]["slot_cleanup"] == {
        "action": "delete_excess_template_elements",
        "planned_excess_slot_count": 1,
        "actual_deleted_slot_count": 1,
        "cleanup_required": True,
        "cleanup_satisfied": True,
        "deleted_text_shape_count": 0,
        "deleted_table_row_count": 1,
        "deleted_media_placeholder_count": 0,
    }
    assert any(target["role"] == "table_cell" for target in targets)
    assert any(target["role"] == "excess_table_row" for target in targets)


def test_brand_clone_edit_fills_inherited_chart_frame(tmp_path: Path) -> None:
    source_file = tmp_path / "brand-chart-source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(
        Inches(0.8),
        Inches(0.55),
        Inches(6.5),
        Inches(0.6),
    ).text = "Template chart title"
    chart_data = CategoryChartData()
    chart_data.categories = ["Sample A", "Sample B"]
    chart_data.add_series("Sample series", (1, 2))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.8),
        Inches(1.45),
        Inches(8.4),
        Inches(3.4),
        chart_data,
    )
    chart.name = "TemplateMetricChart"
    presentation.save(source_file.as_posix())
    template = TemplateProfile(
        id="brand-chart",
        name="Brand Chart",
        type="brand",
        brand=BrandDNA(),
        slides=[
            SlideSpec(
                index=0,
                mode="flexible",
                label="Metric chart frame",
                layout_name="Chart",
                content_category="metrics",
                visual_guidance="title plus inherited chart",
            )
        ],
        source_file=source_file.as_posix(),
        created_at="2026-01-01T00:00:00Z",
        updated_at="2026-01-01T00:00:00Z",
    )
    outline = SlideOutline(
        id="outline-chart",
        job_id="job-brand",
        slide_index=0,
        mode="flexible",
        label="Generated metrics replace the inherited chart data",
        content_json={
            "action_title": "Generated metrics replace the inherited chart data",
            "exhibit_spec": {
                "type": "metric_chart",
                "metrics": [
                    {"label": "Adoption", "value": 7},
                    {"label": "Reliability", "value": 11},
                    {"label": "Coverage", "value": 19},
                ],
            },
        },
        layout_json={"layout": "chart", "archetype": "metric_chart"},
        created_at="2026-01-01T00:00:00Z",
    )
    output_path = tmp_path / "brand-chart-output.pptx"

    warnings = PptxBuilder(node_runner=object(), brand_render_mode="clone").build_deck(
        template,
        [outline],
        output_path,
        tmp_path,
    )

    rel_ns = {"rel": "http://schemas.openxmlformats.org/package/2006/relationships"}
    chart_ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
    with zipfile.ZipFile(output_path, "r") as pptx_zip:
        slide_rels = etree.fromstring(pptx_zip.read("ppt/slides/_rels/slide1.xml.rels"))
        chart_rel = next(
            rel
            for rel in slide_rels.findall("rel:Relationship", namespaces=rel_ns)
            if rel.get("Type", "").endswith("/chart")
        )
        chart_part = f"ppt/charts/{chart_rel.get('Target', '').rsplit('/', 1)[-1]}"
        chart_xml = etree.fromstring(pptx_zip.read(chart_part))
        chart_values = " ".join(chart_xml.xpath(".//c:v/text()", namespaces=chart_ns))
        chart_rels = etree.fromstring(
            pptx_zip.read(f"ppt/charts/_rels/{chart_part.rsplit('/', 1)[-1]}.rels")
        )
        workbook_rel = next(
            rel
            for rel in chart_rels.findall("rel:Relationship", namespaces=rel_ns)
            if rel.get("Type", "").endswith("/package")
        )
        workbook_part = f"ppt/embeddings/{workbook_rel.get('Target', '').rsplit('/', 1)[-1]}"
        workbook = load_workbook(BytesIO(pptx_zip.read(workbook_part)))
        sheet = workbook[workbook.sheetnames[0]]

    artifact = json.loads((tmp_path / "template-clone-edit.json").read_text(encoding="utf-8"))
    targets = artifact["mappings"][0]["editTargets"]

    assert warnings == []
    assert "chart_clone_1_1.xml" in chart_part
    assert "Adoption" in chart_values
    assert "Reliability" in chart_values
    assert "Coverage" in chart_values
    assert "Sample A" not in chart_values
    assert "Sample B" not in chart_values
    assert [sheet.cell(row=row, column=1).value for row in range(2, 5)] == [
        "Adoption",
        "Reliability",
        "Coverage",
    ]
    assert [sheet.cell(row=row, column=2).value for row in range(2, 5)] == [7, 11, 19]
    assert artifact["mappings"][0]["rewritten_chart_count"] == 1
    assert artifact["mappings"][0]["rewritten_chart_point_count"] == 3
    assert artifact["mappings"][0]["duplicated_chart_part_count"] == 1
    assert any(target["role"] == "chart_data" for target in targets)


def test_renderer_borrows_template_frame_chrome_without_source_text(tmp_path: Path) -> None:
    source_file = tmp_path / "brand-source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    rail = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(0),
        Inches(0.38),
        Inches(7.5),
    )
    rail.fill.solid()
    rail.fill.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)
    rail.line.color.rgb = RGBColor(0x12, 0x34, 0x56)
    slide.shapes.add_textbox(
        Inches(1.0),
        Inches(1.0),
        Inches(4.0),
        Inches(0.6),
    ).text = "DO NOT COPY TEMPLATE TEXT"
    presentation.save(source_file.as_posix())
    outline = SlideOutline(
        id="outline-1",
        job_id="job-brand",
        slide_index=0,
        mode="flexible",
        label="Use the frame while replacing content",
        content_json={
            "action_title": "Use the frame while replacing content",
            "bullets": ["Generated evidence replaces template sample copy."],
        },
        layout_json={
            "layout": "two_column",
            "template_frame": {
                "index": 0,
                "source_slide": 1,
                "label": "Branded rail",
                "source_file": source_file.as_posix(),
            },
        },
        created_at="2026-01-01T00:00:00Z",
    )
    output_path = tmp_path / "brand-framed.pptx"

    DeterministicPptxRenderer().render([outline], BrandDNA(), output_path)

    rendered = Presentation(output_path.as_posix())
    text = "\n".join(
        shape.text
        for shape in rendered.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
    )
    fills = []
    for shape in rendered.slides[0].shapes:
        try:
            fills.append(str(shape.fill.fore_color.rgb))
        except Exception:
            continue
    assert "DO NOT COPY TEMPLATE TEXT" not in text
    assert "Use the frame while replacing content" in text
    assert "123456" in fills
    assert outline.layout_json["template_frame"]["chrome_applied"] is True


def test_renderer_clamps_one_line_brand_title_box_height(tmp_path: Path) -> None:
    brand = BrandDNA(
        layout_profile={
            "title_box": {"x": 0.9, "y": 0.81, "w": 9.35, "h": 0.86},
        }
    )
    outline = SlideOutline(
        id="outline-1",
        job_id="job-brand",
        slide_index=0,
        mode="flexible",
        label="Harness design turns workflows into evidence",
        content_json={
            "action_title": "Harness design turns workflows into evidence",
            "bullets": [
                "Existing workflows become evaluation evidence.",
                "Review gates make the benchmark reusable.",
                "Source context keeps the claim grounded.",
            ],
            "sources": ["Bootstrapping Benchmarks"],
            "source_labels": ["Bootstrapping Benchmarks"],
        },
        layout_json={"layout": "callouts", "visual_elements": ["callouts"]},
        created_at="2026-01-01T00:00:00Z",
    )
    output_path = tmp_path / "brand-title-height.pptx"

    DeterministicPptxRenderer().render([outline], brand, output_path)

    rendered = Presentation(output_path.as_posix())
    title_shape = next(
        shape
        for shape in rendered.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
        and "Harness design turns workflows into evidence" in shape.text
    )
    assert title_shape.height / 914400 <= 0.69


def test_renderer_clamps_wrapping_brand_title_before_content_rows(tmp_path: Path) -> None:
    brand = BrandDNA(
        layout_profile={
            "title_box": {"x": 0.9, "y": 0.81, "w": 9.35, "h": 0.86},
        }
    )
    outline = SlideOutline(
        id="outline-1",
        job_id="job-brand",
        slide_index=0,
        mode="flexible",
        label="Harness-centric design turns existing workflows into evaluation evidence",
        content_json={
            "action_title": (
                "Harness-centric design turns existing workflows into evaluation evidence"
            ),
            "bullets": [
                "The harness-centric view treats benchmark generation as a validity problem.",
                "The agent should reason about what makes a test meaningful.",
                "Existing workflows become evaluation evidence when rules and review gates are captured.",
            ],
            "sources": ["Bootstrapping Benchmarks > The Harness-Centric View"],
            "source_labels": ["Bootstrapping Benchmarks > The Harness-Centric View"],
            "exhibit_spec": {
                "type": "callouts",
                "points": [
                    "The harness-centric view treats benchmark generation as a validity problem.",
                    "The agent should reason about what makes a test meaningful.",
                    "Existing workflows become evaluation evidence when rules and review gates are captured.",
                ],
            },
        },
        layout_json={"layout": "callouts", "visual_elements": ["callouts"]},
        created_at="2026-01-01T00:00:00Z",
    )
    output_path = tmp_path / "brand-title-safe-zone.pptx"

    DeterministicPptxRenderer().render([outline], brand, output_path)
    rendered = Presentation(output_path.as_posix())
    title_shape = next(
        shape
        for shape in rendered.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
        and "Harness-centric design turns existing workflows" in shape.text
    )
    _payload, issues = RenderedSlideAudit().inspect(output_path, [outline])

    assert title_shape.height / 914400 <= 0.69
    assert "occluded_text" not in {issue.category for issue in issues}


def test_renderer_places_extracted_brand_logo(tmp_path: Path) -> None:
    logo_path = tmp_path / "logo.png"
    Image.new("RGB", (180, 60), "#111827").save(logo_path)
    brand = BrandDNA(logo={"path": logo_path.as_posix(), "w": 1.2, "h": 0.4})
    outline = SlideOutline(
        id="outline-1",
        job_id="job-1",
        slide_index=0,
        mode="flexible",
        label="Adopt agentic coding to improve production reliability",
        content_json={
            "title": "Adopt agentic coding to improve production reliability",
            "summary": "External memory and specifications improve delivery quality.",
            "bullets": ["Persistent context reduces rework.", "Structured specs improve review."],
            "sources": ["Uploaded source"],
        },
        layout_json={"layout": "two_column", "visual_elements": ["structured_text"]},
        created_at="2026-01-01T00:00:00Z",
    )
    output_path = tmp_path / "logo-rendered.pptx"

    DeterministicPptxRenderer().render([outline], brand, output_path)

    rendered = Presentation(output_path.as_posix())
    assert any(
        shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        for shape in rendered.slides[0].shapes
    )


def test_renderer_omits_subheading_for_wrapped_long_title(tmp_path: Path) -> None:
    long_title = (
        "Replace ephemeral vibe coding with persistent agentic workflows for "
        "scalable software delivery"
    )
    outline = SlideOutline(
        id="outline-1",
        job_id="job-1",
        slide_index=0,
        mode="flexible",
        label=long_title,
        content_json={
            "action_title": long_title,
            "subheading": "Executive Summary: The Shift from Chatting to Managing AI",
            "bullets": ["Persistent context reduces rework."],
            "sources": ["Uploaded source"],
        },
        layout_json={"layout": "two_column", "visual_elements": ["structured_text"]},
        created_at="2026-01-01T00:00:00Z",
    )
    output_path = tmp_path / "long-title-rendered.pptx"

    DeterministicPptxRenderer().render([outline], BrandDNA(), output_path)

    rendered = Presentation(output_path.as_posix())
    text_shapes = [
        shape
        for shape in rendered.slides[0].shapes
        if getattr(shape, "has_text_frame", False) and shape.text
    ]
    title_shape = next(shape for shape in text_shapes if shape.text.startswith("Replace"))

    assert title_shape.text.endswith("delivery")
    assert not any(shape.text.startswith("Executive Summary") for shape in text_shapes)


def test_renderer_adds_editorial_section_chrome(tmp_path: Path) -> None:
    outline = SlideOutline(
        id="outline-1",
        job_id="job-1",
        slide_index=0,
        mode="flexible",
        label="Prevent context rot before AI work scales",
        content_json={
            "action_title": "Prevent context rot before AI work scales",
            "subheading": "Evidence from uploaded source",
            "bullets": ["Persistent memory keeps work reproducible."],
            "sources": ["Uploaded source"],
            "archetype": "anti_patterns",
            "narrative_role": "problem",
        },
        layout_json={
            "layout": "two_column",
            "visual_elements": ["structured_text"],
            "narrative_role": "problem",
            "archetype": "anti_patterns",
        },
        created_at="2026-01-01T00:00:00Z",
    )
    output_path = tmp_path / "editorial-chrome.pptx"

    DeterministicPptxRenderer().render([outline], BrandDNA(), output_path)

    rendered = Presentation(output_path.as_posix())
    text = "\n".join(
        shape.text
        for shape in rendered.slides[0].shapes
        if getattr(shape, "has_text_frame", False) and shape.text
    )

    assert "SECTION 02 - DIAGNOSIS" in text
    assert "Prevent context rot before AI work scales" in text


def test_renderer_adds_executive_summary_proof_rail(tmp_path: Path) -> None:
    outline = SlideOutline(
        id="outline-1",
        job_id="job-1",
        slide_index=1,
        mode="flexible",
        label="Manage AI coding as an operating system, not a prompt habit",
        content_json={
            "action_title": "Manage AI coding as an operating system, not a prompt habit",
            "subheading": "Situation, complication, resolution",
            "bullets": [
                "AI coding works for fast starts.",
                "Production work needs repeatable context and review.",
                "Managed agentic engineering supplies the operating model.",
            ],
            "sources": ["Uploaded source"],
            "archetype": "executive_summary",
            "narrative_role": "executive_summary",
            "exhibit_spec": {
                "type": "executive_summary",
                "messages": [
                    {"label": "Situation", "text": "AI coding works for fast starts."},
                    {
                        "label": "Complication",
                        "text": "Production work needs repeatable context and review.",
                    },
                    {
                        "label": "Resolution",
                        "text": "Managed agentic engineering supplies the operating model.",
                    },
                ],
                "proof_points": [
                    {
                        "label": "AI-generated codebases",
                        "value": 95,
                        "unit": "%",
                        "detail": "Adoption pressure is already visible.",
                    },
                    {
                        "label": "Developers using AI tools",
                        "value": 85,
                        "unit": "%",
                        "detail": "Usage is mainstream enough to manage.",
                    },
                ],
            },
        },
        layout_json={
            "layout": "executive_summary",
            "visual_elements": ["summary"],
            "narrative_role": "executive_summary",
            "archetype": "executive_summary",
        },
        created_at="2026-01-01T00:00:00Z",
    )
    output_path = tmp_path / "executive-summary-proof-rail.pptx"

    DeterministicPptxRenderer().render([outline], BrandDNA(), output_path)

    rendered = Presentation(output_path.as_posix())
    text = "\n".join(
        shape.text
        for shape in rendered.slides[0].shapes
        if getattr(shape, "has_text_frame", False) and shape.text
    )

    assert "PROOF POINTS" in text
    assert "95%" in text
    assert "Decision ask" in text


def test_renderer_does_not_pad_grid_with_generic_filler(tmp_path: Path) -> None:
    outline = SlideOutline(
        id="outline-1",
        job_id="job-1",
        slide_index=0,
        mode="flexible",
        label="Adopt persistent context management",
        content_json={
            "action_title": "Adopt persistent context management",
            "subheading": "Three evidence-backed points",
            "bullets": [
                "Context files reduce repeated explanations.",
                "Specifications improve review quality.",
                "Decision trails support maintainability.",
            ],
            "sources": ["Uploaded source"],
        },
        layout_json={"layout": "icon_grid", "visual_elements": ["icons", "shapes"]},
        created_at="2026-01-01T00:00:00Z",
    )
    output_path = tmp_path / "grid-rendered.pptx"

    DeterministicPptxRenderer().render([outline], BrandDNA(), output_path)

    rendered = Presentation(output_path.as_posix())
    text = "\n".join(
        shape.text
        for shape in rendered.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
    )

    assert "Clarify the implication" not in text
    assert "Decision trails support maintainability." in text


def test_renderer_expands_grid_from_exhibit_points(tmp_path: Path) -> None:
    outline = SlideOutline(
        id="outline-grid-exhibit",
        job_id="job-1",
        slide_index=0,
        mode="flexible",
        label="Use source-grounded harnesses instead of synthetic benchmark shortcuts",
        content_json={
            "action_title": "Use source-grounded harnesses instead of synthetic benchmark shortcuts",
            "subheading": "The harness should expose multiple proof points.",
            "bullets": ["Mapping existing validated assets to their value as benchmark truth sources"],
            "sources": ["Uploaded source"],
            "archetype": "callouts",
            "exhibit_spec": {
                "type": "callouts",
                "points": [
                    "Mapping existing validated assets to their value as benchmark truth sources",
                    "Model contracts specify what success should mean before execution",
                    "Review gates connect generated cases back to operational evidence",
                ],
            },
        },
        layout_json={"layout": "callouts", "visual_elements": ["callouts"]},
        created_at="2026-01-01T00:00:00Z",
    )
    output_path = tmp_path / "grid-exhibit-points.pptx"

    DeterministicPptxRenderer().render([outline], BrandDNA(), output_path)

    rendered = Presentation(output_path.as_posix())
    text = "\n".join(
        shape.text
        for shape in rendered.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
    )
    compact_text = " ".join(text.split())

    assert "Model contracts Specify what success should mean before execution" in compact_text
    assert "Review gates Connect generated cases back to operational evidence" in compact_text


def test_renderer_draws_semantic_icons_as_embedded_artwork(tmp_path: Path) -> None:
    outline = SlideOutline(
        id="outline-1",
        job_id="job-1",
        slide_index=0,
        mode="flexible",
        label="Quantify production readiness before scaling AI development",
        content_json={
            "action_title": "Quantify production readiness before scaling AI development",
            "subheading": "Three evidence-backed points",
            "bullets": [
                "Adoption is high enough to require operating discipline.",
                "Persistent memory reduces repeated context setup.",
                "Review standards protect production reliability.",
            ],
            "sources": ["Uploaded source"],
        },
        layout_json={
            "layout": "icon_grid",
            "visual_elements": ["icons", "shapes"],
            "icons": ["FaChartLine", "FaDatabase", "FaShieldAlt"],
        },
        created_at="2026-01-01T00:00:00Z",
    )
    output_path = tmp_path / "icons-rendered.pptx"

    DeterministicPptxRenderer().render([outline], BrandDNA(), output_path)

    rendered = Presentation(output_path.as_posix())
    pictures = [
        shape
        for shape in rendered.slides[0].shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    icon_cache = output_path.parent / f"{output_path.stem}-icons"

    assert len(pictures) >= 3
    cached_icons = {path.name for path in icon_cache.glob("*.png")}
    assert any(name.startswith("fachartline-") for name in cached_icons)
    assert any(name.startswith("fadatabase-") for name in cached_icons)
    assert any(name.startswith("fashieldalt-") for name in cached_icons)


def test_renderer_does_not_show_diagram_description_placeholders(tmp_path: Path) -> None:
    outline = SlideOutline(
        id="outline-1",
        job_id="job-1",
        slide_index=0,
        mode="flexible",
        label="Map the External Brain architecture",
        content_json={
            "action_title": "Map the External Brain architecture",
            "subheading": "Evidence from uploaded source",
            "bullets": [
                "[Diagram Description: Central Node with arrows] Project context persists.",
                "Rules files guide each AI session.",
                "Memory bank updates after every meaningful change.",
            ],
            "sources": ["Uploaded source"],
        },
        layout_json={"layout": "dependency_map", "visual_elements": ["diagram"]},
        created_at="2026-01-01T00:00:00Z",
    )
    output_path = tmp_path / "cleaned-rendered.pptx"

    DeterministicPptxRenderer().render([outline], BrandDNA(), output_path)

    rendered = Presentation(output_path.as_posix())
    text = "\n".join(
        shape.text
        for shape in rendered.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
    )
    diagram_svgs = list((tmp_path / "cleaned-rendered-diagrams").glob("*.svg"))
    artifact_text = "\n".join(path.read_text(encoding="utf-8") for path in diagram_svgs)

    assert "Diagram Description" not in f"{text}\n{artifact_text}"
    assert "Project context" in f"{text}\n{artifact_text}"
    assert "persists." in f"{text}\n{artifact_text}"


def test_renderer_uses_structured_exhibit_specs_for_authored_archetypes(tmp_path: Path) -> None:
    archetypes = [
        (
            "comparison_table",
            {
                "type": "comparison_table",
                "columns": [
                    {"name": "Dimension", "width": "20%"},
                    {"name": "Before", "width": "40%"},
                    {"name": "After", "width": "40%"},
                ],
                "rows": [{"label": "Context", "values": ["Chat", "Memory bank"]}],
            },
        ),
        (
            "anti_patterns",
            {
                "type": "anti_patterns",
                "patterns": [
                    {
                        "name": "Context rot",
                        "symptom": "Memory disappears",
                        "better_behavior": "Persist context",
                    }
                ],
            },
        ),
        (
            "dependency_map",
            {
                "type": "dependency_map",
                "left_node": "Source packet",
                "middle_nodes": ["Rules", "Architecture", "Decision log"],
                "right_outcome": "Reliable output",
            },
        ),
        (
            "framework_cycle",
            {
                "type": "cycle",
                "center_label": "Agent loop",
                "steps": [
                    {"label": "Frame", "description": "Define the ask"},
                    {"label": "Prime", "description": "Load context"},
                    {"label": "Review", "description": "Check output"},
                ],
            },
        ),
        (
            "code_panel",
            {
                "type": "code_panel",
                "title": "AGENTS.md",
                "lines": ["Require evidence before merge"],
            },
        ),
        (
            "table_reference",
            {
                "type": "reference_table",
                "columns": ["Artifact", "Purpose", "Update trigger"],
                "rows": [
                    [
                        "decision-log.md",
                        "Keeps approvals and assumptions traceable",
                        "Decision changes",
                    ],
                    [
                        "review-gate.md",
                        "Defines the evidence check before rollout",
                        "Pilot milestone",
                    ],
                ],
            },
        ),
        (
            "checklist",
            {
                "type": "checklist",
                "items": [
                    {"action": "Create memory files", "owner": "Lead", "timing": "Week 1"},
                    {"action": "Define review gates", "owner": "Manager", "timing": "Week 1"},
                    {"action": "Run pilot", "owner": "Team", "timing": "Week 2"},
                ],
            },
        ),
    ]
    outlines = [
        SlideOutline(
            id=f"outline-{idx}",
            job_id="job-1",
            slide_index=idx,
            mode="flexible",
            label="Improve AI delivery discipline",
            content_json={
                "action_title": "Improve AI delivery discipline",
                "subheading": "Evidence from uploaded source",
                "bullets": [],
                "sources": ["Uploaded source"],
                "archetype": layout,
                "exhibit_spec": exhibit,
            },
            layout_json={
                "layout": layout,
                "visual_elements": ["structured_exhibit"],
                "icons": ["FaDatabase", "FaShieldAlt", "FaClipboardCheck"],
            },
            created_at="2026-01-01T00:00:00Z",
        )
        for idx, (layout, exhibit) in enumerate(archetypes)
    ]
    output_path = tmp_path / "structured-exhibits.pptx"

    DeterministicPptxRenderer().render(outlines, BrandDNA(), output_path)

    rendered = Presentation(output_path.as_posix())
    text_parts = [
        shape.text
        for slide in rendered.slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    ]
    text_parts.extend(
        cell.text
        for slide in rendered.slides
        for shape in slide.shapes
        if getattr(shape, "has_table", False)
        for row in shape.table.rows
        for cell in row.cells
    )
    text = "\n".join(text_parts)
    diagram_text = "\n".join(
        path.read_text(encoding="utf-8")
        for path in (tmp_path / "structured-exhibits-diagrams").glob("*.svg")
    )
    rendered_surface = f"{text}\n{diagram_text}"

    assert "Memory bank" in text
    assert "COMPARISON LENS" in text
    assert "TARGET TEST" in text
    assert "{'name'" not in text
    assert "Context rot" in text
    assert "Symptom:" in text
    assert "Better move:" in text
    assert "Pattern 1" not in text
    assert "Source packet" in rendered_surface
    assert "project_brief.md" not in text
    assert "Agent loop" not in rendered_surface
    assert "Frame" not in rendered_surface
    assert "Prime" not in rendered_surface
    assert "SOURCE SIGNALS" in rendered_surface
    assert "AGENTS.md" in text
    assert "OPERATING RULES" in text
    assert "REFERENCE ARTIFACT" in text
    assert "REFERENCE MAP" in text
    assert "UPDATE CADENCE" in text
    assert "decision-log.md" in text
    assert "REVIEW GATE" in text
    assert "UPDATE TRIGGER" in text
    assert "Create memory files (Lead / Week 1)" in text


def test_renderer_uses_dark_cadence_for_narrative_beats(tmp_path: Path) -> None:
    outlines = [
        SlideOutline(
            id="cycle-outline",
            job_id="job-1",
            slide_index=0,
            mode="flexible",
            label="Run the operating cycle with explicit review gates",
            content_json={
                "action_title": "Run the operating cycle with explicit review gates",
                "subheading": "A repeatable loop for source-grounded work",
                "bullets": [],
                "sources": ["Uploaded source"],
                "archetype": "framework_cycle",
                "narrative_role": "framework",
                "exhibit_spec": {
                    "type": "cycle",
                    "center_label": "Agent loop",
                    "steps": [
                        {"label": "Load"},
                        {"label": "Plan"},
                        {"label": "Execute"},
                        {"label": "Verify"},
                        {"label": "Update"},
                        {"label": "Reset"},
                    ],
                },
            },
            layout_json={"layout": "framework_cycle", "visual_elements": ["cycle"]},
            created_at="2026-01-01T00:00:00Z",
        ),
        SlideOutline(
            id="quote-outline",
            job_id="job-1",
            slide_index=1,
            mode="flexible",
            label="Adopt the external brain mental model",
            content_json={
                "action_title": "Adopt the external brain mental model",
                "subheading": "Persistent context changes the management problem",
                "bullets": [
                    "Context must survive the chat window.",
                    "Rules must constrain the next run.",
                ],
                "sources": ["Uploaded source"],
                "archetype": "quote_sidebar",
                "narrative_role": "decision",
                "exhibit_spec": {
                    "type": "quote_sidebar",
                    "key_idea": "Treat the agent as capable but memoryless unless context is explicit.",
                    "supporting_points": [
                        "Context must survive the chat window.",
                        "Rules must constrain the next run.",
                    ],
                },
            },
            layout_json={"layout": "quote_sidebar", "visual_elements": ["quote"]},
            created_at="2026-01-01T00:00:00Z",
        ),
        SlideOutline(
            id="closing-outline",
            job_id="job-1",
            slide_index=2,
            mode="flexible",
            label="Commit to governed agentic delivery",
            content_json={
                "action_title": "Commit to governed agentic delivery",
                "subheading": "Recommendation and next steps",
                "bullets": ["Name the owner.", "Pilot the loop.", "Refresh memory."],
                "sources": ["Uploaded source"],
                "archetype": "closing_recommendation",
                "narrative_role": "closing",
                "exhibit_spec": {
                    "type": "recommendation",
                    "recommendation": "Commit to governed agentic delivery",
                    "decision_ask": "Approve the pilot workflow.",
                    "next_steps": ["Name the owner.", "Pilot the loop.", "Refresh memory."],
                },
            },
            layout_json={"layout": "closing_recommendation", "visual_elements": ["decision"]},
            created_at="2026-01-01T00:00:00Z",
        ),
    ]
    output_path = tmp_path / "dark-cadence.pptx"

    DeterministicPptxRenderer().render(outlines, BrandDNA(), output_path)

    rendered = Presentation(output_path.as_posix())
    text = "\n".join(
        shape.text
        for slide in rendered.slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )

    assert "/ OPERATING MODEL" in text
    assert "/ DECISION LENS" in text
    assert "/ FINAL DECISION" in text
    assert "SECTION 04 - OPERATING MODEL" not in text
    assert "SECTION 07 - DECISION" not in text


def test_renderer_text_cleanup_avoids_dangling_truncation() -> None:
    renderer = DeterministicPptxRenderer()

    cleaned = renderer._clean_display_text("Windows range from 32000.0tokens to 1M tokens")
    truncated = renderer._truncate_phrase(
        "A moderately complex project spans dozens of files and services",
        48,
    )

    assert "32000 tokens" in cleaned
    assert not truncated.lower().endswith((" and", " of", " to", " with"))


def test_renderer_avoids_duplicate_closing_recommendation(tmp_path: Path) -> None:
    title = "Commit to the reset habit as the single highest-leverage practice"
    outline = SlideOutline(
        id="closing-outline",
        job_id="job-1",
        slide_index=0,
        mode="flexible",
        label=title,
        content_json={
            "action_title": title,
            "subheading": "Recommendation and decision ask",
            "bullets": ["Name an owner.", "Run the pilot.", "Update memory."],
            "sources": ["Uploaded source"],
            "archetype": "closing_recommendation",
            "narrative_role": "closing",
            "exhibit_spec": {
                "type": "recommendation",
                "recommendation": title,
                "decision_ask": "Approve a governed pilot.",
                "next_steps": ["Name an owner.", "Run the pilot.", "Update memory."],
            },
        },
        layout_json={"layout": "closing_recommendation", "visual_elements": ["decision"]},
        created_at="2026-01-01T00:00:00Z",
    )
    output_path = tmp_path / "closing-deduped.pptx"

    DeterministicPptxRenderer().render([outline], BrandDNA(), output_path)

    rendered = Presentation(output_path.as_posix())
    text = "\n".join(
        shape.text
        for shape in rendered.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
    )

    assert "Make the reset habit the default operating rule before scaling." in text


def test_renderer_completes_sparse_code_panel_rules(tmp_path: Path) -> None:
    outline = SlideOutline(
        id="code-outline",
        job_id="job-1",
        slide_index=0,
        mode="flexible",
        label="Codify memory updates before context goes stale",
        content_json={
            "action_title": "Codify memory updates before context goes stale",
            "subheading": "Memory maintenance rules",
            "bullets": [],
            "sources": ["Uploaded source"],
            "archetype": "code_panel",
            "exhibit_spec": {
                "type": "code_panel",
                "title": "activeContext.md",
                "lines": ["Update active context after material changes."],
            },
        },
        layout_json={"layout": "code_panel", "visual_elements": ["reference_panel"]},
        created_at="2026-01-01T00:00:00Z",
    )
    output_path = tmp_path / "sparse-code-panel.pptx"

    DeterministicPptxRenderer().render([outline], BrandDNA(), output_path)

    rendered = Presentation(output_path.as_posix())
    text = "\n".join(
        shape.text
        for shape in rendered.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
    )

    assert "Update active context after material changes." in text
    assert "Record decisions before starting the next session." in text
    assert "Keep progress notes synchronized with implementation status." in text


def test_renderer_completes_sparse_checklist_items(tmp_path: Path) -> None:
    outline = SlideOutline(
        id="checklist-outline",
        job_id="job-1",
        slide_index=0,
        mode="flexible",
        label="Define acceptance criteria before execution begins",
        content_json={
            "action_title": "Define acceptance criteria before execution begins",
            "subheading": "Specification discipline",
            "bullets": [],
            "sources": ["Uploaded source"],
            "archetype": "checklist",
            "exhibit_spec": {
                "type": "checklist",
                "items": [{"action": "Write acceptance criteria before generation begins"}],
            },
        },
        layout_json={"layout": "checklist", "visual_elements": ["checklist"]},
        created_at="2026-01-01T00:00:00Z",
    )
    output_path = tmp_path / "sparse-checklist.pptx"

    DeterministicPptxRenderer().render([outline], BrandDNA(), output_path)

    rendered = Presentation(output_path.as_posix())
    text = "\n".join(
        shape.text
        for shape in rendered.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
    )

    assert "Write acceptance criteria before generation begins." in text
    assert "Attach source evidence to each material claim." in text
    assert "Review output against the defined test before approval." in text


def test_renderer_splits_metric_units_to_avoid_card_label_overlap(tmp_path: Path) -> None:
    outline = SlideOutline(
        id="metric-outline",
        job_id="job-1",
        slide_index=0,
        mode="flexible",
        label="Quantify capacity before scaling",
        content_json={
            "action_title": "Quantify capacity before scaling",
            "subheading": "Evidence from uploaded source",
            "bullets": [],
            "metrics": [
                {"label": "Context window maximum", "value": 1_000_000, "unit": "tokens"},
                {"label": "Adoption", "value": 95, "unit": "%"},
            ],
            "sources": ["Uploaded source"],
            "archetype": "metric_chart",
            "exhibit_spec": {
                "type": "metric_chart",
                "metrics": [
                    {"label": "Context window maximum", "value": 1_000_000, "unit": "tokens"},
                    {"label": "Adoption", "value": 95, "unit": "%"},
                ],
            },
        },
        layout_json={"layout": "chart", "visual_elements": ["metrics"]},
        created_at="2026-01-01T00:00:00Z",
    )
    output_path = tmp_path / "metric-units.pptx"

    DeterministicPptxRenderer().render([outline], BrandDNA(), output_path)

    rendered = Presentation(output_path.as_posix())
    text = "\n".join(
        shape.text
        for shape in rendered.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
    )

    assert "1M" in text
    assert "tokens" in text
    assert "1M tokens" not in text


def test_stock_office_theme_colors_replaced_by_used_colors(tmp_path) -> None:
    """A deck styled with direct formatting keeps the stock Office palette in
    theme1.xml; the analyzer must derive the brand from colors actually used."""
    from pptx import Presentation as PptxPresentation
    from pptx.dml.color import RGBColor as Rgb
    from pptx.util import Inches as In

    from app.services.template_analyzer import TemplateAnalyzer

    prs = PptxPresentation()  # stock Office theme -> accent1 = 4472C4
    for _ in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bg = slide.shapes.add_shape(1, In(0), In(0), In(10), In(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = Rgb(0x0B, 0x0D, 0x10)  # near-black brand base
        chip = slide.shapes.add_shape(1, In(1), In(1), In(2), In(1))
        chip.fill.solid()
        chip.fill.fore_color.rgb = Rgb(0xD6, 0xA5, 0x29)  # gold accent
    path = tmp_path / "direct-formatted.pptx"
    prs.save(path.as_posix())

    brand = TemplateAnalyzer()._extract_brand(path)
    assert brand.colors.primary.upper() != "4472C4"
    assert brand.colors.primary.upper() == "0B0D10"
    assert brand.colors.accent.upper() == "D6A529"


def test_logo_extraction_returns_none_without_plausible_logo(tmp_path) -> None:
    """Full-bleed background art must never be promoted to a logo."""
    from PIL import Image
    from pptx import Presentation as PptxPresentation
    from pptx.util import Inches as In

    from app.services.template_analyzer import TemplateAnalyzer

    art = tmp_path / "art.png"
    Image.new("RGB", (1600, 900), "navy").save(art)
    prs = PptxPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(art.as_posix(), In(0), In(0), In(10), In(7.5))
    path = tmp_path / "no-logo.pptx"
    prs.save(path.as_posix())

    assert TemplateAnalyzer()._extract_logo(path) is None


def test_logo_extraction_finds_recurring_corner_image(tmp_path) -> None:
    from PIL import Image
    from pptx import Presentation as PptxPresentation
    from pptx.util import Inches as In

    from app.services.template_analyzer import TemplateAnalyzer

    mark = tmp_path / "mark.png"
    Image.new("RGB", (200, 80), "black").save(mark)
    prs = PptxPresentation()
    for _ in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(mark.as_posix(), In(8.6), In(0.3), In(1.0), In(0.4))
    path = tmp_path / "with-logo.pptx"
    prs.save(path.as_posix())

    logo = TemplateAnalyzer()._extract_logo(path)
    assert logo is not None


def test_extract_image_assets_saves_distinct_images(tmp_path) -> None:
    from PIL import Image
    from pptx import Presentation as PptxPresentation
    from pptx.util import Inches as In

    from app.services.template_analyzer import TemplateAnalyzer

    a, b = tmp_path / "a.png", tmp_path / "b.png"
    Image.new("RGB", (400, 300), "navy").save(a)
    Image.new("RGB", (300, 300), "gold").save(b)
    prs = PptxPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(a.as_posix(), In(0), In(0), In(4), In(3))
    slide.shapes.add_picture(b.as_posix(), In(5), In(0), In(3), In(3))
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    slide2.shapes.add_picture(a.as_posix(), In(0), In(0), In(4), In(3))  # duplicate
    path = tmp_path / "imgs.pptx"
    prs.save(path.as_posix())

    names = TemplateAnalyzer().extract_image_assets(path)
    assert len(names) == 2  # deduped by content hash
    assert all((path.parent / "images" / n).exists() for n in names)
