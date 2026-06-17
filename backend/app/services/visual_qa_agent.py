import json
import posixpath
import re
import zipfile
from pathlib import Path
from typing import Optional

from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.clients.bedrock_client import BedrockClient
from app.clients.openai_compatible_client import OpenAICompatibleClient
from app.models.outline import SlideOutline
from app.models.qa import QAEnvelope, QAIssue, QAResult
from app.services.rendering import RenderingError, render_pptx_to_images, render_pptx_to_pdf


EMU_PER_INCH = 914400
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
SAFE_XML_PARSER = etree.XMLParser(resolve_entities=False, no_network=True)


QA_SYSTEM_PROMPT = """You are a visual QA inspector for PowerPoint slides.
Return a JSON object with an array of issues, each with severity (CRITICAL, WARNING, INFO) and message.
Flag text walls, overlaps, cut-off text, low contrast, missing visuals, and repeated layouts.
Return strict JSON only:
{
  "issues": [{"severity":"CRITICAL|WARNING|INFO","message":"...", "category":"...", "slide_index":0}]
}"""


class VisualQAAgent:
    def __init__(
        self,
        bedrock: Optional[BedrockClient] = None,
        openai_client: Optional[OpenAICompatibleClient] = None,
        model_id: Optional[str] = None,
    ) -> None:
        self.bedrock = bedrock
        self.openai_client = openai_client
        self.model_id = model_id or "us.anthropic.claude-sonnet-4-20250514"

    def inspect_deck(
        self,
        pptx_path: Path,
        output_dir: Path,
        outlines: list[SlideOutline],
    ) -> tuple[QAResult, list[Path]]:
        issues: list[QAIssue] = []
        try:
            images = render_pptx_to_images(pptx_path, output_dir)
        except RenderingError as exc:
            issues.append(
                QAIssue(severity="WARNING", message=str(exc), category="render_unavailable")
            )
            images = self._render_pptx_preview_fallback(pptx_path, output_dir)
            if images:
                issues.append(
                    QAIssue(
                        severity="INFO",
                        message="Used approximate PPTX preview rendering for visual QA.",
                        category="render_fallback",
                    )
                )
                if self.bedrock or self.openai_client:
                    for idx, image in enumerate(images):
                        try:
                            report = self._inspect_image(image)
                        except Exception as exc:
                            issues.append(
                                QAIssue(
                                    severity="WARNING",
                                    message=f"Vision QA failed for approximate preview: {exc}",
                                    slide_index=idx,
                                    category="vision_unavailable",
                                )
                            )
                            continue
                        issues.extend(
                            self._downgrade_fallback_vision_issues(
                                self._parse_report(report, idx)
                            )
                        )
            else:
                issues.append(
                    QAIssue(
                        severity="WARNING",
                        message="No preview images could be generated for visual QA.",
                        category="preview_unavailable",
                    )
                )
            issues.extend(self._pptx_structure_checks(pptx_path, outlines))
            issues.extend(self._rule_based_checks(outlines))
            passed = not any(issue.severity == "CRITICAL" for issue in issues)
            return QAResult(issues=issues, passed=passed), images

        if self.bedrock or self.openai_client:
            for idx, image in enumerate(images):
                try:
                    report = self._inspect_image(image)
                except Exception as exc:
                    issues.append(
                        QAIssue(
                            severity="WARNING",
                            message=f"Vision QA failed for slide image: {exc}",
                            slide_index=idx,
                            category="vision_unavailable",
                        )
                    )
                    continue
                issues.extend(self._parse_report(report, idx))
        issues.extend(self._pptx_structure_checks(pptx_path, outlines))
        issues.extend(self._rule_based_checks(outlines))
        passed = not any(issue.severity == "CRITICAL" for issue in issues)
        return QAResult(issues=issues, passed=passed), images

    def _inspect_image(self, image_path: Path) -> str:
        if not self.bedrock:
            if not self.openai_client:
                return ""
            image_format = image_path.suffix.lower().lstrip(".") or "jpeg"
            if image_format == "jpg":
                image_format = "jpeg"
            return self.openai_client.complete_vision(
                system_prompt=QA_SYSTEM_PROMPT,
                user_prompt="Inspect this slide for layout, overlap, text-wall, contrast, and spacing issues. Return JSON only.",
                image_bytes=image_path.read_bytes(),
                image_format=image_format,
                max_tokens=1200,
                temperature=0,
            )
        image_bytes = image_path.read_bytes()
        return self.bedrock.converse_vision(
            model_id=self.model_id,
            system_prompt=QA_SYSTEM_PROMPT,
            user_prompt="Inspect this slide for layout, overlap, or text-wall issues.",
            image_bytes=image_bytes,
        )

    def _parse_report(self, report: str, slide_index: int) -> list[QAIssue]:
        if not report.strip():
            return []
        payload = self._extract_json(report)
        if payload is None:
            return [
                QAIssue(
                    severity="WARNING",
                    message="QA response could not be parsed.",
                    slide_index=slide_index,
                    category="qa_parse",
                )
            ]
        try:
            envelope = QAEnvelope.model_validate(payload)
        except Exception:
            return [
                QAIssue(
                    severity="WARNING",
                    message="QA response JSON did not match expected schema.",
                    slide_index=slide_index,
                    category="qa_schema",
                )
            ]
        parsed: list[QAIssue] = []
        for issue in envelope.issues:
            normalized_severity = issue.severity if issue.severity in {"CRITICAL", "WARNING", "INFO"} else "INFO"
            parsed.append(
                QAIssue(
                    severity=normalized_severity,
                    message=issue.message or "Unspecified issue.",
                    slide_index=slide_index,
                    category=issue.category,
                )
            )
        return parsed

    def _downgrade_fallback_vision_issues(
        self, issues: list[QAIssue]
    ) -> list[QAIssue]:
        downgraded: list[QAIssue] = []
        for issue in issues:
            message = f"Approximate preview finding: {issue.message}"
            if issue.severity == "CRITICAL":
                downgraded.append(
                    issue.model_copy(update={"severity": "WARNING", "message": message})
                )
            else:
                downgraded.append(issue.model_copy(update={"message": message}))
        return downgraded

    def _rule_based_checks(self, outlines: list[SlideOutline]) -> list[QAIssue]:
        issues = []
        for outline in outlines:
            if outline.mode != "flexible":
                continue
            visuals = outline.layout_json.get("visual_elements", [])
            if not visuals:
                issues.append(
                    QAIssue(
                        severity="CRITICAL",
                        message="Text wall detected: no visual elements.",
                        slide_index=outline.slide_index,
                        category="text_wall",
                    )
                )
            sources = outline.content_json.get("sources") or []
            if not sources:
                issues.append(
                    QAIssue(
                        severity="WARNING",
                        message="Slide is missing source coverage metadata.",
                        slide_index=outline.slide_index,
                        category="source_coverage",
                    )
                )
            elif any("source needed" in str(source).lower() for source in sources):
                issues.append(
                    QAIssue(
                        severity="WARNING",
                        message="Slide still contains a [source needed] placeholder.",
                        slide_index=outline.slide_index,
                        category="source_placeholder",
                    )
                )
        issues.extend(self._layout_variety_checks(outlines))
        return issues

    def _layout_variety_checks(self, outlines: list[SlideOutline]) -> list[QAIssue]:
        flexible = [outline for outline in outlines if outline.mode == "flexible"]
        if len(flexible) < 4:
            return []
        layouts = [outline.layout_json.get("layout", "") for outline in flexible]
        issues: list[QAIssue] = []
        if len(set(layouts)) < min(3, len(layouts)):
            issues.append(
                QAIssue(
                    severity="WARNING",
                    message="Deck uses too little layout variety for a consulting narrative.",
                    category="layout_variety",
                )
            )
        run_length = 1
        for idx in range(1, len(layouts)):
            run_length = run_length + 1 if layouts[idx] == layouts[idx - 1] else 1
            if run_length > 2:
                issues.append(
                    QAIssue(
                        severity="WARNING",
                        message="Repeated adjacent layouts may make the deck feel repetitive.",
                        slide_index=flexible[idx].slide_index,
                        category="layout_repetition",
                    )
                )
                break
        return issues

    def _pptx_structure_checks(
        self, pptx_path: Path, outlines: list[SlideOutline]
    ) -> list[QAIssue]:
        if not pptx_path.exists():
            return [
                QAIssue(
                    severity="CRITICAL",
                    message="PPTX file does not exist for QA.",
                    category="office_compatibility",
                )
            ]
        try:
            with zipfile.ZipFile(pptx_path, "r") as pptx_zip:
                bad_entry = pptx_zip.testzip()
                names = set(pptx_zip.namelist())
        except zipfile.BadZipFile:
            return [
                QAIssue(
                    severity="CRITICAL",
                    message="PPTX package is not a valid zip file.",
                    category="office_compatibility",
                )
            ]
        issues: list[QAIssue] = []
        if bad_entry:
            issues.append(
                QAIssue(
                    severity="CRITICAL",
                    message=f"PPTX package has a corrupt entry: {bad_entry}.",
                    category="office_compatibility",
                )
            )
        if "[Content_Types].xml" not in names:
            issues.append(
                QAIssue(
                    severity="CRITICAL",
                    message="PPTX package is missing [Content_Types].xml.",
                    category="office_compatibility",
                )
            )
        else:
            issues.extend(self._content_type_checks(pptx_path, names))
        issues.extend(self._relationship_checks(pptx_path, names))
        try:
            prs = Presentation(pptx_path.as_posix())
        except Exception as exc:
            issues.append(
                QAIssue(
                    severity="CRITICAL",
                    message=f"python-pptx could not open the package: {exc}",
                    category="office_compatibility",
                )
            )
            return issues
        if outlines and len(prs.slides) != len(outlines):
            issues.append(
                QAIssue(
                    severity="WARNING",
                    message="Rendered PPTX slide count does not match planned outlines.",
                    category="slide_count",
                )
            )
        slide_width = self._emu_to_inches(prs.slide_width)
        slide_height = self._emu_to_inches(prs.slide_height)
        for idx, slide in enumerate(prs.slides):
            issues.extend(self._slide_geometry_checks(slide, idx, slide_width, slide_height))
        return issues

    def _content_type_checks(self, pptx_path: Path, names: set[str]) -> list[QAIssue]:
        issues: list[QAIssue] = []
        try:
            with zipfile.ZipFile(pptx_path, "r") as pptx_zip:
                root = etree.fromstring(
                    pptx_zip.read("[Content_Types].xml"),
                    parser=SAFE_XML_PARSER,
                )
        except Exception as exc:
            return [
                QAIssue(
                    severity="CRITICAL",
                    message=f"[Content_Types].xml could not be parsed: {exc}",
                    category="office_compatibility",
                )
            ]
        defaults = {
            node.get("Extension")
            for node in root.findall(f".//{{{CONTENT_TYPES_NS}}}Default")
            if node.get("Extension")
        }
        overrides = {
            node.get("PartName", "").lstrip("/")
            for node in root.findall(f".//{{{CONTENT_TYPES_NS}}}Override")
            if node.get("PartName")
        }
        for name in sorted(names):
            if name == "[Content_Types].xml" or name.endswith("/"):
                continue
            extension = name.rsplit(".", 1)[-1] if "." in name else ""
            if name not in overrides and extension not in defaults:
                issues.append(
                    QAIssue(
                        severity="CRITICAL",
                        message=f"PPTX part has no content type declaration: {name}",
                        category="office_compatibility",
                    )
                )
        return issues

    def _relationship_checks(self, pptx_path: Path, names: set[str]) -> list[QAIssue]:
        issues: list[QAIssue] = []
        try:
            with zipfile.ZipFile(pptx_path, "r") as pptx_zip:
                rel_names = sorted(name for name in names if name.endswith(".rels"))
                for rel_name in rel_names:
                    try:
                        root = etree.fromstring(
                            pptx_zip.read(rel_name),
                            parser=SAFE_XML_PARSER,
                        )
                    except Exception as exc:
                        issues.append(
                            QAIssue(
                                severity="CRITICAL",
                                message=f"Relationship part could not be parsed: {rel_name}: {exc}",
                                category="office_compatibility",
                            )
                        )
                        continue
                    for rel in root.findall(f".//{{{REL_NS}}}Relationship"):
                        target = rel.get("Target")
                        if not target or rel.get("TargetMode") == "External":
                            continue
                        if self._is_external_target(target):
                            continue
                        resolved = self._resolve_relationship_target(rel_name, target)
                        if resolved not in names:
                            issues.append(
                                QAIssue(
                                    severity="CRITICAL",
                                    message=(
                                        f"Relationship target is missing: {rel_name} -> {target}"
                                    ),
                                    category="office_compatibility",
                                )
                            )
        except zipfile.BadZipFile:
            return issues
        return issues

    def _resolve_relationship_target(self, rels_name: str, target: str) -> str:
        if target.startswith("/"):
            return posixpath.normpath(target.lstrip("/"))
        source_part = self._source_part_for_rels(rels_name)
        base_dir = posixpath.dirname(source_part)
        return posixpath.normpath(posixpath.join(base_dir, target))

    def _source_part_for_rels(self, rels_name: str) -> str:
        if rels_name == "_rels/.rels":
            return ""
        return rels_name.replace("/_rels/", "/").removesuffix(".rels")

    def _is_external_target(self, target: str) -> bool:
        return bool(re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", target))

    def _slide_geometry_checks(
        self, slide, slide_index: int, slide_width: float, slide_height: float
    ) -> list[QAIssue]:
        issues: list[QAIssue] = []
        body_chars = 0
        body_text_shape_count = 0
        for shape in slide.shapes:
            left = self._emu_to_inches(shape.left)
            top = self._emu_to_inches(shape.top)
            width = self._emu_to_inches(shape.width)
            height = self._emu_to_inches(shape.height)
            if left < -0.02 or top < -0.02:
                issues.append(
                    QAIssue(
                        severity="CRITICAL",
                        message="Slide shape starts outside the slide bounds.",
                        slide_index=slide_index,
                        category="shape_bounds",
                    )
                )
            if left + width > slide_width + 0.05 or top + height > slide_height + 0.05:
                issues.append(
                    QAIssue(
                        severity="CRITICAL",
                        message="Slide shape extends beyond the slide bounds.",
                        slide_index=slide_index,
                        category="shape_bounds",
                    )
                )
            if not getattr(shape, "has_text_frame", False):
                continue
            text = " ".join(shape.text.split())
            if not text:
                continue
            if top < 0.95 or top > slide_height - 0.7:
                continue
            body_text_shape_count += 1
            body_chars += len(text)
            area = max(width * height, 0.1)
            if len(text) / area > 180:
                issues.append(
                    QAIssue(
                        severity="WARNING",
                        message="Text density creates possible overflow risk.",
                        slide_index=slide_index,
                        category="overflow_risk",
                    )
                )
        if body_chars > 1500:
            issues.append(
                QAIssue(
                    severity="WARNING",
                    message="Slide has high text density; consider splitting or simplifying.",
                    slide_index=slide_index,
                    category="text_density",
                )
            )
        if body_text_shape_count > 10:
            issues.append(
                QAIssue(
                    severity="WARNING",
                    message="Slide has many text objects and may be difficult to scan.",
                    slide_index=slide_index,
                    category="scanability",
                )
            )
        return issues

    def _render_pptx_preview_fallback(self, pptx_path: Path, output_dir: Path) -> list[Path]:
        if not pptx_path.exists():
            return []
        try:
            prs = Presentation(pptx_path.as_posix())
        except Exception:
            return []
        output_dir.mkdir(parents=True, exist_ok=True)
        slide_width = max(int(self._emu_to_inches(prs.slide_width) * 120), 1)
        slide_height = max(int(self._emu_to_inches(prs.slide_height) * 120), 1)
        font = ImageFont.load_default()
        images: list[Path] = []
        for idx, slide in enumerate(prs.slides):
            image = Image.new("RGB", (slide_width, slide_height), "white")
            draw = ImageDraw.Draw(image)
            for shape in slide.shapes:
                self._draw_shape_preview(draw, shape, prs, font)
            image_path = output_dir / f"slide-{idx + 1}.png"
            image.save(image_path)
            images.append(image_path)
        return images

    def _draw_shape_preview(self, draw: ImageDraw.ImageDraw, shape, prs, font) -> None:
        scale_x = (self._emu_to_inches(prs.slide_width) * 120) / int(prs.slide_width)
        scale_y = (self._emu_to_inches(prs.slide_height) * 120) / int(prs.slide_height)
        x = int(shape.left * scale_x)
        y = int(shape.top * scale_y)
        w = max(int(shape.width * scale_x), 1)
        h = max(int(shape.height * scale_y), 1)
        fill = self._shape_color(shape, "fill")
        outline = self._shape_color(shape, "line")
        is_textbox = getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.TEXT_BOX
        if not is_textbox or fill is not None:
            draw.rectangle(
                [x, y, x + w, y + h],
                fill=fill,
                outline=outline or (210, 215, 222),
            )
        if getattr(shape, "has_text_frame", False):
            text = " ".join(shape.text.split())
            if text:
                color = self._first_text_color(shape) or (20, 24, 31)
                text_font = self._font_for_shape(shape)
                self._draw_wrapped_text(
                    draw, text, (x + 6, y + 5, w - 12, h - 10), text_font, color
                )

    def _draw_wrapped_text(
        self,
        draw: ImageDraw.ImageDraw,
        text: str,
        box: tuple[int, int, int, int],
        font,
        color: tuple[int, int, int],
    ) -> None:
        x, y, width, height = box
        if width <= 0 or height <= 0:
            return
        words = text.split()
        lines: list[str] = []
        current = ""
        for word in words:
            candidate = f"{current} {word}".strip()
            if draw.textlength(candidate, font=font) <= width or not current:
                current = candidate
            else:
                lines.append(current)
                current = word
        if current:
            lines.append(current)
        bbox = font.getbbox("Ag")
        line_height = max((bbox[3] - bbox[1]) + 4, 12)
        max_lines = max(height // line_height, 1)
        for line in lines[:max_lines]:
            draw.text((x, y), line[:120], fill=color, font=font)
            y += line_height

    def _font_for_shape(self, shape):
        point_size = self._first_text_point_size(shape) or 12
        pixel_size = max(int(point_size * 120 / 72), 10)
        try:
            return ImageFont.truetype("Arial.ttf", pixel_size)
        except Exception:
            return ImageFont.load_default(size=pixel_size)

    def _first_text_point_size(self, shape) -> float | None:
        try:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.size is not None:
                        return float(run.font.size.pt)
                if paragraph.font.size is not None:
                    return float(paragraph.font.size.pt)
        except Exception:
            return None
        return None

    def _shape_color(self, shape, color_type: str) -> tuple[int, int, int] | None:
        try:
            color = shape.fill.fore_color if color_type == "fill" else shape.line.color
            rgb = color.rgb
        except Exception:
            return None
        if rgb is None:
            return None
        return tuple(int(str(rgb)[idx : idx + 2], 16) for idx in (0, 2, 4))

    def _first_text_color(self, shape) -> tuple[int, int, int] | None:
        try:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    rgb = run.font.color.rgb
                    if rgb is not None:
                        return tuple(int(str(rgb)[idx : idx + 2], 16) for idx in (0, 2, 4))
        except Exception:
            return None
        return None

    def _emu_to_inches(self, value) -> float:
        return int(value) / EMU_PER_INCH

    def export_pdf(self, pptx_path: Path, output_dir: Path) -> Optional[Path]:
        try:
            return render_pptx_to_pdf(pptx_path, output_dir)
        except RenderingError:
            return None

    def _extract_json(self, report: str) -> Optional[dict]:
        try:
            return json.loads(report)
        except json.JSONDecodeError:
            pass
        fenced = re.search(r"```json\s*(\{.*?\})\s*```", report, flags=re.DOTALL)
        if fenced:
            try:
                return json.loads(fenced.group(1))
            except json.JSONDecodeError:
                return None
        first_brace = report.find("{")
        last_brace = report.rfind("}")
        if first_brace != -1 and last_brace > first_brace:
            candidate = report[first_brace : last_brace + 1]
            try:
                return json.loads(candidate)
            except json.JSONDecodeError:
                return None
        return None
