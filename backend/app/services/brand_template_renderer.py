import json
import posixpath
import re
import zipfile
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from typing import Any

from lxml import etree
from openpyxl import load_workbook

from app.models.outline import SlideOutline
from app.models.template import TemplateProfile


PML_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
DML_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
SLIDE_REL_TYPE = f"{REL_NS}/slide"
CHART_REL_TYPE = f"{REL_NS}/chart"
IMAGE_REL_TYPE = f"{REL_NS}/image"
PACKAGE_REL_TYPE = f"{REL_NS}/package"
SLIDE_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.slide+xml"
CHART_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.drawingml.chart+xml"
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SAFE_PARSER = etree.XMLParser(resolve_entities=False, no_network=True)
NSMAP = {"p": PML_NS, "a": DML_NS, "c": CHART_NS, "r": REL_NS}
WEAK_TEMPLATE_FRAME_METHODS = {"cyclic_fallback", "low_confidence_match"}
WEAK_TEMPLATE_FRAME_CONFIDENCES = {"fallback", "low"}


class BrandTemplateCloneRenderer:
    """Render brand decks by duplicating mapped source slides and editing text.

    This is intentionally conservative: it preserves source slide structure,
    masters, layouts, media, and chrome, then rewrites inherited text bodies
    selected as content slots. It does not add a parallel generated design over
    the template.
    """

    def render(
        self,
        template: TemplateProfile,
        outlines: list[SlideOutline],
        output_path: Path,
        working_dir: Path,
    ) -> tuple[bool, list[dict[str, str | int]]]:
        if not (template.source_file or "").strip():
            return False, []
        source_path = Path(template.source_file)
        if not source_path.exists() or not source_path.is_file():
            return False, [self._warning(0, "template_clone_edit", "Brand template source file is missing.")]
        if not outlines:
            return False, [self._warning(0, "template_clone_edit", "No outlines are available to clone/edit.")]

        try:
            with zipfile.ZipFile(source_path, "r") as source_zip:
                entries = {name: source_zip.read(name) for name in source_zip.namelist()}
        except Exception as exc:
            return False, [self._warning(0, "template_clone_edit", f"Template PPTX could not be read: {exc}")]

        required = {"ppt/presentation.xml", "ppt/_rels/presentation.xml.rels", "[Content_Types].xml"}
        if missing := sorted(required - set(entries)):
            return False, [self._warning(0, "template_clone_edit", f"Template PPTX is missing required parts: {', '.join(missing)}")]

        blocked_mappings, blocked_warnings = self._blocked_template_mappings(outlines)
        if blocked_mappings:
            self._write_clone_artifact(
                working_dir,
                template,
                blocked_mappings,
                blocked_warnings,
                status="blocked",
            )
            return False, blocked_warnings

        updates: dict[str, bytes] = {}
        deletions: set[str] = set()
        warnings: list[dict[str, str | int]] = []
        mappings: list[dict[str, Any]] = []
        chart_part_names: set[str] = set()
        cleanup_stats: dict[str, Any] = {
            "deleted_unused_slide_part_count": 0,
            "deleted_unused_slide_rels_count": 0,
            "deleted_notes_comment_tag_part_count": 0,
            "deleted_unreferenced_media_part_count": 0,
            "removed_content_type_override_count": 0,
            "deleted_media_parts": [],
            "deleted_parts": [],
        }

        for position, outline in enumerate(outlines, start=1):
            frame = outline.layout_json.get("template_frame")
            if not isinstance(frame, dict):
                return False, [self._warning(outline.slide_index, "template_clone_edit", "Outline is missing a mapped template frame.")]
            try:
                source_slide_number = int(frame["index"]) + 1
            except (KeyError, TypeError, ValueError):
                return False, [self._warning(outline.slide_index, "template_clone_edit", "Mapped template frame has no valid source slide index.")]

            source_slide = f"ppt/slides/slide{source_slide_number}.xml"
            target_slide = f"ppt/slides/slide{position}.xml"
            source_xml = entries.get(source_slide)
            if source_xml is None:
                return False, [self._warning(outline.slide_index, "template_clone_edit", f"Mapped source slide is missing: {source_slide}")]

            edited_xml, edit_warnings, edit_stats = self._edit_slide_xml(
                source_xml,
                outline,
            )
            updates[target_slide] = edited_xml
            warnings.extend(
                self._warning(outline.slide_index, "template_clone_edit", message)
                for message in edit_warnings
            )

            source_rels = self._slide_rels_name(source_slide)
            target_rels = self._slide_rels_name(target_slide)
            if source_rels in entries:
                rel_xml, rel_warnings, rel_stats = self._clone_slide_rels(
                    entries,
                    updates,
                    entries[source_rels],
                    source_slide,
                    target_slide,
                    source_xml,
                    outline,
                    position,
                    set(edit_stats["deleted_media_relationship_ids"]),
                )
                updates[target_rels] = rel_xml
                warnings.extend(
                    self._warning(outline.slide_index, "template_clone_edit", message)
                    for message in rel_warnings
                )
                edit_stats["rewritten_chart_count"] = rel_stats["rewritten_chart_count"]
                edit_stats["rewritten_chart_point_count"] = rel_stats["rewritten_chart_point_count"]
                edit_stats["duplicated_chart_part_count"] = rel_stats["duplicated_chart_part_count"]
                edit_stats["deleted_media_relationship_count"] = rel_stats["deleted_media_relationship_count"]
                edit_stats["edit_targets"].extend(rel_stats["edit_targets"])
                chart_part_names.update(rel_stats["chart_part_names"])
            else:
                deletions.add(target_rels)

            frame["reuse_mode"] = "duplicate-slide-edit"
            frame["clone_edit_applied"] = True
            frame["output_slide"] = position
            frame["rewritten_text_shape_count"] = edit_stats["rewritten_text_shape_count"]
            frame["rewritten_table_cell_count"] = edit_stats["rewritten_table_cell_count"]
            frame["rewritten_chart_count"] = edit_stats["rewritten_chart_count"]
            frame["rewritten_chart_point_count"] = edit_stats["rewritten_chart_point_count"]
            frame["deleted_media_placeholder_count"] = edit_stats["deleted_media_placeholder_count"]
            frame["bolded_text_run_count"] = edit_stats["bolded_text_run_count"]
            slot_cleanup = self._slot_cleanup_result(frame, edit_stats)
            if (
                slot_cleanup["planned_excess_slot_count"]
                and not slot_cleanup["cleanup_satisfied"]
            ):
                warnings.append(
                    self._warning(
                        outline.slide_index,
                        "template_clone_edit",
                        (
                            "Template slot cleanup was planned but not fully observed; "
                            f"deleted {slot_cleanup['actual_deleted_slot_count']}/"
                            f"{slot_cleanup['planned_excess_slot_count']} inherited slot element(s)."
                        ),
                    )
                )
            frame["slot_cleanup"] = slot_cleanup
            frame["edit_target_count"] = len(edit_stats["edit_targets"])
            mappings.append(
                {
                    "output_slide": position,
                    "outline_slide_index": outline.slide_index,
                    "source_slide": source_slide_number,
                    "label": frame.get("label"),
                    "method": frame.get("method"),
                    "match_confidence": frame.get("match_confidence"),
                    "match_score": frame.get("match_score"),
                    "match_reason": frame.get("match_reason"),
                    "template_category": frame.get("content_category"),
                    "visual_guidance": frame.get("visual_guidance"),
                    "closest_candidates": frame.get("closest_candidates") or [],
                    "clone_edit_blocked": False,
                    "rewritten_text_shape_count": edit_stats["rewritten_text_shape_count"],
                    "rewritten_table_cell_count": edit_stats["rewritten_table_cell_count"],
                    "rewritten_chart_count": edit_stats["rewritten_chart_count"],
                    "rewritten_chart_point_count": edit_stats["rewritten_chart_point_count"],
                    "duplicated_chart_part_count": edit_stats["duplicated_chart_part_count"],
                    "bolded_text_run_count": edit_stats["bolded_text_run_count"],
                    "deleted_excess_text_shape_count": edit_stats["deleted_excess_text_shape_count"],
                    "deleted_table_row_count": edit_stats["deleted_table_row_count"],
                    "deleted_media_placeholder_count": edit_stats["deleted_media_placeholder_count"],
                    "deleted_media_relationship_count": edit_stats["deleted_media_relationship_count"],
                    "slot_plan": frame.get("slot_plan"),
                    "slot_cleanup": slot_cleanup,
                    "edit_target_count": len(edit_stats["edit_targets"]),
                    "editTargets": edit_stats["edit_targets"],
                }
            )

        cleanup_stats = self._mark_unused_package_parts_for_deletion(
            entries,
            deletions,
            len(outlines),
        )
        media_cleanup_stats = self._mark_unreferenced_media_parts_for_deletion(
            entries,
            updates,
            deletions,
        )
        cleanup_stats.update(media_cleanup_stats)
        cleanup_stats["deleted_part_count"] = (
            cleanup_stats["deleted_unused_slide_part_count"]
            + cleanup_stats["deleted_unused_slide_rels_count"]
            + cleanup_stats["deleted_notes_comment_tag_part_count"]
            + cleanup_stats["deleted_unreferenced_media_part_count"]
        )
        updates["ppt/presentation.xml"] = self._rewrite_presentation_xml(
            entries["ppt/presentation.xml"],
            len(outlines),
        )
        updates["ppt/_rels/presentation.xml.rels"] = self._rewrite_presentation_rels(
            entries["ppt/_rels/presentation.xml.rels"],
            len(outlines),
        )
        content_types_xml, removed_override_count = self._rewrite_content_types(
            entries["[Content_Types].xml"],
            len(outlines),
            chart_part_names,
        )
        updates["[Content_Types].xml"] = content_types_xml
        cleanup_stats["removed_content_type_override_count"] = removed_override_count
        if "docProps/app.xml" in entries:
            updates["docProps/app.xml"] = self._rewrite_app_props(entries["docProps/app.xml"], len(outlines))

        output_path.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as output_zip:
            names = set(entries) | set(updates)
            for name in sorted(names):
                if name in deletions and name not in updates:
                    continue
                output_zip.writestr(name, updates.get(name, entries.get(name, b"")))

        placeholder_issues = self._unfilled_placeholder_issues(output_path)
        if placeholder_issues:
            issues_by_slide: dict[int, list[dict[str, Any]]] = {}
            for issue in placeholder_issues:
                output_slide = int(issue.get("output_slide") or 0)
                issues_by_slide.setdefault(output_slide, []).append(issue)
            for mapping in mappings:
                output_slide = int(mapping.get("output_slide") or 0)
                slide_issues = issues_by_slide.get(output_slide, [])
                mapping["unfilled_placeholder_count"] = len(slide_issues)
                mapping["unfilled_placeholders"] = slide_issues[:5]
            for issue in placeholder_issues[:5]:
                warnings.append(
                    self._warning(
                        int(issue.get("slide_index") or 0),
                        "template_clone_edit",
                        (
                            "Final PPTX contains an unfilled inherited "
                            f"placeholder ({issue.get('placeholder_type') or 'body'}); "
                            "fill or delete it."
                        ),
                    )
                )

        self._write_clone_artifact(working_dir, template, mappings, warnings, cleanup_stats)
        return True, warnings

    def _blocked_template_mappings(
        self,
        outlines: list[SlideOutline],
    ) -> tuple[list[dict[str, Any]], list[dict[str, str | int]]]:
        mappings: list[dict[str, Any]] = []
        warnings: list[dict[str, str | int]] = []
        for position, outline in enumerate(outlines, start=1):
            frame = outline.layout_json.get("template_frame")
            if not isinstance(frame, dict):
                continue
            reason = self._weak_template_frame_reason(frame)
            if not reason:
                continue
            source_slide = None
            try:
                source_slide = int(frame["index"]) + 1
            except (KeyError, TypeError, ValueError):
                pass
            message = (
                f"Slide {outline.slide_index + 1} has a weak template-frame match; "
                f"clone/edit was blocked ({reason})."
            )
            warnings.append(self._warning(outline.slide_index, "template_clone_edit", message))
            mappings.append(
                {
                    "output_slide": position,
                    "outline_slide_index": outline.slide_index,
                    "source_slide": source_slide,
                    "label": frame.get("label"),
                    "method": frame.get("method"),
                    "match_confidence": frame.get("match_confidence"),
                    "match_score": frame.get("match_score"),
                    "match_reason": frame.get("match_reason"),
                    "template_category": frame.get("content_category"),
                    "visual_guidance": frame.get("visual_guidance"),
                    "closest_candidates": frame.get("closest_candidates") or [],
                    "clone_edit_blocked": True,
                    "block_reason": reason,
                    "slot_plan": frame.get("slot_plan"),
                    "edit_target_count": 0,
                    "editTargets": [],
                }
            )
        return mappings, warnings

    def _weak_template_frame_reason(self, frame: dict[str, Any]) -> str:
        method = str(frame.get("method") or "").strip().lower()
        confidence = str(frame.get("match_confidence") or "").strip().lower()
        if method in WEAK_TEMPLATE_FRAME_METHODS:
            return f"method={method}"
        if confidence in WEAK_TEMPLATE_FRAME_CONFIDENCES:
            return f"match_confidence={confidence}"
        return ""

    def _slot_cleanup_result(
        self,
        frame: dict[str, Any],
        edit_stats: dict[str, Any],
    ) -> dict[str, Any]:
        slot_plan = frame.get("slot_plan")
        planned = 0
        action = ""
        if isinstance(slot_plan, dict):
            action = str(slot_plan.get("action") or "")
            try:
                planned = int(slot_plan.get("excess_slot_count") or 0)
            except (TypeError, ValueError):
                planned = 0
        actual = sum(
            int(edit_stats.get(key) or 0)
            for key in (
                "deleted_excess_text_shape_count",
                "deleted_table_row_count",
                "deleted_media_placeholder_count",
            )
        )
        cleanup_required = action == "delete_excess_template_elements" and planned > 0
        return {
            "action": action,
            "planned_excess_slot_count": planned,
            "actual_deleted_slot_count": actual,
            "cleanup_required": cleanup_required,
            "cleanup_satisfied": (actual >= planned) if cleanup_required else True,
            "deleted_text_shape_count": int(edit_stats.get("deleted_excess_text_shape_count") or 0),
            "deleted_table_row_count": int(edit_stats.get("deleted_table_row_count") or 0),
            "deleted_media_placeholder_count": int(edit_stats.get("deleted_media_placeholder_count") or 0),
        }

    def _edit_slide_xml(
        self,
        slide_xml: bytes,
        outline: SlideOutline,
    ) -> tuple[bytes, list[str], dict[str, Any]]:
        root = etree.fromstring(slide_xml, parser=SAFE_PARSER)
        shapes = self._text_shapes(root)
        warnings: list[str] = []
        if not shapes:
            warnings.append("Mapped source slide has no editable text bodies.")
            return slide_xml, warnings, {
                "rewritten_text_shape_count": 0,
                "rewritten_table_cell_count": 0,
                "rewritten_chart_count": 0,
                "rewritten_chart_point_count": 0,
                "duplicated_chart_part_count": 0,
                "deleted_table_row_count": 0,
                "deleted_excess_text_shape_count": 0,
                "deleted_media_placeholder_count": 0,
                "deleted_media_relationship_count": 0,
                "deleted_media_relationship_ids": [],
                "bolded_text_run_count": 0,
                "edit_targets": [],
            }

        edit_targets: list[dict[str, Any]] = []
        bolded_text_run_count = 0
        table_edited = False
        title_shape = self._title_shape(shapes)
        if title_shape is None:
            title_shape = shapes[0]
            warnings.append("No clear inherited title slot found; used first editable text body.")
        title_text = self._title_text(outline)
        bolded_text_run_count += self._rewrite_text_body(
            title_shape["element"],
            [self._fit_text(title_text, 118)],
            bold_all=True,
        )
        edit_targets.append(
            self._edit_target(title_shape, "rewrite", "title", [title_text])
        )
        used = {id(title_shape["element"])}

        table_rows = self._table_rows(outline)
        table_frames = self._table_frames(root)
        rewritten_table_count = 0
        deleted_table_row_count = 0
        if table_rows and table_frames:
            table_stats = self._rewrite_table(table_frames[0], table_rows)
            rewritten_table_count = table_stats["rewritten_table_cell_count"]
            deleted_table_row_count = table_stats["deleted_table_row_count"]
            bolded_text_run_count += table_stats["bolded_text_run_count"]
            edit_targets.extend(table_stats["edit_targets"])
            table_edited = True
            if table_stats["truncated_row_count"]:
                warnings.append(
                    f"Inherited table has fewer rows than source data; truncated {table_stats['truncated_row_count']} row(s)."
                )
        elif table_rows:
            warnings.append("Mapped source slide has table-ready evidence but no inherited table frame.")

        chart_ready = bool(self._chart_points(outline))
        body_texts = self._body_texts(
            outline,
            include_exhibit=not table_edited and not chart_ready,
            include_table_blocks=not table_edited,
        )
        body_shapes = [
            shape
            for shape in shapes
            if id(shape["element"]) not in used and self._is_body_slot(shape)
        ]
        rewritten_body_count = 0
        deleted_count = 0
        if body_texts and not body_shapes:
            warnings.append("Mapped source slide has no body slot for generated evidence text.")
        elif body_shapes:
            assignments = self._assign_body_texts(body_texts, body_shapes)
            for shape, texts in assignments:
                bolded_text_run_count += self._rewrite_text_body(
                    shape["element"],
                    texts,
                    bold_first=self._first_body_text_is_subheading(outline, texts),
                )
                used.add(id(shape["element"]))
                rewritten_body_count += 1
                edit_targets.append(self._edit_target(shape, "rewrite", "body", texts))
            for shape in body_shapes:
                if id(shape["element"]) in used:
                    continue
                parent = shape["element"].getparent()
                if parent is not None:
                    parent.remove(shape["element"])
                    deleted_count += 1
                    edit_targets.append(self._edit_target(shape, "delete", "excess_body", []))

        deleted_media_relationship_ids: set[str] = set()
        deleted_media_placeholder_count = 0
        media_placeholders = self._media_placeholders(root)
        has_media_intent = self._has_media_intent(outline)
        for placeholder in media_placeholders:
            if not self._is_media_placeholder(placeholder):
                continue
            if has_media_intent:
                warnings.append(
                    "Mapped source slide has an inherited media placeholder, but generated image replacement is not available yet."
                )
                continue
            parent = placeholder["element"].getparent()
            if parent is None:
                continue
            parent.remove(placeholder["element"])
            deleted_media_placeholder_count += 1
            deleted_media_relationship_ids.update(placeholder["rel_ids"])
            edit_targets.append(self._media_edit_target(placeholder, "delete", "media_placeholder"))

        xml_bytes = etree.tostring(
            root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )
        etree.fromstring(xml_bytes, parser=SAFE_PARSER)
        return xml_bytes, warnings, {
            "rewritten_text_shape_count": 1 + rewritten_body_count,
            "rewritten_table_cell_count": rewritten_table_count,
            "rewritten_chart_count": 0,
            "rewritten_chart_point_count": 0,
            "duplicated_chart_part_count": 0,
            "deleted_excess_text_shape_count": deleted_count,
            "deleted_table_row_count": deleted_table_row_count,
            "deleted_media_placeholder_count": deleted_media_placeholder_count,
            "deleted_media_relationship_count": 0,
            "deleted_media_relationship_ids": sorted(deleted_media_relationship_ids),
            "bolded_text_run_count": bolded_text_run_count,
            "edit_targets": edit_targets,
        }

    def _text_shapes(self, root: etree._Element) -> list[dict[str, Any]]:
        shapes: list[dict[str, Any]] = []
        for sp in root.xpath(".//p:cSld//p:sp", namespaces=NSMAP):
            tx_body = sp.find("p:txBody", namespaces=NSMAP)
            if tx_body is None:
                continue
            bounds = self._shape_bounds(sp)
            text = " ".join(
                (node.text or "").strip()
                for node in sp.xpath(".//a:t", namespaces=NSMAP)
                if (node.text or "").strip()
            )
            ph = sp.find("p:nvSpPr/p:nvPr/p:ph", namespaces=NSMAP)
            placeholder_type = ph.get("type") if ph is not None else ""
            c_nv_pr = sp.find("p:nvSpPr/p:cNvPr", namespaces=NSMAP)
            shape_id = c_nv_pr.get("id") if c_nv_pr is not None else ""
            name = c_nv_pr.get("name") if c_nv_pr is not None else ""
            if self._is_chrome_text(text, placeholder_type, bounds):
                continue
            shapes.append(
                {
                    "element": sp,
                    "shapeId": shape_id,
                    "sourceElementId": shape_id,
                    "name": name,
                    "text": text,
                    "placeholder_type": placeholder_type or "",
                    "bounds": bounds,
                }
            )
        return sorted(shapes, key=lambda item: (item["bounds"]["y"], item["bounds"]["x"]))

    def _table_frames(self, root: etree._Element) -> list[dict[str, Any]]:
        tables: list[dict[str, Any]] = []
        for frame in root.xpath(".//p:cSld//p:graphicFrame[.//a:tbl]", namespaces=NSMAP):
            table = frame.find(".//a:tbl", namespaces=NSMAP)
            if table is None:
                continue
            c_nv_pr = frame.find("p:nvGraphicFramePr/p:cNvPr", namespaces=NSMAP)
            shape_id = c_nv_pr.get("id") if c_nv_pr is not None else ""
            name = c_nv_pr.get("name") if c_nv_pr is not None else ""
            tables.append(
                {
                    "element": frame,
                    "table": table,
                    "shapeId": shape_id,
                    "sourceElementId": shape_id,
                    "name": name,
                    "bounds": self._graphic_bounds(frame),
                }
            )
        return sorted(tables, key=lambda item: (item["bounds"]["y"], item["bounds"]["x"]))

    def _media_placeholders(self, root: etree._Element) -> list[dict[str, Any]]:
        placeholders: list[dict[str, Any]] = []
        for pic in root.xpath(".//p:cSld//p:pic", namespaces=NSMAP):
            c_nv_pr = pic.find("p:nvPicPr/p:cNvPr", namespaces=NSMAP)
            nv_pr = pic.find("p:nvPicPr/p:nvPr", namespaces=NSMAP)
            ph = nv_pr.find("p:ph", namespaces=NSMAP) if nv_pr is not None else None
            rel_ids = []
            for blip in pic.xpath(".//a:blip", namespaces=NSMAP):
                for attr_name in (f"{{{REL_NS}}}embed", f"{{{REL_NS}}}link"):
                    rel_id = blip.get(attr_name)
                    if rel_id:
                        rel_ids.append(rel_id)
            shape_id = c_nv_pr.get("id") if c_nv_pr is not None else ""
            name = c_nv_pr.get("name") if c_nv_pr is not None else ""
            descr = c_nv_pr.get("descr") if c_nv_pr is not None else ""
            placeholders.append(
                {
                    "element": pic,
                    "shapeId": shape_id,
                    "sourceElementId": shape_id,
                    "name": name,
                    "descr": descr,
                    "text": "",
                    "placeholder_type": ph.get("type") if ph is not None else "",
                    "bounds": self._shape_bounds(pic),
                    "rel_ids": sorted(set(rel_ids)),
                }
            )
        return sorted(placeholders, key=lambda item: (item["bounds"]["y"], item["bounds"]["x"]))

    def _shape_bounds(self, shape: etree._Element) -> dict[str, float]:
        xfrm = shape.find("p:spPr/a:xfrm", namespaces=NSMAP)
        off = xfrm.find("a:off", namespaces=NSMAP) if xfrm is not None else None
        ext = xfrm.find("a:ext", namespaces=NSMAP) if xfrm is not None else None
        return self._bounds_from_xfrm(off, ext)

    def _graphic_bounds(self, frame: etree._Element) -> dict[str, float]:
        xfrm = frame.find("p:xfrm", namespaces=NSMAP)
        off = xfrm.find("a:off", namespaces=NSMAP) if xfrm is not None else None
        ext = xfrm.find("a:ext", namespaces=NSMAP) if xfrm is not None else None
        return self._bounds_from_xfrm(off, ext)

    def _bounds_from_xfrm(
        self,
        off: etree._Element | None,
        ext: etree._Element | None,
    ) -> dict[str, float]:
        return {
            "x": self._emu_to_inches(off.get("x") if off is not None else None),
            "y": self._emu_to_inches(off.get("y") if off is not None else None),
            "w": self._emu_to_inches(ext.get("cx") if ext is not None else None),
            "h": self._emu_to_inches(ext.get("cy") if ext is not None else None),
        }

    def _emu_to_inches(self, value: str | None) -> float:
        try:
            return round(int(value or 0) / 914400, 3)
        except (TypeError, ValueError):
            return 0.0

    def _is_chrome_text(
        self,
        text: str,
        placeholder_type: str,
        bounds: dict[str, float],
    ) -> bool:
        if placeholder_type in {"sldNum", "dt", "ftr"}:
            return True
        normalized = " ".join(text.casefold().split())
        if normalized in {"slide number", "date", "footer"}:
            return True
        if bounds["y"] > 6.35:
            return True
        if bounds["h"] <= 0.18 and bounds["w"] <= 2.2:
            return True
        return False

    def _title_shape(self, shapes: list[dict[str, Any]]) -> dict[str, Any] | None:
        for shape in shapes:
            if shape["placeholder_type"] in {"title", "ctrTitle"}:
                return shape
        candidates = [shape for shape in shapes if shape["bounds"]["y"] <= 1.65]
        if candidates:
            return max(candidates, key=lambda item: item["bounds"]["w"] * max(item["bounds"]["h"], 0.1))
        return None

    def _is_body_slot(self, shape: dict[str, Any]) -> bool:
        ph_type = shape["placeholder_type"]
        if ph_type in {"title", "ctrTitle", "sldNum", "dt", "ftr"}:
            return False
        bounds = shape["bounds"]
        return bounds["y"] < 6.35 and bounds["h"] >= 0.18

    def _is_media_placeholder(self, frame: dict[str, Any]) -> bool:
        label = f"{frame.get('name') or ''} {frame.get('descr') or ''}".casefold()
        if any(token in label for token in ("logo", "wordmark", "brand mark", "brandmark")):
            return False
        bounds = frame.get("bounds") or {}
        if float(bounds.get("y") or 0) > 6.25:
            return False
        if float(bounds.get("h") or 0) <= 0.24 and float(bounds.get("w") or 0) <= 2.2:
            return False
        placeholder_type = str(frame.get("placeholder_type") or "").casefold()
        if placeholder_type in {"pic", "obj", "media", "clipart"}:
            return True
        if "placeholder" not in label:
            return False
        return any(token in label for token in ("picture", "image", "media", "photo", "screenshot"))

    def _has_media_intent(self, outline: SlideOutline) -> bool:
        content = outline.content_json or {}
        if self._dict_has_media_intent(content):
            return True
        exhibit = content.get("exhibit_spec")
        if isinstance(exhibit, dict) and self._dict_has_media_intent(exhibit):
            return True
        chart = content.get("chart_spec")
        if isinstance(chart, dict) and self._dict_has_media_intent(chart):
            return True
        layout = " ".join(
            str(value)
            for value in (
                outline.layout_json.get("layout"),
                outline.layout_json.get("archetype"),
                outline.layout_json.get("composition_family"),
                outline.layout_json.get("composition_signature"),
            )
            if value
        ).casefold()
        return any(token in layout for token in ("image", "media", "photo", "screenshot"))

    def _dict_has_media_intent(self, value: dict[str, Any]) -> bool:
        media_keys = {
            "image",
            "images",
            "image_refs",
            "image_assets",
            "media",
            "media_assets",
            "photo",
            "photos",
            "screenshot",
            "screenshots",
            "visual_asset",
            "visual_assets",
        }
        for key, raw in value.items():
            key_text = str(key).casefold()
            if key_text in media_keys and self._has_nonempty_value(raw):
                return True
            if any(token in key_text for token in ("image", "media", "photo", "screenshot")) and self._has_nonempty_value(raw):
                return True
        return False

    def _has_nonempty_value(self, value: Any) -> bool:
        if value is None:
            return False
        if isinstance(value, str):
            return bool(value.strip())
        if isinstance(value, (list, tuple, set, dict)):
            return bool(value)
        return True

    def _assign_body_texts(
        self,
        body_texts: list[str],
        body_shapes: list[dict[str, Any]],
    ) -> list[tuple[dict[str, Any], list[str]]]:
        if not body_texts:
            return []
        if len(body_shapes) <= 1:
            return [(body_shapes[0], [self._fit_text(text, 145) for text in body_texts[:6]])]
        assignments: list[tuple[dict[str, Any], list[str]]] = []
        for idx, shape in enumerate(body_shapes):
            if idx >= len(body_texts):
                break
            if idx == len(body_shapes) - 1:
                texts = body_texts[idx:]
            else:
                texts = [body_texts[idx]]
            assignments.append((shape, [self._fit_text(text, 145) for text in texts[:4]]))
        return assignments

    def _edit_target(
        self,
        shape: dict[str, Any],
        action: str,
        role: str,
        replacement_texts: list[str],
    ) -> dict[str, Any]:
        return {
            "action": action,
            "role": role,
            "shapeId": shape.get("shapeId"),
            "sourceElementId": shape.get("sourceElementId"),
            "name": shape.get("name"),
            "placeholderType": shape.get("placeholder_type"),
            "bounds": shape.get("bounds"),
            "originalText": self._fit_text(str(shape.get("text") or ""), 120),
            "replacementParagraphCount": len(replacement_texts),
        }

    def _media_edit_target(
        self,
        frame: dict[str, Any],
        action: str,
        role: str,
    ) -> dict[str, Any]:
        return {
            "action": action,
            "role": role,
            "shapeId": frame.get("shapeId"),
            "sourceElementId": frame.get("sourceElementId"),
            "name": frame.get("name"),
            "placeholderType": frame.get("placeholder_type"),
            "bounds": frame.get("bounds"),
            "relationshipIds": frame.get("rel_ids") or [],
            "replacementParagraphCount": 0,
        }

    def _unfilled_placeholder_issues(self, pptx_path: Path) -> list[dict[str, Any]]:
        issues: list[dict[str, Any]] = []
        try:
            with zipfile.ZipFile(pptx_path, "r") as pptx_zip:
                slide_names = sorted(
                    name
                    for name in pptx_zip.namelist()
                    if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
                )
                for slide_name in slide_names:
                    root = etree.fromstring(
                        pptx_zip.read(slide_name),
                        parser=SAFE_PARSER,
                    )
                    slide_index = self._slide_index_from_name(slide_name)
                    for shape in root.xpath(
                        ".//p:sp[p:nvSpPr/p:nvPr/p:ph]",
                        namespaces=NSMAP,
                    ):
                        texts = [
                            node.text or ""
                            for node in shape.xpath(".//a:t", namespaces=NSMAP)
                        ]
                        if any(text.strip() for text in texts):
                            continue
                        ph = shape.find("p:nvSpPr/p:nvPr/p:ph", namespaces=NSMAP)
                        c_nv_pr = shape.find("p:nvSpPr/p:cNvPr", namespaces=NSMAP)
                        issues.append(
                            {
                                "output_slide": slide_index + 1,
                                "slide_index": slide_index,
                                "shapeId": c_nv_pr.get("id") if c_nv_pr is not None else "",
                                "name": c_nv_pr.get("name") if c_nv_pr is not None else "",
                                "placeholder_type": ph.get("type") if ph is not None else "body",
                                "bounds": self._shape_bounds(shape),
                            }
                        )
        except (OSError, zipfile.BadZipFile, etree.XMLSyntaxError):
            return issues
        return issues

    def _slide_index_from_name(self, slide_name: str) -> int:
        match = re.search(r"slide(\d+)\.xml$", slide_name)
        if not match:
            return 0
        return max(0, int(match.group(1)) - 1)

    def _table_edit_target(
        self,
        frame: dict[str, Any],
        action: str,
        role: str,
        row_index: int,
        column_index: int | None = None,
        original_text: str = "",
        replacement_text: str = "",
    ) -> dict[str, Any]:
        return {
            "action": action,
            "role": role,
            "shapeId": frame.get("shapeId"),
            "sourceElementId": frame.get("sourceElementId"),
            "name": frame.get("name"),
            "bounds": frame.get("bounds"),
            "row": row_index,
            "column": column_index,
            "originalText": self._fit_text(original_text, 120),
            "replacementParagraphCount": 1 if replacement_text else 0,
        }

    def _rewrite_table(
        self,
        frame: dict[str, Any],
        rows: list[list[str]],
    ) -> dict[str, Any]:
        table = frame["table"]
        existing_rows = list(table.findall("a:tr", namespaces=NSMAP))
        edit_targets: list[dict[str, Any]] = []
        rewritten = 0
        if not existing_rows:
            return {
                "rewritten_table_cell_count": 0,
                "deleted_table_row_count": 0,
                "truncated_row_count": len(rows),
                "bolded_text_run_count": 0,
                "edit_targets": edit_targets,
            }
        writable_rows = min(len(rows), len(existing_rows))
        truncated = max(0, len(rows) - writable_rows)
        bolded_text_run_count = 0
        for row_index in range(writable_rows):
            table_row = existing_rows[row_index]
            cells = list(table_row.findall("a:tc", namespaces=NSMAP))
            if not cells:
                continue
            values = rows[row_index]
            for column_index, cell in enumerate(cells):
                replacement = values[column_index] if column_index < len(values) else ""
                original = self._cell_text(cell)
                bolded_text_run_count += self._rewrite_table_cell(
                    cell,
                    self._fit_text(replacement, 120),
                    bold=row_index == 0,
                )
                rewritten += 1
                edit_targets.append(
                    self._table_edit_target(
                        frame,
                        "rewrite",
                        "table_cell",
                        row_index,
                        column_index,
                        original_text=original,
                        replacement_text=replacement,
                    )
                )
        deleted = 0
        for row_index, table_row in enumerate(existing_rows[writable_rows:], start=writable_rows):
            table.remove(table_row)
            deleted += 1
            edit_targets.append(
                self._table_edit_target(
                    frame,
                    "delete",
                    "excess_table_row",
                    row_index,
                )
            )
        return {
            "rewritten_table_cell_count": rewritten,
            "deleted_table_row_count": deleted,
            "truncated_row_count": truncated,
            "bolded_text_run_count": bolded_text_run_count,
            "edit_targets": edit_targets,
        }

    def _cell_text(self, cell: etree._Element) -> str:
        return " ".join(
            (node.text or "").strip()
            for node in cell.xpath(".//a:t", namespaces=NSMAP)
            if (node.text or "").strip()
        )

    def _rewrite_table_cell(
        self,
        cell: etree._Element,
        text: str,
        bold: bool = False,
    ) -> int:
        tx_body = cell.find("a:txBody", namespaces=NSMAP)
        if tx_body is None:
            tx_body = etree.SubElement(cell, f"{{{DML_NS}}}txBody")
            etree.SubElement(tx_body, f"{{{DML_NS}}}bodyPr")
            etree.SubElement(tx_body, f"{{{DML_NS}}}lstStyle")
        template_p = tx_body.find("a:p", namespaces=NSMAP)
        p_pr = template_p.find("a:pPr", namespaces=NSMAP) if template_p is not None else None
        r_pr = template_p.find(".//a:rPr", namespaces=NSMAP) if template_p is not None else None
        for paragraph in list(tx_body.findall("a:p", namespaces=NSMAP)):
            tx_body.remove(paragraph)
        paragraph = etree.Element(f"{{{DML_NS}}}p")
        if p_pr is not None:
            paragraph.append(deepcopy(p_pr))
        bolded = self._append_text_run(paragraph, text, r_pr, bold=bold)
        tx_body.append(paragraph)
        return bolded

    def _rewrite_text_body(
        self,
        shape: etree._Element,
        texts: list[str],
        bold_all: bool = False,
        bold_first: bool = False,
    ) -> int:
        tx_body = shape.find("p:txBody", namespaces=NSMAP)
        if tx_body is None:
            return 0
        template_p = tx_body.find("a:p", namespaces=NSMAP)
        p_pr = template_p.find("a:pPr", namespaces=NSMAP) if template_p is not None else None
        r_pr = template_p.find(".//a:rPr", namespaces=NSMAP) if template_p is not None else None
        for paragraph in list(tx_body.findall("a:p", namespaces=NSMAP)):
            tx_body.remove(paragraph)
        bolded_text_run_count = 0
        for index, text in enumerate(texts or [""]):
            paragraph = etree.Element(f"{{{DML_NS}}}p")
            if p_pr is not None:
                paragraph.append(deepcopy(p_pr))
            force_bold = bold_all or (bold_first and index == 0)
            bolded_text_run_count += self._append_paragraph_text(
                paragraph,
                str(text),
                r_pr,
                force_bold=force_bold,
            )
            tx_body.append(paragraph)
        return bolded_text_run_count

    def _append_paragraph_text(
        self,
        paragraph: etree._Element,
        text: str,
        r_pr: etree._Element | None,
        force_bold: bool = False,
    ) -> int:
        if force_bold:
            return self._append_text_run(paragraph, text, r_pr, bold=True)
        label_match = re.match(
            r"^([A-Za-z][A-Za-z0-9 /&().-]{1,44}:)(\s+.+)$",
            text,
        )
        if not label_match:
            return self._append_text_run(paragraph, text, r_pr, bold=False)
        bolded = self._append_text_run(paragraph, label_match.group(1), r_pr, bold=True)
        self._append_text_run(paragraph, label_match.group(2), r_pr, bold=False)
        return bolded

    def _append_text_run(
        self,
        paragraph: etree._Element,
        text: str,
        r_pr: etree._Element | None,
        bold: bool = False,
    ) -> int:
        run = etree.SubElement(paragraph, f"{{{DML_NS}}}r")
        run_pr = deepcopy(r_pr) if r_pr is not None else etree.Element(f"{{{DML_NS}}}rPr")
        if bold:
            run_pr.set("b", "1")
        run.append(run_pr)
        node = etree.SubElement(run, f"{{{DML_NS}}}t")
        if text[:1].isspace() or text[-1:].isspace():
            node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        node.text = text
        return 1 if bold else 0

    def _title_text(self, outline: SlideOutline) -> str:
        return str(
            outline.content_json.get("action_title")
            or outline.content_json.get("title")
            or outline.label
        ).strip()

    def _first_body_text_is_subheading(
        self,
        outline: SlideOutline,
        texts: list[str],
    ) -> bool:
        if not texts:
            return False
        first = " ".join(str(texts[0]).split()).casefold()
        content = outline.content_json or {}
        for key in ("subheading", "summary"):
            value = " ".join(str(content.get(key) or "").split()).casefold()
            if value and first == value:
                return True
        return False

    def _body_texts(
        self,
        outline: SlideOutline,
        include_exhibit: bool = True,
        include_table_blocks: bool = True,
    ) -> list[str]:
        content = outline.content_json or {}
        values: list[str] = []
        subheading = str(content.get("subheading") or content.get("summary") or "").strip()
        if subheading and not self._looks_like_meta_text(subheading):
            values.append(subheading)
        for block in content.get("content_blocks") or []:
            if not isinstance(block, dict):
                continue
            body = block.get("body")
            if block.get("type") == "table" and not include_table_blocks:
                continue
            if isinstance(body, list):
                for item in body:
                    if isinstance(item, (list, tuple, dict)):
                        values.append(self._row_text(item))
                    elif str(item).strip():
                        values.append(str(item).strip())
            elif isinstance(body, str) and body.strip():
                values.extend(self._split_multi_item_text(body))
        for item in content.get("bullets") or []:
            if str(item).strip():
                values.append(str(item).strip())
        exhibit = content.get("exhibit_spec")
        if include_exhibit and isinstance(exhibit, dict):
            values.extend(self._exhibit_texts(exhibit))
        deduped: list[str] = []
        seen: set[str] = set()
        for value in values:
            value = self._clean_text(value)
            key = value.casefold()
            if value and key not in seen:
                seen.add(key)
                deduped.append(value)
        return deduped[:8]

    def _table_rows(self, outline: SlideOutline) -> list[list[str]]:
        content = outline.content_json or {}
        rows: list[list[str]] = []
        exhibit = content.get("exhibit_spec")
        if isinstance(exhibit, dict) and exhibit.get("type") in {
            "comparison_table",
            "reference_table",
        }:
            columns = exhibit.get("columns")
            if isinstance(columns, list):
                row = [self._clean_text(str(value)) for value in columns if str(value).strip()]
                if row:
                    rows.append(row)
            for raw_row in exhibit.get("rows") or []:
                row = self._row_values(raw_row)
                if row:
                    rows.append(row)
        if rows:
            return rows[:8]
        for block in content.get("content_blocks") or []:
            if not isinstance(block, dict) or block.get("type") != "table":
                continue
            body = block.get("body")
            if not isinstance(body, list):
                continue
            for raw_row in body:
                row = self._row_values(raw_row)
                if row:
                    rows.append(row)
        return rows[:8]

    def _row_values(self, row: Any) -> list[str]:
        if isinstance(row, dict):
            values = []
            label = str(row.get("label", "")).strip()
            if label:
                values.append(label)
            raw_values = row.get("values")
            if isinstance(raw_values, list):
                values.extend(str(value).strip() for value in raw_values if str(value).strip())
            else:
                values.extend(
                    str(value).strip()
                    for key, value in row.items()
                    if key not in {"label", "values"} and str(value).strip()
                )
        elif isinstance(row, (list, tuple)):
            values = [str(value).strip() for value in row if str(value).strip()]
        elif str(row).strip():
            values = [str(row).strip()]
        else:
            values = []
        return [self._clean_text(value) for value in values]

    def _split_multi_item_text(self, text: str) -> list[str]:
        parts = re.split(r"(?=\b(?:Step|Phase|Stage)\s+\d+\s*[:.)-]?)", text)
        cleaned = [part.strip(" \n\t;") for part in parts if part.strip(" \n\t;")]
        return cleaned or [text.strip()]

    def _exhibit_texts(self, exhibit: dict[str, Any]) -> list[str]:
        values: list[str] = []
        for key in ("items", "steps", "messages", "supporting_points", "next_steps"):
            raw = exhibit.get(key)
            if isinstance(raw, list):
                values.extend(str(item).strip() for item in raw if str(item).strip())
        rows = exhibit.get("rows")
        if isinstance(rows, list):
            for row in rows[:4]:
                values.append(self._row_text(row))
        metrics = exhibit.get("metrics")
        if isinstance(metrics, list):
            for metric in metrics[:4]:
                if isinstance(metric, dict):
                    label = str(metric.get("label") or metric.get("name") or "").strip()
                    value = str(metric.get("value") or "").strip()
                    if label and value:
                        values.append(f"{label}: {value}")
                    elif label:
                        values.append(label)
                elif str(metric).strip():
                    values.append(str(metric).strip())
        return values

    def _row_text(self, row: Any) -> str:
        if isinstance(row, dict):
            parts = [str(value).strip() for value in row.values() if str(value).strip()]
        elif isinstance(row, (list, tuple)):
            parts = [str(value).strip() for value in row if str(value).strip()]
        else:
            parts = [str(row).strip()]
        if not parts:
            return ""
        if len(parts) == 1:
            return parts[0]
        return f"{parts[0]}: {' / '.join(parts[1:])}"

    def _looks_like_meta_text(self, text: str) -> bool:
        normalized = " ".join(text.casefold().split())
        return any(
            marker in normalized
            for marker in (
                "layout instruction",
                "placeholder",
                "cover slide",
                "executive overview deck",
                "visually tied",
            )
        )

    def _clean_text(self, text: str) -> str:
        text = re.sub(r"[\u2022\u25e6\u25aa\u25cf]\s*", "", str(text))
        return " ".join(text.split())

    def _fit_text(self, text: str, max_chars: int) -> str:
        text = self._clean_text(text)
        if len(text) <= max_chars:
            return text
        clipped = text[: max_chars - 1].rsplit(" ", 1)[0].strip()
        return f"{clipped}." if clipped and not clipped.endswith((".", "!", "?")) else clipped

    def _clone_slide_rels(
        self,
        entries: dict[str, bytes],
        updates: dict[str, bytes],
        rel_xml: bytes,
        source_slide: str,
        target_slide: str,
        source_slide_xml: bytes,
        outline: SlideOutline,
        position: int,
        deleted_rel_ids: set[str] | None = None,
    ) -> tuple[bytes, list[str], dict[str, Any]]:
        root = etree.fromstring(rel_xml, parser=SAFE_PARSER)
        warnings: list[str] = []
        deleted_rel_ids = deleted_rel_ids or set()
        stats: dict[str, Any] = {
            "rewritten_chart_count": 0,
            "rewritten_chart_point_count": 0,
            "duplicated_chart_part_count": 0,
            "deleted_media_relationship_count": 0,
            "chart_part_names": set(),
            "edit_targets": [],
        }
        chart_points = self._chart_points(outline)
        chart_frames = self._chart_frames_by_rel_id(source_slide_xml)
        chart_index = 0

        for rel in list(root.findall(f".//{{{PKG_REL_NS}}}Relationship")):
            rel_type = rel.get("Type", "")
            if rel.get("Id") in deleted_rel_ids:
                root.remove(rel)
                if rel_type == IMAGE_REL_TYPE:
                    stats["deleted_media_relationship_count"] += 1
                continue
            if rel_type.endswith(("/notesSlide", "/comments", "/tags")):
                root.remove(rel)
                continue
            if rel_type != CHART_REL_TYPE:
                continue

            chart_index += 1
            rel_id = rel.get("Id") or f"chart{chart_index}"
            target = rel.get("Target")
            if not chart_points:
                warnings.append(
                    "Mapped source slide has inherited chart frame but no chart-ready generated metrics."
                )
                continue
            if not target:
                warnings.append("Mapped source slide has an inherited chart relationship with no target.")
                continue

            source_chart_part = self._resolve_relationship_target(source_slide, target)
            chart_xml = updates.get(source_chart_part, entries.get(source_chart_part))
            if chart_xml is None:
                warnings.append(f"Inherited chart XML was not found: {source_chart_part}")
                continue

            new_chart_part = self._unique_part(
                entries,
                updates,
                f"ppt/charts/chart_clone_{position}_{chart_index}.xml",
            )
            updates[new_chart_part] = self._rewrite_chart_xml(chart_xml, chart_points, outline)
            self._clone_chart_rels(
                entries,
                updates,
                source_chart_part,
                new_chart_part,
                position,
                chart_index,
                chart_points,
            )
            rel.set("Target", self._relative_relationship_target(target_slide, new_chart_part))

            frame = chart_frames.get(rel_id, self._chart_frame_fallback(rel_id))
            stats["rewritten_chart_count"] += 1
            stats["rewritten_chart_point_count"] += len(chart_points)
            stats["duplicated_chart_part_count"] += 1
            stats["chart_part_names"].add(new_chart_part)
            stats["edit_targets"].append(
                self._chart_edit_target(
                    frame,
                    "rewrite",
                    "chart_data",
                    new_chart_part,
                    len(chart_points),
                )
            )

        return (
            etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True),
            warnings,
            stats,
        )

    def _chart_frames_by_rel_id(self, slide_xml: bytes) -> dict[str, dict[str, Any]]:
        root = etree.fromstring(slide_xml, parser=SAFE_PARSER)
        frames: dict[str, dict[str, Any]] = {}
        for frame in root.xpath(".//p:cSld//p:graphicFrame[.//c:chart]", namespaces=NSMAP):
            chart = frame.find(".//c:chart", namespaces=NSMAP)
            rel_id = chart.get(f"{{{REL_NS}}}id") if chart is not None else ""
            if not rel_id:
                continue
            c_nv_pr = frame.find("p:nvGraphicFramePr/p:cNvPr", namespaces=NSMAP)
            shape_id = c_nv_pr.get("id") if c_nv_pr is not None else ""
            name = c_nv_pr.get("name") if c_nv_pr is not None else ""
            frames[rel_id] = {
                "shapeId": shape_id,
                "sourceElementId": shape_id,
                "name": name,
                "placeholder_type": "chart",
                "bounds": self._graphic_bounds(frame),
                "text": "",
            }
        return frames

    def _chart_frame_fallback(self, rel_id: str) -> dict[str, Any]:
        return {
            "shapeId": rel_id,
            "sourceElementId": rel_id,
            "name": rel_id,
            "placeholder_type": "chart",
            "bounds": {"x": 0.0, "y": 0.0, "w": 0.0, "h": 0.0},
            "text": "",
        }

    def _chart_edit_target(
        self,
        frame: dict[str, Any],
        action: str,
        role: str,
        chart_part: str,
        point_count: int,
    ) -> dict[str, Any]:
        return {
            "action": action,
            "role": role,
            "shapeId": frame.get("shapeId"),
            "sourceElementId": frame.get("sourceElementId"),
            "name": frame.get("name"),
            "placeholderType": frame.get("placeholder_type"),
            "bounds": frame.get("bounds"),
            "chartPart": chart_part,
            "replacementParagraphCount": point_count,
        }

    def _chart_points(self, outline: SlideOutline) -> list[dict[str, Any]]:
        content = outline.content_json or {}
        containers: list[dict[str, Any]] = []
        for key in ("exhibit_spec", "chart_spec"):
            value = content.get(key)
            if isinstance(value, dict):
                containers.append(value)
        containers.append(content)

        points: list[dict[str, Any]] = []
        for container in containers:
            points.extend(self._chart_points_from_metric_list(container.get("metrics")))
            points.extend(self._chart_points_from_metric_list(container.get("data")))
            points.extend(self._chart_points_from_metric_list(container.get("points")))
            points.extend(self._chart_points_from_parallel_values(container))
            points.extend(self._chart_points_from_series(container))

        deduped: list[dict[str, Any]] = []
        seen: set[tuple[str, float]] = set()
        for point in points:
            key = (point["label"].casefold(), float(point["value"]))
            if key in seen:
                continue
            seen.add(key)
            deduped.append(point)
        return deduped[:8]

    def _chart_points_from_metric_list(self, raw_items: Any) -> list[dict[str, Any]]:
        if not isinstance(raw_items, list):
            return []
        points: list[dict[str, Any]] = []
        for idx, item in enumerate(raw_items, start=1):
            if isinstance(item, dict):
                label = str(
                    item.get("label")
                    or item.get("name")
                    or item.get("category")
                    or item.get("metric")
                    or f"Metric {idx}"
                ).strip()
                raw_value = (
                    item.get("value")
                    if "value" in item
                    else item.get("amount", item.get("score", item.get("percent")))
                )
            elif isinstance(item, (list, tuple)) and len(item) >= 2:
                label = str(item[0]).strip()
                raw_value = item[1]
            else:
                continue
            value = self._numeric_value(raw_value)
            if label and value is not None:
                points.append({"label": self._fit_text(label, 34), "value": value})
        return points

    def _chart_points_from_parallel_values(self, container: dict[str, Any]) -> list[dict[str, Any]]:
        categories = container.get("categories") or container.get("labels")
        values = container.get("values")
        if not isinstance(categories, list) or not isinstance(values, list):
            return []
        points: list[dict[str, Any]] = []
        for idx, (category, value) in enumerate(zip(categories, values), start=1):
            number = self._numeric_value(value)
            label = str(category or f"Metric {idx}").strip()
            if label and number is not None:
                points.append({"label": self._fit_text(label, 34), "value": number})
        return points

    def _chart_points_from_series(self, container: dict[str, Any]) -> list[dict[str, Any]]:
        categories = container.get("categories") or container.get("labels")
        series = container.get("series")
        if not isinstance(categories, list) or not isinstance(series, list) or not series:
            return []
        first_series = series[0]
        if not isinstance(first_series, dict):
            return []
        values = first_series.get("values") or first_series.get("data")
        if not isinstance(values, list):
            return []
        return self._chart_points_from_parallel_values({"categories": categories, "values": values})

    def _numeric_value(self, raw: Any) -> float | None:
        if isinstance(raw, bool):
            return None
        if isinstance(raw, (int, float)):
            return float(raw)
        text = str(raw or "").strip()
        match = re.search(r"-?\d[\d,]*(?:\.\d+)?", text)
        if match is None:
            return None
        try:
            return float(match.group(0).replace(",", ""))
        except ValueError:
            return None

    def _rewrite_chart_xml(
        self,
        chart_xml: bytes,
        points: list[dict[str, Any]],
        outline: SlideOutline,
    ) -> bytes:
        root = etree.fromstring(chart_xml, parser=SAFE_PARSER)
        series = root.xpath(".//c:ser", namespaces=NSMAP)
        if not series:
            return chart_xml

        primary = series[0]
        for idx_node in primary.findall("c:idx", namespaces=NSMAP):
            idx_node.set("val", "0")
        for order_node in primary.findall("c:order", namespaces=NSMAP):
            order_node.set("val", "0")
        for extra in series[1:]:
            parent = extra.getparent()
            if parent is not None:
                parent.remove(extra)

        self._rewrite_series_name(primary, self._fit_text(self._title_text(outline), 48))
        self._rewrite_cache_points(
            self._ensure_chart_cache(primary, "cat", "strRef", "strCache"),
            [point["label"] for point in points],
        )
        self._rewrite_cache_points(
            self._ensure_chart_cache(primary, "val", "numRef", "numCache"),
            [self._chart_value_text(point["value"]) for point in points],
        )
        self._update_chart_formulas(primary, points)
        updated = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        etree.fromstring(updated, parser=SAFE_PARSER)
        return updated

    def _rewrite_series_name(self, series: etree._Element, title: str) -> None:
        tx = series.find("c:tx", namespaces=NSMAP)
        if tx is None:
            tx = etree.SubElement(series, f"{{{CHART_NS}}}tx")
        cache = tx.find(".//c:strCache", namespaces=NSMAP)
        if cache is None:
            str_ref = tx.find("c:strRef", namespaces=NSMAP)
            if str_ref is None:
                str_ref = etree.SubElement(tx, f"{{{CHART_NS}}}strRef")
            cache = etree.SubElement(str_ref, f"{{{CHART_NS}}}strCache")
        self._rewrite_cache_points(cache, [title])
        value_node = tx.find("c:v", namespaces=NSMAP)
        if value_node is not None:
            value_node.text = title

    def _ensure_chart_cache(
        self,
        series: etree._Element,
        container_name: str,
        ref_name: str,
        cache_name: str,
    ) -> etree._Element:
        container = series.find(f"c:{container_name}", namespaces=NSMAP)
        if container is None:
            container = etree.SubElement(series, f"{{{CHART_NS}}}{container_name}")
        cache = container.find(f".//c:{cache_name}", namespaces=NSMAP)
        if cache is not None:
            return cache
        ref = container.find(f"c:{ref_name}", namespaces=NSMAP)
        if ref is None:
            ref = etree.SubElement(container, f"{{{CHART_NS}}}{ref_name}")
        return etree.SubElement(ref, f"{{{CHART_NS}}}{cache_name}")

    def _rewrite_cache_points(self, cache: etree._Element, values: list[str]) -> None:
        ext_nodes: list[etree._Element] = []
        for child in list(cache):
            local_name = etree.QName(child).localname
            if local_name == "extLst":
                ext_nodes.append(deepcopy(child))
                cache.remove(child)
            elif local_name in {"pt", "ptCount"}:
                cache.remove(child)

        pt_count = etree.SubElement(cache, f"{{{CHART_NS}}}ptCount")
        pt_count.set("val", str(len(values)))
        for idx, value in enumerate(values):
            point = etree.SubElement(cache, f"{{{CHART_NS}}}pt")
            point.set("idx", str(idx))
            value_node = etree.SubElement(point, f"{{{CHART_NS}}}v")
            value_node.text = str(value)
        for ext_node in ext_nodes:
            cache.append(ext_node)

    def _update_chart_formulas(self, series: etree._Element, points: list[dict[str, Any]]) -> None:
        sheet_name = self._chart_sheet_name(series)
        if not sheet_name:
            return
        end_row = len(points) + 1
        replacements = (
            (series.find(".//c:tx//c:f", namespaces=NSMAP), f"{self._quote_sheet_name(sheet_name)}!$B$1"),
            (series.find(".//c:cat//c:f", namespaces=NSMAP), f"{self._quote_sheet_name(sheet_name)}!$A$2:$A${end_row}"),
            (series.find(".//c:val//c:f", namespaces=NSMAP), f"{self._quote_sheet_name(sheet_name)}!$B$2:$B${end_row}"),
        )
        for node, value in replacements:
            if node is not None:
                node.text = value

    def _chart_sheet_name(self, series: etree._Element) -> str:
        for formula in series.xpath(".//c:f", namespaces=NSMAP):
            text = formula.text or ""
            if "!" not in text:
                continue
            return text.split("!", 1)[0].strip("'")
        return "Sheet1"

    def _quote_sheet_name(self, sheet_name: str) -> str:
        escaped = sheet_name.replace("'", "''")
        return f"'{escaped}'"

    def _chart_value_text(self, value: float) -> str:
        return str(int(value)) if float(value).is_integer() else str(value)

    def _clone_chart_rels(
        self,
        entries: dict[str, bytes],
        updates: dict[str, bytes],
        source_chart_part: str,
        new_chart_part: str,
        position: int,
        chart_index: int,
        points: list[dict[str, Any]],
    ) -> None:
        source_rels = self._part_rels_name(source_chart_part)
        rel_xml = entries.get(source_rels)
        if rel_xml is None:
            return
        root = etree.fromstring(rel_xml, parser=SAFE_PARSER)
        workbook_index = 0
        for rel in root.findall(f".//{{{PKG_REL_NS}}}Relationship"):
            if rel.get("Type") != PACKAGE_REL_TYPE or rel.get("TargetMode") == "External":
                continue
            target = rel.get("Target")
            if not target:
                continue
            workbook_part = self._resolve_relationship_target(source_chart_part, target)
            workbook_bytes = updates.get(workbook_part, entries.get(workbook_part))
            if workbook_bytes is None:
                continue
            workbook_index += 1
            new_workbook_part = self._unique_part(
                entries,
                updates,
                f"ppt/embeddings/chart_clone_{position}_{chart_index}_{workbook_index}.xlsx",
            )
            updates[new_workbook_part] = self._rewrite_chart_workbook(workbook_bytes, points)
            rel.set("Target", self._relative_relationship_target(new_chart_part, new_workbook_part))
        updates[self._part_rels_name(new_chart_part)] = etree.tostring(
            root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )

    def _rewrite_chart_workbook(
        self,
        workbook_bytes: bytes,
        points: list[dict[str, Any]],
    ) -> bytes:
        try:
            workbook = load_workbook(BytesIO(workbook_bytes))
        except Exception:
            return workbook_bytes
        if not workbook.sheetnames:
            return workbook_bytes
        sheet = workbook[workbook.sheetnames[0]]
        sheet.cell(row=1, column=1).value = "Category"
        sheet.cell(row=1, column=2).value = "Value"
        for row_index, point in enumerate(points, start=2):
            sheet.cell(row=row_index, column=1).value = point["label"]
            sheet.cell(row=row_index, column=2).value = point["value"]
        desired_rows = len(points) + 1
        if sheet.max_row > desired_rows:
            sheet.delete_rows(desired_rows + 1, sheet.max_row - desired_rows)
        if sheet.max_column > 2:
            sheet.delete_cols(3, sheet.max_column - 2)
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def _unique_part(
        self,
        entries: dict[str, bytes],
        updates: dict[str, bytes],
        preferred: str,
    ) -> str:
        if preferred not in entries and preferred not in updates:
            return preferred
        stem, suffix = preferred.rsplit(".", 1)
        counter = 2
        while True:
            candidate = f"{stem}_{counter}.{suffix}"
            if candidate not in entries and candidate not in updates:
                return candidate
            counter += 1

    def _part_rels_name(self, part_name: str) -> str:
        return (
            f"{posixpath.dirname(part_name)}/_rels/"
            f"{posixpath.basename(part_name)}.rels"
        )

    def _resolve_relationship_target(self, source_part: str, target: str) -> str:
        if target.startswith("/"):
            return posixpath.normpath(target.lstrip("/"))
        return posixpath.normpath(posixpath.join(posixpath.dirname(source_part), target))

    def _relative_relationship_target(self, source_part: str, target_part: str) -> str:
        return posixpath.relpath(target_part, start=posixpath.dirname(source_part))

    def _mark_unused_package_parts_for_deletion(
        self,
        entries: dict[str, bytes],
        deletions: set[str],
        slide_count: int,
    ) -> dict[str, Any]:
        stats: dict[str, Any] = {
            "deleted_unused_slide_part_count": 0,
            "deleted_unused_slide_rels_count": 0,
            "deleted_notes_comment_tag_part_count": 0,
            "deleted_unreferenced_media_part_count": 0,
            "removed_content_type_override_count": 0,
            "deleted_media_parts": [],
            "deleted_parts": [],
        }
        for name in sorted(entries):
            if self._is_unused_slide_part(name, slide_count):
                deletions.add(name)
                stats["deleted_unused_slide_part_count"] += 1
                stats["deleted_parts"].append(name)
                continue
            if self._is_unused_slide_rels_part(name, slide_count):
                deletions.add(name)
                stats["deleted_unused_slide_rels_count"] += 1
                stats["deleted_parts"].append(name)
                continue
            if self._is_notes_comment_or_tag_part(name):
                deletions.add(name)
                stats["deleted_notes_comment_tag_part_count"] += 1
                stats["deleted_parts"].append(name)
        stats["deleted_part_count"] = (
            stats["deleted_unused_slide_part_count"]
            + stats["deleted_unused_slide_rels_count"]
            + stats["deleted_notes_comment_tag_part_count"]
        )
        return stats

    def _mark_unreferenced_media_parts_for_deletion(
        self,
        entries: dict[str, bytes],
        updates: dict[str, bytes],
        deletions: set[str],
    ) -> dict[str, Any]:
        referenced_media_parts: set[str] = set()
        package_names = set(entries) | set(updates)
        for rels_name in sorted(name for name in package_names if name.endswith(".rels")):
            if rels_name in deletions:
                continue
            rels_xml = updates.get(rels_name, entries.get(rels_name))
            if rels_xml is None:
                continue
            source_part = self._source_part_from_rels_name(rels_name)
            if not source_part:
                continue
            try:
                root = etree.fromstring(rels_xml, parser=SAFE_PARSER)
            except Exception:
                continue
            for rel in root.findall(f".//{{{PKG_REL_NS}}}Relationship"):
                if rel.get("TargetMode") == "External":
                    continue
                target = rel.get("Target")
                if not target:
                    continue
                target_part = self._resolve_relationship_target(source_part, target)
                if target_part.startswith("ppt/media/"):
                    referenced_media_parts.add(target_part)

        deleted_media_parts: list[str] = []
        for media_part in sorted(name for name in package_names if name.startswith("ppt/media/")):
            if media_part in deletions or media_part in referenced_media_parts:
                continue
            deletions.add(media_part)
            deleted_media_parts.append(media_part)

        deleted_parts = sorted(
            part
            for part in deletions
            if (
                self._is_notes_comment_or_tag_part(part)
                or part.startswith("ppt/slides/")
                or part.startswith("ppt/media/")
            )
        )
        return {
            "deleted_unreferenced_media_part_count": len(deleted_media_parts),
            "deleted_media_parts": deleted_media_parts,
            "deleted_parts": deleted_parts,
        }

    def _source_part_from_rels_name(self, rels_name: str) -> str:
        marker = "/_rels/"
        if marker not in rels_name or not rels_name.endswith(".rels"):
            return ""
        prefix, filename = rels_name.split(marker, 1)
        base_name = filename[: -len(".rels")]
        if not base_name:
            return ""
        return f"{prefix}/{base_name}" if prefix else base_name

    def _is_unused_slide_part(self, name: str, slide_count: int) -> bool:
        match = re.fullmatch(r"ppt/slides/slide(\d+)\.xml", name)
        return bool(match and int(match.group(1)) > slide_count)

    def _is_unused_slide_rels_part(self, name: str, slide_count: int) -> bool:
        match = re.fullmatch(r"ppt/slides/_rels/slide(\d+)\.xml\.rels", name)
        return bool(match and int(match.group(1)) > slide_count)

    def _is_notes_comment_or_tag_part(self, name: str) -> bool:
        return bool(
            re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", name)
            or re.fullmatch(r"ppt/notesSlides/_rels/notesSlide\d+\.xml\.rels", name)
            or re.fullmatch(r"ppt/comments/comment\d+\.xml", name)
            or re.fullmatch(r"ppt/comments/_rels/comment\d+\.xml\.rels", name)
            or re.fullmatch(r"ppt/tags/tag\d+\.xml", name)
            or re.fullmatch(r"ppt/tags/_rels/tag\d+\.xml\.rels", name)
        )

    def _clean_slide_rels(self, rel_xml: bytes) -> bytes:
        root = etree.fromstring(rel_xml, parser=SAFE_PARSER)
        for rel in list(root.findall(f".//{{{PKG_REL_NS}}}Relationship")):
            rel_type = rel.get("Type", "")
            if rel_type.endswith(("/notesSlide", "/comments", "/tags")):
                root.remove(rel)
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    def _rewrite_presentation_xml(self, xml: bytes, slide_count: int) -> bytes:
        root = etree.fromstring(xml, parser=SAFE_PARSER)
        sld_id_list = root.find("p:sldIdLst", namespaces=NSMAP)
        if sld_id_list is None:
            sld_id_list = etree.SubElement(root, f"{{{PML_NS}}}sldIdLst")
        for child in list(sld_id_list):
            sld_id_list.remove(child)
        for idx in range(1, slide_count + 1):
            slide_id = etree.SubElement(sld_id_list, f"{{{PML_NS}}}sldId")
            slide_id.set("id", str(255 + idx))
            slide_id.set(f"{{{REL_NS}}}id", f"rIdClone{idx}")
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    def _rewrite_presentation_rels(self, rel_xml: bytes, slide_count: int) -> bytes:
        root = etree.fromstring(rel_xml, parser=SAFE_PARSER)
        for rel in list(root.findall(f".//{{{PKG_REL_NS}}}Relationship")):
            if rel.get("Type") == SLIDE_REL_TYPE:
                root.remove(rel)
        for idx in range(1, slide_count + 1):
            rel = etree.SubElement(root, f"{{{PKG_REL_NS}}}Relationship")
            rel.set("Id", f"rIdClone{idx}")
            rel.set("Type", SLIDE_REL_TYPE)
            rel.set("Target", f"slides/slide{idx}.xml")
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    def _rewrite_content_types(
        self,
        xml: bytes,
        slide_count: int,
        chart_part_names: set[str] | None = None,
    ) -> tuple[bytes, int]:
        root = etree.fromstring(xml, parser=SAFE_PARSER)
        removed_overrides = 0
        for node in list(root.findall(f".//{{{CONTENT_TYPES_NS}}}Override")):
            part_name = node.get("PartName") or ""
            normalized = part_name.lstrip("/")
            if self._is_unused_slide_part(normalized, slide_count) or self._is_notes_comment_or_tag_part(normalized):
                root.remove(node)
                removed_overrides += 1
        existing = {
            node.get("PartName")
            for node in root.findall(f".//{{{CONTENT_TYPES_NS}}}Override")
            if node.get("PartName")
        }
        default_extensions = {
            node.get("Extension")
            for node in root.findall(f".//{{{CONTENT_TYPES_NS}}}Default")
            if node.get("Extension")
        }
        for idx in range(1, slide_count + 1):
            part_name = f"/ppt/slides/slide{idx}.xml"
            if part_name in existing:
                continue
            override = etree.SubElement(root, f"{{{CONTENT_TYPES_NS}}}Override")
            override.set("PartName", part_name)
            override.set("ContentType", SLIDE_CONTENT_TYPE)
        for chart_part in sorted(chart_part_names or set()):
            part_name = f"/{chart_part}"
            if part_name in existing:
                continue
            override = etree.SubElement(root, f"{{{CONTENT_TYPES_NS}}}Override")
            override.set("PartName", part_name)
            override.set("ContentType", CHART_CONTENT_TYPE)
        if chart_part_names and "xlsx" not in default_extensions:
            default = etree.SubElement(root, f"{{{CONTENT_TYPES_NS}}}Default")
            default.set("Extension", "xlsx")
            default.set("ContentType", XLSX_CONTENT_TYPE)
        return (
            etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True),
            removed_overrides,
        )

    def _rewrite_app_props(self, xml: bytes, slide_count: int) -> bytes:
        try:
            root = etree.fromstring(xml, parser=SAFE_PARSER)
        except Exception:
            return xml
        for node in root.iter():
            if etree.QName(node).localname == "Slides":
                node.text = str(slide_count)
                break
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    def _slide_rels_name(self, slide_name: str) -> str:
        return self._part_rels_name(slide_name)

    def _write_clone_artifact(
        self,
        working_dir: Path,
        template: TemplateProfile,
        mappings: list[dict[str, Any]],
        warnings: list[dict[str, str | int]],
        cleanup_stats: dict[str, Any] | None = None,
        status: str | None = None,
    ) -> None:
        payload = {
            "artifact": "template-clone-edit",
            "template_id": template.id,
            "source_file": Path(template.source_file).name if template.source_file else None,
            "status": status or ("warning" if warnings else "pass"),
            "slide_count": len(mappings),
            "mappings": mappings,
            "package_cleanup": cleanup_stats
            or {
                "deleted_part_count": 0,
                "deleted_unused_slide_part_count": 0,
                "deleted_unused_slide_rels_count": 0,
                "deleted_notes_comment_tag_part_count": 0,
                "deleted_unreferenced_media_part_count": 0,
                "removed_content_type_override_count": 0,
                "deleted_media_parts": [],
                "deleted_parts": [],
            },
            "warnings": warnings,
        }
        working_dir.mkdir(parents=True, exist_ok=True)
        (working_dir / "template-clone-edit.json").write_text(
            json.dumps(payload, indent=2, ensure_ascii=True),
            encoding="utf-8",
        )
        frame_map = self._template_frame_map(template, payload)
        (working_dir / "template-frame-map.json").write_text(
            json.dumps(frame_map, indent=2, ensure_ascii=True),
            encoding="utf-8",
        )
        deviation_log = self._template_deviation_log(payload)
        (working_dir / "template-deviation-log.json").write_text(
            json.dumps(deviation_log, indent=2, ensure_ascii=True),
            encoding="utf-8",
        )

    def _template_frame_map(
        self,
        template: TemplateProfile,
        clone_payload: dict[str, Any],
    ) -> dict[str, Any]:
        mappings = [
            mapping
            for mapping in clone_payload.get("mappings") or []
            if isinstance(mapping, dict)
        ]
        used_source_slides = {
            int(mapping["source_slide"])
            for mapping in mappings
            if str(mapping.get("source_slide") or "").isdigit()
        }
        output_slides: list[dict[str, Any]] = []
        for mapping in mappings:
            output_slides.append(
                {
                    "outputSlide": mapping.get("output_slide"),
                    "sourceSlide": mapping.get("source_slide"),
                    "outlineSlideIndex": mapping.get("outline_slide_index"),
                    "narrativeRole": mapping.get("narrative_role"),
                    "reuseMode": (
                        "blocked"
                        if mapping.get("clone_edit_blocked")
                        else "duplicate-slide-edit"
                    ),
                    "method": mapping.get("method"),
                    "matchConfidence": mapping.get("match_confidence"),
                    "matchScore": mapping.get("match_score"),
                    "matchReason": mapping.get("match_reason"),
                    "blockReason": mapping.get("block_reason"),
                    "editTargets": mapping.get("editTargets") or [],
                    "closestCandidates": mapping.get("closest_candidates") or [],
                }
            )
        omitted_source_slides = []
        for slide in template.slides:
            source_slide = slide.index + 1
            if source_slide in used_source_slides:
                continue
            omitted_source_slides.append(
                {
                    "sourceSlide": source_slide,
                    "label": slide.label,
                    "contentCategory": slide.content_category,
                    "reason": "not selected for generated narrative",
                }
            )
        return {
            "artifact": "template-frame-map",
            "standard": "claude-pptx-template-following",
            "source": "https://github.com/anthropics/skills/blob/main/skills/pptx/editing.md",
            "templateId": template.id,
            "sourceFile": Path(template.source_file).name if template.source_file else None,
            "status": clone_payload.get("status"),
            "outputSlideCount": len(output_slides),
            "sourceSlideCount": len(template.slides),
            "omittedSourceSlideCount": len(omitted_source_slides),
            "outputSlides": output_slides,
            "omittedSourceSlides": omitted_source_slides,
        }

    def _template_deviation_log(self, clone_payload: dict[str, Any]) -> dict[str, Any]:
        deviations: list[dict[str, Any]] = []
        for mapping in clone_payload.get("mappings") or []:
            if not isinstance(mapping, dict):
                continue
            output_slide = mapping.get("output_slide")
            source_slide = mapping.get("source_slide")
            method = str(mapping.get("method") or "").strip().lower()
            confidence = str(mapping.get("match_confidence") or "").strip().lower()
            if mapping.get("clone_edit_blocked"):
                deviations.append(
                    {
                        "type": "blocked_clone_edit",
                        "severity": "warning",
                        "output_slide": output_slide,
                        "source_slide": source_slide,
                        "reason": mapping.get("block_reason") or "weak template-frame match",
                        "closest_candidates": mapping.get("closest_candidates") or [],
                    }
                )
            elif method in WEAK_TEMPLATE_FRAME_METHODS or confidence in WEAK_TEMPLATE_FRAME_CONFIDENCES:
                deviations.append(
                    {
                        "type": "weak_template_mapping",
                        "severity": "warning",
                        "output_slide": output_slide,
                        "source_slide": source_slide,
                        "reason": mapping.get("match_reason") or method or confidence,
                        "closest_candidates": mapping.get("closest_candidates") or [],
                    }
                )
            slot_cleanup = mapping.get("slot_cleanup")
            if (
                isinstance(slot_cleanup, dict)
                and slot_cleanup.get("cleanup_required")
                and not slot_cleanup.get("cleanup_satisfied")
            ):
                deviations.append(
                    {
                        "type": "unsatisfied_slot_cleanup",
                        "severity": "warning",
                        "output_slide": output_slide,
                        "source_slide": source_slide,
                        "reason": (
                            f"deleted {slot_cleanup.get('actual_deleted_slot_count', 0)}/"
                            f"{slot_cleanup.get('planned_excess_slot_count', 0)} planned excess slot(s)"
                        ),
                    }
                )
            if int(mapping.get("unfilled_placeholder_count") or 0) > 0:
                deviations.append(
                    {
                        "type": "unfilled_inherited_placeholder",
                        "severity": "critical",
                        "output_slide": output_slide,
                        "source_slide": source_slide,
                        "reason": (
                            f"{mapping.get('unfilled_placeholder_count')} inherited "
                            "placeholder(s) remain empty in final PPTX XML"
                        ),
                        "placeholders": mapping.get("unfilled_placeholders") or [],
                    }
                )
        for warning in clone_payload.get("warnings") or []:
            if not isinstance(warning, dict):
                continue
            deviations.append(
                {
                    "type": "template_warning",
                    "severity": str(warning.get("severity") or "WARNING").lower(),
                    "output_slide": int(warning.get("slide_index") or 0) + 1,
                    "source_slide": None,
                    "reason": warning.get("message") or "",
                }
            )
        status = "pass" if not deviations else "warning"
        if clone_payload.get("status") == "blocked":
            status = "blocked"
        return {
            "artifact": "template-deviation-log",
            "standard": "claude-pptx-template-following",
            "source": "https://github.com/anthropics/skills/blob/main/skills/pptx/editing.md",
            "status": status,
            "deviation_count": len(deviations),
            "deviations": deviations,
        }

    def _warning(self, slide_index: int, field: str, message: str) -> dict[str, str | int]:
        return {
            "slide_index": slide_index,
            "field": field,
            "message": message,
            "severity": "WARNING",
        }
