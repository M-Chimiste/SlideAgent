"""HTML slide renderer: outlines -> HTML -> headless-Chrome PDF -> images -> PPTX.

This is the polished generated-deck path. The model still emits the same
structured specs; this renderer turns them into a cohesive, content-fitted deck
using a single CSS design system rendered by headless Chrome. Slides are
embedded full-bleed as high-resolution images (the codex-skill/Gamma approach),
with speaker notes preserved as real text.
"""

from __future__ import annotations

import os
import shutil
import signal
import subprocess
import tempfile
import time
from pathlib import Path

from app.models.brand import BrandDNA
from app.models.outline import SlideOutline

from .css import build_css
from .design_system import resolve_theme, resolve_modes
from .templates import esc, render_slide_html

_CHROME_CANDIDATES = [
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    "/Applications/Chromium.app/Contents/MacOS/Chromium",
    "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
    "/usr/bin/google-chrome",
    "/usr/bin/google-chrome-stable",
    "/usr/bin/chromium",
    "/usr/bin/chromium-browser",
    "/opt/google/chrome/chrome",
]
_PDFTOPPM_CANDIDATES = ["/opt/homebrew/bin/pdftoppm", "/usr/local/bin/pdftoppm", "/usr/bin/pdftoppm"]


def find_chrome() -> str | None:
    explicit = os.environ.get("SLIDEFORGE_CHROME_BINARY")
    if explicit and Path(explicit).exists():
        return explicit
    for name in ("google-chrome", "google-chrome-stable", "chromium", "chromium-browser", "chrome"):
        found = shutil.which(name)
        if found:
            return found
    for path in _CHROME_CANDIDATES:
        if Path(path).exists():
            return path
    return None


def _find_pdftoppm() -> str | None:
    found = shutil.which("pdftoppm")
    if found:
        return found
    for path in _PDFTOPPM_CANDIDATES:
        if Path(path).exists():
            return path
    return None


class HtmlRenderError(RuntimeError):
    """Raised when the HTML rendering toolchain is unavailable or fails."""


class HtmlSlideRenderer:
    """Renders generated freeform/brand decks via HTML + headless Chrome."""

    def __init__(self, *, dpi: int | None = None) -> None:
        self.dpi = dpi or int(os.environ.get("SLIDEFORGE_HTML_DPI", "160"))

    # ---- public renderer interface (matches DeterministicPptxRenderer) ---- #
    def render(
        self,
        outlines: list[SlideOutline],
        brand: BrandDNA,
        output_path: Path,
        enable_diagrams: bool = True,
    ) -> list[dict[str, str | int]]:
        output_path = Path(output_path)
        work = output_path.parent
        work.mkdir(parents=True, exist_ok=True)
        warnings: list[dict[str, str | int]] = []

        html_doc = self.build_html(outlines, brand, warnings)
        html_path = work / f"{output_path.stem}.html"
        html_path.write_text(html_doc, encoding="utf-8")

        pdf_path = work / f"{output_path.stem}.deck.pdf"
        self._html_to_pdf(html_path, pdf_path)

        images = self._pdf_to_images(pdf_path, work / f"{output_path.stem}-pages")
        if len(images) < len(outlines):
            warnings.append({
                "slide_index": -1,
                "field": "render",
                "severity": "warning",
                "message": f"rendered {len(images)} of {len(outlines)} slide images",
            })
        self._assemble_pptx(images, outlines, output_path)
        return warnings

    # ---- HTML assembly ---- #
    def build_html(
        self,
        outlines: list[SlideOutline],
        brand: BrandDNA,
        warnings: list[dict[str, str | int]] | None = None,
    ) -> str:
        warnings = warnings if warnings is not None else []
        theme = resolve_theme(brand)
        css = build_css(theme)
        keyed = [
            (
                (o.content_json or {}).get("narrative_role") or (o.content_json or {}).get("slide_type"),
                (o.content_json or {}).get("composition_family") or (o.layout_json or {}).get("composition_family"),
            )
            for o in outlines
        ]
        modes = resolve_modes(keyed)
        total = len(outlines)
        sections = []
        history: list[str] = []
        for i, outline in enumerate(outlines):
            try:
                sections.append(render_slide_html(outline, theme, modes[i], i, total, history))
            except Exception as exc:  # never let one slide break the deck
                warnings.append({
                    "slide_index": i,
                    "field": "render",
                    "severity": "warning",
                    "message": f"slide render fallback: {exc}",
                })
                sections.append(self._fallback_section(outline, modes[i], i, total))
        font_face = self._font_face_css()
        return (
            "<!doctype html><html><head><meta charset='utf-8'>"
            f"<style>{font_face}{css}</style></head>"
            f"<body><div class='deck'>{''.join(sections)}</div></body></html>"
        )

    def _font_face_css(self) -> str:
        """Embed bundled web fonts if present; otherwise rely on the system stack."""
        fonts_dir = Path(__file__).parent / "assets" / "fonts"
        if not fonts_dir.exists():
            return ""
        faces = []
        mapping = {
            "Newsreader": ("Newsreader", "normal"),
            "HankenGrotesk": ("Hanken Grotesk", "normal"),
        }
        for file in sorted(fonts_dir.glob("*.woff2")):
            stem = file.stem  # e.g. Newsreader-600 or HankenGrotesk-700i
            for prefix, (family, _style) in mapping.items():
                if stem.startswith(prefix):
                    rest = stem[len(prefix):].lstrip("-")
                    italic = rest.endswith("i")
                    weight = "".join(c for c in rest if c.isdigit()) or "400"
                    style = "italic" if italic else "normal"
                    faces.append(
                        f"@font-face{{font-family:'{family}';font-style:{style};"
                        f"font-weight:{weight};font-display:swap;"
                        f"src:url('file://{file.as_posix()}') format('woff2');}}"
                    )
        return "".join(faces)

    def _fallback_section(self, outline, mode: str, index: int, total: int) -> str:
        content = outline.content_json or {}
        title = esc(content.get("action_title") or content.get("title") or outline.label or "")
        bullets = content.get("bullets") or []
        items = "".join(f"<li>{esc(b if isinstance(b, str) else b.get('text',''))}</li>" for b in bullets[:5])
        return (
            f'<section class="slide {mode}"><div class="head-block">'
            f'<h2 class="headline h-md">{title}</h2></div>'
            f'<ul style="font-size:18px;line-height:1.6;color:var(--fg-muted);margin-left:20px">{items}</ul>'
            f'<div class="foot"><div class="src"></div><div class="pg">{index+1:02d} / {total:02d}</div></div></section>'
        )

    # ---- toolchain ---- #
    def _html_to_pdf(self, html_path: Path, pdf_path: Path, *, max_wait: float = 45.0) -> None:
        """Render HTML to PDF with headless Chrome.

        Headless Chrome on macOS frequently writes the PDF quickly but then
        refuses to exit, so we poll for a size-stable PDF and kill the process
        once it is complete rather than waiting on process exit.
        """
        chrome = find_chrome()
        if not chrome:
            raise HtmlRenderError("no Chrome/Chromium binary found for HTML rendering")
        if pdf_path.exists():
            pdf_path.unlink()
        profile = tempfile.mkdtemp(prefix="sf-chrome-")
        cmd = [
            chrome,
            "--headless=new",
            "--disable-gpu",
            "--no-sandbox",
            "--no-first-run",
            "--no-default-browser-check",
            "--hide-scrollbars",
            "--disable-extensions",
            "--disable-dev-shm-usage",
            "--disable-background-networking",
            "--disable-sync",
            f"--user-data-dir={profile}",
            "--no-pdf-header-footer",
            f"--print-to-pdf={pdf_path.as_posix()}",
            html_path.as_uri(),
        ]
        proc = subprocess.Popen(
            cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, start_new_session=True
        )
        start = time.monotonic()
        last_size = -1
        stable = 0
        try:
            while time.monotonic() - start < max_wait:
                if proc.poll() is not None:
                    break
                if pdf_path.exists():
                    size = pdf_path.stat().st_size
                    if size > 2000 and size == last_size:
                        stable += 1
                        if stable >= 2:
                            break
                    else:
                        stable = 0
                    last_size = size
                time.sleep(0.25)
        finally:
            self._kill(proc)
            shutil.rmtree(profile, ignore_errors=True)
        if not pdf_path.exists() or pdf_path.stat().st_size < 2000:
            raise HtmlRenderError("Chrome did not produce a valid PDF")

    @staticmethod
    def _kill(proc: subprocess.Popen) -> None:
        if proc.poll() is not None:
            return
        try:
            os.killpg(os.getpgid(proc.pid), signal.SIGKILL)
        except Exception:
            try:
                proc.kill()
            except Exception:
                pass

    def _pdf_to_images(self, pdf_path: Path, prefix: Path) -> list[Path]:
        pdftoppm = _find_pdftoppm()
        if not pdftoppm:
            raise HtmlRenderError("pdftoppm (poppler) not found for image rasterization")
        prefix.parent.mkdir(parents=True, exist_ok=True)
        # clear stale pages
        for old in prefix.parent.glob(f"{prefix.name}-*.png"):
            old.unlink()
        cmd = [pdftoppm, "-png", "-r", str(self.dpi), pdf_path.as_posix(), prefix.as_posix()]
        subprocess.run(cmd, capture_output=True, text=True, timeout=120, check=False)
        images = sorted(
            prefix.parent.glob(f"{prefix.name}-*.png"),
            key=lambda p: int("".join(c for c in p.stem.split("-")[-1] if c.isdigit()) or 0),
        )
        if not images:
            raise HtmlRenderError("pdftoppm produced no page images")
        return images

    def _assemble_pptx(self, images: list[Path], outlines: list[SlideOutline], output_path: Path) -> None:
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        blank = prs.slide_layouts[6]
        for i, image in enumerate(images):
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(
                image.as_posix(), 0, 0, width=prs.slide_width, height=prs.slide_height
            )
            if i < len(outlines):
                notes = (outlines[i].content_json or {}).get("speaker_notes") or ""
                if notes:
                    slide.notes_slide.notes_text_frame.text = str(notes)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(output_path.as_posix())
