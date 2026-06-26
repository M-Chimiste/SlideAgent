from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.util import Inches

from app.models.brand import BrandDNA
from app.models.outline import SlideOutline
from app.services.concept_diagram_renderer import ConceptDiagramRenderer
from app.services.pptx_rendering.assets import AssetRenderingMixin
from app.services.pptx_rendering.chrome import ChromeRenderingMixin
from app.services.pptx_rendering.constants import SLIDE_H, SLIDE_W
from app.services.pptx_rendering.core_layouts import CoreLayoutRenderingMixin
from app.services.pptx_rendering.drawing import DrawingMixin
from app.services.pptx_rendering.exhibit_layouts import ExhibitLayoutRenderingMixin
from app.services.pptx_rendering.immersive_layouts import ImmersiveLayoutRenderingMixin
from app.services.pptx_rendering.table_layouts import TableLayoutRenderingMixin


class DeterministicPptxRenderer(
    AssetRenderingMixin,
    ChromeRenderingMixin,
    CoreLayoutRenderingMixin,
    TableLayoutRenderingMixin,
    ImmersiveLayoutRenderingMixin,
    ExhibitLayoutRenderingMixin,
    DrawingMixin,
):
    def render(
        self,
        outlines: list[SlideOutline],
        brand: BrandDNA,
        output_path: Path,
        enable_diagrams: bool = True,
    ) -> list[dict[str, str | int]]:
        self._warnings: list[dict[str, str | int]] = []
        self._icon_cache_dir = output_path.parent / f"{output_path.stem}-icons"
        self._icon_cache_dir.mkdir(parents=True, exist_ok=True)
        self._diagram_cache_dir = (
            output_path.parent / f"{output_path.stem}-diagrams"
            if enable_diagrams
            else None
        )
        self._diagram_renderer = ConceptDiagramRenderer() if enable_diagrams else None
        self._template_frame_presentations: dict[str, Presentation] = {}
        prs = Presentation()
        prs.slide_width = Inches(SLIDE_W)
        prs.slide_height = Inches(SLIDE_H)
        blank_layout = prs.slide_layouts[6]
        for index, outline in enumerate(outlines):
            slide = prs.slides.add_slide(blank_layout)
            self._render_slide(slide, outline, brand, index + 1, len(outlines))
        prs.save(output_path.as_posix())
        return self._warnings

    def _render_slide(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
        slide_number: int,
        total_slides: int,
    ) -> None:
        title = outline.content_json.get("action_title") or outline.content_json.get("title") or outline.label
        subheading = outline.content_json.get("subheading") or outline.content_json.get("summary", "")
        layout = outline.layout_json.get("layout", "two_column")
        if layout == "cover":
            self._add_cover(slide, outline, brand)
            self._add_logo(slide, brand)
            return
        if layout == "section_divider":
            self._add_section_divider(slide, outline, brand, slide_number)
            self._add_logo(slide, brand)
            return
        if layout == "framework_cycle":
            self._add_framework_cycle_immersive(
                slide, outline, brand, slide_number, total_slides
            )
            self._add_logo(slide, brand)
            return
        if layout == "quote_sidebar":
            self._add_quote_sidebar_immersive(
                slide, outline, brand, slide_number, total_slides
            )
            self._add_logo(slide, brand)
            return
        if layout == "closing_recommendation":
            self._add_closing_recommendation_immersive(
                slide, outline, brand, slide_number, total_slides
            )
            self._add_logo(slide, brand)
            return
        composition_family = str(outline.layout_json.get("composition_family") or "")
        if self._add_full_bleed_authored_composition(
            slide,
            outline,
            brand,
            composition_family,
            slide_number,
            total_slides,
        ):
            self._add_logo(slide, brand)
            return
        frame_applied = self._add_template_frame_background(slide, outline)
        if not frame_applied:
            self._add_profile_background(slide, brand)
        self._add_header(slide, title, subheading, brand, outline, slide_number)
        self._add_logo(slide, brand)
        self._add_footer(slide, outline, brand, slide_number, total_slides)

        if self._add_authored_composition(slide, outline, brand, composition_family):
            return

        if layout == "executive_summary":
            self._add_executive_summary(slide, outline, brand)
        elif layout == "chart":
            self._add_metric_chart(slide, outline, brand)
        elif layout == "comparison_table":
            self._add_comparison_table(slide, outline, brand)
        elif layout == "matrix_2x2":
            self._add_matrix_2x2(slide, outline, brand)
        elif layout == "callouts":
            self._add_callouts(slide, outline, brand)
        elif layout == "process":
            self._add_table_or_process(slide, outline, brand)
        elif layout == "quote_sidebar":
            self._add_quote_sidebar(slide, outline, brand)
        elif layout == "framework_cycle":
            self._add_framework_cycle(slide, outline, brand)
        elif layout == "dependency_map":
            self._add_dependency_map(slide, outline, brand)
        elif layout == "checklist":
            self._add_checklist(slide, outline, brand)
        elif layout == "code_panel":
            self._add_code_panel(slide, outline, brand)
        elif layout == "anti_patterns":
            self._add_anti_patterns(slide, outline, brand)
        elif layout == "table_reference":
            self._add_reference_table(slide, outline, brand)
        elif layout == "closing_recommendation":
            self._add_closing_recommendation(slide, outline, brand)
        elif layout == "icon_grid":
            self._add_grid(slide, outline, brand)
        elif layout == "icon_rows":
            self._add_icon_rows(slide, outline, brand)
        else:
            self._add_two_column(slide, outline, brand)

    def _add_profile_background(self, slide, brand: BrandDNA) -> None:
        fill = (brand.layout_profile or {}).get("dominant_fill")
        if not fill or str(fill).upper() in {"FFFFFF", "FFF"}:
            return
        try:
            if self._relative_luminance(str(fill)) < 0.72:
                return
        except Exception:
            return
        background = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0),
            Inches(0),
            Inches(SLIDE_W),
            Inches(SLIDE_H),
        )
        background.fill.solid()
        background.fill.fore_color.rgb = self._rgb(str(fill))
        background.line.color.rgb = background.fill.fore_color.rgb

    def _add_template_frame_background(self, slide, outline: SlideOutline) -> bool:
        frame = outline.layout_json.get("template_frame")
        if not isinstance(frame, dict):
            return False
        source_file = str(frame.get("source_file") or "").strip()
        try:
            source_index = int(frame.get("index"))
        except (TypeError, ValueError):
            return False
        if not source_file or not Path(source_file).exists():
            return False
        try:
            source_prs = self._template_frame_presentations.get(source_file)
            if source_prs is None:
                source_prs = Presentation(source_file)
                self._template_frame_presentations[source_file] = source_prs
            source_slide = source_prs.slides[source_index]
        except Exception:
            return False
        copied = 0
        for source_shape in source_slide.shapes:
            if not self._safe_template_frame_shape(source_shape):
                continue
            try:
                new_element = deepcopy(source_shape.element)
                slide.shapes._spTree.insert_element_before(new_element, "p:extLst")
                copied += 1
            except Exception:
                continue
            if copied >= 48:
                break
        if copied:
            frame["chrome_shape_count"] = copied
            frame["chrome_applied"] = True
        return copied > 0

    def _safe_template_frame_shape(self, shape) -> bool:
        shape_type = getattr(shape, "shape_type", None)
        if shape_type in {MSO_SHAPE_TYPE.PLACEHOLDER, MSO_SHAPE_TYPE.TEXT_BOX}:
            return False
        if getattr(shape, "has_text_frame", False) and str(getattr(shape, "text", "")).strip():
            return False
        if getattr(shape, "has_table", False) or getattr(shape, "has_chart", False):
            return False
        if shape_type not in {
            MSO_SHAPE_TYPE.AUTO_SHAPE,
            MSO_SHAPE_TYPE.FREEFORM,
            MSO_SHAPE_TYPE.LINE,
        }:
            return False
        try:
            left = int(shape.left) / 914400
            top = int(shape.top) / 914400
            width = int(shape.width) / 914400
            height = int(shape.height) / 914400
        except Exception:
            return False
        if left < -0.01 or top < -0.01:
            return False
        if left + width > SLIDE_W + 0.01 or top + height > SLIDE_H + 0.01:
            return False
        return width > 0.01 and height > 0.01
