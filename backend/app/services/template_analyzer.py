import re
import uuid
import zipfile
import json
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Optional

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from app.clients.bedrock_client import BedrockClient
from app.models.brand import BrandDNA, BrandFonts, BrandLogo
from app.models.template import (
    LayoutSpec,
    PlaceholderSpec,
    SlideField,
    SlideSchema,
    SlideSpec,
    TemplateProfile,
)
from app.services.brand_layout_renderer import role_for_layout
from app.services.rendering import render_pptx_to_images


SAFE_PARSER = etree.XMLParser(resolve_entities=False, no_network=True)


class TemplateAnalyzer:
    def __init__(self, bedrock: Optional[BedrockClient] = None) -> None:
        self.bedrock = bedrock

    def analyze(
        self,
        template_path: Path,
        template_name: str,
        template_type: str,
        template_id: Optional[str] = None,
    ) -> tuple[TemplateProfile, list[Path]]:
        brand = self._extract_brand(template_path)
        slides = self._extract_slides(template_path, template_type)
        try:
            layout_library = self._extract_layout_library(template_path)
        except Exception:
            layout_library = []
        profile = TemplateProfile(
            id=template_id or self._new_id(),
            name=template_name,
            type=template_type,
            brand=brand,
            slides=slides,
            source_file=template_path.as_posix(),
            created_at=self._timestamp(),
            updated_at=self._timestamp(),
            layout_library=layout_library,
        )
        thumbnails = []
        try:
            thumbnails = render_pptx_to_images(
                template_path, template_path.parent / "thumbnails"
            )
        except Exception:
            thumbnails = []
        try:
            self._write_frame_map(template_path, profile)
        except Exception:
            pass
        return profile, thumbnails

    def _write_frame_map(
        self,
        template_path: Path,
        profile: TemplateProfile,
    ) -> Path:
        frame_map = self._frame_map(template_path, profile)
        target = template_path.parent / "frame-map.json"
        target.write_text(
            json.dumps(frame_map, indent=2, ensure_ascii=True),
            encoding="utf-8",
        )
        return target

    def _frame_map(
        self,
        template_path: Path,
        profile: TemplateProfile,
    ) -> dict[str, Any]:
        presentation = Presentation(template_path.as_posix())
        slide_specs = {slide.index: slide for slide in profile.slides}
        slides: list[dict[str, Any]] = []
        total_slots = 0
        schema_slide_count = 0
        for index, slide in enumerate(presentation.slides):
            spec = slide_specs.get(index)
            text_slots: list[dict[str, Any]] = []
            media_slots: list[dict[str, Any]] = []
            table_count = 0
            chart_count = 0
            picture_count = 0
            for shape in slide.shapes:
                bounds = self._shape_bounds(shape)
                shape_name = shape.name or f"Shape{len(text_slots) + len(media_slots) + 1}"
                if getattr(shape, "has_chart", False):
                    chart_count += 1
                    media_slots.append(
                        {
                            "name": shape_name,
                            "kind": "chart",
                            "bounds": bounds,
                            "placeholder": bool(getattr(shape, "is_placeholder", False)),
                        }
                    )
                    continue
                if getattr(shape, "has_table", False):
                    table_count += 1
                    text_slots.append(
                        {
                            "name": shape_name,
                            "kind": "table",
                            "bounds": bounds,
                            "placeholder": bool(getattr(shape, "is_placeholder", False)),
                            "text_preview": self._truncate_text(
                                self._shape_text(shape),
                                160,
                            ),
                            "cell_count": self._table_cell_count(shape),
                        }
                    )
                    continue
                if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                    picture_count += 1
                    media_slots.append(
                        {
                            "name": shape_name,
                            "kind": "picture",
                            "bounds": bounds,
                            "placeholder": bool(getattr(shape, "is_placeholder", False)),
                        }
                    )
                    continue
                if getattr(shape, "has_text_frame", False):
                    text = self._shape_text(shape)
                    text_slots.append(
                        {
                            "name": shape_name,
                            "kind": self._text_slot_kind(bounds, text),
                            "bounds": bounds,
                            "placeholder": bool(getattr(shape, "is_placeholder", False)),
                            "text_preview": self._truncate_text(text, 160),
                            "paragraph_count": len(
                                [
                                    paragraph
                                    for paragraph in text.splitlines()
                                    if paragraph.strip()
                                ]
                            ),
                        }
                    )
            schema_fields = (
                len(spec.slide_schema.fields)
                if spec and spec.slide_schema
                else 0
            )
            if schema_fields:
                schema_slide_count += 1
            slot_count = len(text_slots) + len(media_slots)
            total_slots += slot_count
            slides.append(
                {
                    "slide_index": index,
                    "label": spec.label if spec else f"Slide {index + 1}",
                    "layout_name": spec.layout_name if spec else slide.slide_layout.name,
                    "mode": spec.mode if spec else "flexible",
                    "content_category": spec.content_category if spec else self._slide_content_category(slide, ""),
                    "visual_guidance": spec.visual_guidance if spec else self._slide_visual_guidance(slide),
                    "classification_reason": spec.classification_reason if spec else None,
                    "schema_field_count": schema_fields,
                    "slot_count": slot_count,
                    "text_slot_count": len(text_slots),
                    "media_slot_count": len(media_slots),
                    "table_count": table_count,
                    "chart_count": chart_count,
                    "picture_count": picture_count,
                    "text_inventory": self._truncate_text(
                        " ".join(
                            slot.get("text_preview", "")
                            for slot in text_slots
                            if slot.get("text_preview")
                        ),
                        260,
                    ),
                    "text_slots": text_slots,
                    "media_slots": media_slots,
                }
            )
        return {
            "artifact": "template-frame-map",
            "standard": "claude-pptx-editing-v1",
            "source": "https://github.com/anthropics/skills/blob/main/skills/pptx/editing.md",
            "template_id": profile.id,
            "template_type": profile.type,
            "slide_count": len(slides),
            "schema_bearing_slide_count": schema_slide_count,
            "slot_count": total_slots,
            "slides": slides,
        }

    def _text_slot_kind(self, bounds: dict[str, float], text: str) -> str:
        lowered = text.casefold()
        if bounds["y"] < 1.4 and bounds["h"] <= 1.4:
            return "title"
        if bounds["y"] > 6.2:
            return "footer"
        if any(token in lowered for token in ("source", "footnote", "copyright")):
            return "caption"
        return "body"

    def _table_cell_count(self, shape) -> int:
        try:
            return len(shape.table.rows) * len(shape.table.columns)
        except Exception:
            return 0

    def _truncate_text(self, text: str, limit: int) -> str:
        cleaned = " ".join(str(text or "").split())
        if len(cleaned) <= limit:
            return cleaned
        return cleaned[:limit].rsplit(" ", 1)[0].rstrip(".,;:")

    def _extract_brand(self, template_path: Path) -> BrandDNA:
        brand = BrandDNA()
        with zipfile.ZipFile(template_path, "r") as zip_ref:
            theme_path = "ppt/theme/theme1.xml"
            if theme_path in zip_ref.namelist():
                xml = zip_ref.read(theme_path)
                tree = etree.fromstring(xml, parser=SAFE_PARSER)
                ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                clr_scheme = tree.find(".//a:clrScheme", namespaces=ns)
                colors = {}
                if clr_scheme is not None:
                    for child in clr_scheme:
                        tag = etree.QName(child).localname
                        color_node = child.find(".//a:srgbClr", namespaces=ns)
                        if color_node is None:
                            color_node = child.find(".//a:sysClr", namespaces=ns)
                        if color_node is not None:
                            color_value = self._theme_color_value(color_node)
                            if color_value:
                                colors[tag] = color_value
                font_scheme = tree.find(".//a:fontScheme", namespaces=ns)
                fonts = BrandFonts()
                if font_scheme is not None:
                    major = font_scheme.find(".//a:majorFont", namespaces=ns)
                    minor = font_scheme.find(".//a:minorFont", namespaces=ns)
                    if major is not None:
                        major_latn = major.find(".//a:latin", namespaces=ns)
                        if major_latn is not None and major_latn.get("typeface"):
                            fonts.heading = major_latn.get("typeface")
                    if minor is not None:
                        minor_latn = minor.find(".//a:latin", namespaces=ns)
                        if minor_latn is not None and minor_latn.get("typeface"):
                            fonts.body = minor_latn.get("typeface")
                brand.fonts = fonts
                if colors:
                    brand.colors.primary = colors.get("accent1", brand.colors.primary)
                    brand.colors.secondary = colors.get("accent2", brand.colors.secondary)
                    brand.colors.accent = colors.get("accent3", brand.colors.accent)
                    brand.colors.text_dark = colors.get("dk1", brand.colors.text_dark)
                    brand.colors.text_light = colors.get("lt1", brand.colors.text_light)
                    brand.colors.background_dark = colors.get(
                        "dk2", brand.colors.background_dark
                    )
                    brand.colors.background_light = colors.get(
                        "lt2", brand.colors.background_light
                    )
        logo = self._extract_logo(template_path)
        if logo:
            brand.logo = logo
        brand.design_notes = self._design_notes(template_path)
        brand.layout_profile = self._layout_profile(template_path)
        return brand

    def _theme_color_value(self, color_node) -> str | None:
        last_color = str(color_node.get("lastClr") or "").strip()
        value = str(color_node.get("val") or "").strip()
        for candidate in (last_color, value):
            cleaned = candidate.replace("#", "").strip()
            if re.fullmatch(r"[0-9A-Fa-f]{6}", cleaned):
                return cleaned.upper()
        return None

    def _layout_profile(self, template_path: Path) -> dict[str, Any]:
        presentation = Presentation(template_path.as_posix())
        title_boxes = []
        footer_boxes = []
        body_boxes = []
        logo_boxes = []
        fills: list[str] = []
        table_fills: list[str] = []
        chart_fills: list[str] = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                bounds = self._shape_bounds(shape)
                fill = self._shape_fill(shape)
                if fill:
                    fills.append(fill)
                if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                    logo_boxes.append(bounds)
                    continue
                if getattr(shape, "has_chart", False):
                    chart_fills.extend(self._chart_colors(shape))
                    continue
                if getattr(shape, "has_table", False):
                    table_fills.extend(self._table_colors(shape))
                text = self._shape_text(shape).strip()
                if not text:
                    continue
                if bounds["y"] < 1.4 and bounds["h"] <= 1.4:
                    title_boxes.append(bounds)
                elif bounds["y"] > 6.2:
                    footer_boxes.append(bounds)
                else:
                    body_boxes.append(bounds)
        profile: dict[str, Any] = {}
        if title_boxes:
            profile["title_box"] = self._average_box(title_boxes)
        if footer_boxes:
            profile["footer_box"] = self._average_box(footer_boxes)
        if body_boxes:
            profile["body_box"] = self._bounding_box(body_boxes)
        if logo_boxes:
            profile["logo_box"] = self._average_box(logo_boxes)
        dominant_fill = self._most_common(fills)
        if dominant_fill:
            profile["dominant_fill"] = dominant_fill
        if table_fills:
            profile["table_colors"] = self._unique_colors(table_fills)[:4]
        if chart_fills:
            profile["chart_colors"] = self._unique_colors(chart_fills)[:4]
        return profile

    def _shape_bounds(self, shape) -> dict[str, float]:
        return {
            "x": round(int(shape.left) / 914400, 3),
            "y": round(int(shape.top) / 914400, 3),
            "w": round(int(shape.width) / 914400, 3),
            "h": round(int(shape.height) / 914400, 3),
        }

    def _shape_fill(self, shape) -> str | None:
        try:
            rgb = shape.fill.fore_color.rgb
        except Exception:
            return None
        return str(rgb) if rgb else None

    def _table_colors(self, shape) -> list[str]:
        colors = []
        try:
            for row in shape.table.rows:
                for cell in row.cells:
                    rgb = cell.fill.fore_color.rgb
                    if rgb:
                        colors.append(str(rgb))
        except Exception:
            return []
        return colors

    def _chart_colors(self, shape) -> list[str]:
        colors = []
        try:
            for series in shape.chart.series:
                rgb = series.format.fill.fore_color.rgb
                if rgb:
                    colors.append(str(rgb))
        except Exception:
            return []
        return colors

    def _average_box(self, boxes: list[dict[str, float]]) -> dict[str, float]:
        return {
            key: round(sum(box[key] for box in boxes) / len(boxes), 3)
            for key in ("x", "y", "w", "h")
        }

    def _bounding_box(self, boxes: list[dict[str, float]]) -> dict[str, float]:
        left = min(box["x"] for box in boxes)
        top = min(box["y"] for box in boxes)
        right = max(box["x"] + box["w"] for box in boxes)
        bottom = max(box["y"] + box["h"] for box in boxes)
        return {
            "x": round(left, 3),
            "y": round(top, 3),
            "w": round(right - left, 3),
            "h": round(bottom - top, 3),
        }

    def _most_common(self, values: list[str]) -> str | None:
        if not values:
            return None
        return max(set(values), key=values.count)

    def _unique_colors(self, values: list[str]) -> list[str]:
        unique = []
        for value in values:
            if value not in unique:
                unique.append(value)
        return unique

    def _extract_logo(self, template_path: Path) -> BrandLogo | None:
        presentation = Presentation(template_path.as_posix())
        candidates = []
        slide_area = int(presentation.slide_width) * int(presentation.slide_height)
        for slide in presentation.slides:
            for shape in slide.shapes:
                if getattr(shape, "shape_type", None) != MSO_SHAPE_TYPE.PICTURE:
                    continue
                area = int(shape.width) * int(shape.height)
                candidates.append((area / max(slide_area, 1), shape))
        if not candidates:
            return None
        candidates.sort(key=lambda item: item[0])
        _, shape = candidates[0]
        image = shape.image
        extension = image.ext or "png"
        logo_path = template_path.parent / f"brand-logo.{extension}"
        logo_path.write_bytes(image.blob)
        width_inches = max(min(int(shape.width) / 914400, 1.6), 0.45)
        height_inches = max(min(int(shape.height) / 914400, 0.8), 0.25)
        return BrandLogo(path=logo_path.as_posix(), w=width_inches, h=height_inches)

    def _design_notes(self, template_path: Path) -> str:
        presentation = Presentation(template_path.as_posix())
        layout_names = []
        for slide in presentation.slides:
            layout_name = slide.slide_layout.name
            if layout_name not in layout_names:
                layout_names.append(layout_name)
        if not layout_names:
            return "No reusable layout examples detected."
        return "Reusable layout examples: " + ", ".join(layout_names[:8])

    def _extract_layout_library(self, template_path: Path) -> list[LayoutSpec]:
        presentation = Presentation(template_path.as_posix())
        library: list[LayoutSpec] = []
        for index, layout in enumerate(presentation.slide_layouts):
            placeholders: list[PlaceholderSpec] = []
            text_slots = 0
            media_slots = 0
            has_title = False
            for ph in layout.placeholders:
                try:
                    ptype = ph.placeholder_format.type
                    idx = ph.placeholder_format.idx
                except Exception:
                    continue
                if ptype in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                    has_title = True
                    text_slots += 1
                elif ptype in (PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT):
                    text_slots += 1
                elif ptype == PP_PLACEHOLDER.PICTURE:
                    media_slots += 1
                try:
                    bounds = self._shape_bounds(ph)
                except Exception:
                    bounds = None
                placeholders.append(
                    PlaceholderSpec(
                        idx=idx,
                        ph_type=getattr(ptype, "name", str(ptype)) if ptype is not None else "UNKNOWN",
                        name=ph.name or "",
                        bounds=bounds,
                    )
                )
            library.append(
                LayoutSpec(
                    index=index,
                    name=layout.name or f"Layout {index}",
                    role=role_for_layout(layout),
                    placeholders=placeholders,
                    text_slot_count=text_slots,
                    media_slot_count=media_slots,
                    has_title=has_title,
                )
            )
        return library

    def _extract_slides(self, template_path: Path, template_type: str) -> list[SlideSpec]:
        presentation = Presentation(template_path.as_posix())
        layouts = list(presentation.slide_layouts)
        slides = []
        for index, slide in enumerate(presentation.slides):
            title = self._slide_title(slide)
            label = title or f"Slide {index + 1}"
            content_category = self._slide_content_category(slide, label)
            visual_guidance = self._slide_visual_guidance(slide)
            mode = "flexible"
            reason = None
            schema = None
            intent = self._slide_intent(label, content_category, visual_guidance)
            if template_type == "strict":
                mode, reason = self._classify_slide(slide, label)
                if mode == "strict":
                    schema = self._extract_schema(slide)
                else:
                    intent = label
            try:
                layout_index = layouts.index(slide.slide_layout)
            except ValueError:
                layout_index = None
            slides.append(
                SlideSpec(
                    index=index,
                    mode=mode,
                    label=label,
                    layout_name=slide.slide_layout.name,
                    layout_index=layout_index,
                    schema=schema,
                    intent=intent,
                    content_category=content_category,
                    visual_guidance=visual_guidance,
                    classification_reason=reason,
                )
            )
        return slides

    def _slide_content_category(self, slide, label: str) -> str:
        text = self._slide_text(slide).casefold()
        has_chart = any(getattr(shape, "has_chart", False) for shape in slide.shapes)
        has_table = any(getattr(shape, "has_table", False) for shape in slide.shapes)
        picture_count = sum(
            1
            for shape in slide.shapes
            if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE
        )
        body_boxes = [
            self._shape_bounds(shape)
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False)
            and self._shape_text(shape).strip()
            and self._shape_bounds(shape)["y"] >= 1.2
            and self._shape_bounds(shape)["y"] <= 6.2
        ]
        lowered_label = label.casefold()
        metric_like = bool(
            re.search(r"\b(?:kpi|scorecard|dashboard)\b", text)
            or re.search(r"\b\d+(?:\.\d+)?\s*(?:%|x|/10)\b", text)
        )
        if has_chart:
            return "metric_chart"
        if has_table:
            if any(token in text for token in ("current", "target", "before", "after", "versus", " vs ")):
                return "comparison"
            return "structured_table"
        if any(
            token in text or token in lowered_label
            for token in (
                "compare",
                "comparison",
                "versus",
                "before",
                "after",
                "current state",
                "target state",
            )
        ):
            return "comparison"
        if any(token in text for token in ("step", "phase", "timeline", "roadmap", "sequence", "process")):
            return "process"
        if picture_count and len(body_boxes) <= 3:
            return "media_story"
        if any(token in text for token in ("quote", "key idea", "decision", "recommendation")):
            return "quote_callout"
        if metric_like:
            return "metric_chart"
        if len(body_boxes) >= 3:
            return "evidence_points"
        if len(self._slide_text(slide).split()) <= 12:
            return "section_or_cover"
        return "content"

    def _slide_visual_guidance(self, slide) -> str:
        text_slots = []
        media_count = 0
        table_count = 0
        chart_count = 0
        for shape in slide.shapes:
            if getattr(shape, "has_chart", False):
                chart_count += 1
                continue
            if getattr(shape, "has_table", False):
                table_count += 1
                continue
            if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                media_count += 1
                continue
            if getattr(shape, "has_text_frame", False) and self._shape_text(shape).strip():
                bounds = self._shape_bounds(shape)
                text_slots.append(self._text_slot_kind(bounds, self._shape_text(shape)))
        parts = []
        if text_slots:
            parts.append(
                f"{len(text_slots)} text slot(s): {', '.join(text_slots[:5])}"
            )
        if table_count:
            parts.append(f"{table_count} table frame(s)")
        if chart_count:
            parts.append(f"{chart_count} chart frame(s)")
        if media_count:
            parts.append(f"{media_count} media slot(s)")
        return "; ".join(parts) or "Open canvas frame"

    def _slide_intent(
        self,
        label: str,
        content_category: str,
        visual_guidance: str,
    ) -> str:
        category = content_category.replace("_", " ")
        if label:
            return f"{label} frame for {category}; {visual_guidance}"
        return f"Reusable {category} frame; {visual_guidance}"

    def _slide_text(self, slide) -> str:
        return " ".join(self._shape_text(shape).strip() for shape in slide.shapes)

    def _slide_title(self, slide) -> str:
        for shape in slide.shapes:
            text = self._shape_text(shape).strip()
            if text:
                return text.splitlines()[0][:80]
        return ""

    def _classify_slide(self, slide, label: str) -> tuple[str, str]:
        text = " ".join(self._shape_text(shape).strip() for shape in slide.shapes).lower()
        if any(keyword in text for keyword in ["status", "dashboard", "kpi", "score"]):
            return "strict", "Detected structured dashboard keywords."
        try:
            if any(getattr(shape, "has_table", False) for shape in slide.shapes):
                return "strict", "Detected table element."
            if any(getattr(shape, "has_chart", False) for shape in slide.shapes):
                return "strict", "Detected chart element."
        except Exception:
            pass
        if "title" in label.lower() and len(text.split()) < 8:
            return "strict", "Likely title slide."
        return "flexible", "Defaulted to flexible layout."

    def _extract_schema(self, slide) -> SlideSchema:
        fields = []
        for shape in slide.shapes:
            if getattr(shape, "has_chart", False):
                fields.extend(self._extract_chart_fields(shape))
                continue
            if getattr(shape, "has_table", False):
                fields.extend(self._extract_table_fields(shape, len(fields)))
                continue
            if not shape.has_text_frame:
                continue
            name = shape.name or f"Shape{len(fields) + 1}"
            shape_text = shape.text.strip()
            field_id = self._sanitize_id(name)
            field_type = self._infer_field_type(shape_text)
            constraints = self._infer_constraints(field_id, field_type, shape_text)
            is_placeholder = bool(getattr(shape, "is_placeholder", False))
            fields.append(
                SlideField(
                    id=field_id,
                    type=field_type,
                    location=f"shape:{name}",
                    required=is_placeholder or bool(shape_text),
                    max_chars=constraints.get("max_chars", 160),
                    values=constraints.get("values"),
                    render=constraints.get("render"),
                    color_map=constraints.get("color_map"),
                    max_items=constraints.get("max_items"),
                    max_chars_per_item=constraints.get("max_chars_per_item"),
                )
            )
        return SlideSchema(fields=fields)

    def _extract_chart_fields(self, shape) -> list[SlideField]:
        fields: list[SlideField] = []
        name = shape.name or "Chart"
        chart = shape.chart
        for series_idx, series in enumerate(chart.series):
            fields.append(
                SlideField(
                    id=self._sanitize_id(f"{name}_series_{series_idx + 1}_name"),
                    type="text",
                    location=f"chart:{name}:series_name:{series_idx}",
                    required=True,
                    max_chars=80,
                )
            )
            for point_idx, value in enumerate(series.values):
                fields.append(
                    SlideField(
                        id=self._sanitize_id(
                            f"{name}_series_{series_idx + 1}_value_{point_idx + 1}"
                        ),
                        type="number",
                        location=f"chart:{name}:value:{series_idx}:{point_idx}",
                        required=True,
                        max_chars=20,
                    )
                )
        try:
            categories = list(chart.plots[0].categories)
        except Exception:
            categories = []
        for point_idx, category in enumerate(categories):
            fields.append(
                SlideField(
                    id=self._sanitize_id(f"{name}_category_{point_idx + 1}"),
                    type="text",
                    location=f"chart:{name}:category:{point_idx}",
                    required=True,
                    max_chars=max(len(str(category)) + 20, 60),
                )
            )
        return fields

    def _extract_table_fields(self, shape, existing_count: int) -> list[SlideField]:
        fields: list[SlideField] = []
        name = shape.name or f"Table{existing_count + 1}"
        for row_idx, row in enumerate(shape.table.rows):
            for col_idx, cell in enumerate(row.cells):
                cell_text = cell.text.strip()
                if not cell_text:
                    continue
                field_id = self._sanitize_id(f"{name}_{row_idx + 1}_{col_idx + 1}")
                field_type = self._infer_field_type(cell_text)
                constraints = self._infer_constraints(field_id, field_type, cell_text)
                fields.append(
                    SlideField(
                        id=field_id,
                        type=field_type,
                        location=f"table:{name}:{row_idx}:{col_idx}",
                        required=True,
                        max_chars=constraints.get("max_chars", 160),
                        values=constraints.get("values"),
                        render=constraints.get("render"),
                        color_map=constraints.get("color_map"),
                        max_items=constraints.get("max_items"),
                        max_chars_per_item=constraints.get("max_chars_per_item"),
                    )
                )
        return fields

    def _shape_text(self, shape) -> str:
        if getattr(shape, "has_text_frame", False):
            return shape.text or ""
        if getattr(shape, "has_table", False):
            cell_text = []
            for row in shape.table.rows:
                for cell in row.cells:
                    if cell.text:
                        cell_text.append(cell.text)
            return " ".join(cell_text)
        return ""

    def _infer_field_type(self, text: str) -> str:
        lowered = text.lower()
        if any(token in lowered for token in ["rag", "status", "traffic light"]):
            return "enum"
        if any(token in lowered for token in ["risks", "actions", "items"]) and "\n" in text:
            return "text_list"
        if re.search(r"\d+%", lowered):
            return "number"
        if any(token in lowered for token in ["date", "month", "year"]):
            return "date"
        return "text"

    def _infer_constraints(
        self, field_id: str, field_type: str, shape_text: str
    ) -> dict[str, object]:
        lowered = f"{field_id} {shape_text}".lower()
        constraints: dict[str, object] = {}
        if field_type == "enum" or "rag" in lowered or "status" in lowered:
            constraints["values"] = ["green", "amber", "red"]
            constraints["render"] = "fill_color"
            constraints["color_map"] = {
                "green": "00B050",
                "amber": "FFC000",
                "red": "FF0000",
            }
        if field_type == "text_list":
            constraints["max_items"] = 4
            constraints["max_chars_per_item"] = 80
        if "title" in lowered:
            constraints["max_chars"] = 80
        elif "summary" in lowered or "description" in lowered:
            constraints["max_chars"] = 220
        return constraints

    def _sanitize_id(self, name: str) -> str:
        cleaned = re.sub(r"[^a-zA-Z0-9]+", "_", name.strip().lower())
        return cleaned.strip("_") or "field"

    def _timestamp(self) -> str:
        return datetime.now(UTC).isoformat().replace("+00:00", "Z")

    def _new_id(self) -> str:
        return str(uuid.uuid4())
