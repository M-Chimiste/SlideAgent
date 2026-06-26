from collections import Counter
import json
import math
import re
from pathlib import Path
from typing import Any

from pptx import Presentation

from app.models.outline import SlideOutline
from app.models.qa import QAIssue


TRAILING_FRAGMENT_RE = re.compile(
    r"\b(?:a|an|and|as|by|for|from|in|into|of|or|that|the|their|through|to|which|with)\.?$",
    re.IGNORECASE,
)
CONVERT_ARTIFACT_RE = re.compile(
    r"\bconvert\b.{0,180}\binto an(?: owned)?(?: action)?",
    re.IGNORECASE,
)
OWNER_NEXT_RE = re.compile(r"\(\s*owner\s*/\s*next\s*\)|\bowner\s*/\s*next\b", re.IGNORECASE)
GENERIC_MATRIX_RE = re.compile(
    r"\b(?:high|low) impact\s*/\s*(?:high|low) readiness\b",
    re.IGNORECASE,
)
RAW_TABLE_RE = re.compile(r"\w\s*\|\s*\w")
META_RENDERED_RE = re.compile(
    r"\b(?:cover slide|executive overview deck)\b",
    re.IGNORECASE,
)
GENERIC_COMPARISON_COPY_RE = re.compile(
    r"\b(?:current readout|source claim|operating implication)\b",
    re.IGNORECASE,
)
NONSENSICAL_RENDERED_RE = re.compile(
    r"\b(?:cannot|can|make|makes|around|through|from|to)\s+is\b|"
    r"\bmodel contracts make is\b|"
    r"\bevaluation systems around is\b|"
    r"\bsynthetic benchmarks cannot is\b",
    re.IGNORECASE,
)
ORPHAN_META_LABELS = {
    "architecture overview",
    "closing remarks",
    "executive summary",
    "historically",
    "however",
    "introduction",
    "the model contract",
    "the problem",
    "the solution",
}
SOURCE_FRAGMENT_RE = re.compile(
    r"\b(?:use this evidence to decide|review evidence for|instead of looking|"
    r"frame the request with an explicit outcome|prime the agent with memory|"
    r"generate bounded changes from the spec|review output before updating memory|"
    r"this white paper has presented|the key contributions are|"
    r"a reframing of the benchmark problem|rather than treating)\b",
    re.IGNORECASE,
)
AUTHORED_FILLER_RE = re.compile(
    r"\b(?:tie the claim to a source-backed evaluation artifact|"
    r"tie the claim.{0,80}source-backed evaluation artifact|"
    r"connect the source evidence to the decision before scaling|"
    r"make the handoff inspectable before the decision moves|"
    r"each step should pair evidence,?\s+ownership,?\s+and timing|"
    r"connect.{0,80}to an explicit review gate|"
    r"make.{0,80}visible before execution starts|"
    r"name the review gate before expanding the benchmark|"
    r"update the benchmark when source evidence changes|"
    r"adopt the operating model through a named pilot and review gate|"
    r"confirm owner,?\s+scope,?\s+and timing|confirm ownership and timing|"
    r"make the next move visible enough|use the model to decide|"
    r"evidence tension|reliability risk|operating move|"
    r"evidence gap|validation risk|operating choice|"
    r"make the operating implication explicit|assign the next review gate|"
    r"name the owner before the workflow expands|"
    r"inspect unresolved assumptions before commitment|"
    r"deterministic beat for)\b",
    re.IGNORECASE,
)
REPEATED_TRIGGER_RE = re.compile(
    r"\bwhen (?:conditions|assumptions) change\b",
    re.IGNORECASE,
)
MULTI_ITEM_MARKER_RE = re.compile(
    r"(?<!\w)(?:Step|Phase|Stage)\s+\d+\s*[:.)-]?|\b\d{1,2}[.)](?=\s+[A-Z])",
    re.IGNORECASE,
)
UNICODE_BULLET_RE = re.compile(r"[\u2022\u25e6\u25aa\u25cf]")

GENERIC_DIAGRAM_LABELS = {
    "frame",
    "ground",
    "build",
    "prime",
    "generate",
    "review",
    "update",
    "reset",
    "persist",
    "source context",
    "rules",
    "memory",
    "reliable output",
    "operating loop",
}
EMU_PER_INCH = 914400


class RenderedSlideAudit:
    def inspect(
        self,
        pptx_path: Path,
        outlines: list[SlideOutline],
        artifact_dir: Path | None = None,
    ) -> tuple[dict[str, Any], list[QAIssue]]:
        prs = Presentation(pptx_path.as_posix())
        diagram_labels = self._diagram_labels_by_slide(pptx_path)
        slides: list[dict[str, Any]] = []
        issues: list[QAIssue] = []
        slide_width = prs.slide_width / EMU_PER_INCH
        slide_height = prs.slide_height / EMU_PER_INCH
        for slide_index, slide in enumerate(prs.slides):
            outline = outlines[slide_index] if slide_index < len(outlines) else None
            texts = self._slide_texts(slide)
            paragraph_texts = self._slide_paragraph_texts(slide)
            slide_issues = self._text_issues(slide_index, texts, paragraph_texts)
            layout_diagnostics, layout_issues = self._layout_diagnostics(
                slide,
                slide_index,
                slide_width,
                slide_height,
            )
            slide_issues.extend(layout_issues)
            labels = diagram_labels.get(slide_index, [])
            if labels and self._is_generic_diagram(labels):
                slide_issues.append(
                    QAIssue(
                        severity="CRITICAL",
                        message=(
                            "Rendered diagram uses generic fallback labels instead "
                            "of source-specific explanatory nodes."
                        ),
                        slide_index=slide_index,
                        category="diagram_semantic_fit",
                    )
                )
            if outline:
                slide_issues.extend(self._outline_semantic_issues(slide_index, outline))
            issues.extend(slide_issues)
            slides.append(
                {
                    "slide_index": slide_index,
                    "text": texts,
                    "paragraph_text": paragraph_texts,
                    "layout_diagnostics": layout_diagnostics,
                    "diagram_labels": labels,
                    "composition_signature": self._composition_signature(outline),
                    "template_frame": self._template_frame(outline),
                    "visual_degradation": (
                        (outline.content_json or {}).get("visual_degradation")
                        if outline
                        else None
                    ),
                    "source_refs": (outline.content_json.get("source_refs") if outline else []) or [],
                    "issues": [issue.model_dump() for issue in slide_issues],
                }
            )
        visual_rhythm = self._visual_rhythm_summary(slides, outlines)
        issues.extend(self._visual_rhythm_issues(visual_rhythm))
        payload = {
            "passed": not any(issue.severity in {"CRITICAL", "WARNING"} for issue in issues),
            "issue_count": len(issues),
            "critical_count": sum(1 for issue in issues if issue.severity == "CRITICAL"),
            "warning_count": sum(1 for issue in issues if issue.severity == "WARNING"),
            "issues": [issue.model_dump() for issue in issues],
            "visual_rhythm": visual_rhythm,
            "slides": slides,
        }
        if artifact_dir is not None:
            artifact_dir.mkdir(parents=True, exist_ok=True)
            (artifact_dir / "rendered-slide-audit.json").write_text(
                json.dumps(payload, indent=2, ensure_ascii=True),
                encoding="utf-8",
            )
        return payload, issues

    def _slide_texts(self, slide) -> list[str]:
        texts: list[str] = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = " ".join(shape.text.split())
            if text:
                texts.append(text)
        return texts

    def _slide_paragraph_texts(self, slide) -> list[str]:
        texts: list[str] = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                text = " ".join(paragraph.text.split())
                if text:
                    texts.append(text)
        return texts

    def _template_frame(self, outline: SlideOutline | None) -> dict[str, Any] | None:
        if outline is None:
            return None
        frame = (outline.layout_json or {}).get("template_frame")
        if not isinstance(frame, dict):
            return None
        return {
            key: value
            for key, value in frame.items()
            if key
            in {
                "index",
                "source_slide",
                "label",
                "layout_name",
                "mode",
                "method",
                "match_score",
                "match_confidence",
                "match_reason",
                "intent",
                "content_category",
                "visual_guidance",
                "schema_field_count",
                "item_slot_count",
                "reuse_mode",
                "chrome_shape_count",
                "chrome_applied",
            }
        }

    def _text_issues(
        self,
        slide_index: int,
        texts: list[str],
        paragraph_texts: list[str],
    ) -> list[QAIssue]:
        issues: list[QAIssue] = []
        joined = " || ".join(texts)
        if CONVERT_ARTIFACT_RE.search(joined):
            issues.append(
                QAIssue(
                    severity="CRITICAL",
                    message="Rendered slide contains generated boilerplate such as 'Convert ... into an owned action'.",
                    slide_index=slide_index,
                    category="content_quality",
                )
            )
        if OWNER_NEXT_RE.search(joined):
            issues.append(
                QAIssue(
                    severity="CRITICAL",
                    message="Rendered slide contains unresolved owner/timing placeholder text.",
                    slide_index=slide_index,
                    category="placeholder_text",
                )
            )
        if GENERIC_MATRIX_RE.search(joined):
            hits = set(match.group(0).casefold() for match in GENERIC_MATRIX_RE.finditer(joined))
            if len(hits) >= 2:
                issues.append(
                    QAIssue(
                        severity="CRITICAL",
                        message="Rendered 2x2 uses generic impact/readiness quadrant labels without a source-specific axis.",
                        slide_index=slide_index,
                        category="semantic_visual_fit",
                    )
                )
        if any(RAW_TABLE_RE.search(text) for text in texts):
            issues.append(
                QAIssue(
                    severity="CRITICAL",
                    message="Rendered slide exposes raw table or markdown delimiter text.",
                    slide_index=slide_index,
                    category="raw_artifact",
                )
            )
        if META_RENDERED_RE.search(joined):
            issues.append(
                QAIssue(
                    severity="WARNING",
                    message="Rendered slide exposes meta slide-planning labels such as 'Cover Slide'.",
                    slide_index=slide_index,
                    category="meta_copy",
                )
            )
        if len(GENERIC_COMPARISON_COPY_RE.findall(joined)) >= 2:
            issues.append(
                QAIssue(
                    severity="WARNING",
                    message="Rendered slide uses generic comparison filler instead of topic-specific row language.",
                    slide_index=slide_index,
                    category="generic_comparison_copy",
                )
            )
        if NONSENSICAL_RENDERED_RE.search(joined):
            issues.append(
                QAIssue(
                    severity="CRITICAL",
                    message="Rendered slide contains a nonsensical title-derived phrase.",
                    slide_index=slide_index,
                    category="nonsensical_copy",
                )
            )
        orphan_counts = Counter(
            " ".join(text.casefold().split()).strip(" .:-")
            for text in texts
            if " ".join(text.casefold().split()).strip(" .:-") in ORPHAN_META_LABELS
        )
        repeated_orphans = [
            label for label, count in orphan_counts.items() if count >= 2
        ]
        if repeated_orphans:
            issues.append(
                QAIssue(
                    severity="WARNING",
                    message=(
                        "Rendered slide repeats source section/meta labels as body "
                        f"copy: {', '.join(repeated_orphans)}."
                    ),
                    slide_index=slide_index,
                    category="meta_copy",
                )
            )
        if self._has_prefix_repetition(texts):
            issues.append(
                QAIssue(
                    severity="WARNING",
                    message="Rendered slide repeats the same source phrase as both label and body copy.",
                    slide_index=slide_index,
                    category="repeated_source_fragment",
                )
            )
        if any(len(MULTI_ITEM_MARKER_RE.findall(text)) >= 2 for text in paragraph_texts):
            issues.append(
                QAIssue(
                    severity="WARNING",
                    message="Rendered paragraph concatenates multiple numbered/list items instead of separate paragraphs.",
                    slide_index=slide_index,
                    category="multi_item_concatenation",
                )
            )
        if any(UNICODE_BULLET_RE.search(text) for text in texts):
            issues.append(
                QAIssue(
                    severity="WARNING",
                    message=(
                        "Rendered slide contains unicode bullet glyphs; use native "
                        "PowerPoint list formatting or separate item shapes instead."
                    ),
                    slide_index=slide_index,
                    category="unicode_bullets",
                )
            )
        if SOURCE_FRAGMENT_RE.search(joined):
            issues.append(
                QAIssue(
                    severity="CRITICAL",
                    message="Rendered slide contains source-fragment or demo-boilerplate text instead of authored copy.",
                    slide_index=slide_index,
                    category="content_quality",
                )
            )
        if self._has_authored_filler_copy(joined):
            issues.append(
                QAIssue(
                    severity="CRITICAL",
                    message="Rendered slide contains generic renderer-authored filler copy.",
                    slide_index=slide_index,
                    category="renderer_filler_copy",
                )
            )
        trigger_hits = REPEATED_TRIGGER_RE.findall(joined)
        if len(trigger_hits) >= 3:
            issues.append(
                QAIssue(
                    severity="WARNING",
                    message="Rendered reference exhibit repeats the same generic update trigger across rows.",
                    slide_index=slide_index,
                    category="repeated_placeholder",
                )
            )
        dangling = [
            text
            for index, text in enumerate(texts)
            if self._looks_dangling(text)
            and not self._next_text_completes_fragment(
                text,
                texts[index + 1] if index + 1 < len(texts) else "",
            )
        ]
        if dangling:
            issues.append(
                QAIssue(
                    severity="CRITICAL",
                    message="Rendered slide contains incomplete sentence fragments.",
                    slide_index=slide_index,
                    category="incomplete_content",
                )
            )
        return issues

    def _has_authored_filler_copy(self, text: str) -> bool:
        if AUTHORED_FILLER_RE.search(text):
            return True
        normalized = re.sub(r"[^a-z0-9]+", " ", text.casefold())
        normalized = " ".join(normalized.split())
        return any(
            phrase in normalized
            for phrase in (
                "tie the claim source backed evaluation artifact",
                "tie the claim to a source backed evaluation artifact",
                "name the review gate before expanding the benchmark",
                "update the benchmark when source evidence changes",
                "adopt the operating model through a named pilot and review gate",
                "confirm owner scope and timing",
                "confirm ownership and timing",
            )
        ) or bool(
            re.search(
                r"\bconnect\b.{0,90}\b(?:to\s+)?an explicit review gate\b|"
                r"\bmake\b.{0,90}\bvisible before execution starts\b",
                normalized,
            )
        )

    def _layout_diagnostics(
        self,
        slide,
        slide_index: int,
        slide_width: float,
        slide_height: float,
    ) -> tuple[dict[str, Any], list[QAIssue]]:
        boxes = self._text_shape_boxes(slide)
        opaque_shapes = self._opaque_shape_boxes(slide, slide_width, slide_height)
        overflow_risks = [
            box for box in boxes if self._text_box_has_overflow_risk(box)
        ]
        small_text_risks = [
            box for box in boxes if self._text_box_has_small_text_risk(box)
        ]
        overlap_pairs = self._text_overlap_pairs(boxes, slide_width, slide_height)
        occlusion_pairs = self._text_occlusion_pairs(
            boxes,
            opaque_shapes,
            slide_width,
            slide_height,
        )
        issues: list[QAIssue] = []
        if overflow_risks:
            issues.append(
                QAIssue(
                    severity="WARNING",
                    message=(
                        "Rendered slide has text boxes with cut-off/overflow risk "
                        f"after fitting: {len(overflow_risks)} region(s)."
                    ),
                    slide_index=slide_index,
                    category="cut-off-text",
                )
            )
        if small_text_risks:
            issues.append(
                QAIssue(
                    severity="WARNING",
                    message=(
                        "Rendered slide has body/support text below the readable "
                        f"font-size floor: {len(small_text_risks)} region(s)."
                    ),
                    slide_index=slide_index,
                    category="small_text",
                )
            )
        if overlap_pairs:
            issues.append(
                QAIssue(
                    severity="WARNING",
                    message=(
                        "Rendered slide has overlapping text-bearing regions: "
                        f"{len(overlap_pairs)} pair(s)."
                    ),
                    slide_index=slide_index,
                    category="overlap",
                )
            )
        if occlusion_pairs:
            issues.append(
                QAIssue(
                    severity="WARNING",
                    message=(
                        "Rendered slide has text partially covered by opaque "
                        f"non-text shapes: {len(occlusion_pairs)} region(s)."
                    ),
                    slide_index=slide_index,
                    category="occluded_text",
                )
            )
        diagnostics = {
            "text_box_count": len(boxes),
            "opaque_shape_count": len(opaque_shapes),
            "overflow_risk_count": len(overflow_risks),
            "small_text_risk_count": len(small_text_risks),
            "overlap_pair_count": len(overlap_pairs),
            "occlusion_pair_count": len(occlusion_pairs),
            "overflow_risks": [
                self._layout_box_payload(box) for box in overflow_risks[:6]
            ],
            "small_text_risks": [
                self._layout_box_payload(box) for box in small_text_risks[:6]
            ],
            "overlap_pairs": overlap_pairs[:6],
            "occlusion_pairs": occlusion_pairs[:6],
        }
        return diagnostics, issues

    def _text_shape_boxes(self, slide) -> list[dict[str, Any]]:
        boxes: list[dict[str, Any]] = []
        for shape_index, shape in enumerate(slide.shapes):
            if not getattr(shape, "has_text_frame", False):
                continue
            paragraphs = [
                " ".join(paragraph.text.split())
                for paragraph in shape.text_frame.paragraphs
                if " ".join(paragraph.text.split())
            ]
            text = " ".join(paragraphs)
            if not text:
                continue
            width = shape.width / EMU_PER_INCH
            height = shape.height / EMU_PER_INCH
            if width <= 0 or height <= 0:
                continue
            font_size = self._shape_font_size(shape)
            box = {
                "shape_index": shape_index,
                "text": text,
                "paragraphs": paragraphs,
                "x": shape.left / EMU_PER_INCH,
                "y": shape.top / EMU_PER_INCH,
                "w": width,
                "h": height,
                "font_size": font_size,
                "is_chrome": self._is_chrome_text(text, height, font_size),
            }
            boxes.append(box)
        return boxes

    def _opaque_shape_boxes(
        self,
        slide,
        slide_width: float,
        slide_height: float,
    ) -> list[dict[str, Any]]:
        boxes: list[dict[str, Any]] = []
        for shape_index, shape in enumerate(slide.shapes):
            if getattr(shape, "has_text_frame", False) and " ".join(shape.text.split()):
                continue
            width = shape.width / EMU_PER_INCH
            height = shape.height / EMU_PER_INCH
            if width <= 0.08 or height <= 0.08:
                continue
            box = {
                "shape_index": shape_index,
                "x": shape.left / EMU_PER_INCH,
                "y": shape.top / EMU_PER_INCH,
                "w": width,
                "h": height,
            }
            if self._is_full_slide_text_region(box, slide_width, slide_height):
                continue
            if not self._shape_has_visible_fill(shape):
                continue
            boxes.append(box)
        return boxes

    def _shape_has_visible_fill(self, shape) -> bool:
        fill = getattr(shape, "fill", None)
        if fill is None:
            return False
        try:
            fill_type = fill.type
        except Exception:
            return False
        return fill_type is not None

    def _shape_font_size(self, shape) -> float:
        sizes: list[float] = []
        for paragraph in shape.text_frame.paragraphs:
            if paragraph.font.size is not None:
                sizes.append(float(paragraph.font.size.pt))
            for run in paragraph.runs:
                if run.font.size is not None:
                    sizes.append(float(run.font.size.pt))
        if not sizes:
            return 12.0
        return max(6.0, min(sizes))

    def _is_chrome_text(self, text: str, height: float, font_size: float) -> bool:
        cleaned = " ".join(str(text).split())
        if cleaned.startswith("Source:"):
            return True
        if re.match(r"^\d+\s*/\s*\d+$", cleaned):
            return True
        if re.match(r"^SECTION\s+\d+", cleaned, re.IGNORECASE):
            return True
        if ">" in cleaned and height <= 0.35 and font_size <= 9:
            return True
        if height <= 0.26 and len(cleaned) <= 90 and font_size <= 10:
            return True
        return False

    def _text_box_has_overflow_risk(self, box: dict[str, Any]) -> bool:
        if box["is_chrome"]:
            return False
        text = str(box["text"])
        if len(text) < 32:
            return False
        font_size = float(box["font_size"])
        width_points = float(box["w"]) * 72
        height_points = float(box["h"]) * 72
        chars_per_line = max(8, int(width_points / max(font_size * 0.52, 1)))
        estimated_lines = sum(
            max(1, math.ceil(len(paragraph) / chars_per_line))
            for paragraph in box["paragraphs"]
        )
        capacity_lines = max(1.0, height_points / max(font_size * 1.18, 1))
        box["estimated_lines"] = estimated_lines
        box["capacity_lines"] = round(capacity_lines, 2)
        box["chars_per_line"] = chars_per_line
        return estimated_lines > capacity_lines + 0.8

    def _text_box_has_small_text_risk(self, box: dict[str, Any]) -> bool:
        if box["is_chrome"]:
            return False
        text = str(box["text"]).strip()
        if len(text) < 34 or len(text.split()) < 4:
            return False
        return float(box["font_size"]) < 10

    def _text_overlap_pairs(
        self,
        boxes: list[dict[str, Any]],
        slide_width: float,
        slide_height: float,
    ) -> list[dict[str, Any]]:
        pairs: list[dict[str, Any]] = []
        content_boxes = [
            box
            for box in boxes
            if not box["is_chrome"]
            and not self._is_full_slide_text_region(box, slide_width, slide_height)
        ]
        for left_index, left in enumerate(content_boxes):
            for right in content_boxes[left_index + 1 :]:
                overlap = self._overlap_ratio(left, right)
                if overlap < 0.35:
                    continue
                pairs.append(
                    {
                        "shape_indexes": [left["shape_index"], right["shape_index"]],
                        "overlap_ratio": round(overlap, 3),
                        "texts": [
                            self._truncate_text(str(left["text"]), 64),
                            self._truncate_text(str(right["text"]), 64),
                        ],
                    }
                )
        return pairs

    def _text_occlusion_pairs(
        self,
        boxes: list[dict[str, Any]],
        opaque_shapes: list[dict[str, Any]],
        slide_width: float,
        slide_height: float,
    ) -> list[dict[str, Any]]:
        pairs: list[dict[str, Any]] = []
        content_boxes = [
            box
            for box in boxes
            if not box["is_chrome"]
            and not self._is_full_slide_text_region(box, slide_width, slide_height)
        ]
        for text_box in content_boxes:
            for shape_box in opaque_shapes:
                overlap = self._overlap_ratio(text_box, shape_box)
                # Text fully inside its card background is expected. Partial
                # overlap by a neighboring card is how Office-visible text gets
                # covered even when text extraction still sees the full string.
                if overlap < 0.045 or overlap > 0.88:
                    continue
                pairs.append(
                    {
                        "shape_indexes": [
                            text_box["shape_index"],
                            shape_box["shape_index"],
                        ],
                        "overlap_ratio": round(overlap, 3),
                        "text": self._truncate_text(str(text_box["text"]), 80),
                    }
                )
        return pairs

    def _is_full_slide_text_region(
        self,
        box: dict[str, Any],
        slide_width: float,
        slide_height: float,
    ) -> bool:
        return box["w"] >= slide_width * 0.92 and box["h"] >= slide_height * 0.88

    def _overlap_ratio(self, left: dict[str, Any], right: dict[str, Any]) -> float:
        x_overlap = max(
            0.0,
            min(left["x"] + left["w"], right["x"] + right["w"])
            - max(left["x"], right["x"]),
        )
        y_overlap = max(
            0.0,
            min(left["y"] + left["h"], right["y"] + right["h"])
            - max(left["y"], right["y"]),
        )
        overlap_area = x_overlap * y_overlap
        if overlap_area <= 0:
            return 0.0
        left_area = left["w"] * left["h"]
        right_area = right["w"] * right["h"]
        return overlap_area / max(min(left_area, right_area), 0.001)

    def _layout_box_payload(self, box: dict[str, Any]) -> dict[str, Any]:
        return {
            "shape_index": box["shape_index"],
            "text": self._truncate_text(str(box["text"]), 96),
            "bounds": {
                "x": round(float(box["x"]), 3),
                "y": round(float(box["y"]), 3),
                "w": round(float(box["w"]), 3),
                "h": round(float(box["h"]), 3),
            },
            "font_size": round(float(box["font_size"]), 1),
            "estimated_lines": box.get("estimated_lines"),
            "capacity_lines": box.get("capacity_lines"),
        }

    def _truncate_text(self, text: str, limit: int) -> str:
        cleaned = " ".join(str(text).split())
        if len(cleaned) <= limit:
            return cleaned
        return cleaned[: max(0, limit - 3)].rsplit(" ", 1)[0].rstrip(" ,.;:") + "..."

    def _has_prefix_repetition(self, texts: list[str]) -> bool:
        normalized = [
            re.sub(r"[^a-z0-9 ]+", "", " ".join(text.casefold().split())).strip()
            for text in texts
        ]
        normalized = [text for text in normalized if len(text.split()) >= 4]
        hits = 0
        for index, left in enumerate(normalized):
            for right in normalized[index + 1 :]:
                shorter, longer = sorted([left, right], key=len)
                shorter_words = len(shorter.split())
                longer_words = len(longer.split())
                if (
                    4 <= shorter_words <= 6
                    and longer.startswith(shorter)
                    and longer_words >= shorter_words + 3
                ):
                    # Many authored compositions intentionally use a short
                    # lead beside the full claim. That is not the repeated
                    # source-fragment bug this check is meant to catch.
                    continue
                if shorter_words >= 4 and longer.startswith(shorter):
                    hits += 1
                    if hits >= 2:
                        return True
        return False

    def _outline_semantic_issues(
        self, slide_index: int, outline: SlideOutline
    ) -> list[QAIssue]:
        issues: list[QAIssue] = []
        content = outline.content_json or {}
        if content.get("visual_degradation"):
            return issues
        exhibit = content.get("exhibit_spec")
        if not isinstance(exhibit, dict):
            return issues
        exhibit_type = str(exhibit.get("type") or "").lower()
        if exhibit_type == "matrix_2x2":
            quadrants = exhibit.get("quadrants") or []
            labels = [
                str(item.get("label") or "").casefold()
                for item in quadrants
                if isinstance(item, dict)
            ]
            generic = [label for label in labels if GENERIC_MATRIX_RE.search(label)]
            if len(generic) >= 2:
                issues.append(
                    QAIssue(
                        severity="WARNING",
                        message="Planned matrix appears to be a generic fallback rather than a real source-backed tradeoff.",
                        slide_index=slide_index,
                        category="semantic_visual_fit",
                    )
                )
        if exhibit_type in {"dependency_map", "cycle"}:
            flattened = self._flatten(exhibit).casefold()
            generic_hits = [
                label for label in GENERIC_DIAGRAM_LABELS if re.search(rf"\b{re.escape(label)}\b", flattened)
            ]
            if len(generic_hits) >= 3:
                issues.append(
                    QAIssue(
                        severity="WARNING",
                        message="Planned diagram contains too many generic fallback labels.",
                        slide_index=slide_index,
                        category="diagram_semantic_fit",
                    )
                )
        return issues

    def _looks_dangling(self, text: str) -> bool:
        cleaned = " ".join(text.split()).strip()
        if len(cleaned) < 16:
            return False
        if len(cleaned.split()) < 4:
            return bool(
                re.search(
                    r"\b(?:can|cannot|could|should|must|will|would|make|makes|"
                    r"create|creates|turn|turns|give|gives|evaluate|evaluates|"
                    r"specify|specifies|determine|determines|standardize|"
                    r"standardizes)\.?$",
                    cleaned,
                    re.IGNORECASE,
                )
            )
        if cleaned.endswith("..."):
            return True
        if TRAILING_FRAGMENT_RE.search(cleaned):
            return True
        if re.search(
            r"\b(?:a core|to help|enabling|managed|operating|reproducible|"
            r"varying|development of evaluation|reflect their actual use)\.?$",
            cleaned,
            re.IGNORECASE,
        ):
            return True
        if re.search(
            r"\b(?:can|cannot|could|should|must|will|would|make|makes|"
            r"create|creates|turn|turns|give|gives|evaluate|evaluates|"
            r"specify|specifies|determine|determines|standardize|"
            r"standardizes)\.?$",
            cleaned,
            re.IGNORECASE,
        ):
            return True
        if re.search(r"\benabling scalable\.?$", cleaned, re.IGNORECASE):
            return True
        lowered = cleaned.casefold()
        if lowered in ORPHAN_META_LABELS:
            return True
        if re.search(
            r"\b(?:cannot|can|make|makes|around|through|from|to)\s+is\b",
            lowered,
        ):
            return True
        if re.search(
            r"\b(?:because|while|when|where|that|which)\s+(?:it|they|the)$",
            lowered,
        ):
            return True
        return bool(re.search(r"\bdevelopment of\.?$", cleaned, re.IGNORECASE))

    def _next_text_completes_fragment(self, text: str, next_text: str) -> bool:
        if not next_text:
            return False
        current = " ".join(str(text).split()).rstrip(".,;:")
        following = " ".join(str(next_text).split()).strip()
        if len(following.split()) > 10:
            return False
        if not re.match(r"^(?:a|an|the|of|to|for|with|in|on|by|as|[a-z]\w*)\b", following):
            return False
        combined = f"{current} {following}"
        return not self._looks_dangling(combined)

    def _diagram_labels_by_slide(self, pptx_path: Path) -> dict[int, list[str]]:
        labels_by_slide: dict[int, list[str]] = {}
        diagram_dir = pptx_path.parent / f"{pptx_path.stem}-diagrams"
        if not diagram_dir.exists():
            return labels_by_slide
        for svg_path in sorted(diagram_dir.glob("*.svg")):
            match = re.match(r"(\d+)-", svg_path.name)
            if not match:
                continue
            slide_index = max(0, int(match.group(1)) - 1)
            try:
                text = svg_path.read_text(encoding="utf-8")
            except Exception:
                continue
            labels = [
                " ".join(label.split())
                for label in re.findall(r">([^<>]+)</text>", text)
                if " ".join(label.split())
            ]
            labels_by_slide.setdefault(slide_index, []).extend(labels)
        return labels_by_slide

    def _is_generic_diagram(self, labels: list[str]) -> bool:
        normalized = {" ".join(label.casefold().split()) for label in labels}
        hits = normalized.intersection(GENERIC_DIAGRAM_LABELS)
        return len(hits) >= 3

    def _composition_signature(self, outline: SlideOutline | None) -> str | None:
        if outline is None:
            return None
        content = outline.content_json or {}
        layout = outline.layout_json or {}
        exhibit = content.get("exhibit_spec")
        exhibit_type = exhibit.get("type") if isinstance(exhibit, dict) else ""
        return "|".join(
            str(part or "")
            for part in (
                layout.get("composition_family") or layout.get("layout"),
                layout.get("composition_variant"),
                content.get("narrative_role") or layout.get("narrative_role"),
                exhibit_type,
            )
        )

    def _visual_rhythm_summary(
        self,
        slides: list[dict[str, Any]],
        outlines: list[SlideOutline],
    ) -> dict[str, Any]:
        families: list[str] = []
        signatures: list[str] = []
        for outline in outlines:
            if outline.mode != "flexible":
                continue
            layout = outline.layout_json or {}
            family = str(layout.get("composition_family") or layout.get("layout") or "")
            if family in {"editorial_cover", "path_forward_close", "section_divider"}:
                continue
            if family:
                families.append(family)
            signature = str(
                layout.get("composition_signature")
                or self._composition_signature(outline)
                or ""
            )
            if signature:
                signatures.append(signature)
        family_counts = Counter(families)
        signature_counts = Counter(signatures)
        card_like_families = {
            "evidence_wall",
            "proof_strip",
            "toolkit_grid",
            "challenge_cards",
            "why_it_matters_cards",
            "source_repair_cards",
        }
        marker_counts: Counter[str] = Counter()
        for slide in slides:
            text = " || ".join(str(item) for item in slide.get("text", []))
            for marker in (
                "COMPARISON LENS",
                "PROOF",
                "OPERATING KIT",
                "WHY IT MATTERS",
            ):
                if marker in text:
                    marker_counts[marker] += 1
        slide_count = len(families)
        card_like_count = sum(
            count
            for family, count in family_counts.items()
            if family in card_like_families
        )
        return {
            "slide_count": slide_count,
            "unique_family_count": len(family_counts),
            "family_counts": dict(sorted(family_counts.items())),
            "signature_counts": dict(sorted(signature_counts.items())),
            "card_like_count": card_like_count,
            "card_like_ratio": round(card_like_count / slide_count, 3)
            if slide_count
            else 0,
            "marker_counts": dict(sorted(marker_counts.items())),
        }

    def _visual_rhythm_issues(self, rhythm: dict[str, Any]) -> list[QAIssue]:
        slide_count = int(rhythm.get("slide_count") or 0)
        if slide_count < 8:
            return []
        issues: list[QAIssue] = []
        unique_family_count = int(rhythm.get("unique_family_count") or 0)
        minimum_unique = min(6, max(4, slide_count // 3))
        if unique_family_count < minimum_unique:
            issues.append(
                QAIssue(
                    severity="WARNING",
                    message=(
                        "Deck visual rhythm is too narrow: too few composition "
                        "families for the number of generated slides."
                    ),
                    category="visual_rhythm",
                )
            )
        family_counts = rhythm.get("family_counts") or {}
        if isinstance(family_counts, dict) and family_counts:
            most_common = max(int(count) for count in family_counts.values())
            repeat_limit = 2 if slide_count >= 12 else 3
            if most_common > repeat_limit:
                issues.append(
                    QAIssue(
                        severity="WARNING",
                        message=(
                            "Deck repeats one authored composition family often "
                            "enough to feel like layout cycling."
                        ),
                        category="visual_rhythm",
                    )
                )
            process_like_count = sum(
                int(family_counts.get(family, 0) or 0)
                for family in (
                    "lifecycle_timeline",
                    "decision_ladder",
                    "operating_map",
                )
            )
            if process_like_count > max(3, int(slide_count * 0.28)):
                issues.append(
                    QAIssue(
                        severity="WARNING",
                        message=(
                            "Deck overuses process-like compositions; mix in "
                            "editorial spreads, proof views, contrasts, or "
                            "system views before marking the deck ready."
                        ),
                        category="visual_rhythm",
                    )
                )
        if float(rhythm.get("card_like_ratio") or 0) > 0.46:
            issues.append(
                QAIssue(
                    severity="WARNING",
                    message=(
                        "Deck relies on card/grid/strip compositions for too many "
                        "slides; add full-canvas editorial, timeline, or system views."
                    ),
                    category="visual_rhythm",
                )
            )
        marker_counts = rhythm.get("marker_counts") or {}
        if isinstance(marker_counts, dict):
            repeated_markers = [
                marker for marker, count in marker_counts.items() if int(count) >= 3
            ]
            if repeated_markers:
                issues.append(
                    QAIssue(
                        severity="WARNING",
                        message=(
                            "Rendered slides repeat the same visible composition "
                            f"markers: {', '.join(repeated_markers)}."
                        ),
                        category="visual_rhythm",
                    )
                )
        return issues

    def _flatten(self, value: Any) -> str:
        if isinstance(value, dict):
            return " ".join(self._flatten(item) for item in value.values())
        if isinstance(value, list):
            return " ".join(self._flatten(item) for item in value)
        return str(value or "")
