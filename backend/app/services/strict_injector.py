import zipfile
import posixpath
import re
from io import BytesIO
from pathlib import Path
from typing import Any

from lxml import etree
from openpyxl import load_workbook

from app.models.outline import SlideOutline
from app.models.template import SlideSpec, TemplateProfile
from app.services.strict_validator import StrictSchemaValidator


NSMAP = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"

SAFE_PARSER = etree.XMLParser(resolve_entities=False, no_network=True)


class StrictSlideInjector:
    def __init__(self) -> None:
        self.validator = StrictSchemaValidator()

    def inject(
        self,
        template_path: Path,
        template: TemplateProfile,
        outlines: list[SlideOutline],
        output_path: Path,
    ) -> list[dict[str, str | int]]:
        outline_map = {outline.slide_index: outline for outline in outlines}
        warnings: list[dict[str, str | int]] = []
        with zipfile.ZipFile(template_path, "r") as source_zip:
            entries = {name: source_zip.read(name) for name in source_zip.namelist()}

        updated_entries: dict[str, bytes] = {}
        for slide_spec in template.slides:
            if slide_spec.mode != "strict":
                continue
            outline = outline_map.get(slide_spec.index)
            if not outline:
                continue
            slide_name = f"ppt/slides/slide{slide_spec.index + 1}.xml"
            slide_xml = entries.get(slide_name)
            if slide_xml is None:
                warnings.append(
                    {
                        "slide_index": slide_spec.index,
                        "field": "slide",
                        "message": f"Slide XML not found: {slide_name}",
                    }
                )
                continue
            updated_xml, chart_updates, slide_warnings = self._apply_outline_xml(
                slide_xml, slide_name, entries, slide_spec, outline
            )
            updated_entries[slide_name] = updated_xml
            updated_entries.update(chart_updates)
            warnings.extend(slide_warnings)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as output_zip:
            for name, data in entries.items():
                output_zip.writestr(name, updated_entries.get(name, data))
        return warnings

    def _apply_outline_xml(
        self,
        slide_xml: bytes,
        slide_name: str,
        entries: dict[str, bytes],
        slide_spec: SlideSpec,
        outline: SlideOutline,
    ) -> tuple[bytes, dict[str, bytes], list[dict[str, str | int]]]:
        values = outline.content_json.get("fields", {})
        if not slide_spec.slide_schema:
            return slide_xml, {}, []
        root = etree.fromstring(slide_xml, parser=SAFE_PARSER)
        chart_updates: dict[str, bytes] = {}
        warnings: list[dict[str, str | int]] = []
        for field in slide_spec.slide_schema.fields:
            value = values.get(field.id)
            normalized_value, field_warnings = self.validator.validate_field_value(
                field, value
            )
            for warning in field_warnings:
                warnings.append(
                    {"slide_index": slide_spec.index, "field": field.id, "message": warning}
                )
            if value is None:
                continue
            if field.location.startswith("chart:"):
                chart_warning = self._apply_chart_value(
                    entries,
                    chart_updates,
                    slide_name,
                    root,
                    field.location,
                    normalized_value,
                )
                if chart_warning:
                    warnings.append(
                        {
                            "slide_index": slide_spec.index,
                            "field": field.id,
                            "message": chart_warning,
                        }
                    )
                continue
            target = self._find_table_cell(root, field.location)
            if target is None:
                target = self._find_shape(root, field.location)
            if target is None:
                warnings.append(
                    {
                        "slide_index": slide_spec.index,
                        "field": field.id,
                        "message": f"Target not found for location {field.location}",
                    }
                )
                continue
            self._apply_value(target, field, normalized_value)
        xml_bytes = etree.tostring(
            root, xml_declaration=True, encoding="UTF-8", standalone=True
        )
        etree.fromstring(xml_bytes, parser=SAFE_PARSER)
        return xml_bytes, chart_updates, warnings

    def _apply_chart_value(
        self,
        entries: dict[str, bytes],
        chart_updates: dict[str, bytes],
        slide_name: str,
        slide_root: etree._Element,
        location: str,
        value: Any,
    ) -> str | None:
        parsed = self._parse_chart_location(location)
        if parsed is None:
            return f"Invalid chart location {location}"
        chart_name, target_type, series_idx, point_idx = parsed
        chart_part = self._find_chart_part(entries, slide_name, slide_root, chart_name)
        if chart_part is None:
            return f"Chart not found for location {location}"
        chart_xml = chart_updates.get(chart_part, entries.get(chart_part))
        if chart_xml is None:
            return f"Chart XML not found for location {location}"
        chart_root = etree.fromstring(chart_xml, parser=SAFE_PARSER)
        if target_type == "series_name":
            warning = self._replace_chart_series_name(chart_root, series_idx, str(value))
        elif target_type == "category":
            warning = self._replace_chart_category(chart_root, point_idx, str(value))
        else:
            warning = self._replace_chart_value(chart_root, series_idx, point_idx, value)
        if warning:
            return warning
        workbook_warning = self._sync_chart_workbook(
            entries,
            chart_updates,
            chart_part,
            chart_root,
            target_type,
            series_idx,
            point_idx,
            value,
        )
        if workbook_warning:
            return workbook_warning
        updated = etree.tostring(
            chart_root, xml_declaration=True, encoding="UTF-8", standalone=True
        )
        etree.fromstring(updated, parser=SAFE_PARSER)
        chart_updates[chart_part] = updated
        return None

    def _parse_chart_location(
        self, location: str
    ) -> tuple[str, str, int, int | None] | None:
        parts = location.split(":")
        if len(parts) < 3 or parts[0] != "chart":
            return None
        chart_name = parts[1]
        target_type = parts[2]
        try:
            if target_type == "series_name" and len(parts) == 4:
                return chart_name, target_type, int(parts[3]), None
            if target_type == "category" and len(parts) == 4:
                return chart_name, target_type, 0, int(parts[3])
            if target_type == "value" and len(parts) == 5:
                return chart_name, target_type, int(parts[3]), int(parts[4])
        except ValueError:
            return None
        return None

    def _find_chart_part(
        self,
        entries: dict[str, bytes],
        slide_name: str,
        slide_root: etree._Element,
        chart_name: str,
    ) -> str | None:
        rel_id = None
        for frame in slide_root.xpath(".//p:graphicFrame", namespaces=NSMAP):
            c_nv_pr = frame.find(".//p:cNvPr", namespaces=NSMAP)
            if c_nv_pr is None or c_nv_pr.get("name") != chart_name:
                continue
            chart = frame.find(".//c:chart", namespaces=NSMAP)
            if chart is not None:
                rel_id = chart.get(f"{{{NSMAP['r']}}}id")
                break
        if not rel_id:
            return None
        rel_name = self._slide_rels_name(slide_name)
        rel_xml = entries.get(rel_name)
        if rel_xml is None:
            return None
        rel_root = etree.fromstring(rel_xml, parser=SAFE_PARSER)
        for rel in rel_root.findall(f".//{{{REL_NS}}}Relationship"):
            if rel.get("Id") != rel_id:
                continue
            target = rel.get("Target")
            if not target:
                return None
            return self._resolve_relationship_target(slide_name, target)
        return None

    def _slide_rels_name(self, slide_name: str) -> str:
        return (
            f"{posixpath.dirname(slide_name)}/_rels/"
            f"{posixpath.basename(slide_name)}.rels"
        )

    def _resolve_relationship_target(self, source_part: str, target: str) -> str:
        if target.startswith("/"):
            return posixpath.normpath(target.lstrip("/"))
        return posixpath.normpath(posixpath.join(posixpath.dirname(source_part), target))

    def _chart_series(self, chart_root: etree._Element, series_idx: int) -> etree._Element | None:
        series = chart_root.xpath(".//c:ser", namespaces=NSMAP)
        return series[series_idx] if 0 <= series_idx < len(series) else None

    def _replace_chart_series_name(
        self, chart_root: etree._Element, series_idx: int, value: str
    ) -> str | None:
        series = self._chart_series(chart_root, series_idx)
        if series is None:
            return f"Chart series index not found: {series_idx}"
        node = series.find(".//c:tx//c:strCache/c:pt[@idx='0']/c:v", namespaces=NSMAP)
        if node is None:
            return f"Chart series name cache not found: {series_idx}"
        node.text = value
        return None

    def _replace_chart_category(
        self, chart_root: etree._Element, point_idx: int | None, value: str
    ) -> str | None:
        if point_idx is None:
            return "Chart category point index missing."
        node = chart_root.find(
            f".//c:ser[1]/c:cat//c:strCache/c:pt[@idx='{point_idx}']/c:v",
            namespaces=NSMAP,
        )
        if node is None:
            return f"Chart category point not found: {point_idx}"
        node.text = value
        return None

    def _replace_chart_value(
        self,
        chart_root: etree._Element,
        series_idx: int,
        point_idx: int | None,
        value: Any,
    ) -> str | None:
        if point_idx is None:
            return "Chart value point index missing."
        series = self._chart_series(chart_root, series_idx)
        if series is None:
            return f"Chart series index not found: {series_idx}"
        node = series.find(
            f".//c:val//c:numCache/c:pt[@idx='{point_idx}']/c:v",
            namespaces=NSMAP,
        )
        if node is None:
            return f"Chart value point not found: {series_idx}:{point_idx}"
        node.text = str(value)
        return None

    def _sync_chart_workbook(
        self,
        entries: dict[str, bytes],
        updates: dict[str, bytes],
        chart_part: str,
        chart_root: etree._Element,
        target_type: str,
        series_idx: int,
        point_idx: int | None,
        value: Any,
    ) -> str | None:
        workbook_part = self._chart_workbook_part(entries, chart_part)
        if workbook_part is None:
            return None
        workbook_bytes = updates.get(workbook_part, entries.get(workbook_part))
        if workbook_bytes is None:
            return f"Embedded chart workbook not found: {workbook_part}"
        cell_ref = self._chart_workbook_cell(
            chart_root, target_type, series_idx, point_idx
        )
        if cell_ref is None:
            return None
        sheet_name, cell_address = cell_ref
        workbook = load_workbook(BytesIO(workbook_bytes))
        if sheet_name not in workbook.sheetnames:
            return f"Embedded chart workbook sheet not found: {sheet_name}"
        sheet = workbook[sheet_name]
        sheet[cell_address] = value
        output = BytesIO()
        workbook.save(output)
        updates[workbook_part] = output.getvalue()
        return None

    def _chart_workbook_part(
        self, entries: dict[str, bytes], chart_part: str
    ) -> str | None:
        rel_name = (
            f"{posixpath.dirname(chart_part)}/_rels/"
            f"{posixpath.basename(chart_part)}.rels"
        )
        rel_xml = entries.get(rel_name)
        if rel_xml is None:
            return None
        rel_root = etree.fromstring(rel_xml, parser=SAFE_PARSER)
        for rel in rel_root.findall(f".//{{{REL_NS}}}Relationship"):
            rel_type = rel.get("Type", "")
            target = rel.get("Target")
            if not target or not rel_type.endswith("/package"):
                continue
            return self._resolve_relationship_target(chart_part, target)
        return None

    def _chart_workbook_cell(
        self,
        chart_root: etree._Element,
        target_type: str,
        series_idx: int,
        point_idx: int | None,
    ) -> tuple[str, str] | None:
        series = self._chart_series(chart_root, series_idx)
        if series is None:
            return None
        if target_type == "series_name":
            formula = series.find(".//c:tx//c:f", namespaces=NSMAP)
            return self._first_cell_from_formula(formula.text if formula is not None else "")
        if target_type == "category":
            if point_idx is None:
                return None
            formula = series.find(".//c:cat//c:f", namespaces=NSMAP)
            return self._cell_from_formula_point(formula.text if formula is not None else "", point_idx)
        if target_type == "value":
            if point_idx is None:
                return None
            formula = series.find(".//c:val//c:f", namespaces=NSMAP)
            return self._cell_from_formula_point(formula.text if formula is not None else "", point_idx)
        return None

    def _first_cell_from_formula(self, formula: str) -> tuple[str, str] | None:
        parsed = self._parse_sheet_range(formula)
        if parsed is None:
            return None
        sheet_name, start_cell, _ = parsed
        return sheet_name, start_cell

    def _cell_from_formula_point(
        self, formula: str, point_idx: int
    ) -> tuple[str, str] | None:
        parsed = self._parse_sheet_range(formula)
        if parsed is None:
            return None
        sheet_name, start_cell, end_cell = parsed
        cells = self._expand_cell_range(start_cell, end_cell)
        if point_idx >= len(cells):
            return None
        return sheet_name, cells[point_idx]

    def _parse_sheet_range(self, formula: str) -> tuple[str, str, str] | None:
        if "!" not in formula:
            return None
        raw_sheet, raw_range = formula.split("!", 1)
        sheet_name = raw_sheet.strip("'")
        cells = raw_range.replace("$", "").split(":")
        if not cells or not cells[0]:
            return None
        start_cell = cells[0]
        end_cell = cells[-1] if len(cells) > 1 else start_cell
        return sheet_name, start_cell, end_cell

    def _expand_cell_range(self, start_cell: str, end_cell: str) -> list[str]:
        start_col, start_row = self._split_cell(start_cell)
        end_col, end_row = self._split_cell(end_cell)
        if start_col == end_col:
            return [f"{start_col}{row}" for row in range(start_row, end_row + 1)]
        if start_row == end_row:
            return [
                f"{self._column_name(col)}{start_row}"
                for col in range(self._column_index(start_col), self._column_index(end_col) + 1)
            ]
        return [start_cell]

    def _split_cell(self, cell: str) -> tuple[str, int]:
        match = re.match(r"^([A-Z]+)(\d+)$", cell.upper())
        if not match:
            return cell.upper(), 1
        return match.group(1), int(match.group(2))

    def _column_index(self, column: str) -> int:
        value = 0
        for char in column.upper():
            value = value * 26 + (ord(char) - ord("A") + 1)
        return value

    def _column_name(self, index: int) -> str:
        chars = []
        while index:
            index, remainder = divmod(index - 1, 26)
            chars.append(chr(ord("A") + remainder))
        return "".join(reversed(chars)) or "A"

    def _find_shape(self, root: etree._Element, location: str) -> etree._Element | None:
        if location.startswith("shape:"):
            target_name = location.split(":", 1)[1]
            for shape in root.xpath(".//p:sp", namespaces=NSMAP):
                c_nv_pr = shape.find(".//p:cNvPr", namespaces=NSMAP)
                if c_nv_pr is not None and c_nv_pr.get("name") == target_name:
                    return shape
            for shape in root.xpath(".//p:graphicFrame", namespaces=NSMAP):
                c_nv_pr = shape.find(".//p:cNvPr", namespaces=NSMAP)
                if c_nv_pr is not None and c_nv_pr.get("name") == target_name:
                    return shape
        if location.startswith("shape_id:"):
            target_id = location.split(":", 1)[1]
            for shape in root.xpath(".//p:sp", namespaces=NSMAP):
                c_nv_pr = shape.find(".//p:cNvPr", namespaces=NSMAP)
                if c_nv_pr is not None and c_nv_pr.get("id") == target_id:
                    return shape
        return None

    def _find_table_cell(self, root: etree._Element, location: str) -> etree._Element | None:
        if not location.startswith("table:"):
            return None
        parts = location.split(":")
        if len(parts) != 4:
            return None
        _, table_name, row_text, col_text = parts
        try:
            row_idx = int(row_text)
            col_idx = int(col_text)
        except ValueError:
            return None
        for frame in root.xpath(".//p:graphicFrame", namespaces=NSMAP):
            c_nv_pr = frame.find(".//p:cNvPr", namespaces=NSMAP)
            if c_nv_pr is None or c_nv_pr.get("name") != table_name:
                continue
            rows = frame.xpath(".//a:tbl/a:tr", namespaces=NSMAP)
            if row_idx >= len(rows):
                return None
            cells = rows[row_idx].xpath("./a:tc", namespaces=NSMAP)
            if col_idx >= len(cells):
                return None
            return cells[col_idx]
        return None

    def _apply_value(self, shape: etree._Element, field, value: Any) -> None:
        if field.type == "text_list" and isinstance(value, list):
            text_value = "\n".join(str(item) for item in value)
        else:
            text_value = str(value)

        if field.render == "fill_color" and field.color_map:
            color = field.color_map.get(str(value).lower())
            if color:
                self._apply_fill_color(shape, color)

        self._replace_text(shape, text_value)

    def _replace_text(self, shape: etree._Element, text: str) -> None:
        tx_body = shape.find(".//p:txBody", namespaces=NSMAP)
        if tx_body is None:
            tx_body = shape.find(".//a:txBody", namespaces=NSMAP)
        if tx_body is None:
            return
        paragraphs = tx_body.findall("a:p", namespaces=NSMAP)
        if not paragraphs:
            return
        first_para = paragraphs[0]
        runs = first_para.findall("a:r", namespaces=NSMAP)
        if not runs:
            run = etree.SubElement(first_para, f"{{{NSMAP['a']}}}r")
            text_node = etree.SubElement(run, f"{{{NSMAP['a']}}}t")
            text_node.text = text
        else:
            text_node = runs[0].find("a:t", namespaces=NSMAP)
            if text_node is None:
                text_node = etree.SubElement(runs[0], f"{{{NSMAP['a']}}}t")
            text_node.text = text
            for run in runs[1:]:
                first_para.remove(run)
        for para in paragraphs[1:]:
            tx_body.remove(para)

    def _apply_fill_color(self, shape: etree._Element, color: str) -> None:
        sp_pr = shape.find("p:spPr", namespaces=NSMAP)
        if sp_pr is None:
            return
        for fill in sp_pr.findall("a:solidFill", namespaces=NSMAP):
            sp_pr.remove(fill)
        solid = etree.SubElement(sp_pr, f"{{{NSMAP['a']}}}solidFill")
        srgb = etree.SubElement(solid, f"{{{NSMAP['a']}}}srgbClr")
        srgb.set("val", color.replace("#", "")[:6])
