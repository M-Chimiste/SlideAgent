# ruff: noqa: F401
from pathlib import Path
import math
import re
import shutil
import subprocess
from typing import Any

from PIL import Image, ImageDraw
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt

from app.models.brand import BrandDNA
from app.models.outline import SlideOutline
from app.services.concept_diagram_renderer import ConceptDiagramRenderer, DiagramRenderError
from app.services.pptx_rendering.constants import ICON_SCALE, SLIDE_H, SLIDE_W


class CoreLayoutRenderingMixin:
    def _add_cover(self, slide, outline: SlideOutline, brand: BrandDNA) -> None:
        background = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0),
            Inches(0),
            Inches(SLIDE_W),
            Inches(SLIDE_H),
        )
        background.fill.solid()
        background.fill.fore_color.rgb = self._rgb(brand.colors.primary)
        background.line.color.rgb = self._rgb(brand.colors.primary)
        title = (
            outline.content_json.get("deck_title")
            or outline.content_json.get("title")
            or outline.content_json.get("action_title")
            or outline.label
        )
        exhibit = self._exhibit(outline)
        thesis = (
            exhibit.get("thesis")
            or outline.content_json.get("subheading")
            or outline.content_json.get("summary")
            or "A practical operating model for executive action."
        )
        if self._looks_like_meta_subheading(str(thesis)):
            thesis = "Source-grounded operating choices for the leadership review."
        self._add_dark_text(
            slide,
            "EXECUTIVE PERSPECTIVE",
            0.78,
            0.74,
            2.8,
            0.26,
            brand,
            size=10,
            bold=True,
            color=self._tint(brand.colors.primary, 0.78),
        )
        rail = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.78),
            Inches(1.08),
            Inches(1.15),
            Inches(0.04),
        )
        rail.fill.solid()
        rail.fill.fore_color.rgb = self._rgb(brand.colors.accent)
        rail.line.color.rgb = self._rgb(brand.colors.accent)
        rail.line.width = Pt(0)
        title_field = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.0),
            Inches(1.34),
            Inches(0.18),
            Inches(2.25),
        )
        title_field.fill.solid()
        title_field.fill.fore_color.rgb = self._rgb(brand.colors.accent)
        title_field.line.color.rgb = title_field.fill.fore_color.rgb
        self._add_dark_text(slide, title, 0.78, 1.48, 8.05, 1.56, brand, size=34, bold=True)
        self._add_dark_text(
            slide,
            self._truncate_at_word(thesis, 175),
            0.82,
            3.25,
            7.0,
            0.8,
            brand,
            size=14,
            color=self._tint(brand.colors.primary, 0.76),
        )
        signal_panel = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(8.62),
            Inches(1.18),
            Inches(3.74),
            Inches(4.38),
        )
        signal_panel.fill.solid()
        signal_panel.fill.fore_color.rgb = self._rgb(self._tint(brand.colors.primary, 0.12))
        signal_panel.line.color.rgb = signal_panel.fill.fore_color.rgb
        stack_x, stack_y = 8.95, 1.5
        stack_items = self._cover_stack_items(exhibit, outline)
        for idx, (number, label, detail) in enumerate(stack_items):
            y = stack_y + idx * 1.2
            divider = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(stack_x + 0.06),
                Inches(y + 0.94),
                Inches(2.72),
                Inches(0.015),
            )
            divider.fill.solid()
            divider.fill.fore_color.rgb = self._rgb(self._tint(brand.colors.primary, 0.38))
            divider.line.color.rgb = divider.fill.fore_color.rgb
            self._add_dark_text(
                slide,
                number,
                stack_x + 0.08,
                y + 0.04,
                0.56,
                0.32,
                brand,
                size=13,
                bold=True,
                color=brand.colors.accent,
            )
            self._add_dark_text(
                slide,
                label,
                stack_x + 0.08,
                y + 0.42,
                2.52,
                0.28,
                brand,
                size=12,
                bold=True,
                color=brand.colors.text_light,
            )
            self._add_dark_text(
                slide,
                detail,
                stack_x + 0.08,
                y + 0.73,
                2.56,
                0.22,
                brand,
                size=8,
                color=self._tint(brand.colors.primary, 0.82),
            )

    def _cover_stack_items(
        self, exhibit: dict[str, Any], outline: SlideOutline
    ) -> list[tuple[str, str, str]]:
        raw_signals = exhibit.get("signals") if isinstance(exhibit, dict) else []
        if not isinstance(raw_signals, list):
            raw_signals = []
        signals = [
            self._truncate_at_word(self._clean_display_text(str(item)), 28)
            for item in raw_signals
            if self._clean_display_text(str(item))
        ]
        if not signals:
            signals = [
                self._truncate_at_word(text, 28)
                for text in self._bullets(outline)[:3]
                if text
            ]
        if not signals:
            signals = self._default_cover_signals(outline)
        defaults = ["Context", "Decision", "Execution"]
        items: list[tuple[str, str, str]] = []
        for idx in range(3):
            signal = signals[idx] if idx < len(signals) else defaults[idx]
            items.append((f"{idx + 1:02d}", signal, self._cover_signal_detail(signal, idx)))
        return items

    def _default_cover_signals(self, outline: SlideOutline) -> list[str]:
        text = " ".join(
            [
                self._authored_title(outline),
                str(outline.content_json.get("summary") or ""),
                str(outline.content_json.get("speaker_notes") or ""),
            ]
        ).lower()
        if any(
            token in text
            for token in ("benchmark", "harness", "ground truth", "model contract")
        ):
            return ["Ground truth", "Model contract", "Harness execution"]
        if any(token in text for token in ("memory", "context", "agent", "coding")):
            return ["Persistent context", "Explicit rules", "Review loop"]
        return ["Decision context", "Evidence standard", "Operating move"]

    def _cover_signal_detail(self, signal: str, index: int) -> str:
        lowered = signal.lower()
        if any(token in lowered for token in ("memory", "context", "brain")):
            return "Keep knowledge persistent"
        if "ground truth" in lowered:
            return "Find validated evidence"
        if "model contract" in lowered:
            return "Define success criteria"
        if "harness" in lowered:
            return "Run repeatable tests"
        if any(token in lowered for token in ("rule", "spec", "standard")):
            return "Make criteria explicit"
        if any(token in lowered for token in ("review", "quality", "evidence")):
            return "Verify before scale"
        if any(token in lowered for token in ("cycle", "workflow", "loop")):
            return "Run the operating cadence"
        return [
            "Frame the decision",
            "Pressure-test the change",
            "Commit to next actions",
        ][index % 3]

    def _add_executive_summary(self, slide, outline: SlideOutline, brand: BrandDNA) -> None:
        exhibit = self._exhibit(outline)
        messages = exhibit.get("messages") if exhibit.get("type") == "executive_summary" else None
        bullets = [
            str(item.get("text", ""))
            for item in messages
            if isinstance(item, dict) and str(item.get("text", "")).strip()
        ] if isinstance(messages, list) else self._bullets(outline)
        y = 1.55
        labels = [
            str(item.get("label", ""))
            for item in messages
            if isinstance(item, dict) and str(item.get("label", "")).strip()
        ] if isinstance(messages, list) else ["Situation", "Complication", "Resolution"]
        while len(labels) < 3:
            labels.append(["Situation", "Complication", "Resolution"][len(labels)])
        colors = [brand.colors.primary, brand.colors.secondary, brand.colors.accent]
        panel = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(8.42),
            Inches(1.48),
            Inches(3.72),
            Inches(3.95),
        )
        panel.fill.solid()
        panel.fill.fore_color.rgb = self._rgb(brand.colors.primary)
        panel.line.color.rgb = self._rgb(brand.colors.primary)
        self._add_dark_text(
            slide,
            "PROOF POINTS",
            8.72,
            1.82,
            1.5,
            0.22,
            brand,
            size=8,
            bold=True,
            color=brand.colors.accent,
        )
        proof_points = self._summary_proof_points(exhibit, outline)
        for idx, proof in enumerate(proof_points[:3]):
            y_pos = 2.24 + idx * 0.88
            value = str(proof.get("value", ""))
            unit = ""
            if isinstance(proof.get("value"), (int, float)):
                value, unit = self._metric_card_value_parts(proof)
            self._add_dark_text(
                slide,
                value,
                8.72,
                y_pos,
                0.86,
                0.34,
                brand,
                size=18,
                bold=True,
                color=brand.colors.text_light,
            )
            label = str(proof.get("label") or "Signal")
            if unit and unit.lower() not in label.lower():
                label = f"{unit} / {label}"
            self._add_dark_text(
                slide,
                self._truncate_at_word(label, 32),
                9.68,
                y_pos + 0.02,
                2.06,
                0.2,
                brand,
                size=8,
                bold=True,
                color=self._tint(brand.colors.primary, 0.84),
            )
            self._add_dark_text(
                slide,
                self._truncate_phrase(str(proof.get("detail") or ""), 62),
                9.68,
                y_pos + 0.28,
                2.0,
                0.35,
                brand,
                size=7,
                color=self._tint(brand.colors.primary, 0.7),
            )
        rail = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.86),
            Inches(1.55),
            Inches(0.04),
            Inches(3.78),
        )
        rail.fill.solid()
        rail.fill.fore_color.rgb = self._rgb(self._tint(brand.colors.primary, 0.58))
        rail.line.color.rgb = rail.fill.fore_color.rgb
        for idx, label in enumerate(labels[:3]):
            row_y = y + idx * 1.24
            row = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(1.05),
                Inches(row_y),
                Inches(6.95),
                Inches(0.96),
            )
            row.fill.solid()
            row.fill.fore_color.rgb = self._rgb("FFFFFF" if idx % 2 == 0 else "F4F6F8")
            row.line.color.rgb = self._rgb(brand.colors.background_light)
            strip = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(1.05),
                Inches(row_y),
                Inches(0.12),
                Inches(0.96),
            )
            strip.fill.solid()
            strip.fill.fore_color.rgb = self._rgb(colors[idx % len(colors)])
            strip.line.color.rgb = strip.fill.fore_color.rgb
            self._add_dark_text(
                slide,
                f"0{idx + 1}",
                1.34,
                row_y + 0.29,
                0.45,
                0.24,
                brand,
                size=10,
                bold=True,
                color=colors[idx % len(colors)],
            )
            self._add_dark_text(
                slide,
                label,
                2.0,
                row_y + 0.19,
                2.35,
                0.28,
                brand,
                size=13,
                bold=True,
                color=brand.colors.primary,
            )
            text = bullets[idx] if idx < len(bullets) else "Define the critical implication for leadership."
            self._add_body_text(
                slide,
                self._truncate_at_word(text, 128),
                4.42,
                row_y + 0.2,
                3.1,
                0.48,
                brand,
                size=10,
        )
        recommendation = (
            exhibit.get("recommendation")
            or "Manage the work with persistent context, explicit rules, and evidence-backed review."
        )
        rec = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(1.05),
            Inches(5.62),
            Inches(11.1),
            Inches(0.72),
        )
        rec.fill.solid()
        rec.fill.fore_color.rgb = self._rgb(self._tint(brand.colors.accent, 0.86))
        rec.line.color.rgb = self._rgb(self._tint(brand.colors.accent, 0.7))
        self._add_body_text(
            slide,
            "Decision ask",
            1.34,
            5.84,
            1.2,
            0.24,
            brand,
            size=8,
        )
        self._add_body_text(
            slide,
            self._truncate_at_word(str(recommendation), 150),
            2.58,
            5.78,
            8.8,
            0.32,
            brand,
            size=11,
        )

    def _summary_proof_points(
        self, exhibit: dict[str, Any], outline: SlideOutline
    ) -> list[dict[str, Any]]:
        proof_points = exhibit.get("proof_points", [])
        if isinstance(proof_points, list) and proof_points:
            strong = [
                item
                for item in proof_points
                if isinstance(item, dict)
                and str(item.get("label") or item.get("value") or "").strip()
                and self._is_strong_proof_point(item)
            ]
            if len(strong) >= 2:
                return strong
        metrics = [
            metric
            for metric in outline.content_json.get("metrics", [])
            if isinstance(metric, dict) and self._is_strong_proof_point(metric)
        ]
        if len(metrics) >= 2:
            return metrics
        qualitative = self._qualitative_summary_proof_points(exhibit, outline)
        if qualitative:
            return qualitative
        return [
            {"label": "Context", "value": "1", "detail": "Persistent memory keeps work reproducible."},
            {"label": "Rules", "value": "2", "detail": "Acceptance criteria constrain generation."},
            {"label": "Review", "value": "3", "detail": "Evidence gates protect delivery quality."},
        ]

    def _is_strong_proof_point(self, proof: dict[str, Any]) -> bool:
        unit = str(proof.get("unit") or "").strip()
        value = proof.get("value")
        if isinstance(value, str) and not value.strip():
            return False
        if unit:
            try:
                numeric = float(str(value).replace(",", "").rstrip("%"))
            except (TypeError, ValueError):
                return True
            label = str(proof.get("label") or "").strip().lower()
            if numeric <= 2 and label in {"prediction", "predictions", "item", "items"}:
                return False
            return True
        try:
            numeric = float(str(value).replace(",", "").rstrip("%"))
        except (TypeError, ValueError):
            return bool(str(value).strip())
        return abs(numeric) >= 10

    def _qualitative_summary_proof_points(
        self,
        exhibit: dict[str, Any],
        outline: SlideOutline,
    ) -> list[dict[str, Any]]:
        messages = exhibit.get("messages")
        points: list[dict[str, Any]] = []
        if isinstance(messages, list):
            for idx, item in enumerate(messages[:3], start=1):
                if not isinstance(item, dict):
                    continue
                label = str(item.get("label") or f"Signal {idx}").strip()
                detail = str(
                    item.get("detail")
                    or item.get("body")
                    or item.get("text")
                    or item.get("description")
                    or ""
                ).strip()
                if not detail:
                    continue
                points.append(
                    {
                        "label": label,
                        "value": f"{idx:02d}",
                        "detail": detail,
                    }
                )
        if points:
            return points
        bullets = self._bullets(outline)
        return [
            {
                "label": f"Signal {idx}",
                "value": f"{idx:02d}",
                "detail": bullet,
            }
            for idx, bullet in enumerate(bullets[:3], start=1)
            if str(bullet).strip()
        ]

    def _add_two_column(self, slide, outline: SlideOutline, brand: BrandDNA) -> None:
        bullets = self._bullets(outline)
        left = bullets[:2] or ["Clarify the decision context and why it matters now."]
        right = bullets[2:5] or ["Focus leadership attention on the highest-value next moves."]
        icons = self._icons(outline)
        self._add_card(slide, 0.85, 1.55, 5.75, 4.75, "FFFFFF", brand.colors.background_light)
        self._add_card(slide, 6.9, 1.55, 5.75, 4.75, "FFFFFF", brand.colors.background_light)
        self._add_icon(slide, icons[0], 1.12, 1.74, 0.78, brand, brand.colors.primary)
        self._add_label(slide, "Implication", 2.06, 1.95, 4.2, brand, bold=True)
        self._add_bullets(slide, left, 1.15, 2.52, 4.95, 2.82, brand)
        self._add_icon(slide, icons[1], 7.17, 1.74, 0.78, brand, brand.colors.secondary)
        self._add_label(slide, "Evidence", 8.11, 1.95, 4.2, brand, bold=True)
        self._add_bullets(slide, right, 7.2, 2.52, 4.95, 2.82, brand)

    def _add_authored_composition(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
        family: str,
    ) -> bool:
        if family == "proof_strip":
            self._add_authored_proof_strip(slide, outline, brand)
            return True
        if family == "toolkit_grid":
            self._add_authored_toolkit_grid(slide, outline, brand)
            return True
        if family == "challenge_cards":
            self._add_authored_challenge_cards(slide, outline, brand)
            return True
        if family == "why_it_matters_cards":
            self._add_authored_why_matters(slide, outline, brand)
            return True
        if family == "circular_trap":
            self._add_anti_patterns(slide, outline, brand)
            return True
        if family == "reframe_split":
            self._add_quote_sidebar(slide, outline, brand)
            return True
        return False

    def _add_full_bleed_authored_composition(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
        family: str,
        slide_number: int,
        total_slides: int,
    ) -> bool:
        if family == "editorial_spread":
            self._add_authored_editorial_spread(
                slide, outline, brand, slide_number, total_slides
            )
            return True
        if family == "evidence_wall":
            self._add_authored_evidence_wall(
                slide, outline, brand, slide_number, total_slides
            )
            return True
        if family == "decision_ladder":
            self._add_authored_decision_ladder(
                slide, outline, brand, slide_number, total_slides
            )
            return True
        if family == "architecture_layers":
            self._add_authored_architecture_layers(
                slide, outline, brand, slide_number, total_slides
            )
            return True
        if family == "reframe_comparison":
            self._add_authored_reframe_comparison(
                slide, outline, brand, slide_number, total_slides
            )
            return True
        if family == "lifecycle_timeline":
            self._add_authored_lifecycle_timeline(
                slide, outline, brand, slide_number, total_slides
            )
            return True
        if family == "operating_map":
            self._add_authored_operating_map(
                slide, outline, brand, slide_number, total_slides
            )
            return True
        if family == "statement_canvas":
            self._add_authored_statement_canvas(
                slide, outline, brand, slide_number, total_slides
            )
            return True
        if family == "spotlight_quote":
            self._add_authored_spotlight_quote(
                slide, outline, brand, slide_number, total_slides
            )
            return True
        if family == "metric_signal":
            self._add_authored_metric_signal(
                slide, outline, brand, slide_number, total_slides
            )
            return True
        return False

    def _add_authored_canvas_chrome(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
        slide_number: int,
        total_slides: int,
        label: str,
        dark: bool = True,
    ) -> None:
        fill = brand.colors.background_dark if dark else brand.colors.background_light
        background = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0),
            Inches(0),
            Inches(SLIDE_W),
            Inches(SLIDE_H),
        )
        background.fill.solid()
        background.fill.fore_color.rgb = self._rgb(fill)
        background.line.color.rgb = background.fill.fore_color.rgb
        section_number, _ = self._section_marker(outline, slide_number)
        text_color = brand.colors.text_light if dark else brand.colors.text_dark
        secondary = (
            self._tint(brand.colors.primary, 0.74)
            if dark
            else brand.colors.secondary
        )
        self._add_dark_text(
            slide,
            f"SECTION {section_number} / {label}",
            0.72,
            0.34,
            4.2,
            0.2,
            brand,
            size=8,
            bold=True,
            color=brand.colors.accent,
        )
        self._add_dark_text(
            slide,
            f"{slide_number}/{total_slides}",
            11.72,
            0.34,
            0.74,
            0.2,
            brand,
            size=8,
            bold=True,
            color=secondary,
        )
        source = self._footer_source_text(outline)
        if source:
            self._add_dark_text(
                slide,
                self._truncate_at_word(source, 150),
                0.72,
                7.0,
                8.8,
                0.22,
                brand,
                size=7,
                color=secondary,
            )
        rail = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.72),
            Inches(0.68),
            Inches(1.0),
            Inches(0.035),
        )
        rail.fill.solid()
        rail.fill.fore_color.rgb = self._rgb(brand.colors.accent)
        rail.line.color.rgb = rail.fill.fore_color.rgb
        # Return values are not needed, but local names clarify the color intent
        # for the full-canvas renderers below.
        _ = text_color

    def _add_authored_editorial_spread(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
        slide_number: int,
        total_slides: int,
    ) -> None:
        self._add_authored_canvas_chrome(
            slide, outline, brand, slide_number, total_slides, "EXECUTIVE READ"
        )
        title = self._authored_title(outline)
        items = self._authored_items(outline, limit=3)
        thesis = self._authored_support_text(outline, items[0])
        self._add_dark_text(
            slide,
            self._truncate_at_word(title, 115),
            0.82,
            1.25,
            6.25,
            1.55,
            brand,
            size=28,
            bold=True,
            color=brand.colors.text_light,
        )
        self._add_dark_text(
            slide,
            self._truncate_phrase(thesis, 190),
            0.86,
            3.16,
            5.45,
            0.82,
            brand,
            size=15,
            color=self._tint(brand.colors.primary, 0.78),
        )
        accent_panel = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(7.34),
            Inches(1.16),
            Inches(4.55),
            Inches(5.18),
        )
        accent_panel.fill.solid()
        accent_panel.fill.fore_color.rgb = self._rgb(
            self._tint(brand.colors.primary, 0.13)
        )
        accent_panel.line.color.rgb = self._rgb(self._tint(brand.colors.primary, 0.32))
        for idx, item in enumerate(items[:4]):
            y = 1.58 + idx * 1.08
            self._add_dark_text(
                slide,
                f"0{idx + 1}",
                7.72,
                y,
                0.46,
                0.24,
                brand,
                size=10,
                bold=True,
                color=brand.colors.accent,
            )
            lead, rest = self._split_lead(item)
            self._add_dark_text(
                slide,
                lead,
                8.36,
                y - 0.02,
                2.78,
                0.24,
                brand,
                size=12,
                bold=True,
                color=brand.colors.text_light,
            )
            self._add_dark_text(
                slide,
                rest or item,
                8.36,
                y + 0.32,
                2.72,
                0.44,
                brand,
                size=10,
                color=self._tint(brand.colors.primary, 0.78),
            )

    def _add_authored_evidence_wall(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
        slide_number: int,
        total_slides: int,
    ) -> None:
        self._add_authored_canvas_chrome(
            slide, outline, brand, slide_number, total_slides, "EVIDENCE WALL", dark=False
        )
        title = self._authored_title(outline)
        items = self._evidence_wall_items(outline)
        rail = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.72),
            Inches(1.25),
            Inches(3.38),
            Inches(5.36),
        )
        rail.fill.solid()
        rail.fill.fore_color.rgb = self._rgb(brand.colors.primary)
        rail.line.color.rgb = rail.fill.fore_color.rgb
        self._add_dark_text(
            slide,
            "SOURCE READ",
            1.06,
            1.66,
            1.8,
            0.22,
            brand,
            size=8,
            bold=True,
            color=brand.colors.accent,
        )
        self._add_dark_text(
            slide,
            self._truncate_at_word(title, 96),
            1.06,
            2.14,
            2.86,
            1.34,
            brand,
            size=19,
            bold=True,
            color=brand.colors.text_light,
        )
        self._add_dark_text(
            slide,
            self._truncate_phrase(
                self._authored_distinct_support_text(outline, items, items[0]),
                150,
            ),
            1.08,
            4.24,
            2.34,
            0.74,
            brand,
            size=11,
            color=self._tint(brand.colors.primary, 0.76),
        )
        cells = [
            (4.58, 1.32, 3.05, 1.48),
            (8.0, 1.62, 3.66, 1.08),
            (4.98, 3.34, 2.45, 1.88),
            (7.78, 3.02, 3.88, 1.58),
            (8.22, 5.02, 2.92, 1.2),
        ]
        used_details: list[str] = []
        for idx, item in enumerate(items[:5]):
            x, y, w, h = cells[idx]
            accent = self._icon_fill(brand, idx)
            fill = "FFFFFF" if idx % 2 == 0 else self._tint(accent, 0.9)
            self._add_card(slide, x, y, w, h, fill, self._tint(accent, 0.72))
            self._add_dark_text(
                slide,
                f"{idx + 1}",
                x + 0.24,
                y + 0.22,
                0.34,
                0.2,
                brand,
                size=9,
                bold=True,
                color=accent,
            )
            lead, rest = self._split_lead(item)
            self._add_body_text(
                slide,
                self._safe_authored_heading(outline, lead, rest, idx, 28),
                x + 0.74,
                y + 0.22,
                w - 1.02,
                0.4,
                brand,
                size=9,
            )
            detail = self._authored_detail_text(
                outline,
                item,
                idx,
                rest,
                avoid=[*items, *used_details],
            )
            if detail:
                used_details.append(detail)
            detail_size = 8 if h < 1.35 else 9
            self._add_body_text(
                slide,
                detail,
                x + 0.28,
                y + 0.62,
                w - 0.56,
                max(0.52, h - 0.82),
                brand,
                size=detail_size,
            )

    def _evidence_wall_items(self, outline: SlideOutline) -> list[str]:
        source_items = self._unique_display_items(self._grid_items(outline), limit=5)
        items = [
            item
            for item in source_items
            if self._evidence_wall_item_has_detail(item)
        ]
        if len(items) < 3:
            items = [
                item
                for item in self._authored_items(outline, limit=5)
                if self._evidence_wall_item_has_detail(item)
            ]
        if len(items) < 3:
            items = self._benchmark_evidence_wall_defaults(outline)
        return self._unique_display_items(items, limit=4)

    def _evidence_wall_item_has_detail(self, item: str) -> bool:
        lead, rest = self._split_lead(item)
        if not lead or not rest:
            return False
        if self._is_incomplete_display_fragment(lead) or self._is_incomplete_display_fragment(rest):
            return False
        return len(rest.split()) >= 4

    def _benchmark_evidence_wall_defaults(self, outline: SlideOutline) -> list[str]:
        text = " ".join(
            [
                self._authored_title(outline),
                str(outline.content_json.get("subheading") or ""),
                str(outline.content_json.get("summary") or ""),
                str(outline.content_json.get("speaker_notes") or ""),
                " ".join(str(item) for item in outline.content_json.get("sources", []) if str(item).strip()),
            ]
        ).casefold()
        if any(token in text for token in ("benchmark", "harness", "ground truth", "evaluation")):
            return [
                "Validated evidence anchors each benchmark case.",
                "Agents discover candidate cases from real workflows.",
                "Harness reviews keep failures inspectable before scaling.",
            ]
        return [
            "Source evidence defines the operating signal.",
            "Review gates turn the signal into a managed decision.",
            "Owners update the workflow when assumptions change.",
        ]

    def _add_authored_reframe_comparison(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
        slide_number: int,
        total_slides: int,
    ) -> None:
        self._add_authored_canvas_chrome(
            slide, outline, brand, slide_number, total_slides, "REFRAME"
        )
        exhibit = self._exhibit(outline)
        columns, rows = self._comparison_rows(exhibit, outline)
        rows = rows[:3] or [[item, item, "Managed move"] for item in self._authored_items(outline, 3)]
        title = self._authored_title(outline)
        left_header = str(columns[1] if len(columns) > 1 else "Current readout")
        right_header = str(columns[2] if len(columns) > 2 else "Target move")
        self._add_dark_text(
            slide,
            self._truncate_at_word(title, 92),
            0.82,
            1.16,
            5.48,
            0.98,
            brand,
            size=24,
            bold=True,
            color=brand.colors.text_light,
        )
        self._add_dark_text(
            slide,
            self._truncate_phrase(
                self._comparison_panel_insight(columns, rows, outline),
                135,
            ),
            0.86,
            2.46,
            4.6,
            0.58,
            brand,
            size=11,
            color=self._tint(brand.colors.primary, 0.76),
        )
        left = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.9),
            Inches(3.42),
            Inches(4.38),
            Inches(2.78),
        )
        left.fill.solid()
        left.fill.fore_color.rgb = self._rgb(self._tint(brand.colors.primary, 0.13))
        left.line.color.rgb = self._rgb(self._tint(brand.colors.primary, 0.32))
        right = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(7.92),
            Inches(1.42),
            Inches(4.26),
            Inches(4.78),
        )
        right.fill.solid()
        right.fill.fore_color.rgb = self._rgb("FFFFFF")
        right.line.color.rgb = self._rgb(self._tint(brand.colors.accent, 0.62))
        self._add_dark_text(
            slide,
            left_header.upper(),
            1.22,
            3.74,
            2.6,
            0.22,
            brand,
            size=8,
            bold=True,
            color=brand.colors.accent,
        )
        for idx, row in enumerate(rows):
            current_text = self._truncate_at_word((row + [""])[1], 64)
            self._add_dark_text(
                slide,
                current_text,
                1.24,
                4.14 + idx * 0.48,
                3.46,
                0.26,
                brand,
                size=9,
                color=self._tint(brand.colors.primary, 0.84),
            )
        self._add_dark_text(
            slide,
            "SHIFT",
            5.94,
            3.34,
            0.82,
            0.22,
            brand,
            size=8,
            bold=True,
            color=brand.colors.accent,
        )
        for idx, row in enumerate(rows):
            y = 3.86 + idx * 0.58
            self._add_arrow(slide, 5.26, y, 7.78, y - 1.36 + idx * 0.8, brand.colors.accent, width=1.6)
            self._add_arrowhead(slide, 7.78, y - 1.36 + idx * 0.8, 0.2, 0, brand.colors.accent)
        self._add_body_text(
            slide,
            right_header.upper(),
            8.24,
            1.82,
            2.56,
            0.2,
            brand,
            size=8,
        )
        for idx, row in enumerate(rows):
            cells = (row + ["", ""])[:3]
            y = 2.34 + idx * 1.08
            self._add_badge(
                slide,
                str(idx + 1),
                8.22,
                y,
                0.42,
                brand,
                self._icon_fill(brand, idx),
            )
            self._add_body_text(
                slide,
                self._truncate_at_word(cells[2] or cells[0], 86),
                8.82,
                y,
                2.64,
                0.44,
                brand,
                size=10,
            )

    def _add_authored_decision_ladder(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
        slide_number: int,
        total_slides: int,
    ) -> None:
        self._add_authored_canvas_chrome(
            slide, outline, brand, slide_number, total_slides, "DECISION LADDER", dark=False
        )
        title = self._authored_title(outline)
        items = self._authored_items(outline, limit=4)
        if len(items) < 3:
            for filler in self._authored_filler_items(outline):
                if len(items) >= 3:
                    break
                if filler not in items and not self._authored_text_overlaps(filler, title):
                    items.append(filler)
        decision_support = self._authored_distinct_support_text(
            outline,
            items,
            self._topic_support_sentence(outline),
        )
        if self._authored_text_overlaps(decision_support, title):
            decision_support = ""
        self._add_dark_text(
            slide,
            self._truncate_at_word(title, 98),
            0.82,
            1.12,
            5.6,
            0.9,
            brand,
            size=23,
            bold=True,
            color=brand.colors.text_dark,
        )
        self._add_body_text(
            slide,
            decision_support,
            0.88,
            2.28,
            4.6,
            0.74,
            brand,
            size=12,
        )
        left_rule = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.86),
            Inches(3.48),
            Inches(4.2),
            Inches(0.04),
        )
        left_rule.fill.solid()
        left_rule.fill.fore_color.rgb = self._rgb(brand.colors.accent)
        left_rule.line.color.rgb = left_rule.fill.fore_color.rgb
        self._add_dark_text(
            slide,
            "DECISION LOGIC",
            0.88,
            3.72,
            1.75,
            0.2,
            brand,
            size=8,
            bold=True,
            color=brand.colors.accent,
        )
        self._add_body_text(
            slide,
            self._truncate_at_word(items[0], 110) if items else "",
            0.88,
            4.08,
            4.14,
            0.92,
            brand,
            size=14,
        )
        used_details: list[str] = []
        for idx, item in enumerate(items[:4]):
            x = 6.18
            y = 2.16 + idx * 1.02
            accent = self._icon_fill(brand, idx)
            band = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(x),
                Inches(y),
                Inches(5.78),
                Inches(0.88),
            )
            band.fill.solid()
            band.fill.fore_color.rgb = self._rgb("FFFFFF")
            band.line.color.rgb = self._rgb(self._tint(accent, 0.72))
            stripe = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(x),
                Inches(y),
                Inches(0.08),
                Inches(0.88),
            )
            stripe.fill.solid()
            stripe.fill.fore_color.rgb = self._rgb(accent)
            stripe.line.color.rgb = stripe.fill.fore_color.rgb
            self._add_badge(slide, str(idx + 1), x + 0.28, y + 0.2, 0.44, brand, accent)
            lead, rest = self._split_lead(item)
            if self._authored_heading_reads_like_fragment(lead, rest):
                lead = self._fallback_heading_from_title(outline, idx)
                rest = item
            self._add_dark_text(
                slide,
                self._safe_authored_heading(outline, lead, rest, idx, 52),
                x + 0.9,
                y + 0.16,
                2.08,
                0.28,
                brand,
                size=12,
                bold=True,
                color=brand.colors.primary,
            )
            detail = self._authored_detail_text(
                outline,
                item,
                idx,
                rest,
                avoid=[*items, *used_details],
            )
            if detail:
                used_details.append(detail)
                self._add_body_text(
                    slide,
                    self._truncate_at_word(detail, 88),
                    x + 3.04,
                    y + 0.15,
                    2.18,
                    0.42,
                    brand,
                    size=11,
                )

    def _add_authored_architecture_layers(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
        slide_number: int,
        total_slides: int,
    ) -> None:
        self._add_authored_canvas_chrome(
            slide, outline, brand, slide_number, total_slides, "ARCHITECTURE"
        )
        title = self._authored_title(outline)
        rows = self._architecture_layer_rows(outline)
        title_w = 5.25
        title_size = self._fit_font_size(title, title_w, [24, 22, 20, 18, 16], max_lines=3)
        self._add_dark_text(
            slide,
            self._truncate_at_word(title, 98),
            0.82,
            1.12,
            title_w,
            1.24,
            brand,
            size=title_size,
            bold=True,
            color=brand.colors.text_light,
        )
        self._add_dark_text(
            slide,
            self._authored_distinct_support_text(
                outline,
                [item for row in rows for item in row],
                (self._authored_filler_items(outline) or [rows[0][1]])[0],
            ),
            0.86,
            2.78,
            4.58,
            0.74,
            brand,
            size=12,
            color=self._tint(brand.colors.primary, 0.76),
        )
        layer_field = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(6.46),
            Inches(1.54),
            Inches(5.14),
            Inches(4.48),
        )
        layer_field.fill.solid()
        layer_field.fill.fore_color.rgb = self._rgb(self._tint(brand.colors.primary, 0.12))
        layer_field.line.color.rgb = layer_field.fill.fore_color.rgb
        guide = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(6.66),
            Inches(1.82),
            Inches(0.04),
            Inches(3.7),
        )
        guide.fill.solid()
        guide.fill.fore_color.rgb = self._rgb(brand.colors.accent)
        guide.line.color.rgb = guide.fill.fore_color.rgb
        for idx, (label, detail) in enumerate(rows[:4]):
            x = 6.88 + idx * 0.08
            y = 1.9 + idx * 0.84
            w = 4.42 - idx * 0.12
            safe_detail = self._complete_display_item(detail, limit=112)
            if not safe_detail:
                safe_detail = self._architecture_detail_fallback(label, outline, idx)
            divider = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(x),
                Inches(y + 0.66),
                Inches(w),
                Inches(0.018),
            )
            divider.fill.solid()
            divider.fill.fore_color.rgb = self._rgb(self._tint(brand.colors.primary, 0.36))
            divider.line.color.rgb = divider.fill.fore_color.rgb
            self._add_dark_text(
                slide,
                self._truncate_at_word(label, 30),
                x + 0.02,
                y + 0.09,
                1.52,
                0.26,
                brand,
                size=11,
                bold=True,
                color=brand.colors.text_light,
            )
            self._add_dark_text(
                slide,
                safe_detail,
                x + 1.72,
                y + 0.08,
                w - 1.86,
                0.52,
                brand,
                size=10,
                color=self._tint(brand.colors.primary, 0.9),
            )
        bottom_note = self._authored_distinct_support_text(
            outline,
            [item for row in rows for item in row],
            self._topic_support_sentence(outline),
        )
        if bottom_note and not self._authored_text_overlaps(bottom_note, title):
            self._add_dark_text(
                slide,
                self._truncate_at_word(bottom_note, 92),
                6.68,
                6.12,
                4.5,
                0.42,
                brand,
                size=11,
                color=self._tint(brand.colors.primary, 0.76),
            )

    def _architecture_detail_fallback(
        self, label: str, outline: SlideOutline, index: int
    ) -> str:
        text = " ".join([label, self._authored_title(outline)]).lower()
        if "contract" in text:
            return [
                "Make success criteria explicit before execution.",
                "Define what must be tested by the harness.",
                "Turn operational questions into valid benchmark cases.",
            ][index % 3]
        if "harness" in text or "architecture" in text or "layer" in text:
            return [
                "Assign the handoff from contract to evidence.",
                "Preserve run evidence for review and replay.",
                "Connect execution back to the source standard.",
            ][index % 3]
        return [
            "Anchor the layer in source-backed evidence.",
            "Name the review rule before scaling.",
            "Keep the benchmark update path explicit.",
        ][index % 3]

    def _add_authored_lifecycle_timeline(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
        slide_number: int,
        total_slides: int,
    ) -> None:
        self._add_authored_canvas_chrome(
            slide, outline, brand, slide_number, total_slides, "LIFECYCLE", dark=False
        )
        title = self._authored_title(outline)
        items = self._authored_items(outline, limit=4)
        title_w = 6.4
        title_size = self._fit_font_size(title, title_w, [23, 21, 19, 17], max_lines=3)
        self._add_dark_text(
            slide,
            self._truncate_at_word(title, 104),
            0.82,
            1.04,
            title_w,
            1.05,
            brand,
            size=title_size,
            bold=True,
            color=brand.colors.text_dark,
        )
        y = 3.62
        self._add_arrow(slide, 1.18, y, 11.64, y, brand.colors.primary, width=1.6)
        self._add_arrowhead(slide, 11.64, y, 0.3, 0, brand.colors.primary)
        timeline_items = items[:4]
        step_gap = 10.18 / max(len(timeline_items) - 1, 1)
        used_details: list[str] = []
        for idx, item in enumerate(timeline_items):
            x = 1.16 + idx * step_gap
            accent = self._icon_fill(brand, idx)
            marker = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                Inches(x),
                Inches(y - 0.34),
                Inches(0.68),
                Inches(0.68),
            )
            marker.fill.solid()
            marker.fill.fore_color.rgb = self._rgb(accent)
            marker.line.color.rgb = marker.fill.fore_color.rgb
            self._add_dark_text(
                slide,
                str(idx + 1),
                x + 0.21,
                y - 0.16,
                0.24,
                0.2,
                brand,
                size=8,
                bold=True,
                color=self._readable_text_color(accent, brand),
            )
            lead, rest = self._split_lead(item)
            text_y = 2.2 if idx % 2 == 0 else 4.34
            self._add_body_text(slide, lead, x - 0.66, text_y, 2.28, 0.46, brand, size=11)
            detail = self._authored_detail_text(
                outline,
                item,
                idx,
                rest,
                avoid=[*items, *used_details],
            )
            if detail:
                used_details.append(detail)
            self._add_body_text(
                slide,
                self._truncate_at_word(detail, 58),
                x - 0.66,
                text_y + 0.52,
                2.28,
                0.56,
                brand,
                size=10,
            )
        self._add_body_text(
            slide,
            self._authored_distinct_support_text(outline, items, items[0]),
            7.62,
            1.12,
            3.82,
            0.64,
            brand,
            size=12,
        )

    def _add_authored_operating_map(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
        slide_number: int,
        total_slides: int,
    ) -> None:
        self._add_authored_canvas_chrome(
            slide, outline, brand, slide_number, total_slides, "OPERATING MAP"
        )
        title = self._authored_title(outline)
        items = self._authored_items(outline, limit=4)
        title_size = self._fit_font_size(title, 4.62, [19, 17, 15], max_lines=4)
        self._add_dark_text(
            slide,
            self._truncate_at_word(title, 108),
            0.82,
            1.06,
            4.62,
            1.38,
            brand,
            size=title_size,
            bold=True,
            color=brand.colors.text_light,
        )
        support = self._authored_distinct_support_text(outline, items, items[0])
        self._add_dark_text(
            slide,
            self._truncate_phrase(support, 150),
            0.86,
            2.72,
            4.36,
            0.72,
            brand,
            size=12,
            color=self._tint(brand.colors.primary, 0.76),
        )
        map_panel = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(5.72),
            Inches(1.08),
            Inches(6.12),
            Inches(5.42),
        )
        map_panel.fill.solid()
        map_panel.fill.fore_color.rgb = self._rgb(self._tint(brand.colors.primary, 0.12))
        map_panel.line.color.rgb = self._rgb(self._tint(brand.colors.primary, 0.28))
        if items:
            self._add_dark_text(
                slide,
                "OPERATING SHIFT",
                1.0,
                4.54,
                1.95,
                0.18,
                brand,
                size=8,
                bold=True,
                color=brand.colors.accent,
            )
            self._add_dark_text(
                slide,
                self._truncate_phrase(items[0], 118),
                1.0,
                4.92,
                3.9,
                0.62,
                brand,
                size=12,
                color=self._tint(brand.colors.primary, 0.82),
            )
        node_positions = [
            (6.18, 1.54, 2.18, 1.02),
            (9.02, 1.94, 2.18, 1.02),
            (6.62, 3.68, 2.18, 1.02),
            (9.36, 4.34, 2.0, 1.2),
        ]
        labels = ["FRAME", "TEST", "REVIEW", "COMMIT"]
        used_details: list[str] = []
        active_positions = node_positions[: max(1, min(len(items), len(node_positions)))]
        for idx, (x, y, w, h) in enumerate(active_positions):
            item = items[idx]
            accent = self._icon_fill(brand, idx)
            fill = self._tint(accent, 0.88) if idx in {1, 3} else "FFFFFF"
            node = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(x),
                Inches(y),
                Inches(w),
                Inches(h),
            )
            node.fill.solid()
            node.fill.fore_color.rgb = self._rgb(fill)
            node.line.color.rgb = self._rgb(self._tint(accent, 0.46))
            self._add_dark_text(
                slide,
                labels[idx],
                x + 0.2,
                y + 0.18,
                0.82,
                0.16,
                brand,
                size=7,
                bold=True,
                color=accent,
            )
            lead, rest = self._split_lead(item)
            if self._authored_heading_reads_like_fragment(lead, rest):
                lead = self._fallback_heading_from_title(outline, idx)
                rest = item
            self._add_dark_text(
                slide,
                self._safe_authored_heading(outline, lead, rest, idx, 40),
                x + 0.2,
                y + 0.46,
                w - 0.4,
                0.26,
                brand,
                size=10,
                bold=True,
                color=brand.colors.primary,
            )
            detail = self._authored_detail_text(
                outline,
                item,
                idx,
                rest,
                avoid=[*items, *used_details],
            )
            if detail:
                used_details.append(detail)
                detail = self._complete_display_item(detail, limit=62)
            if detail:
                self._add_dark_text(
                    slide,
                    detail,
                    x + 0.2,
                    y + 0.74,
                    w - 0.42,
                    0.24,
                    brand,
                    size=8,
                    color=brand.colors.primary,
                )
        connectors = [
            ((8.36, 2.05), (9.02, 2.34)),
            ((10.08, 2.96), (7.9, 3.68)),
            ((8.8, 4.22), (9.36, 4.92)),
        ]
        for start, end in connectors[: max(0, len(active_positions) - 1)]:
            self._add_arrow(
                slide,
                start[0],
                start[1],
                end[0],
                end[1],
                brand.colors.accent,
                width=1.3,
            )
            self._add_arrowhead(
                slide,
                end[0],
                end[1],
                end[0] - start[0],
                end[1] - start[1],
                brand.colors.accent,
            )

    def _add_authored_statement_canvas(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
        slide_number: int,
        total_slides: int,
    ) -> None:
        self._add_authored_canvas_chrome(
            slide, outline, brand, slide_number, total_slides, "DECISION POINT"
        )
        title = self._authored_title(outline)
        items = self._authored_items(outline, limit=3)
        title_size = self._fit_font_size(title, 6.05, [30, 27, 24, 21], max_lines=4)
        self._add_dark_text(
            slide,
            self._truncate_at_word(title, 120),
            0.82,
            1.08,
            6.05,
            2.02,
            brand,
            size=title_size,
            bold=True,
            color=brand.colors.text_light,
        )
        support = self._authored_distinct_support_text(outline, items, "")
        if support:
            self._add_dark_text(
                slide,
                self._truncate_phrase(support, 150),
                0.88,
                3.58,
                4.9,
                0.62,
                brand,
                size=13,
                color=self._tint(brand.colors.primary, 0.78),
            )
        accent_bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.88),
            Inches(4.62),
            Inches(4.92),
            Inches(0.12),
        )
        accent_bar.fill.solid()
        accent_bar.fill.fore_color.rgb = self._rgb(brand.colors.accent)
        accent_bar.line.color.rgb = accent_bar.fill.fore_color.rgb
        self._add_dark_text(
            slide,
            self._truncate_at_word(self._source_label(outline), 86),
            0.9,
            5.04,
            4.4,
            0.28,
            brand,
            size=8,
            bold=True,
            color=self._tint(brand.colors.primary, 0.68),
        )
        field = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(7.0),
            Inches(1.18),
            Inches(4.82),
            Inches(5.26),
        )
        field.fill.solid()
        field.fill.fore_color.rgb = self._rgb(self._tint(brand.colors.primary, 0.13))
        field.line.color.rgb = self._rgb(self._tint(brand.colors.primary, 0.34))
        used_details: list[str] = []
        for idx, item in enumerate(items[:3]):
            y = 1.72 + idx * 1.48
            accent = self._icon_fill(brand, idx)
            rule = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(7.42),
                Inches(y - 0.2),
                Inches(3.76),
                Inches(0.025),
            )
            rule.fill.solid()
            rule.fill.fore_color.rgb = self._rgb(self._tint(accent, 0.48))
            rule.line.color.rgb = rule.fill.fore_color.rgb
            self._add_dark_text(
                slide,
                f"0{idx + 1}",
                7.42,
                y,
                0.42,
                0.2,
                brand,
                size=9,
                bold=True,
                color=accent,
            )
            lead, rest = self._split_lead(item)
            self._add_dark_text(
                slide,
                lead,
                8.08,
                y - 0.02,
                3.24,
                0.48,
                brand,
                size=13,
                bold=True,
                color=brand.colors.text_light,
            )
            detail = self._authored_detail_text(
                outline,
                item,
                idx,
                rest,
                avoid=[support, *used_details],
            )
            if detail:
                used_details.append(detail)
                self._add_dark_text(
                    slide,
                    detail,
                    8.08,
                    y + 0.52,
                    3.12,
                    0.62,
                    brand,
                    size=10,
                    color=self._tint(brand.colors.primary, 0.82),
                )

    def _add_authored_spotlight_quote(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
        slide_number: int,
        total_slides: int,
    ) -> None:
        self._add_authored_canvas_chrome(
            slide, outline, brand, slide_number, total_slides, "CALLOUT"
        )
        title = self._authored_title(outline)
        items = self._authored_items(outline, limit=3)
        if len(items) < 3:
            for fallback in self._spotlight_callout_fallback_items(outline):
                if len(items) >= 3:
                    break
                if fallback in items:
                    continue
                if any(self._authored_text_overlaps(fallback, item) for item in items):
                    continue
                items.append(fallback)
        support = self._authored_distinct_support_text(outline, items, "")
        spotlight = self._truncate_at_word(title, 118)
        title_size = self._fit_font_size(spotlight, 7.1, [34, 31, 28, 25], max_lines=4)
        band = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0),
            Inches(1.04),
            Inches(0.18),
            Inches(5.36),
        )
        band.fill.solid()
        band.fill.fore_color.rgb = self._rgb(brand.colors.accent)
        band.line.color.rgb = band.fill.fore_color.rgb
        self._add_dark_text(
            slide,
            spotlight,
            0.82,
            1.24,
            7.1,
            1.82,
            brand,
            size=title_size,
            bold=True,
            color=brand.colors.text_light,
        )
        if support:
            self._add_dark_text(
                slide,
                self._truncate_phrase(support, 156),
                0.88,
                3.38,
                5.9,
                0.66,
                brand,
                size=14,
                color=self._tint(brand.colors.primary, 0.78),
            )
        self._add_dark_text(
            slide,
            "WHAT CHANGES",
            8.2,
            1.25,
            2.4,
            0.22,
            brand,
            size=8,
            bold=True,
            color=brand.colors.accent,
        )
        used_details: list[str] = [support]
        for idx, item in enumerate(items[:3]):
            y = 1.74 + idx * 1.42
            lead, rest = self._split_lead(item)
            if self._authored_heading_reads_like_fragment(lead, rest):
                lead = self._fallback_heading_from_title(outline, idx)
            self._add_dark_text(
                slide,
                lead,
                8.18,
                y,
                3.18,
                0.32,
                brand,
                size=13,
                bold=True,
                color=brand.colors.text_light,
            )
            detail = self._authored_detail_text(
                outline,
                item,
                idx,
                rest,
                avoid=used_details,
            )
            if detail:
                used_details.append(detail)
                self._add_dark_text(
                    slide,
                    detail,
                    8.2,
                    y + 0.42,
                    3.1,
                    0.42,
                    brand,
                    size=10,
                    color=self._tint(brand.colors.primary, 0.82),
                )
            rule = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(8.18),
                Inches(y + 1.02),
                Inches(3.02),
                Inches(0.02),
            )
            rule.fill.solid()
            rule.fill.fore_color.rgb = self._rgb(self._tint(brand.colors.primary, 0.42))
            rule.line.color.rgb = rule.fill.fore_color.rgb
        source = self._source_label(outline)
        if source:
            self._add_dark_text(
                slide,
                self._truncate_at_word(source, 88),
                0.9,
                5.78,
                5.6,
                0.28,
                brand,
                size=8,
                bold=True,
                color=self._tint(brand.colors.primary, 0.68),
            )

    def _spotlight_callout_fallback_items(self, outline: SlideOutline) -> list[str]:
        context = " ".join(
            [
                self._authored_title(outline),
                str(outline.content_json.get("subheading") or ""),
                self._source_label(outline),
            ]
        ).lower()
        if "executive summary" in context or "benchmark" in context:
            return [
                "Benchmark saturation: Public tests are crowded and weakly tied to enterprise workflows.",
                "Implicit ground truth: Existing documents and decisions can become benchmark evidence.",
                "Harness-centric discovery: Agents convert operating evidence into reusable tests.",
            ]
        if "contract" in context:
            return [
                "Explicit criteria: Inputs, outputs, and success rules are defined before execution.",
                "Shared standard: Humans and agents inspect the same evaluation contract.",
                "Valid cases: Operational questions become benchmark requirements.",
            ]
        return self._authored_filler_items(outline)

    def _add_authored_metric_signal(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
        slide_number: int,
        total_slides: int,
    ) -> None:
        self._add_authored_canvas_chrome(
            slide, outline, brand, slide_number, total_slides, "NUMBER SIGNAL"
        )
        metrics = self._authored_quantitative_metrics(outline)
        if not metrics:
            self._add_authored_spotlight_quote(slide, outline, brand, slide_number, total_slides)
            return
        primary = metrics[0]
        value = self._format_metric_value(primary)
        label = self._metric_label_text(primary.get("label", "Metric"))
        title = self._authored_title(outline)
        insight = self._metric_insight(outline, metrics)
        accent_panel = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.78),
            Inches(1.18),
            Inches(5.0),
            Inches(4.84),
        )
        accent_panel.fill.solid()
        accent_panel.fill.fore_color.rgb = self._rgb(self._tint(brand.colors.primary, 0.13))
        accent_panel.line.color.rgb = self._rgb(self._tint(brand.colors.primary, 0.32))
        value_size = self._fit_font_size(value, 4.32, [66, 58, 50, 42], max_lines=1)
        self._add_dark_text(
            slide,
            value,
            1.1,
            1.72,
            4.32,
            0.98,
            brand,
            size=value_size,
            bold=True,
            color=brand.colors.text_light,
        )
        self._add_dark_text(
            slide,
            label,
            1.16,
            2.86,
            3.84,
            0.36,
            brand,
            size=14,
            bold=True,
            color=brand.colors.accent,
        )
        self._add_dark_text(
            slide,
            self._truncate_phrase(insight, 150),
            1.16,
            3.44,
            3.96,
            0.72,
            brand,
            size=12,
            color=self._tint(brand.colors.primary, 0.8),
        )
        self._add_dark_text(
            slide,
            self._truncate_at_word(title, 112),
            6.36,
            1.3,
            5.08,
            1.42,
            brand,
            size=self._fit_font_size(title, 5.08, [27, 24, 21], max_lines=3),
            bold=True,
            color=brand.colors.text_light,
        )
        self._add_dark_text(
            slide,
            "SUPPORTING SIGNALS",
            6.42,
            3.18,
            2.6,
            0.2,
            brand,
            size=8,
            bold=True,
            color=brand.colors.accent,
        )
        for idx, metric in enumerate(metrics[1:4]):
            y = 3.62 + idx * 0.82
            mini_value = self._format_metric_value(metric)
            mini_label = self._metric_label_text(metric.get("label", "Metric"))
            self._add_dark_text(
                slide,
                mini_value,
                6.42,
                y,
                1.12,
                0.28,
                brand,
                size=15,
                bold=True,
                color=brand.colors.text_light,
            )
            self._add_dark_text(
                slide,
                mini_label,
                7.74,
                y + 0.03,
                3.28,
                0.24,
                brand,
                size=10,
                color=self._tint(brand.colors.primary, 0.8),
            )
            rule = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(6.42),
                Inches(y + 0.48),
                Inches(4.36),
                Inches(0.018),
            )
            rule.fill.solid()
            rule.fill.fore_color.rgb = self._rgb(self._tint(brand.colors.primary, 0.4))
            rule.line.color.rgb = rule.fill.fore_color.rgb
        source = self._source_label(outline)
        if source:
            self._add_dark_text(
                slide,
                self._truncate_at_word(source, 92),
                6.42,
                6.1,
                4.7,
                0.24,
                brand,
                size=8,
                bold=True,
                color=self._tint(brand.colors.primary, 0.68),
            )

    def _architecture_layer_rows(self, outline: SlideOutline) -> list[tuple[str, str]]:
        exhibit = self._exhibit(outline)
        rows = exhibit.get("rows")
        layer_rows: list[tuple[str, str]] = []
        if isinstance(rows, list):
            for row in rows:
                if isinstance(row, list):
                    cells = [
                        self._complete_display_item(str(cell))
                        for cell in row
                        if str(cell).strip()
                    ]
                    cells = [cell for cell in cells if cell]
                    if cells:
                        label = cells[0]
                        detail = cells[1] if len(cells) > 1 else cells[0]
                        label = self._architecture_label_text(label, outline, len(layer_rows))
                        detail = self._architecture_detail_text(detail, outline, len(layer_rows))
                        if self._authored_text_overlaps(label, detail):
                            detail = ""
                        layer_rows.append((label, detail))
                elif isinstance(row, dict):
                    label = self._complete_display_item(
                        str(row.get("label") or row.get("artifact") or row.get("name") or "")
                    )
                    detail = self._complete_display_item(
                        str(row.get("purpose") or row.get("description") or row.get("text") or "")
                    )
                    if label or detail:
                        label = self._architecture_label_text(
                            label or "Layer", outline, len(layer_rows)
                        )
                        detail = self._architecture_detail_text(
                            detail, outline, len(layer_rows)
                        )
                        if label and detail and self._authored_text_overlaps(label, detail):
                            detail = ""
                        layer_rows.append((label, detail))
        if not layer_rows:
            for item in self._authored_items(outline, limit=5):
                lead, rest = self._split_lead(item)
                lead = self._architecture_label_text(lead or "Layer", outline, len(layer_rows))
                detail = rest if rest and not self._authored_text_overlaps(lead, rest) else ""
                detail = self._architecture_detail_text(detail, outline, len(layer_rows))
                layer_rows.append((lead or "Layer", detail))
        return layer_rows

    def _architecture_label_text(
        self, label: str, outline: SlideOutline, index: int
    ) -> str:
        cleaned = self._truncate_at_word(str(label or "Layer"), 34)
        normalized = " ".join(cleaned.casefold().split())
        if re.search(
            r"\b(?:validation risk|evidence gap|operating choice|operating move|"
            r"reliability risk)\b",
            normalized,
        ):
            return self._architecture_label_fallback(outline, index)
        return cleaned

    def _architecture_detail_text(
        self, detail: str, outline: SlideOutline, index: int
    ) -> str:
        cleaned = self._complete_display_item(str(detail or ""), limit=112)
        if not cleaned or self._is_renderer_filler_copy(cleaned):
            return self._architecture_detail_fallback("", outline, index)
        return cleaned

    def _architecture_label_fallback(self, outline: SlideOutline, index: int) -> str:
        text = " ".join([self._authored_title(outline), self._source_label(outline)]).lower()
        if "harness" in text or "benchmark" in text:
            return [
                "Synthetic benchmark loop",
                "Harness reasoning shift",
                "Coverage mismatch",
                "Source evidence path",
            ][index % 4]
        return ["Layer input", "Layer handoff", "Layer output", "Layer review"][index % 4]

    def _add_authored_proof_strip(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
    ) -> None:
        source_items = self._grid_items(outline)
        items = source_items[:4] if len(source_items) >= 3 else self._authored_items(outline, limit=4)
        panel = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.86),
            Inches(1.56),
            Inches(3.35),
            Inches(4.78),
        )
        panel.fill.solid()
        panel.fill.fore_color.rgb = self._rgb(brand.colors.primary)
        panel.line.color.rgb = self._rgb(brand.colors.primary)
        self._add_dark_text(
            slide,
            self._authored_proof_marker_label(outline),
            1.16,
            1.92,
            1.75,
            0.22,
            brand,
            size=9,
            bold=True,
            color=brand.colors.accent,
        )
        self._add_dark_text(
            slide,
            f"{len(items)} source signals",
            1.16,
            2.42,
            2.35,
            0.5,
            brand,
            size=22,
            bold=True,
            color=brand.colors.text_light,
        )
        support = self._authored_distinct_support_text(outline, items, "")
        if support:
            self._add_dark_text(
                slide,
                support,
                1.18,
                3.25,
                2.45,
                1.28,
                brand,
                size=12,
                color=self._tint(brand.colors.primary, 0.78),
            )
        source_label = self._source_label(outline)
        if source_label:
            self._add_dark_text(
                slide,
                self._truncate_phrase(source_label, 72),
                1.18,
                5.54,
                2.45,
                0.24,
                brand,
                size=7,
                color=self._tint(brand.colors.primary, 0.66),
            )
        for idx, item in enumerate(items[:4]):
            y = 1.58 + idx * 1.12
            accent = self._icon_fill(brand, idx)
            fill = "FFFFFF" if idx % 2 == 0 else self._tint(brand.colors.secondary, 0.92)
            self._add_card(slide, 4.72, y, 7.62, 0.88, fill, brand.colors.background_light)
            marker = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(4.72),
                Inches(y),
                Inches(0.1),
                Inches(0.88),
            )
            marker.fill.solid()
            marker.fill.fore_color.rgb = self._rgb(accent)
            marker.line.color.rgb = marker.fill.fore_color.rgb
            self._add_dark_text(
                slide,
                f"{idx + 1:02d}",
                5.03,
                y + 0.28,
                0.46,
                0.2,
                brand,
                size=8,
                bold=True,
                color=accent,
            )
            self._add_body_text(
                slide,
                item,
                5.72,
                y + 0.18,
                5.82,
                0.36,
                brand,
                size=12,
            )

    def _authored_proof_marker_label(self, outline: SlideOutline) -> str:
        role = str(
            outline.content_json.get("narrative_role")
            or outline.layout_json.get("narrative_role")
            or ""
        ).strip().lower()
        if role == "reference":
            return "SOURCE READ"
        if role == "decision":
            return "DECISION BASIS"
        if role in {"risk", "problem", "challenge"}:
            return "RISK SIGNALS"
        return "SOURCE SIGNALS"

    def _add_authored_toolkit_grid(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
    ) -> None:
        items = self._authored_items(outline, limit=4)
        hero = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.9),
            Inches(1.6),
            Inches(4.28),
            Inches(4.74),
        )
        hero.fill.solid()
        hero.fill.fore_color.rgb = self._rgb(self._tint(brand.colors.primary, 0.9))
        hero.line.color.rgb = self._rgb(self._tint(brand.colors.primary, 0.74))
        self._add_dark_text(
            slide,
            "OPERATING KIT",
            1.24,
            1.98,
            2.2,
            0.22,
            brand,
            size=8,
            bold=True,
            color=brand.colors.accent,
        )
        primary_support = self._authored_support_text(outline, items[0])
        self._add_body_text(
            slide,
            primary_support,
            1.24,
            2.58,
            3.42,
            1.56,
            brand,
            size=16,
        )
        secondary_support = self._authored_distinct_support_text(
            outline,
            [*items, primary_support],
            items[0],
        )
        self._add_body_text(
            slide,
            secondary_support,
            1.26,
            5.15,
            3.35,
            0.5,
            brand,
            size=10,
        )
        positions = [
            (5.72, 1.52),
            (9.0, 1.82),
            (5.72, 4.0),
            (9.0, 4.3),
        ]
        for idx, item in enumerate(items[:4]):
            x, y = positions[idx]
            accent = self._icon_fill(brand, idx)
            self._add_card(slide, x, y, 2.96, 1.84, "FFFFFF", brand.colors.background_light)
            self._add_badge(slide, str(idx + 1), x + 0.24, y + 0.25, 0.38, brand, accent)
            lead, rest = self._split_lead(item)
            self._add_label(
                slide,
                self._truncate_at_word(lead, 36),
                x + 0.78,
                y + 0.28,
                1.88,
                brand,
                bold=True,
            )
            self._add_body_text(
                slide,
                self._truncate_at_word(rest or item, 82),
                x + 0.3,
                y + 0.82,
                2.28,
                0.78,
                brand,
                size=9,
            )

    def _add_authored_challenge_cards(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
    ) -> None:
        items = self._authored_items(outline, limit=3)
        labels = ["TENSION", "RISK", "MOVE"]
        used_details: list[str] = []
        for idx, label in enumerate(labels):
            item = items[idx] if idx < len(items) else items[-1]
            x = 0.96 + idx * 4.02
            y = 1.72 + (0.28 if idx == 1 else 0)
            accent = self._icon_fill(brand, idx)
            fill = self._tint(accent, 0.88) if idx == 0 else "FFFFFF"
            self._add_card(slide, x, y, 3.42, 3.78, fill, brand.colors.background_light)
            strip = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(x),
                Inches(y),
                Inches(3.42),
                Inches(0.14),
            )
            strip.fill.solid()
            strip.fill.fore_color.rgb = self._rgb(accent)
            strip.line.color.rgb = strip.fill.fore_color.rgb
            self._add_dark_text(
                slide,
                label,
                x + 0.34,
                y + 0.5,
                1.28,
                0.18,
                brand,
                size=8,
                bold=True,
                color=accent,
            )
            lead, rest = self._split_lead(item)
            if self._authored_heading_reads_like_fragment(lead, rest):
                rest = item
                lead = self._fallback_heading_from_title(outline, idx)
            self._add_body_text(
                slide,
                lead,
                x + 0.34,
                y + 1.1,
                2.72,
                0.54,
                brand,
                size=15,
            )
            detail = (
                self._truncate_at_word(self._clean_display_text(str(item)), 112)
                if outline.content_json.get("visual_qa_source_repair")
                else self._authored_detail_text(
                    outline,
                    item,
                    idx,
                    rest,
                    avoid=[*items, *used_details],
                )
            )
            if detail:
                used_details.append(detail)
            self._add_body_text(
                slide,
                detail,
                x + 0.34,
                y + 2.05,
                2.64,
                0.9,
                brand,
                size=11,
            )

    def _add_authored_why_matters(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
    ) -> None:
        items = self._authored_items(outline, limit=3)
        support = self._authored_distinct_support_text(
            outline,
            items,
            "",
        )
        if not support and items:
            support = self._complete_display_item(items[0], limit=118)
        band = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.9),
            Inches(1.6),
            Inches(11.52),
            Inches(1.35),
        )
        band.fill.solid()
        band.fill.fore_color.rgb = self._rgb(self._tint(brand.colors.accent, 0.86))
        band.line.color.rgb = self._rgb(self._tint(brand.colors.accent, 0.7))
        self._add_dark_text(
            slide,
            "WHY IT MATTERS",
            1.22,
            1.9,
            2.1,
            0.2,
            brand,
            size=8,
            bold=True,
            color=brand.colors.primary,
        )
        self._add_body_text(
            slide,
            self._truncate_phrase(support, 118),
            3.18,
            1.88,
            7.78,
            0.46,
            brand,
            center=False,
            size=16,
        )
        used_details: list[str] = []
        for idx, item in enumerate(items[:3]):
            x = 1.0 + idx * 3.92
            accent = self._icon_fill(brand, idx)
            self._add_card(slide, x, 3.42, 3.34, 2.28, "FFFFFF", brand.colors.background_light)
            self._add_icon(slide, self._icons(outline)[idx], x + 0.28, 3.75, 0.62, brand, accent)
            lead, rest = self._split_lead(item)
            if self._authored_heading_reads_like_fragment(lead, rest):
                rest = item
                lead = self._fallback_heading_from_title(outline, idx)
            self._add_dark_text(
                slide,
                self._safe_authored_heading(outline, lead, rest, idx, 42),
                x + 1.02,
                3.78,
                2.08,
                0.5,
                brand,
                size=10,
                bold=True,
                color=brand.colors.primary,
            )
            detail = self._authored_detail_text(
                outline,
                item,
                idx,
                rest,
                avoid=[support, *items, *used_details],
            )
            if detail:
                used_details.append(detail)
            self._add_body_text(
                slide,
                detail,
                x + 0.34,
                4.62,
                2.58,
                0.68,
                brand,
                size=11,
            )

    def _authored_heading_reads_like_fragment(self, lead: str, rest: str) -> bool:
        cleaned = " ".join(str(lead).split()).strip()
        if not cleaned:
            return True
        if not rest and len(cleaned.split()) > 4:
            return True
        if re.match(r"^(?:a|an|the)\b", cleaned, re.IGNORECASE) and not rest:
            return True
        if re.search(
            r"\b(?:a|an|and|as|by|for|from|in|into|of|or|the|their|through|to|with)\.?$",
            cleaned,
            re.IGNORECASE,
        ):
            return True
        return False

    def _safe_authored_heading(
        self,
        outline: SlideOutline,
        lead: str,
        rest: str,
        index: int,
        limit: int,
    ) -> str:
        heading = self._truncate_at_word(lead, limit)
        if (
            self._authored_heading_reads_like_fragment(heading, rest)
            or self._is_incomplete_display_fragment(heading)
            or self._is_renderer_filler_copy(heading)
        ):
            heading = self._truncate_at_word(
                self._fallback_heading_from_title(outline, index),
                limit,
            )
        if self._is_incomplete_display_fragment(heading):
            words = [
                word
                for word in re.findall(r"[A-Za-z][A-Za-z0-9'-]*", lead)
                if word.casefold() not in {"a", "an", "and", "of", "the", "to"}
            ]
            heading = " ".join(words[:3]) or self._compact_title_subject(
                self._authored_title(outline)
            )
        return heading

    def _authored_items(self, outline: SlideOutline, limit: int) -> list[str]:
        items = self._grid_items(outline)
        if not items:
            fallback = (
                outline.content_json.get("subheading")
                or outline.content_json.get("summary")
                or outline.label
            )
            items = [str(fallback)]
        cleaned: list[str] = []
        seen: set[str] = set()
        seen_leads: set[str] = set()
        for item in items:
            complete = self._complete_display_item(str(item), limit=132)
            if not complete:
                continue
            text = self._truncate_at_word(
                complete.replace("...", ""),
                132,
            )
            if not text:
                continue
            if self._is_authored_placeholder_item(text):
                continue
            key = text.casefold()
            if key in seen:
                continue
            lead, _ = self._split_lead(text)
            lead_key = re.sub(r"[^a-z0-9 ]+", "", lead.casefold()).strip()
            if len(lead_key.split()) >= 3 and self._authored_lead_seen(
                lead_key,
                seen_leads,
            ):
                continue
            seen.add(key)
            if len(lead_key.split()) >= 3:
                seen_leads.add(lead_key)
            cleaned.append(text)
            if len(cleaned) >= limit:
                break
        if not cleaned:
            fallback = self._complete_display_item(self._authored_title(outline), limit=132)
            if fallback:
                cleaned.append(fallback)
        return cleaned[:limit]

    def _authored_lead_seen(self, lead_key: str, seen_leads: set[str]) -> bool:
        return any(
            lead_key == seen
            or lead_key.startswith(f"{seen} ")
            or seen.startswith(f"{lead_key} ")
            for seen in seen_leads
        )

    def _authored_filler_items(self, outline: SlideOutline) -> list[str]:
        title = self._authored_title(outline)
        candidates: list[str] = []
        notes = str(outline.content_json.get("speaker_notes") or "")
        for sentence in re.split(r"(?<=[.!?])\s+", notes):
            if sentence.strip():
                candidates.append(sentence)
        source_label = self._source_label(outline)
        if ">" in source_label:
            source_label = source_label.rsplit(">", 1)[-1].strip()
        if (
            source_label
            and not self._looks_like_meta_subheading(source_label)
            and not self._is_placeholder_bullet(source_label)
        ):
            candidates.append(source_label)
        candidates.extend(self._title_facet_items(title))
        candidates.append(self._topic_support_sentence(outline))
        return self._unique_display_items(candidates, limit=6)

    def _title_facet_items(self, title: str) -> list[str]:
        normalized = self._clean_display_text(title)
        facets: list[str] = []
        if normalized and not self._is_incomplete_display_fragment(normalized):
            facets.append(self._truncate_at_word(normalized, 54))
        return facets

    def _fallback_heading_from_title(self, outline: SlideOutline, index: int) -> str:
        facets = self._authored_filler_items(outline)
        if facets:
            return self._truncate_at_word(facets[index % len(facets)], 42)
        return self._truncate_at_word(self._authored_title(outline), 42)

    def _topic_support_sentence(self, outline: SlideOutline) -> str:
        support = str(
            outline.content_json.get("subheading")
            or outline.content_json.get("summary")
            or ""
        )
        if (
            support
            and not self._looks_like_meta_subheading(support)
            and not self._looks_like_section_title(support)
        ):
            return self._truncate_at_word(self._clean_display_text(support), 92)
        subject = self._compact_title_subject(self._authored_title(outline))
        if subject and not self._is_incomplete_display_fragment(subject):
            return self._truncate_at_word(subject.capitalize(), 92)
        return ""

    def _looks_like_section_title(self, text: str) -> bool:
        cleaned = " ".join(str(text).split()).strip(" .:-")
        if not cleaned:
            return False
        lowered = cleaned.casefold()
        if lowered in {
            "architecture overview",
            "closing remarks",
            "conclusion and future directions",
            "the case for implicit ground truth discovery",
            "the model contract",
        }:
            return True
        verbs = {
            "are",
            "build",
            "builds",
            "create",
            "creates",
            "define",
            "defines",
            "discover",
            "discovers",
            "enable",
            "enables",
            "improve",
            "improves",
            "is",
            "make",
            "makes",
            "must",
            "need",
            "needs",
            "shift",
            "shifts",
            "use",
            "uses",
        }
        words = re.findall(r"[A-Za-z][A-Za-z'-]*", cleaned)
        if len(words) <= 8 and not any(word.casefold() in verbs for word in words):
            return True
        return False

    def _compact_title_subject(self, title: str) -> str:
        words = [
            word
            for word in re.findall(r"[A-Za-z][A-Za-z0-9'-]*", title)
            if word.casefold()
            not in {
                "a",
                "an",
                "and",
                "are",
                "as",
                "be",
                "because",
                "before",
                "build",
                "by",
                "for",
                "from",
                "in",
                "into",
                "is",
                "must",
                "of",
                "or",
                "rather",
                "than",
                "the",
                "through",
                "to",
                "use",
                "with",
            }
        ]
        if not words:
            return ""
        return self._truncate_at_word(" ".join(words[:3]), 38)

    def _unique_display_items(self, items: list[str], limit: int) -> list[str]:
        cleaned: list[str] = []
        seen: set[str] = set()
        for item in items:
            text = self._complete_display_item(str(item), limit=132)
            if not text:
                continue
            key = re.sub(r"[^a-z0-9]+", " ", text.casefold()).strip()
            if key in seen:
                continue
            if any(self._authored_text_overlaps(existing, text) for existing in cleaned):
                continue
            seen.add(key)
            cleaned.append(text)
            if len(cleaned) >= limit:
                break
        return cleaned

    def _authored_distinct_support_text(
        self,
        outline: SlideOutline,
        items: list[str],
        fallback: str,
    ) -> str:
        raw_support = str(
            outline.content_json.get("subheading")
            or outline.content_json.get("summary")
            or ""
        )
        if self._looks_like_meta_subheading(raw_support) or self._looks_like_section_title(raw_support):
            raw_support = ""
        candidates = [raw_support, fallback]
        for candidate in candidates:
            text = self._complete_display_item(str(candidate), limit=128).replace(
                "...",
                "",
            )
            if not text:
                continue
            if any(self._authored_text_overlaps(item, text) for item in items[:5]):
                continue
            return text
        return ""

    def _authored_canvas_support(self, outline: SlideOutline) -> str:
        return self._topic_support_sentence(outline)

    def _is_authored_placeholder_item(self, text: str) -> bool:
        normalized = " ".join(str(text).casefold().split()).strip(" .:-")
        return normalized in {
            "current readout",
            "target move",
            "signal",
            "managed behavior",
            "source signal",
            "source read",
        }

    def _authored_detail_text(
        self,
        outline: SlideOutline,
        item: str,
        index: int,
        preferred: str = "",
        avoid: list[str] | None = None,
    ) -> str:
        lead, _ = self._split_lead(item)
        candidates = [preferred]
        for candidate in candidates:
            detail = self._clean_display_text(str(candidate)).replace("...", "")
            if not detail or self._is_authored_placeholder_item(detail):
                continue
            detail = re.sub(
                r"^(?:and|but|or|because)\s+",
                "",
                detail,
                flags=re.IGNORECASE,
            )
            if detail[:1].islower():
                detail = f"{detail[:1].upper()}{detail[1:]}"
            if self._is_incomplete_display_fragment(detail):
                continue
            if self._authored_text_overlaps(lead, detail):
                continue
            if any(self._authored_text_overlaps(other, detail) for other in avoid or []):
                continue
            return self._truncate_at_word(detail, 96)
        return ""

    def _authored_text_overlaps(self, left: str, right: str) -> bool:
        left_norm = re.sub(r"[^a-z0-9 ]+", "", str(left).casefold()).strip()
        right_norm = re.sub(r"[^a-z0-9 ]+", "", str(right).casefold()).strip()
        if not left_norm or not right_norm:
            return False
        shorter, longer = sorted([left_norm, right_norm], key=len)
        if len(shorter.split()) < 3:
            return shorter == longer
        return longer.startswith(shorter) or shorter.startswith(longer)

    def _authored_title(self, outline: SlideOutline) -> str:
        return self._clean_display_text(
            str(
                outline.content_json.get("action_title")
                or outline.content_json.get("title")
                or outline.label
            )
        )

    def _authored_support_text(self, outline: SlideOutline, fallback: str) -> str:
        support = str(
            outline.content_json.get("subheading")
            or outline.content_json.get("summary")
            or fallback
        )
        if self._looks_like_meta_subheading(support) or self._looks_like_section_title(support):
            support = fallback
        return self._truncate_phrase(support, 132)

    def _source_label(self, outline: SlideOutline) -> str:
        sources = outline.content_json.get("sources") or []
        if isinstance(sources, list) and sources:
            return str(sources[0])
        refs = outline.content_json.get("source_refs") or []
        if isinstance(refs, list) and refs:
            return str(refs[0])
        return ""

    def _add_callouts(self, slide, outline: SlideOutline, brand: BrandDNA) -> None:
        metrics = [
            metric
            for metric in (outline.content_json.get("metrics") or [])
            if isinstance(metric, dict)
        ]
        quantitative = [metric for metric in metrics if self._is_quantitative_metric(metric)]
        if len(quantitative) < 2:
            # A "1 / 2 / 3" sequence rendered as big numbers reads as filler;
            # show the points as labeled cards instead of fake KPIs.
            self._add_grid(slide, outline, brand)
            return
        metrics = quantitative
        icons = self._icons(outline)
        for idx, metric in enumerate(metrics[:3]):
            x = 0.95 + idx * 4.05
            self._add_card(slide, x, 1.65, 3.45, 3.75, "FFFFFF", brand.colors.background_light)
            strip = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(x),
                Inches(1.65),
                Inches(3.45),
                Inches(0.12),
            )
            strip.fill.solid()
            strip.fill.fore_color.rgb = self._rgb(brand.colors.accent if idx == 0 else brand.colors.secondary)
            strip.line.color.rgb = strip.fill.fore_color.rgb
            self._add_icon(slide, icons[idx], x + 2.54, 1.86, 0.74, brand, self._icon_fill(brand, idx))
            self._add_big_number(slide, self._format_metric_value(metric), x + 0.2, 2.05, 3.05, brand)
            self._add_body_text(slide, self._metric_label_text(metric.get("label", "Metric")), x + 0.35, 3.35, 2.75, 1.35, brand, center=True, size=14)

    def _authored_quantitative_metrics(self, outline: SlideOutline) -> list[dict[str, Any]]:
        metrics = [
            metric
            for metric in (outline.content_json.get("metrics") or [])
            if isinstance(metric, dict)
        ]
        exhibit = self._exhibit(outline)
        exhibit_metrics = exhibit.get("metrics")
        if isinstance(exhibit_metrics, list):
            metrics.extend(metric for metric in exhibit_metrics if isinstance(metric, dict))
        cleaned: list[dict[str, Any]] = []
        seen: set[tuple[str, str, str]] = set()
        for metric in metrics:
            if not self._is_quantitative_metric(metric):
                continue
            key = (
                str(metric.get("label") or "").casefold(),
                str(metric.get("value") or ""),
                str(metric.get("unit") or "").casefold(),
            )
            if key in seen:
                continue
            seen.add(key)
            cleaned.append(metric)
        return cleaned

    def _add_metric_chart(self, slide, outline: SlideOutline, brand: BrandDNA) -> None:
        metrics = outline.content_json.get("metrics") or []
        exhibit = self._exhibit(outline)
        if not metrics and exhibit.get("type") == "metric_chart":
            metrics = exhibit.get("metrics") or []
        if not metrics and exhibit.get("type") == "line_chart":
            metrics = exhibit.get("metrics") or []
        metrics = [metric for metric in metrics if isinstance(metric, dict)]
        if not metrics:
            self._add_two_column(slide, outline, brand)
            return
        chart_type = str((outline.content_json.get("chart_spec") or {}).get("type") or exhibit.get("type") or "")
        if chart_type == "line" or exhibit.get("type") == "line_chart":
            series = self._line_series(metrics[:8])
            if series:
                self._add_native_line_chart(slide, outline, brand, series)
                return
        series = self._chart_series(metrics[:6])
        if series:
            self._add_native_bar_chart(slide, outline, brand, series)
            return
        display_metrics = metrics[:4]
        count = len(display_metrics)
        gap = 0.28
        card_w = (11.75 - gap * (count - 1)) / count
        self._add_label(slide, "SOURCED SIGNALS", 0.9, 1.45, 2.2, brand, bold=True)
        for idx, metric in enumerate(display_metrics):
            x = 0.85 + idx * (card_w + gap)
            accent = self._icon_fill(brand, idx)
            value_text, unit_text = self._metric_card_value_parts(metric)
            self._add_card(slide, x, 1.86, card_w, 2.6, "FFFFFF", brand.colors.background_light)
            strip = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(x),
                Inches(1.86),
                Inches(card_w),
                Inches(0.12),
            )
            strip.fill.solid()
            strip.fill.fore_color.rgb = self._rgb(accent)
            strip.line.color.rgb = strip.fill.fore_color.rgb
            self._add_big_number(
                slide,
                value_text,
                x + 0.22,
                2.26,
                card_w - 0.44,
                brand,
            )
            label_y = 3.28
            if unit_text:
                self._add_dark_text(
                    slide,
                    unit_text,
                    x + 0.3,
                    3.02,
                    card_w - 0.6,
                    0.24,
                    brand,
                    size=16,
                    bold=True,
                    color=brand.colors.primary,
                )
                label_y = 3.58
            self._add_body_text(
                slide,
                self._metric_label_text(metric.get("label", "Metric")),
                x + 0.28,
                label_y,
                card_w - 0.56,
                0.72,
                brand,
                center=True,
                size=12,
            )

        insight = self._metric_insight(outline, display_metrics)
        self._add_card(slide, 1.2, 4.92, 10.9, 0.86, brand.colors.background_light, brand.colors.background_light)
        self._add_body_text(slide, insight, 1.52, 5.14, 10.25, 0.34, brand, center=True, size=13)

    def _metric_card_value_parts(self, metric: dict[str, Any]) -> tuple[str, str]:
        unit = str(metric.get("unit") or "")
        try:
            numeric = float(metric.get("value", ""))
        except (TypeError, ValueError):
            return self._truncate_at_word(str(metric.get("value", "")), 10), ""
        if unit == "%":
            return f"{self._format_compact_number(numeric)}%", ""
        if unit.lower() == "tokens":
            return self._format_compact_number(numeric), "tokens"
        if unit:
            return self._format_compact_number(numeric), self._truncate_at_word(unit, 12)
        return self._format_compact_number(numeric), ""

    def _is_quantitative_metric(self, metric: dict[str, Any]) -> bool:
        """True when a metric carries a real measurement worth a big number, as
        opposed to a sequence index (1, 2, 3) used as a list ordinal."""
        unit = str(metric.get("unit") or "").strip()
        label = str(metric.get("label") or metric.get("name") or "").strip().casefold()
        try:
            number = float(str(metric.get("value", "")).replace(",", "").rstrip("%"))
        except (TypeError, ValueError):
            number = None
        if unit and number is not None and abs(number) <= 2 and label in {
            "prediction",
            "predictions",
            "item",
            "items",
        }:
            return False
        if unit:
            return True
        if number is None:
            return False
        return abs(number) >= 10

    def _format_metric_value(self, metric: dict[str, Any]) -> str:
        value = metric.get("value", "")
        unit = str(metric.get("unit") or "")
        try:
            numeric = float(value)
        except (TypeError, ValueError):
            return str(value)[:12]
        if unit == "%":
            return f"{self._format_compact_number(numeric)}%"
        if unit.lower() == "tokens":
            return f"{self._format_compact_number(numeric)} tokens"
        return self._format_compact_number(numeric)

    def _format_compact_number(self, value: float) -> str:
        if abs(value) >= 1_000_000:
            compact = value / 1_000_000
            suffix = "M"
        elif abs(value) >= 1_000:
            compact = value / 1_000
            suffix = "k"
        else:
            compact = value
            suffix = ""
        if float(compact).is_integer():
            return f"{int(compact)}{suffix}"
        return f"{compact:.1f}".rstrip("0").rstrip(".") + suffix

    def _metric_label_text(self, label: Any) -> str:
        words = " ".join(str(label).split()).split()
        if len(words) <= 8:
            return " ".join(words)
        return " ".join(words[:8])

    def _metric_insight(self, outline: SlideOutline, metrics: list[dict[str, Any]]) -> str:
        units = {str(metric.get("unit") or "").lower() for metric in metrics}
        if "%" in units and "tokens" in units:
            return "Adoption and capacity are high enough to make operating discipline the constraint."
        if "%" in units:
            return "Treat high adoption as a signal to formalize review, ownership, and quality gates."
        if "tokens" in units:
            return "Larger context windows help, but persistent project memory still carries the workflow."
        bullets = self._bullets(outline)
        text = bullets[0] if bullets else outline.content_json.get("subheading", "")
        if str(text).lower().startswith("evidence from"):
            text = outline.content_json.get("subheading", "")
        if not text:
            return "Use these figures as pressure signals, then manage the operating model."
        words = " ".join(str(text).split()).split()
        return " ".join(words[:18])

    def _chart_series(
        self, metrics: list[dict[str, Any]]
    ) -> list[tuple[str, float, str]] | None:
        """Return a chartable (label, value, unit) series, or None.

        A native bar chart is only meaningful when at least three values share a
        single unit; mixed-unit KPIs (e.g. 95% next to 1M tokens) stay as cards.
        """
        series: list[tuple[str, float, str]] = []
        for metric in metrics:
            text = str(metric.get("value", "")).strip()
            had_pct = text.endswith("%")
            try:
                value = float(text.replace(",", "").rstrip("%").strip())
            except (TypeError, ValueError):
                return None
            label = self._metric_label_text(metric.get("label", "")) or "Item"
            unit = str(metric.get("unit") or "").strip() or ("%" if had_pct else "")
            series.append((label, value, unit))
        if len(series) < 3:
            return None
        if len({unit.lower() for _, _, unit in series}) != 1:
            return None
        return series

    def _line_series(
        self, metrics: list[dict[str, Any]]
    ) -> list[tuple[str, float, str]] | None:
        series: list[tuple[str, float, str]] = []
        for index, metric in enumerate(metrics):
            try:
                value = float(str(metric.get("value", "")).replace(",", "").rstrip("%"))
            except (TypeError, ValueError):
                continue
            label = self._metric_label_text(metric.get("label", "")) or f"Point {index + 1}"
            unit = str(metric.get("unit") or "")
            series.append((label, value, unit))
        if len(series) < 3:
            return None
        units = {unit.lower() for _, _, unit in series if unit}
        if len(units) > 1:
            return None
        return series

    def _add_native_line_chart(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
        series: list[tuple[str, float, str]],
    ) -> None:
        labels = [self._truncate_at_word(label, 20) for label, _, _ in series]
        values = [value for _, value, _ in series]
        unit = series[0][2]
        self._add_label(slide, "TREND SIGNAL", 0.9, 1.45, 3.0, brand, bold=True)
        chart_data = CategoryChartData()
        chart_data.categories = labels
        chart_data.add_series("trend", tuple(values))
        frame = slide.shapes.add_chart(
            XL_CHART_TYPE.LINE_MARKERS,
            Inches(0.9),
            Inches(1.9),
            Inches(11.35),
            Inches(3.6),
            chart_data,
        )
        chart = frame.chart
        chart.has_legend = False
        chart.has_title = False
        plot = chart.plots[0]
        plot.has_data_labels = True
        data_labels = plot.data_labels
        data_labels.number_format = '0"%"' if unit == "%" else "#,##0"
        data_labels.number_format_is_linked = False
        data_labels.position = XL_LABEL_POSITION.ABOVE
        data_labels.font.size = Pt(10)
        data_labels.font.bold = True
        data_labels.font.name = brand.fonts.body
        data_labels.font.color.rgb = self._rgb(brand.colors.text_dark)
        line_series = plot.series[0]
        line_series.format.line.color.rgb = self._rgb(brand.colors.accent)
        line_series.format.line.width = Pt(2.25)
        value_axis = chart.value_axis
        value_axis.has_major_gridlines = True
        value_axis.major_gridlines.format.line.color.rgb = self._rgb(
            self._tint(brand.colors.secondary, 0.78)
        )
        value_axis.tick_labels.font.size = Pt(9)
        category_axis = chart.category_axis
        category_axis.tick_labels.font.size = Pt(10)
        category_axis.tick_labels.font.name = brand.fonts.body
        category_axis.tick_labels.font.color.rgb = self._rgb(brand.colors.text_dark)
        insight = self._metric_insight(
            outline,
            [{"label": label, "value": value, "unit": unit} for label, value, unit in series],
        )
        self._add_card(
            slide,
            1.2,
            5.74,
            10.9,
            0.74,
            brand.colors.background_light,
            brand.colors.background_light,
        )
        self._add_body_text(
            slide, insight, 1.52, 5.92, 10.25, 0.4, brand, center=True, size=13
        )

    def _add_native_bar_chart(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
        series: list[tuple[str, float, str]],
    ) -> None:
        labels = [self._truncate_at_word(label, 22) for label, _, _ in series]
        values = [value for _, value, _ in series]
        unit = series[0][2]
        self._add_label(slide, "SOURCED SIGNALS", 0.9, 1.45, 3.0, brand, bold=True)

        chart_data = CategoryChartData()
        chart_data.categories = labels
        chart_data.add_series("signal", tuple(values))
        frame = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(0.9),
            Inches(1.98),
            Inches(11.5),
            Inches(3.5),
            chart_data,
        )
        chart = frame.chart
        chart.has_legend = False
        chart.has_title = False

        plot = chart.plots[0]
        plot.gap_width = 80
        plot.has_data_labels = True
        data_labels = plot.data_labels
        data_labels.number_format = '0"%"' if unit == "%" else "#,##0"
        data_labels.number_format_is_linked = False
        data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        data_labels.font.size = Pt(12)
        data_labels.font.bold = True
        data_labels.font.name = brand.fonts.body
        data_labels.font.color.rgb = self._rgb(brand.colors.text_dark)

        # One accent bar carries the message; the rest stay a neutral gray.
        highlight = max(range(len(values)), key=lambda index: values[index])
        muted = self._tint(brand.colors.secondary, 0.5)
        bar_series = plot.series[0]
        for index, point in enumerate(bar_series.points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = self._rgb(
                brand.colors.accent if index == highlight else muted
            )

        value_axis = chart.value_axis
        value_axis.has_major_gridlines = False
        value_axis.visible = False
        value_axis.minimum_scale = 0
        value_axis.maximum_scale = max(values) * 1.18 if values else 1
        category_axis = chart.category_axis
        category_axis.has_major_gridlines = False
        category_axis.tick_labels.font.size = Pt(11)
        category_axis.tick_labels.font.name = brand.fonts.body
        category_axis.tick_labels.font.color.rgb = self._rgb(brand.colors.text_dark)

        insight = self._metric_insight(
            outline,
            [{"label": label, "value": value, "unit": unit} for label, value, unit in series],
        )
        self._add_card(
            slide, 1.2, 5.74, 10.9, 0.74, brand.colors.background_light, brand.colors.background_light
        )
        self._add_body_text(slide, insight, 1.52, 5.92, 10.25, 0.4, brand, center=True, size=13)
