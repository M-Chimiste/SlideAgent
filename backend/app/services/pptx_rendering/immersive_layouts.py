# ruff: noqa: F401
from pathlib import Path
import math
import re
import shutil
import subprocess
from typing import Any

from PIL import Image, ImageDraw
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt

from app.models.brand import BrandDNA
from app.models.outline import SlideOutline
from app.services.concept_diagram_renderer import ConceptDiagramRenderer, DiagramRenderError
from app.services.pptx_rendering.constants import ICON_SCALE, SLIDE_H, SLIDE_W


class ImmersiveLayoutRenderingMixin:
    def _add_framework_cycle_immersive(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
        slide_number: int,
        total_slides: int,
    ) -> None:
        self._add_dark_slide_chrome(
            slide, outline, brand, slide_number, total_slides, "OPERATING MODEL"
        )
        if self._try_add_diagram_asset(
            slide, outline, brand, 0.92, 1.76, 11.5, 4.92, dark=True
        ):
            return
        exhibit = self._exhibit(outline)
        steps = exhibit.get("steps") if exhibit.get("type") == "cycle" else None
        if isinstance(steps, list) and steps:
            bullets = [
                str(step.get("label") or step.get("description") or "").strip()
                for step in steps
                if isinstance(step, dict)
            ][:6]
        else:
            bullets = self._bullets(outline)[:6]
        if len(bullets) < 4 or self._labels_are_generic_visual_fallbacks(bullets):
            self._add_authored_proof_strip(slide, outline, brand)
            return
        while len(bullets) < 6:
            bullets.append(bullets[-1])
        center_x, center_y = 6.66, 4.18
        loop = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(center_x - 1.04),
            Inches(center_y - 1.04),
            Inches(2.08),
            Inches(2.08),
        )
        loop.fill.solid()
        loop.fill.fore_color.rgb = self._rgb(self._tint(brand.colors.primary, 0.1))
        loop.line.color.rgb = self._rgb(brand.colors.accent)
        loop.line.width = Pt(1.4)
        self._add_dark_text(
            slide,
            self._truncate_at_word(str(exhibit.get("center_label") or "Operating loop"), 28),
            center_x - 0.62,
            center_y - 0.22,
            1.24,
            0.44,
            brand,
            size=12,
            bold=True,
            color=brand.colors.text_light,
        )
        positions = [
            (3.35, 2.72),
            (6.66, 2.38),
            (9.98, 2.72),
            (9.98, 5.06),
            (6.66, 5.38),
            (3.35, 5.06),
        ]
        icons = self._icons(outline)
        for idx, (x1, y1) in enumerate(positions):
            x2, y2 = positions[(idx + 1) % len(positions)]
            self._add_arrow(slide, x1, y1, x2, y2, brand.colors.accent, width=1.7)
            mid_x = (x1 + x2) / 2
            mid_y = (y1 + y2) / 2
            self._add_arrowhead(slide, mid_x, mid_y, x2 - x1, y2 - y1, brand.colors.accent)
        for idx, (x, y) in enumerate(positions):
            card = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(x - 1.18),
                Inches(y - 0.42),
                Inches(2.36),
                Inches(0.86),
            )
            card.fill.solid()
            card.fill.fore_color.rgb = self._rgb(self._tint(brand.colors.primary, 0.16 + (idx % 2) * 0.05))
            card.line.color.rgb = self._rgb(self._tint(brand.colors.primary, 0.34))
            self._add_icon(slide, icons[idx % len(icons)], x - 1.0, y - 0.26, 0.52, brand, self._icon_fill(brand, idx))
            self._add_dark_text(
                slide,
                self._truncate_at_word(bullets[idx], 34),
                x - 0.32,
                y - 0.19,
                1.2,
                0.34,
                brand,
                size=9,
                bold=True,
                color=self._tint(brand.colors.primary, 0.86),
            )

    def _add_quote_sidebar_immersive(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
        slide_number: int,
        total_slides: int,
    ) -> None:
        self._add_dark_slide_chrome(
            slide, outline, brand, slide_number, total_slides, "DECISION LENS"
        )
        exhibit = self._exhibit(outline)
        bullets = [
            str(item)
            for item in exhibit.get("supporting_points", [])
            if str(item).strip()
        ][:4] or self._bullets(outline)[:4]
        if not bullets:
            bullets = ["Shift the operating model from ad hoc execution to managed discipline."]
        quote = self._first_content_text(
            [
                exhibit.get("quote"),
                exhibit.get("key_idea"),
                outline.content_json.get("summary"),
                outline.content_json.get("subheading"),
                bullets[0],
            ],
            fallback=bullets[0],
        )
        callout = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.86),
            Inches(2.54),
            Inches(5.86),
            Inches(2.8),
        )
        callout.fill.solid()
        callout.fill.fore_color.rgb = self._rgb(self._tint(brand.colors.primary, 0.1))
        callout.line.color.rgb = self._rgb(self._tint(brand.colors.primary, 0.3))
        self._add_dark_text(slide, "THE MENTAL MODEL", 1.16, 2.88, 2.2, 0.22, brand, size=8, bold=True, color=brand.colors.accent)
        self._add_dark_text(
            slide,
            self._truncate_phrase(str(quote), 165),
            1.16,
            3.35,
            4.92,
            1.22,
            brand,
            size=17,
            bold=True,
            color=brand.colors.text_light,
        )
        self._add_dark_text(
            slide,
            self._topic_support_sentence(outline),
            1.18,
            4.72,
            4.8,
            0.34,
            brand,
            size=11,
            color=self._tint(brand.colors.primary, 0.68),
        )
        icons = self._icons(outline)
        for idx, text in enumerate(bullets[:4]):
            y = 2.42 + idx * 0.86
            card = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(7.35),
                Inches(y),
                Inches(4.35),
                Inches(0.76),
            )
            card.fill.solid()
            card.fill.fore_color.rgb = self._rgb(self._tint(brand.colors.primary, 0.15 + idx * 0.03))
            card.line.color.rgb = self._rgb(self._tint(brand.colors.primary, 0.34))
            self._add_icon(slide, icons[idx % len(icons)], 7.58, y + 0.1, 0.42, brand, self._icon_fill(brand, idx))
            self._add_dark_text(
                slide,
                self._truncate_phrase(text, 78),
                8.16,
                y + 0.12,
                3.1,
                0.44,
                brand,
                size=10,
                color=self._tint(brand.colors.primary, 0.84),
            )
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(6.74),
            Inches(3.94),
            Inches(7.34),
            Inches(3.94),
        )
        connector.line.color.rgb = self._rgb(brand.colors.accent)
        connector.line.width = Pt(1.4)

    def _add_closing_recommendation_immersive(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
        slide_number: int,
        total_slides: int,
    ) -> None:
        outline = self._closing_display_outline(outline)
        self._add_dark_slide_chrome(
            slide, outline, brand, slide_number, total_slides, "FINAL DECISION"
        )
        exhibit = self._exhibit(outline)
        recommendation = (
            exhibit.get("recommendation")
            or outline.content_json.get("action_title")
            or outline.label
        )
        recommendation = self._nonduplicate_closing_recommendation(
            str(recommendation),
            str(outline.content_json.get("action_title") or outline.label),
        )
        if self._closing_text_is_source_excerpt(recommendation) or self._is_incomplete_display_fragment(recommendation):
            recommendation = self._default_closing_recommendation(outline)
        next_steps = [
            str(item)
            for item in exhibit.get("next_steps", [])
            if str(item).strip() and self._closing_step_is_actionable(str(item))
        ][:3]
        if len(next_steps) < 2:
            next_steps = self._closing_next_steps(outline)
        if not next_steps:
            next_steps = [
                "Select the first governed benchmark pilot.",
                "Bind the pilot to validated source evidence.",
                "Review failures before expanding coverage.",
            ]
        ask = self._closing_decision_ask(
            str(exhibit.get("decision_ask") or ""),
            outline,
        )
        self._add_dark_text(slide, "RECOMMENDATION", 0.92, 2.5, 2.2, 0.24, brand, size=9, bold=True, color=brand.colors.accent)
        self._add_dark_text(
            slide,
            self._truncate_phrase(recommendation, 155),
            0.92,
            2.96,
            6.1,
            1.05,
            brand,
            size=24,
            bold=True,
            color=brand.colors.text_light,
        )
        ask_box = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.92),
            Inches(4.82),
            Inches(5.86),
            Inches(0.66),
        )
        ask_box.fill.solid()
        ask_box.fill.fore_color.rgb = self._rgb(self._tint(brand.colors.accent, 0.18))
        ask_box.line.color.rgb = ask_box.fill.fore_color.rgb
        self._add_dark_text(slide, "DECISION ASK", 1.14, 5.0, 1.3, 0.18, brand, size=7, bold=True, color=brand.colors.text_dark)
        self._add_dark_text(
            slide,
            self._truncate_at_word(ask, 96),
            2.48,
            4.94,
            3.82,
            0.34,
            brand,
            size=10,
            color=brand.colors.text_dark,
        )
        rail = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(7.86),
            Inches(2.56),
            Inches(0.035),
            Inches(3.34),
        )
        rail.fill.solid()
        rail.fill.fore_color.rgb = self._rgb(brand.colors.accent)
        rail.line.color.rgb = rail.fill.fore_color.rgb
        for idx, step in enumerate(next_steps[:3]):
            y = 2.54 + idx * 1.14
            self._add_badge(slide, str(idx + 1), 7.65, y + 0.05, 0.48, brand, self._icon_fill(brand, idx))
            rule = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(8.38),
                Inches(y + 0.94),
                Inches(3.12),
                Inches(0.018),
            )
            rule.fill.solid()
            rule.fill.fore_color.rgb = self._rgb(self._tint(brand.colors.primary, 0.4))
            rule.line.color.rgb = rule.fill.fore_color.rgb
            self._add_dark_text(
                slide,
                self._truncate_phrase(step, 118),
                8.38,
                y,
                3.62,
                0.84,
                brand,
                size=10,
                color=self._tint(brand.colors.primary, 0.86),
            )

    def _add_closing_recommendation(
        self, slide, outline: SlideOutline, brand: BrandDNA
    ) -> None:
        outline = self._closing_display_outline(outline)
        exhibit = self._exhibit(outline)
        recommendation = (
            exhibit.get("recommendation")
            or outline.content_json.get("action_title")
            or outline.label
        )
        if self._closing_text_is_source_excerpt(str(recommendation)) or self._is_incomplete_display_fragment(str(recommendation)):
            recommendation = self._default_closing_recommendation(outline)
        next_steps = [
            str(item)
            for item in exhibit.get("next_steps", [])
            if str(item).strip() and self._closing_step_is_actionable(str(item))
        ][:3]
        if len(next_steps) < 2:
            next_steps = self._closing_next_steps(outline)
        ask = self._closing_decision_ask(str(exhibit.get("decision_ask") or ""), outline)
        panel = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.85),
            Inches(1.6),
            Inches(5.25),
            Inches(4.65),
        )
        panel.fill.solid()
        panel.fill.fore_color.rgb = self._rgb(brand.colors.primary)
        panel.line.color.rgb = self._rgb(brand.colors.primary)
        self._add_dark_text(slide, "RECOMMENDATION", 1.18, 2.05, 2.2, 0.3, brand, size=10, color=brand.colors.accent)
        self._add_dark_text(
            slide,
            self._truncate_at_word(recommendation, 140),
            1.18,
            2.58,
            4.45,
            1.3,
            brand,
            size=18,
            bold=True,
        )
        self._add_dark_text(
            slide,
            self._truncate_at_word(str(ask), 120),
            1.2,
            4.55,
            4.25,
            0.66,
            brand,
            size=11,
            color=self._tint(brand.colors.primary, 0.75),
        )
        for idx, step in enumerate(next_steps[:3]):
            y = 1.92 + idx * 1.26
            self._add_badge(slide, str(idx + 1), 6.85, y + 0.08, 0.44, brand, self._icon_fill(brand, idx))
            self._add_body_text(
                slide,
                self._truncate_at_word(step, 108),
                7.5,
                y,
                4.5,
                0.58,
                brand,
                size=13,
            )

    def _closing_display_outline(self, outline: SlideOutline) -> SlideOutline:
        title = str(
            outline.content_json.get("action_title")
            or outline.content_json.get("title")
            or outline.label
        )
        if not self._closing_text_is_source_excerpt(title) and not self._is_incomplete_display_fragment(title):
            return outline
        title = self._default_closing_title(outline)
        content = dict(outline.content_json)
        content["action_title"] = title
        return outline.model_copy(update={"label": title, "content_json": content})

    def _default_closing_title(self, outline: SlideOutline) -> str:
        text = self._outline_text_for_closing(outline)
        if any(token in text for token in ("benchmark", "harness", "ground truth", "evaluation")):
            return "Commit to a governed benchmark pilot"
        return "Commit to the next governed operating move"

    def _default_closing_recommendation(self, outline: SlideOutline) -> str:
        text = self._outline_text_for_closing(outline)
        if any(token in text for token in ("benchmark", "harness", "ground truth", "evaluation")):
            return "Launch a governed benchmark pilot tied to validated source evidence."
        return "Approve the first governed pilot and make the evidence standard explicit."

    def _closing_text_is_source_excerpt(self, text: str) -> bool:
        lowered = " ".join(str(text).casefold().split())
        return lowered.startswith(
            (
                "benchmarks have been",
                "commit to models are only as trustworthy",
                "historically,",
                "histopathology",
                "models are only as trustworthy",
                "organizations deploying",
                "the proliferation",
                "this white paper",
                "with the proliferation",
            )
        )

    def _closing_decision_ask(self, ask: str, outline: SlideOutline) -> str:
        cleaned = self._clean_display_text(ask)
        normalized = re.sub(r"[^a-z0-9]+", " ", cleaned.casefold()).strip()
        if normalized in {
            "",
            "confirm owner scope and timing",
            "confirm ownership and timing",
            "confirm owner timing and success measure",
        }:
            text = self._outline_text_for_closing(outline)
            if any(token in text for token in ("benchmark", "harness", "ground truth", "evaluation")):
                return "Approve the first governed benchmark pilot."
            return "Approve the first governed pilot."
        return cleaned

    def _closing_next_steps(self, outline: SlideOutline) -> list[str]:
        bullets = [
            bullet
            for bullet in self._bullets(outline)
            if self._closing_step_is_actionable(bullet)
        ][:3]
        if len(bullets) >= 2:
            return bullets
        text = self._outline_text_for_closing(outline)
        if any(token in text for token in ("benchmark", "harness", "ground truth", "evaluation")):
            return [
                "Select one consequential internal benchmark pilot.",
                "Bind test cases to validated source evidence.",
                "Review failures before expanding coverage.",
            ]
        if "memory" in text or "context" in text:
            return [
                "Name the shared context owner.",
                "Run one workflow with persistent memory.",
                "Refresh the record after each milestone.",
            ]
        return [
            "Select the first governed pilot.",
            "Define the evidence standard.",
            "Review results before scaling.",
        ]

    def _closing_step_is_actionable(self, text: str) -> bool:
        cleaned = self._clean_display_text(text)
        if not cleaned or self._is_incomplete_display_fragment(cleaned):
            return False
        if self._closing_text_is_source_excerpt(cleaned):
            return False
        first = re.findall(r"[A-Za-z][A-Za-z'-]*", cleaned[:32])
        if not first:
            return False
        return first[0].casefold() in {
            "approve",
            "assign",
            "bind",
            "calibrate",
            "commit",
            "define",
            "launch",
            "name",
            "review",
            "run",
            "select",
        }

    def _outline_text_for_closing(self, outline: SlideOutline) -> str:
        content = outline.content_json
        parts = [
            str(content.get("action_title") or ""),
            str(content.get("subheading") or ""),
            str(content.get("summary") or ""),
            str(content.get("speaker_notes") or ""),
            " ".join(str(item) for item in content.get("sources", []) if str(item).strip()),
        ]
        return " ".join(parts).casefold()

    def _add_section_divider(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
        slide_number: int,
    ) -> None:
        background = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0),
            Inches(0),
            Inches(SLIDE_W),
            Inches(SLIDE_H),
        )
        background.fill.solid()
        background.fill.fore_color.rgb = self._rgb(brand.colors.primary)
        background.line.color.rgb = self._rgb(brand.colors.primary)
        title = outline.content_json.get("action_title") or outline.content_json.get("title") or outline.label
        subheading = outline.content_json.get("subheading") or outline.content_json.get("summary", "")
        self._add_dark_text(slide, "SECTION", 0.8, 0.82, 2.0, 0.25, brand, size=10, color=brand.colors.accent)
        self._add_dark_text(slide, title, 0.8, 1.7, 7.4, 1.1, brand, size=30, bold=True)
        self._add_dark_text(slide, subheading, 0.82, 3.12, 6.8, 0.7, brand, size=13, color=self._tint(brand.colors.primary, 0.72))
        rail = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(9.2),
            Inches(0.92),
            Inches(0.12),
            Inches(4.95),
        )
        rail.fill.solid()
        rail.fill.fore_color.rgb = self._rgb(brand.colors.accent)
        rail.line.color.rgb = rail.fill.fore_color.rgb
        for idx, label in enumerate(["Diagnose", "Design", "Execute"]):
            y = 1.0 + idx * 1.52
            marker = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(9.62),
                Inches(y),
                Inches(2.35),
                Inches(0.68),
            )
            marker.fill.solid()
            marker.fill.fore_color.rgb = self._rgb(self._tint(brand.colors.primary, 0.18 + idx * 0.1))
            marker.line.color.rgb = self._rgb(self._tint(brand.colors.primary, 0.34))
            self._add_dark_text(slide, label, 9.86, y + 0.2, 1.6, 0.2, brand, size=10, bold=True)
        number = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(9.62),
            Inches(5.04),
            Inches(1.0),
            Inches(0.62),
        )
        number.fill.solid()
        number.fill.fore_color.rgb = self._rgb(brand.colors.accent)
        number.line.color.rgb = self._rgb(brand.colors.accent)
        self._add_text_in_shape(number, f"{slide_number:02d}", brand, size=15, bold=True, color=brand.colors.text_light, center=True)
        self._add_dark_text(
            slide,
            self._footer_source_text(outline),
            0.8,
            6.92,
            6.0,
            0.25,
            brand,
            size=8,
            color=self._tint(brand.colors.primary, 0.72),
        )
