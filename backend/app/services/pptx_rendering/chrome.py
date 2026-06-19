# ruff: noqa: F401
from pathlib import Path
import math
import re
import shutil
import subprocess
from typing import Any

from PIL import Image, ImageDraw
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt

from app.models.brand import BrandDNA
from app.models.outline import SlideOutline
from app.services.concept_diagram_renderer import ConceptDiagramRenderer, DiagramRenderError
from app.services.pptx_rendering.constants import ICON_SCALE, SLIDE_H, SLIDE_W


class ChromeRenderingMixin:
    def _add_header(
        self,
        slide,
        title: str,
        subheading: str,
        brand: BrandDNA,
        outline: SlideOutline,
        slide_number: int,
    ) -> None:
        subheading = "" if self._looks_like_meta_subheading(subheading) else subheading
        reserved_logo_width = brand.logo.size_w + 0.35 if brand.logo else 0
        section_number, section_label = self._section_marker(outline, slide_number)
        kicker = slide.shapes.add_textbox(Inches(0.68), Inches(0.18), Inches(5.4), Inches(0.2))
        kicker_frame = kicker.text_frame
        kicker_frame.clear()
        kicker_run = kicker_frame.paragraphs[0].add_run()
        kicker_run.text = f"SECTION {section_number} - {section_label}"
        kicker_run.font.name = brand.fonts.body
        kicker_run.font.size = Pt(7.5)
        kicker_run.font.bold = True
        kicker_run.font.color.rgb = self._rgb(brand.colors.accent)

        number_x = max(9.95, SLIDE_W - 0.7 - reserved_logo_width - 0.78)
        number_box = slide.shapes.add_textbox(
            Inches(number_x),
            Inches(0.16),
            Inches(0.72),
            Inches(0.34),
        )
        number_frame = number_box.text_frame
        number_frame.clear()
        number_para = number_frame.paragraphs[0]
        number_para.alignment = 2
        number_run = number_para.add_run()
        number_run.text = section_number
        number_run.font.name = brand.fonts.heading
        number_run.font.size = Pt(18)
        number_run.font.bold = True
        number_run.font.color.rgb = self._rgb(self._tint(brand.colors.primary, 0.84))

        accent = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.68),
            Inches(0.46),
            Inches(0.62),
            Inches(0.035),
        )
        accent.fill.solid()
        accent.fill.fore_color.rgb = self._rgb(brand.colors.accent)
        accent.line.color.rgb = self._rgb(brand.colors.accent)
        rule = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.68),
            Inches(1.12),
            Inches(11.95),
            Inches(0.012),
        )
        rule.fill.solid()
        rule.fill.fore_color.rgb = self._rgb(self._tint(brand.colors.secondary, 0.82))
        rule.line.color.rgb = rule.fill.fore_color.rgb
        title_profile = self._profile_box(brand, "title_box")
        title_x = title_profile.get("x", 0.68)
        title_y = title_profile.get("y", 0.48)
        title_width = title_profile.get("w", 12.1 - reserved_logo_width)
        title_font_size = self._fit_font_size(title, title_width, [24, 20, 18], max_lines=2)
        one_line_chars = max(1, int((title_width * 72) / (0.52 * title_font_size)))
        wraps_two_lines = len(" ".join(title.split())) > one_line_chars
        if not wraps_two_lines:
            title_height, show_subheading = 0.55, True
        elif title_font_size >= 20:
            title_height, show_subheading = 0.86, False
        else:
            title_height, show_subheading = 0.98, False
        title_height = max(title_height, min(float(title_profile.get("h", title_height)), 1.1))
        title_box = slide.shapes.add_textbox(
            Inches(title_x),
            Inches(title_y),
            Inches(title_width),
            Inches(title_height),
        )
        frame = title_box.text_frame
        frame.clear()
        self._autofit(frame)
        para = frame.paragraphs[0]
        para.line_spacing = 1.04
        run = para.add_run()
        run.text = self._truncate_at_word(title, 120)
        run.font.name = brand.fonts.heading
        run.font.size = Pt(title_font_size)
        run.font.bold = True
        run.font.color.rgb = self._rgb(brand.colors.text_dark)
        if subheading and show_subheading:
            sub_y = (
                title_y + title_height + 0.04
                if title_profile
                else 0.88 if title_height <= 0.55 else 0.48 + title_height + 0.04
            )
            sub_box = slide.shapes.add_textbox(
                Inches(title_x), Inches(sub_y), Inches(min(10.6, title_width)), Inches(0.22)
            )
            sub_frame = sub_box.text_frame
            sub_frame.clear()
            sub = sub_frame.paragraphs[0].add_run()
            sub.text = self._truncate_at_word(subheading, 150)
            sub.font.name = brand.fonts.body
            sub.font.size = Pt(9)
            sub.font.color.rgb = self._rgb(brand.colors.secondary)

    def _looks_like_meta_subheading(self, text: str) -> bool:
        normalized = " ".join(str(text or "").lower().split())
        return any(
            marker in normalized
            for marker in (
                "quote sidebar",
                "layout instruction",
                "diagram description",
                "placeholder",
                "visually tied",
                "compact reference block",
                "distinct exhibit",
            )
        )

    def _section_marker(self, outline: SlideOutline, slide_number: int) -> tuple[str, str]:
        role = str(
            outline.layout_json.get("narrative_role")
            or outline.content_json.get("narrative_role")
            or ""
        ).lower()
        archetype = str(
            outline.layout_json.get("archetype")
            or outline.content_json.get("archetype")
            or outline.layout_json.get("layout")
            or ""
        ).lower()
        role_sections = {
            "executive_summary": ("01", "EXECUTIVE SUMMARY"),
            "problem": ("02", "DIAGNOSIS"),
            "evidence": ("03", "EVIDENCE"),
            "framework": ("04", "OPERATING MODEL"),
            "reference": ("05", "REFERENCE SYSTEM"),
            "implementation": ("06", "IMPLEMENTATION"),
            "decision": ("07", "DECISION"),
            "closing": ("08", "RECOMMENDATION"),
        }
        if role in role_sections:
            return role_sections[role]
        archetype_sections = {
            "executive_summary": ("01", "EXECUTIVE SUMMARY"),
            "anti_patterns": ("02", "DIAGNOSIS"),
            "dependency_map": ("03", "EVIDENCE"),
            "chart": ("03", "EVIDENCE"),
            "metric_chart": ("03", "EVIDENCE"),
            "comparison_table": ("03", "EVIDENCE"),
            "framework_cycle": ("04", "OPERATING MODEL"),
            "code_panel": ("05", "REFERENCE SYSTEM"),
            "table_reference": ("05", "REFERENCE SYSTEM"),
            "checklist": ("06", "IMPLEMENTATION"),
            "quote_sidebar": ("07", "DECISION"),
            "closing_recommendation": ("08", "RECOMMENDATION"),
        }
        if archetype in archetype_sections:
            return archetype_sections[archetype]
        coarse_section = min(8, max(1, (slide_number + 1) // 2))
        return f"{coarse_section:02d}", "ANALYSIS"

    def _add_footer(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
        slide_number: int,
        total_slides: int,
    ) -> None:
        source_text = self._footer_source_text(outline)
        footer_box = self._profile_box(brand, "footer_box")
        footer = slide.shapes.add_textbox(
            Inches(footer_box.get("x", 0.68)),
            Inches(footer_box.get("y", 7.0)),
            Inches(footer_box.get("w", 10.7)),
            Inches(footer_box.get("h", 0.26)),
        )
        frame = footer.text_frame
        frame.clear()
        run = frame.paragraphs[0].add_run()
        run.text = source_text[:180]
        run.font.name = brand.fonts.body
        run.font.size = Pt(10)
        run.font.color.rgb = self._rgb(brand.colors.secondary)
        page = slide.shapes.add_textbox(
            Inches(11.52),
            Inches(footer_box.get("y", 7.0)),
            Inches(0.9),
            Inches(footer_box.get("h", 0.26)),
        )
        page_frame = page.text_frame
        page_frame.clear()
        page_para = page_frame.paragraphs[0]
        page_para.alignment = 2
        page_run = page_para.add_run()
        page_run.text = f"{slide_number}/{total_slides}"
        page_run.font.name = brand.fonts.body
        page_run.font.size = Pt(10)
        page_run.font.bold = True
        page_run.font.color.rgb = self._rgb(brand.colors.secondary)

    def _add_dark_slide_chrome(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
        slide_number: int,
        total_slides: int,
        label: str,
    ) -> None:
        background = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0),
            Inches(0),
            Inches(SLIDE_W),
            Inches(SLIDE_H),
        )
        background.fill.solid()
        background.fill.fore_color.rgb = self._rgb(brand.colors.primary)
        background.line.color.rgb = self._rgb(brand.colors.primary)
        title = (
            outline.content_json.get("action_title")
            or outline.content_json.get("title")
            or outline.label
        )
        subheading = outline.content_json.get("subheading") or outline.content_json.get("summary", "")
        subheading = "" if self._looks_like_meta_subheading(subheading) else subheading
        section_number, _section_label = self._section_marker(outline, slide_number)
        self._add_dark_text(
            slide,
            f"SECTION {section_number} / {label}",
            0.78,
            0.48,
            3.8,
            0.24,
            brand,
            size=8,
            bold=True,
            color=brand.colors.accent,
        )
        rail = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.78),
            Inches(0.83),
            Inches(1.08),
            Inches(0.045),
        )
        rail.fill.solid()
        rail.fill.fore_color.rgb = self._rgb(brand.colors.accent)
        rail.line.color.rgb = rail.fill.fore_color.rgb
        title_size = self._fit_font_size(title, 8.8, [27, 23, 20], max_lines=2)
        one_line_chars = max(1, int((8.8 * 72) / (0.52 * title_size)))
        title_height = 0.72 if len(" ".join(title.split())) <= one_line_chars else 1.04
        self._add_dark_text(
            slide,
            self._truncate_at_word(title, 130),
            0.78,
            1.02,
            8.8,
            title_height,
            brand,
            size=title_size,
            bold=True,
        )
        if subheading and len(title) <= 96:
            self._add_dark_text(
                slide,
                self._truncate_at_word(subheading, 150),
                0.8,
                1.92,
                7.6,
                0.36,
                brand,
                size=10,
                color=self._tint(brand.colors.primary, 0.74),
            )
        number = slide.shapes.add_textbox(Inches(11.55), Inches(0.46), Inches(0.95), Inches(0.44))
        number_frame = number.text_frame
        number_frame.clear()
        number_para = number_frame.paragraphs[0]
        number_para.alignment = 2
        number_run = number_para.add_run()
        number_run.text = f"{slide_number:02d}"
        number_run.font.name = brand.fonts.heading
        number_run.font.size = Pt(20)
        number_run.font.bold = True
        number_run.font.color.rgb = self._rgb(self._tint(brand.colors.primary, 0.46))
        self._add_dark_footer(slide, outline, brand, slide_number, total_slides)

    def _add_dark_footer(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
        slide_number: int,
        total_slides: int,
    ) -> None:
        source_text = self._footer_source_text(outline)
        self._add_dark_text(
            slide,
            source_text[:180],
            0.78,
            6.95,
            8.4,
            0.24,
            brand,
            size=9,
            color=self._tint(brand.colors.primary, 0.66),
        )
        self._add_dark_text(
            slide,
            f"{slide_number}/{total_slides}",
            11.45,
            6.95,
            0.88,
            0.24,
            brand,
            size=9,
            bold=True,
            color=self._tint(brand.colors.primary, 0.66),
        )

    def _footer_source_text(self, outline: SlideOutline) -> str:
        sources = (
            outline.content_json.get("source_labels")
            or outline.content_json.get("sources")
            or []
        )
        cleaned = [
            str(source)
            for source in sources
            if str(source).strip()
        ]
        return "Source: " + "; ".join(cleaned[:2]) if cleaned else "Source: [source needed]"

    def _profile_box(self, brand: BrandDNA, key: str) -> dict[str, float]:
        box = (brand.layout_profile or {}).get(key)
        if not isinstance(box, dict):
            return {}
        try:
            return {
                item: float(box[item])
                for item in ("x", "y", "w", "h")
                if item in box
            }
        except (TypeError, ValueError):
            return {}
