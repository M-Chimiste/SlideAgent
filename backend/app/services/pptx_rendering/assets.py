# ruff: noqa: F401
from pathlib import Path
import math
import re
import shutil
import subprocess
from typing import Any

from PIL import Image, ImageDraw
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt

from app.models.brand import BrandDNA
from app.models.outline import SlideOutline
from app.services.concept_diagram_renderer import ConceptDiagramRenderer, DiagramRenderError
from app.services.pptx_rendering.constants import ICON_SCALE, SLIDE_H, SLIDE_W


class AssetRenderingMixin:
    def _try_add_diagram_asset(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
        x: float,
        y: float,
        width: float,
        height: float,
        dark: bool = False,
    ) -> bool:
        renderer = getattr(self, "_diagram_renderer", None)
        output_dir = getattr(self, "_diagram_cache_dir", None)
        if renderer is None or output_dir is None:
            return False
        try:
            layout = str(outline.layout_json.get("layout") or "diagram")
            artifact = renderer.render(
                outline,
                brand,
                output_dir,
                self._diagram_stem(outline, layout),
                dark=dark,
            )
            slide.shapes.add_picture(
                artifact.png_path.as_posix(),
                Inches(x),
                Inches(y),
                Inches(width),
                Inches(height),
            )
            return True
        except (DiagramRenderError, OSError, RuntimeError, ValueError) as exc:
            self._warnings.append(
                {
                    "slide_index": outline.slide_index,
                    "field": "diagram_render",
                    "message": self._truncate_phrase(
                        f"Diagram asset render failed; used native fallback. Reason: {exc}",
                        220,
                    ),
                }
            )
            return False

    def _diagram_stem(self, outline: SlideOutline, layout: str) -> str:
        label = outline.label or outline.content_json.get("action_title") or layout
        slug = re.sub(r"[^a-z0-9]+", "-", str(label).lower()).strip("-")
        return f"{outline.slide_index + 1:02d}-{layout}-{slug[:42] or 'diagram'}"

    def _add_logo(self, slide, brand: BrandDNA) -> None:
        if not brand.logo:
            return
        logo_path = Path(brand.logo.path)
        if not logo_path.exists():
            return
        x = SLIDE_W - 0.6 - brand.logo.size_w
        y = 0.3
        slide.shapes.add_picture(
            logo_path.as_posix(),
            Inches(x),
            Inches(y),
            width=Inches(brand.logo.size_w),
            height=Inches(brand.logo.size_h),
        )

    def _add_icon(
        self,
        slide,
        icon_name: str,
        x: float,
        y: float,
        size: float,
        brand: BrandDNA,
        fill: str,
    ) -> None:
        halo = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(x),
            Inches(y),
            Inches(size),
            Inches(size),
        )
        halo.fill.solid()
        halo.fill.fore_color.rgb = self._rgb(self._tint(fill, 0.84))
        halo.line.color.rgb = self._rgb(self._tint(fill, 0.62))
        inner_inset = size * 0.08
        badge = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(x + inner_inset),
            Inches(y + inner_inset),
            Inches(size - inner_inset * 2),
            Inches(size - inner_inset * 2),
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = self._rgb(fill)
        badge.line.color.rgb = self._rgb(fill)
        icon_size = size * 0.43
        inset = (size - icon_size) / 2
        icon_path = self._icon_image_path(icon_name, "FFFFFF")
        slide.shapes.add_picture(
            icon_path.as_posix(),
            Inches(x + inset),
            Inches(y + inset),
            Inches(icon_size),
            Inches(icon_size),
        )

    def _icon_fill(self, brand: BrandDNA, index: int) -> str:
        palette = [brand.colors.primary, brand.colors.secondary, brand.colors.accent]
        return palette[index % len(palette)]

    def _icon_image_path(self, icon_name: str, color: str) -> Path:
        safe_name = re.sub(r"[^a-z0-9]+", "-", icon_name.lower()).strip("-") or "default"
        safe_color = self._clean_hex(color).lower()
        path = self._icon_cache_dir / f"{safe_name}-{safe_color}.png"
        if not path.exists():
            if not self._draw_icon_with_node(icon_name, color, path):
                self._draw_icon_png(icon_name, color, path)
        return path

    def _draw_icon_with_node(self, icon_name: str, color: str, path: Path) -> bool:
        script_path = Path(__file__).resolve().parents[1] / "workers" / "icon_renderer.js"
        if shutil.which("node") is None or not script_path.exists():
            return False
        try:
            result = subprocess.run(
                [
                    "node",
                    script_path.as_posix(),
                    icon_name,
                    self._clean_hex(color),
                    path.as_posix(),
                ],
                check=False,
                capture_output=True,
                text=True,
                timeout=10,
            )
        except (OSError, subprocess.SubprocessError):
            return False
        return result.returncode == 0 and path.exists() and path.stat().st_size > 0

    def _draw_icon_png(self, icon_name: str, color: str, path: Path) -> None:
        normalized = icon_name.lower()
        size = 128 * ICON_SCALE
        image = Image.new("RGBA", (size, size), (255, 255, 255, 0))
        draw = ImageDraw.Draw(image)
        stroke = self._rgba(color)
        width = 8
        if "chart" in normalized or "growth" in normalized:
            self._draw_chart_icon(draw, stroke, width)
        elif "risk" in normalized or "exclamation" in normalized:
            self._draw_warning_icon(draw, stroke, width)
        elif (
            "database" in normalized
            or "data" in normalized
            or "brain" in normalized
            or "book" in normalized
            or "file" in normalized
        ):
            self._draw_database_icon(draw, stroke, width)
        elif "users" in normalized or "customer" in normalized or "comments" in normalized or "usertie" in normalized:
            self._draw_users_icon(draw, stroke, width)
        elif (
            "shield" in normalized
            or "security" in normalized
            or "clipboard" in normalized
            or "tasks" in normalized
            or "check" in normalized
        ):
            self._draw_shield_icon(draw, stroke, width)
        elif (
            "timeline" in normalized
            or "project" in normalized
            or "robot" in normalized
            or "code" in normalized
            or "sitemap" in normalized
            or "tools" in normalized
        ):
            self._draw_process_icon(draw, stroke, width)
        elif "finance" in normalized or "dollar" in normalized:
            self._draw_finance_icon(draw, stroke, width)
        elif "strategy" in normalized or "chess" in normalized:
            self._draw_strategy_icon(draw, stroke, width)
        elif "status" in normalized or "tachometer" in normalized:
            self._draw_status_icon(draw, stroke, width)
        else:
            self._draw_lightbulb_icon(draw, stroke, width)
        image.save(path)

    def _draw_chart_icon(self, draw: ImageDraw.ImageDraw, color, width: int) -> None:
        self._line(draw, [(24, 96), (24, 30)], color, width)
        self._line(draw, [(24, 96), (104, 96)], color, width)
        self._line(draw, [(36, 78), (54, 58), (72, 66), (100, 34)], color, width)
        self._dot(draw, 36, 78, 4.5, color)
        self._dot(draw, 54, 58, 4.5, color)
        self._dot(draw, 72, 66, 4.5, color)
        self._polygon(draw, [(100, 34), (86, 38), (96, 50)], fill=color)

    def _draw_warning_icon(self, draw: ImageDraw.ImageDraw, color, width: int) -> None:
        self._line(draw, [(64, 22), (108, 100), (20, 100), (64, 22)], color, width, rounded=False)
        self._line(draw, [(64, 48), (64, 76)], color, width)
        self._dot(draw, 64, 92, 5.5, color)

    def _draw_database_icon(self, draw: ImageDraw.ImageDraw, color, width: int) -> None:
        self._ellipse(draw, (26, 22, 102, 50), outline=color, width=width)
        self._line(draw, [(26, 36), (26, 92)], color, width)
        self._line(draw, [(102, 36), (102, 92)], color, width)
        self._arc(draw, (26, 48, 102, 76), 0, 180, color, width)
        self._arc(draw, (26, 68, 102, 96), 0, 180, color, width)
        self._arc(draw, (26, 78, 102, 106), 0, 180, color, width)

    def _draw_users_icon(self, draw: ImageDraw.ImageDraw, color, width: int) -> None:
        self._ellipse(draw, (46, 22, 82, 58), outline=color, width=width)
        self._arc(draw, (34, 54, 94, 116), 205, 335, color, width)
        self._ellipse(draw, (18, 40, 48, 70), outline=color, width=max(6, width - 2))
        self._ellipse(draw, (80, 40, 110, 70), outline=color, width=max(6, width - 2))
        self._arc(draw, (12, 66, 58, 116), 210, 320, color, max(5, width - 3))
        self._arc(draw, (70, 66, 116, 116), 220, 330, color, max(5, width - 3))

    def _draw_shield_icon(self, draw: ImageDraw.ImageDraw, color, width: int) -> None:
        points = [(64, 18), (104, 34), (98, 76), (64, 108), (30, 76), (24, 34)]
        self._line(draw, points + [points[0]], color, width, rounded=False)
        self._line(draw, [(48, 64), (60, 76), (86, 48)], color, width)

    def _draw_process_icon(self, draw: ImageDraw.ImageDraw, color, width: int) -> None:
        self._line(draw, [(42, 42), (86, 42)], color, max(5, width - 2))
        self._line(draw, [(64, 58), (64, 82)], color, max(5, width - 2))
        self._line(draw, [(42, 92), (86, 92)], color, max(5, width - 2))
        self._rounded_rectangle(draw, (18, 24, 50, 60), outline=color, width=width, radius=7)
        self._rounded_rectangle(draw, (78, 24, 110, 60), outline=color, width=width, radius=7)
        self._rounded_rectangle(draw, (48, 76, 80, 112), outline=color, width=width, radius=7)
        self._dot(draw, 34, 42, 3.5, color)
        self._dot(draw, 94, 42, 3.5, color)
        self._dot(draw, 64, 94, 3.5, color)

    def _draw_finance_icon(self, draw: ImageDraw.ImageDraw, color, width: int) -> None:
        self._ellipse(draw, (26, 20, 102, 100), outline=color, width=width)
        self._line(draw, [(64, 34), (64, 90)], color, width)
        self._arc(draw, (44, 36, 82, 68), 90, 270, color, width)
        self._arc(draw, (46, 58, 84, 90), 270, 90, color, width)

    def _draw_strategy_icon(self, draw: ImageDraw.ImageDraw, color, width: int) -> None:
        self._ellipse(draw, (30, 30, 98, 98), outline=color, width=width)
        self._ellipse(draw, (48, 48, 80, 80), outline=color, width=max(5, width - 2))
        self._line(draw, [(64, 18), (64, 42)], color, max(5, width - 2))
        self._line(draw, [(64, 86), (64, 110)], color, max(5, width - 2))
        self._line(draw, [(18, 64), (42, 64)], color, max(5, width - 2))
        self._line(draw, [(86, 64), (110, 64)], color, max(5, width - 2))
        self._line(draw, [(68, 60), (102, 26)], color, width)
        self._polygon(draw, [(102, 26), (92, 28), (100, 38)], fill=color)

    def _draw_status_icon(self, draw: ImageDraw.ImageDraw, color, width: int) -> None:
        self._arc(draw, (22, 26, 106, 110), 200, 340, color, width)
        self._line(draw, [(38, 82), (30, 90)], color, max(5, width - 2))
        self._line(draw, [(90, 82), (100, 90)], color, max(5, width - 2))
        self._line(draw, [(64, 76), (90, 50)], color, width)
        self._dot(draw, 64, 76, 7, color)

    def _draw_lightbulb_icon(self, draw: ImageDraw.ImageDraw, color, width: int) -> None:
        self._ellipse(draw, (40, 22, 88, 76), outline=color, width=width)
        self._line(draw, [(50, 78), (78, 78)], color, width)
        self._line(draw, [(54, 92), (74, 92)], color, width)
        self._line(draw, [(58, 106), (70, 106)], color, width)
        self._line(draw, [(64, 12), (64, 4)], color, max(5, width - 2))
        self._line(draw, [(32, 28), (24, 20)], color, max(5, width - 2))
        self._line(draw, [(96, 28), (104, 20)], color, max(5, width - 2))

