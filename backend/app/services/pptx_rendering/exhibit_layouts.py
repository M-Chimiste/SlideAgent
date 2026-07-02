# ruff: noqa: F401
from pathlib import Path
import math
import re
import shutil
import subprocess
from typing import Any

from PIL import Image, ImageDraw
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt

from app.models.brand import BrandDNA
from app.models.outline import SlideOutline
from app.services.concept_diagram_renderer import ConceptDiagramRenderer, DiagramRenderError
from app.services.pptx_rendering.constants import ICON_SCALE, SLIDE_H, SLIDE_W


class ExhibitLayoutRenderingMixin:
    def _add_matrix_2x2(self, slide, outline: SlideOutline, brand: BrandDNA) -> None:
        exhibit = self._exhibit(outline)
        quadrants = exhibit.get("quadrants") if exhibit.get("type") == "matrix_2x2" else []
        if not self._usable_matrix_quadrants(quadrants):
            self._add_authored_proof_strip(slide, outline, brand)
            return
        x0, y0, w, h = 0.92, 1.58, 10.95, 4.55
        x_mid, y_mid = x0 + w / 2, y0 + h / 2
        fills = [
            self._tint(brand.colors.accent, 0.82),
            "FFFFFF",
            "FFFFFF",
            self._tint(brand.colors.secondary, 0.86),
        ]
        positions = [
            (x0, y0, w / 2, h / 2),
            (x_mid, y0, w / 2, h / 2),
            (x0, y_mid, w / 2, h / 2),
            (x_mid, y_mid, w / 2, h / 2),
        ]
        for idx, item in enumerate(quadrants[:4]):
            x, y, qw, qh = positions[idx]
            rect = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(x),
                Inches(y),
                Inches(qw),
                Inches(qh),
            )
            rect.fill.solid()
            rect.fill.fore_color.rgb = self._rgb(fills[idx])
            rect.line.color.rgb = self._rgb(brand.colors.background_light)
            label = str(item.get("label") if isinstance(item, dict) else item)
            description = str(item.get("description", "") if isinstance(item, dict) else "")
            self._add_body_text(
                slide,
                self._truncate_at_word(label, 46),
                x + 0.28,
                y + 0.32,
                qw - 0.55,
                0.32,
                brand,
                size=12,
            )
            self._add_body_text(
                slide,
                self._truncate_at_word(description, 96),
                x + 0.28,
                y + 0.86,
                qw - 0.55,
                0.78,
                brand,
                size=10,
            )
        self._add_label(slide, str(exhibit.get("y_axis") or "Impact"), 0.88, 1.26, 1.4, brand, bold=True)
        self._add_label(slide, str(exhibit.get("x_axis") or "Readiness"), 5.25, 6.28, 1.8, brand, bold=True)

    def _add_table_or_process(self, slide, outline: SlideOutline, brand: BrandDNA) -> None:
        rows = self._table_rows(outline)
        x, y, w = 0.85, 1.55, 11.75
        col_w = w / max(len(rows[0]), 1)
        row_h = 0.68
        for r_idx, row in enumerate(rows[:7]):
            for c_idx, cell in enumerate(row):
                fill = brand.colors.primary if r_idx == 0 else ("F4F6F8" if r_idx % 2 else "FFFFFF")
                rect = slide.shapes.add_shape(
                    MSO_SHAPE.RECTANGLE,
                    Inches(x + c_idx * col_w),
                    Inches(y + r_idx * row_h),
                    Inches(col_w),
                    Inches(row_h),
                )
                rect.fill.solid()
                rect.fill.fore_color.rgb = self._rgb(fill)
                rect.line.color.rgb = self._rgb(brand.colors.background_light)
                text_color = brand.colors.text_light if r_idx == 0 else brand.colors.text_dark
                self._add_text_in_shape(rect, str(cell), brand, size=9 if r_idx else 10, bold=r_idx == 0, color=text_color)

    def _add_quote_sidebar(self, slide, outline: SlideOutline, brand: BrandDNA) -> None:
        exhibit = self._exhibit(outline)
        bullets = [
            str(item)
            for item in exhibit.get("supporting_points", [])
            if str(item).strip()
        ][:4] or self._bullets(outline)[:4]
        if not bullets:
            bullets = ["Shift the operating model from ad hoc execution to managed discipline."]
        header_texts = [
            self._authored_title(outline),
            str(outline.content_json.get("subheading") or ""),
        ]
        quote = self._first_distinct_sidebar_text(
            [
                exhibit.get("quote"),
                exhibit.get("key_idea"),
                outline.content_json.get("summary"),
                bullets[0],
                self._authored_canvas_support(outline),
            ],
            header_texts,
            fallback=self._authored_canvas_support(outline),
        )
        panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.4), Inches(1.48), Inches(3.75), Inches(4.85))
        panel.fill.solid()
        panel.fill.fore_color.rgb = self._rgb(brand.colors.primary)
        panel.line.color.rgb = self._rgb(brand.colors.primary)
        self._add_dark_text(slide, "KEY IDEA", 8.74, 1.92, 1.3, 0.25, brand, size=9, color=brand.colors.accent)
        self._add_dark_text(
            slide,
            self._truncate_phrase(str(quote), 135),
            8.74,
            2.35,
            3.0,
            1.3,
            brand,
            size=16,
            bold=True,
        )
        support = self._first_distinct_sidebar_text(
            [self._authored_canvas_support(outline), *bullets[1:], *bullets[:1]],
            [*header_texts, str(quote)],
            fallback=self._topic_support_sentence(outline),
        )
        self._add_dark_text(
            slide,
            self._truncate_at_word(support, 105),
            8.76,
            4.25,
            2.9,
            0.9,
            brand,
            size=10,
            color=self._tint(brand.colors.primary, 0.72),
        )
        icons = self._icons(outline)
        sidebar_items = [
            text
            for text in bullets
            if not self._authored_text_overlaps(str(text), str(quote))
            and not self._authored_text_overlaps(str(text), str(support))
        ][:4]
        if not sidebar_items:
            sidebar_items = [
                text
                for text in self._authored_filler_items(outline)
                if not self._authored_text_overlaps(str(text), str(quote))
                and not self._authored_text_overlaps(str(text), str(support))
            ][:3]
        for idx, text in enumerate(sidebar_items):
            y = 1.55 + idx * 1.08
            self._add_icon(slide, icons[idx], 0.95, y, 0.7, brand, self._icon_fill(brand, idx))
            connector = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(7.58),
                Inches(y + 0.36),
                Inches(8.4),
                Inches(3.1),
            )
            connector.line.color.rgb = self._rgb(self._tint(brand.colors.primary, 0.58))
            connector.line.width = Pt(0.8)
            self._add_body_text(
                slide,
                self._truncate_phrase(text, 120),
                1.88,
                y + 0.02,
                5.55,
                0.72,
                brand,
                size=13,
            )

    def _first_distinct_sidebar_text(
        self,
        candidates: list[Any],
        existing: list[str],
        fallback: str,
    ) -> str:
        for candidate in candidates:
            text = self._clean_display_text(str(candidate or ""))
            if not text:
                continue
            if any(self._authored_text_overlaps(text, value) for value in existing):
                continue
            return text
        return fallback

    def _add_framework_cycle(self, slide, outline: SlideOutline, brand: BrandDNA) -> None:
        if self._try_add_diagram_asset(
            slide, outline, brand, 0.78, 1.36, 11.75, 5.08
        ):
            return
        exhibit = self._exhibit(outline)
        steps = exhibit.get("steps") if exhibit.get("type") == "cycle" else None
        if isinstance(steps, list) and steps:
            bullets = [
                str(step.get("label") or step.get("description") or "")
                for step in steps
                if isinstance(step, dict)
            ][:6]
        else:
            bullets = self._bullets(outline)[:6]
        if len(bullets) < 4 or self._labels_are_generic_visual_fallbacks(bullets):
            self._add_authored_proof_strip(slide, outline, brand)
            return
        center_x, center_y = 6.65, 3.85
        center = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(center_x - 0.78), Inches(center_y - 0.78), Inches(1.56), Inches(1.56))
        center.fill.solid()
        center.fill.fore_color.rgb = self._rgb(brand.colors.primary)
        center.line.color.rgb = self._rgb(brand.colors.primary)
        self._add_text_in_shape(
            center,
            self._truncate_at_word(str(exhibit.get("center_label") or "CYCLE"), 24),
            brand,
            size=13,
            bold=True,
            color=brand.colors.text_light,
            center=True,
        )
        icons = self._icons(outline)
        points: list[tuple[float, float]] = []
        for idx, text in enumerate(bullets):
            angle = -math.pi / 2 + idx * (2 * math.pi / len(bullets))
            x = center_x + math.cos(angle) * 3.65
            y = center_y + math.sin(angle) * 2.0
            points.append((x, y))
            self._add_card(slide, x - 1.3, y - 0.45, 2.6, 0.9, "FFFFFF", brand.colors.background_light)
            self._add_icon(slide, icons[idx % len(icons)], x - 1.14, y - 0.31, 0.52, brand, self._icon_fill(brand, idx))
            self._add_body_text(
                slide,
                self._truncate_at_word(text, 42),
                x - 0.46,
                y - 0.28,
                1.55,
                0.45,
                brand,
                size=10,
            )
        for idx, (x1, y1) in enumerate(points):
            x2, y2 = points[(idx + 1) % len(points)]
            self._add_arrow(slide, x1, y1, x2, y2, brand.colors.secondary, width=1.8)
            mid_x = (x1 + x2) / 2
            mid_y = (y1 + y2) / 2
            self._add_arrowhead(slide, mid_x, mid_y, x2 - x1, y2 - y1, brand.colors.secondary)

    def _add_dependency_map(self, slide, outline: SlideOutline, brand: BrandDNA) -> None:
        if self._try_add_diagram_asset(
            slide, outline, brand, 0.78, 1.38, 11.75, 5.05
        ):
            return
        exhibit = self._exhibit(outline)
        middle_nodes = [
            str(item)
            for item in exhibit.get("middle_nodes", [])
            if str(item).strip()
        ][:3]
        bullets = self._bullets(outline)[:5]
        left_label = str(exhibit.get("left_node") or "Source context")
        right_label = str(exhibit.get("right_outcome") or "Reliable next session")
        if (
            len(middle_nodes) < 2
            or self._labels_are_generic_visual_fallbacks(
                [left_label, right_label, *middle_nodes, *bullets[:3]]
            )
        ):
            self._add_authored_proof_strip(slide, outline, brand)
            return
        if len(middle_nodes) < 3:
            middle_nodes.extend(bullets[: 3 - len(middle_nodes)])
        if len(middle_nodes) < 3:
            self._add_authored_proof_strip(slide, outline, brand)
            return
        connector_labels = [
            str(item)
            for item in exhibit.get("connector_labels", [])
            if str(item).strip()
        ]
        left = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.78), Inches(3.0), Inches(2.55), Inches(0.88))
        left.fill.solid()
        left.fill.fore_color.rgb = self._rgb(brand.colors.primary)
        left.line.color.rgb = self._rgb(brand.colors.primary)
        self._add_text_in_shape(
            left,
            self._truncate_at_word(left_label, 42),
            brand,
            size=10,
            bold=True,
            color=brand.colors.text_light,
            center=True,
        )
        mid_positions = [(4.15, 1.95), (4.15, 3.22), (4.15, 4.49)]
        for idx, (x, y) in enumerate(mid_positions):
            self._add_card(slide, x, y, 2.9, 0.92, "FFFFFF", brand.colors.secondary)
            self._add_body_text(
                slide,
                self._truncate_phrase(middle_nodes[idx], 48),
                x + 0.22,
                y + 0.2,
                2.45,
                0.36,
                brand,
                size=10,
            )
            self._add_arrow(slide, 3.36, 3.44, x - 0.05, y + 0.46, brand.colors.secondary, width=2.0)
            if idx < len(connector_labels):
                self._add_body_text(
                    slide,
                    connector_labels[idx],
                    3.42,
                    y + 0.06,
                    0.62,
                    0.24,
                    brand,
                    center=True,
                    size=7,
                )
        right = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.25), Inches(2.96), Inches(3.55), Inches(0.96))
        right.fill.solid()
        right.fill.fore_color.rgb = self._rgb(self._tint(brand.colors.accent, 0.84))
        right.line.color.rgb = self._rgb(brand.colors.accent)
        self._add_body_text(
            slide,
            self._truncate_phrase(right_label, 58),
            8.46,
            3.18,
            3.1,
            0.42,
            brand,
            size=10,
            center=True,
        )
        for _idx, (_x, y) in enumerate(mid_positions):
            self._add_arrow(slide, 7.08, y + 0.46, 8.18, 3.44, brand.colors.accent, width=2.0)

    def _usable_matrix_quadrants(self, quadrants) -> bool:
        if not isinstance(quadrants, list) or len(quadrants) < 4:
            return False
        labels = [
            str(item.get("label") if isinstance(item, dict) else item).strip()
            for item in quadrants[:4]
        ]
        if len([label for label in labels if label]) < 4:
            return False
        return not self._labels_are_generic_matrix(labels)

    def _labels_are_generic_matrix(self, labels: list[str]) -> bool:
        generic = {
            "high impact / high readiness",
            "high impact / low readiness",
            "low impact / high readiness",
            "low impact / low readiness",
        }
        normalized = {" ".join(label.casefold().split()) for label in labels if label.strip()}
        return len(normalized.intersection(generic)) >= 2

    def _labels_are_generic_visual_fallbacks(self, labels: list[str]) -> bool:
        generic = {
            "frame",
            "ground",
            "build",
            "prime",
            "generate",
            "review",
            "update",
            "reset",
            "persist",
            "source context",
            "rules",
            "memory",
            "reliable output",
            "operating loop",
            "reliable next session",
            "product context",
            "system patterns",
            "active decisions",
        }
        normalized = {" ".join(str(label).casefold().split()) for label in labels if str(label).strip()}
        return len(normalized.intersection(generic)) >= 3

    def _add_checklist(self, slide, outline: SlideOutline, brand: BrandDNA) -> None:
        exhibit = self._exhibit(outline)
        items = exhibit.get("items") if exhibit.get("type") == "checklist" else None
        if isinstance(items, list) and items:
            bullets = [
                str(item.get("action") or "")
                for item in items
                if isinstance(item, dict) and str(item.get("action") or "").strip()
            ][:6]
            owners = [
                str(item.get("owner") or "")
                for item in items
                if isinstance(item, dict)
            ][:6]
            timings = [
                str(item.get("timing") or "")
                for item in items
                if isinstance(item, dict)
            ][:6]
        else:
            bullets = self._bullets(outline)[:6]
            owners = []
            timings = []
        bullets = self._complete_operating_rules(outline, bullets, minimum=4)[:6]
        if not bullets:
            bullets = self._complete_operating_rules(outline, [], minimum=4)[:4]
        for idx, text in enumerate(bullets[:6]):
            y = 1.5 + idx * 0.76
            badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.98), Inches(y), Inches(0.38), Inches(0.38))
            badge.fill.solid()
            badge.fill.fore_color.rgb = self._rgb(self._icon_fill(brand, idx))
            badge.line.color.rgb = badge.fill.fore_color.rgb
            self._add_text_in_shape(
                badge,
                str(idx + 1),
                brand,
                size=8,
                bold=True,
                color=brand.colors.text_light,
                center=True,
            )
            detail = ""
            if idx < len(owners) and owners[idx]:
                detail = owners[idx]
            if idx < len(timings) and timings[idx]:
                detail = f"{detail} / {timings[idx]}".strip(" /")
            if self._is_placeholder_owner_timing(detail):
                detail = ""
            body = self._checklist_body_text(text, detail)
            self._add_body_text(
                slide,
                body,
                1.58,
                y - 0.03,
                6.2,
                0.43,
                brand,
                size=12,
            )
        panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.45), Inches(1.62), Inches(3.55), Inches(3.92))
        panel.fill.solid()
        panel.fill.fore_color.rgb = self._rgb(brand.colors.primary)
        panel.line.color.rgb = self._rgb(brand.colors.primary)
        kicker, headline, support = self._emphasis_panel_text(
            outline,
            [],
            default_headline="Use the checklist as an operating cadence, not a one-time readout.",
            default_support="Assign ownership, timing, and evidence checks before scaling.",
        )
        self._add_dark_text(slide, kicker, 8.82, 2.0, 2.9, 0.28, brand, size=10, color=brand.colors.accent)
        self._add_dark_text(slide, headline, 8.82, 2.52, 2.8, 1.5, brand, size=16, bold=True)
        self._add_dark_text(slide, support, 8.84, 4.55, 2.8, 0.9, brand, size=9, color=self._tint(brand.colors.primary, 0.72))

    def _is_placeholder_owner_timing(self, detail: str) -> bool:
        normalized = " ".join(str(detail).casefold().split())
        return normalized in {"owner", "next", "owner / next", "owner/next"}

    def _checklist_body_text(self, text: str, detail: str) -> str:
        if not detail:
            return self._truncate_at_word(text, 118)
        suffix = f" ({detail})"
        action_limit = max(48, 118 - len(suffix))
        action = self._truncate_at_word(text.rstrip("."), action_limit).rstrip(".")
        return f"{action}{suffix}"

    def _add_code_panel(self, slide, outline: SlideOutline, brand: BrandDNA) -> None:
        exhibit = self._exhibit(outline)
        artifact_title = str(exhibit.get("title") or "operating-rules.md")
        artifact_label = self._truncate_at_word(artifact_title.rsplit("/", 1)[-1], 28)
        lines = [
            str(item)
            for item in exhibit.get("lines", [])
            if str(item).strip()
        ][:5]
        bullets = lines[:4] or self._bullets(outline)[:4]
        bullets = self._complete_operating_rules(outline, bullets, minimum=3)[:4]
        if not bullets:
            bullets = self._complete_operating_rules(outline, [], minimum=3)[:3]
        self._add_label(slide, "OPERATING RULES", 0.94, 1.45, 2.5, brand, bold=True)
        for idx, text in enumerate(bullets[:3]):
            y = 1.82 + idx * 1.04
            self._add_card(slide, 0.88, y, 5.68, 0.82, "FFFFFF", brand.colors.background_light)
            self._add_badge(slide, str(idx + 1), 1.12, y + 0.2, 0.42, brand, self._icon_fill(brand, idx))
            self._add_body_text(
                slide,
                self._truncate_at_word(text, 112),
                1.72,
                y + 0.18,
                4.45,
                0.38,
                brand,
                size=11,
            )
        key_rule = self._supporting_rule_text(outline, bullets)
        self._add_card(slide, 0.88, 5.12, 5.68, 0.74, self._tint(brand.colors.accent, 0.86), self._tint(brand.colors.accent, 0.7))
        self._add_body_text(slide, "OPERATING TEST", 1.16, 5.22, 1.65, 0.2, brand, size=8)
        self._add_body_text(
            slide,
            self._truncate_at_word(key_rule, 88),
            1.16,
            5.42,
            5.0,
            0.24,
            brand,
            size=8,
        )

        panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.02), Inches(1.5), Inches(5.18), Inches(4.18))
        panel.fill.solid()
        panel.fill.fore_color.rgb = self._rgb("111827")
        panel.line.color.rgb = self._rgb("111827")
        header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.02), Inches(1.5), Inches(5.18), Inches(0.52))
        header.fill.solid()
        header.fill.fore_color.rgb = self._rgb("1F2937")
        header.line.color.rgb = header.fill.fore_color.rgb
        self._add_dark_text(
            slide,
            "REFERENCE ARTIFACT",
            7.28,
            1.66,
            1.75,
            0.2,
            brand,
            size=8,
            color=brand.colors.accent,
        )
        self._add_dark_text(
            slide,
            self._truncate_at_word(artifact_title, 42),
            9.1,
            1.66,
            2.55,
            0.22,
            brand,
            size=8,
            color=self._tint("111827", 0.74),
        )
        code_lines = [f"# {artifact_title}"] + [
            "- " + self._truncate_at_word(str(bullet), 52) for bullet in bullets[:5]
        ]
        self._add_code_text(slide, code_lines, 7.34, 2.22, 4.34, 2.5, brand)

        chip_specs = [
            ("ARTIFACT", artifact_label),
            ("REVIEW GATE", "Evidence check"),
            ("UPDATE TRIGGER", "Material change"),
        ]
        for idx, (label, value) in enumerate(chip_specs):
            x = 7.18 + idx * 1.62
            chip = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(x),
                Inches(4.88),
                Inches(1.42),
                Inches(0.52),
            )
            chip.fill.solid()
            chip.fill.fore_color.rgb = self._rgb("1F2937")
            chip.line.color.rgb = self._rgb(self._tint("111827", 0.22))
            self._add_dark_text(slide, label, x + 0.12, 4.97, 1.02, 0.14, brand, size=6, color=brand.colors.accent)
            self._add_dark_text(slide, value, x + 0.12, 5.13, 1.16, 0.16, brand, size=7, color=self._tint("111827", 0.78))

    def _supporting_rule_text(
        self,
        outline: SlideOutline,
        bullets: list[str],
    ) -> str:
        if len(bullets) > 3:
            return bullets[3]
        if len(bullets) > 1:
            return bullets[-1]
        title = str(
            outline.content_json.get("action_title")
            or outline.content_json.get("title")
            or outline.label
        )
        return f"Test every rule against the slide decision: {title}"

    def _first_content_text(self, values: list, fallback: str) -> str:
        for value in values:
            text = str(value or "").strip()
            if text and not self._looks_like_meta_instruction(text):
                return text
        return fallback

    def _looks_like_meta_instruction(self, text: str) -> bool:
        normalized = " ".join(text.lower().split())
        return any(
            marker in normalized
            for marker in (
                "quote sidebar",
                "layout instruction",
                "diagram description",
                "placeholder",
                "visually tied",
                "source-grounded evidence",
                "highlighting",
            )
        )

    def _complete_operating_rules(
        self,
        outline: SlideOutline,
        rules: list[str],
        minimum: int,
    ) -> list[str]:
        completed: list[str] = []
        seen: set[str] = set()
        for rule in rules:
            cleaned = self._clean_display_text(str(rule)).rstrip(".")
            if not cleaned:
                continue
            key = cleaned.casefold()
            if key in seen:
                continue
            seen.add(key)
            completed.append(cleaned + ".")
        if len(completed) >= minimum:
            return completed
        intent = " ".join(
            str(value)
            for value in [
                outline.label,
                outline.content_json.get("action_title", ""),
                outline.content_json.get("subheading", ""),
                outline.content_json.get("summary", ""),
                outline.layout_json.get("layout", ""),
            ]
        ).lower()
        if any(token in intent for token in ("memory", "context", "external brain")):
            candidates = [
                "Update active context after material changes.",
                "Record decisions before starting the next session.",
                "Keep progress notes synchronized with implementation status.",
                "Refresh rules when ownership, risk, or scope changes.",
            ]
        elif any(token in intent for token in ("spec", "acceptance", "criteria")):
            candidates = [
                "Write acceptance criteria before generation begins.",
                "Attach source evidence to each material claim.",
                "Review output against the defined test before approval.",
                "Revise the specification when assumptions change.",
            ]
        elif any(token in intent for token in ("cycle", "loop", "workflow", "reset")):
            candidates = [
                "Load source context before planning the work.",
                "Verify output against the acceptance test.",
                "Update the shared record before resetting context.",
                "Repeat the loop only after evidence is captured.",
            ]
        else:
            candidates = [
                "State the decision before execution begins.",
                "Attach source context to the working artifact.",
                "Run review before expanding the workflow.",
                "Name the owner for the next operating step.",
            ]
        for candidate in candidates:
            key = candidate.rstrip(".").casefold()
            if key in seen:
                continue
            seen.add(key)
            completed.append(candidate)
            if len(completed) >= minimum:
                break
        return completed

    def _add_anti_patterns(self, slide, outline: SlideOutline, brand: BrandDNA) -> None:
        exhibit = self._exhibit(outline)
        patterns = exhibit.get("patterns") if exhibit.get("type") == "anti_patterns" else None
        if not isinstance(patterns, list) or not patterns:
            bullets = self._bullets(outline)[:4]
            patterns = [
                {
                    "name": self._truncate_at_word(text, 28),
                    "symptom": text,
                    "consequence": "Reliability drops.",
                    "better_behavior": "Add a durable rule.",
                }
                for text in bullets
            ]
        if not patterns:
            patterns = [
                {"name": "Chat drift", "symptom": "Relying on chat history", "better_behavior": "Persist context"},
                {"name": "Thin sourcing", "symptom": "Skipping source checks", "better_behavior": "Cite evidence"},
                {"name": "Blind accept", "symptom": "Accepting code blindly", "better_behavior": "Run QA gates"},
            ]
        patterns = [pattern for pattern in patterns if isinstance(pattern, dict)]
        patterns = self._safe_anti_pattern_patterns(patterns)
        if not patterns:
            self._add_grid(slide, self._anti_pattern_fallback_outline(outline), brand)
            return
        self._pad_safe_anti_patterns(patterns)
        icons = self._icons(outline)
        count = min(len(patterns), 4)
        card_w = 11.4 / count - 0.24
        for idx, pattern in enumerate(patterns[:4]):
            if not isinstance(pattern, dict):
                continue
            x = 0.86 + idx * (card_w + 0.32)
            self._add_card(slide, x, 1.72, card_w, 3.85, "FFFFFF", brand.colors.background_light)
            self._add_icon(slide, icons[idx], x + 0.22, 1.98, 0.68, brand, self._icon_fill(brand, idx))
            self._add_label(
                slide,
                self._truncate_at_word(str(pattern.get("name") or f"Risk {idx + 1}"), 26),
                x + 1.02,
                2.14,
                card_w - 1.22,
                brand,
                bold=True,
            )
            symptom = self._truncate_phrase(str(pattern.get("symptom") or ""), 76)
            behavior = self._truncate_phrase(str(pattern.get("better_behavior") or ""), 74)
            body_parts = []
            if symptom:
                body_parts.append(f"Symptom: {symptom}")
            if behavior:
                body_parts.append(f"Better move: {behavior}")
            body = "\n".join(body_parts)
            self._add_body_text(slide, body, x + 0.32, 2.88, card_w - 0.58, 1.72, brand, size=10)

    def _safe_anti_pattern_patterns(self, patterns: list[dict[str, Any]]) -> list[dict[str, Any]]:
        safe: list[dict[str, Any]] = []
        for pattern in patterns:
            name = self._clean_display_text(str(pattern.get("name") or ""))
            symptom = self._clean_display_text(str(pattern.get("symptom") or ""))
            behavior = self._clean_display_text(str(pattern.get("better_behavior") or ""))
            if self._anti_pattern_text_is_broken(name, max_words=7):
                continue
            if symptom and self._anti_pattern_text_is_broken(symptom, max_words=9):
                continue
            if behavior and self._anti_pattern_text_is_broken(behavior, max_words=9):
                continue
            if not symptom and not behavior:
                continue
            safe.append(
                {
                    **pattern,
                    "name": name,
                    "symptom": symptom,
                    "better_behavior": behavior,
                }
            )
        return safe

    def _pad_safe_anti_patterns(self, patterns: list[dict[str, Any]]) -> None:
        defaults = [
            {"name": "Chat drift", "symptom": "Relying on chat history", "better_behavior": "Persist context"},
            {"name": "Thin sourcing", "symptom": "Skipping source checks", "better_behavior": "Cite evidence"},
            {"name": "Blind accept", "symptom": "Accepting code blindly", "better_behavior": "Run QA gates"},
        ]
        seen = {str(pattern.get("name") or "").casefold() for pattern in patterns}
        for default in defaults:
            if len(patterns) >= 3:
                return
            key = str(default["name"]).casefold()
            if key in seen:
                continue
            seen.add(key)
            patterns.append(default)

    def _anti_pattern_text_is_broken(self, text: str, max_words: int) -> bool:
        cleaned = " ".join(str(text).split()).strip()
        if not cleaned:
            return True
        if self._is_incomplete_display_fragment(cleaned):
            return True
        lowered = cleaned.casefold()
        if "..." in cleaned or "…" in cleaned:
            return True
        if cleaned.count("(") != cleaned.count(")"):
            return True
        if re.match(r"^(?:and|or|but|to|of|with|while|instead)\b", lowered):
            return True
        if re.search(
            r"\b(?:a|an|already|and|as|been|being|by|for|from|in|into|of|or|rather|single|the|their|through|to|with|without)\.?$",
            cleaned,
            re.IGNORECASE,
        ):
            return True
        if len(cleaned.split()) > max_words and not re.search(r"[.!?)]$", cleaned):
            return True
        source_prose_starts = (
            "benchmarks have been",
            "historically,",
            "with the proliferation",
            "while this approach",
            "if we have capable",
        )
        return lowered.startswith(source_prose_starts)

    def _anti_pattern_fallback_outline(self, outline: SlideOutline) -> SlideOutline:
        content = dict(outline.content_json)
        filler_items = self._authored_filler_items(outline)[:4]
        content.update(
            {
                "bullets": filler_items,
                "exhibit_spec": {"type": "callouts", "points": filler_items},
                "visual_degradation": {
                    "from": "anti_patterns",
                    "to": "callouts",
                    "reason": "Anti-pattern cards require complete named failure modes.",
                },
            }
        )
        return outline.model_copy(update={"content_json": content})

    def _add_grid(self, slide, outline: SlideOutline, brand: BrandDNA) -> None:
        if (
            outline.content_json.get("visual_qa_source_repair")
            and outline.layout_json.get("composition_family") == "source_repair_cards"
        ):
            self._add_source_repair_cards(slide, outline, brand)
            return
        bullets = self._grid_items(outline)[:4]
        if not bullets:
            fallback = (
                outline.content_json.get("subheading")
                or outline.content_json.get("summary")
                or outline.content_json.get("speaker_notes")
                or outline.label
            )
            bullets = [str(fallback)]
        icons = self._icons(outline)
        if len(bullets) == 1:
            self._add_card(slide, 1.1, 2.0, 11.2, 2.5, brand.colors.background_light, brand.colors.secondary)
            self._add_icon(slide, icons[0], 1.48, 2.42, 1.05, brand, brand.colors.primary)
            self._add_body_text(slide, bullets[0], 2.78, 2.7, 8.45, 1.0, brand, center=False, size=16)
            return
        if len(bullets) in {2, 3}:
            card_w = 5.35 if len(bullets) == 2 else 3.55
            gap = 0.55
            card_h = 3.05
            card_y = 2.05
            total_w = len(bullets) * card_w + (len(bullets) - 1) * gap
            start_x = (SLIDE_W - total_w) / 2
            for idx, text in enumerate(bullets):
                x = start_x + idx * (card_w + gap)
                accent = self._icon_fill(brand, idx)
                self._add_card(slide, x, card_y, card_w, card_h, "FFFFFF", brand.colors.background_light)
                strip = slide.shapes.add_shape(
                    MSO_SHAPE.RECTANGLE,
                    Inches(x),
                    Inches(card_y),
                    Inches(card_w),
                    Inches(0.1),
                )
                strip.fill.solid()
                strip.fill.fore_color.rgb = self._rgb(accent)
                strip.line.color.rgb = strip.fill.fore_color.rgb
                self._add_icon(slide, icons[idx], x + 0.34, card_y + 0.32, 0.86, brand, accent)
                lead, rest = self._split_lead(text)
                self._add_label(slide, lead, x + 1.36, card_y + 0.52, card_w - 1.6, brand, bold=True)
                self._add_body_text(
                    slide,
                    rest or lead,
                    x + 0.4,
                    card_y + 1.42,
                    card_w - 0.8,
                    card_h - 1.6,
                    brand,
                    center=False,
                    size=13,
                )
            return
        for idx, text in enumerate(bullets):
            row = idx // 2
            col = idx % 2
            x = 0.9 + col * 6.0
            y = 1.55 + row * 2.25
            self._add_card(slide, x, y, 5.25, 1.85, "FFFFFF", brand.colors.background_light)
            self._add_icon(slide, icons[idx], x + 0.28, y + 0.25, 0.86, brand, self._icon_fill(brand, idx))
            self._add_body_text(slide, text, x + 1.34, y + 0.36, 3.52, 1.0, brand, size=13)

    def _grid_items(self, outline: SlideOutline) -> list[str]:
        exhibit = self._exhibit(outline)
        items: list[str] = []
        structured_items = exhibit.get("items")
        if isinstance(structured_items, list):
            items.extend(self._coerce_item_text(item) for item in structured_items)
        items.extend(self._bullets(outline))
        for key in ("points", "supporting_points", "next_steps", "lines", "rules"):
            value = exhibit.get(key)
            if isinstance(value, list):
                items.extend(self._coerce_item_text(item) for item in value)
        rows = exhibit.get("rows")
        if isinstance(rows, list):
            items.extend(self._coerce_item_text(row) for row in rows)
        deduped: list[str] = []
        seen: set[str] = set()
        for item in items:
            cleaned = self._complete_display_item(str(item))
            if not cleaned or self._is_placeholder_bullet(cleaned):
                continue
            key = cleaned.casefold()
            if key in seen:
                continue
            seen.add(key)
            deduped.append(cleaned)
        return deduped

    def _add_icon_rows(self, slide, outline: SlideOutline, brand: BrandDNA) -> None:
        if (
            outline.content_json.get("visual_qa_source_repair")
            and outline.layout_json.get("composition_family") == "source_repair_cards"
        ):
            self._add_source_repair_cards(slide, outline, brand)
            return
        bullets = self._grid_items(outline)[:4]
        if len(bullets) < 2:
            # A single point reads as a near-empty slide in a multi-row layout;
            # render it as one deliberate statement panel instead.
            self._add_grid(slide, outline, brand)
            return
        icons = self._icons(outline)
        for idx, text in enumerate(bullets[:4]):
            y = 1.45 + idx * 1.25
            self._add_card(slide, 0.9, y, 11.7, 0.98, "FFFFFF", brand.colors.background_light)
            self._add_icon(
                slide,
                icons[idx],
                1.2,
                y + 0.12,
                0.82,
                brand,
                self._icon_fill(brand, idx),
            )
            self._add_body_text(slide, text, 2.24, y + 0.25, 9.55, 0.48, brand, size=14)

    def _add_source_repair_cards(
        self,
        slide,
        outline: SlideOutline,
        brand: BrandDNA,
    ) -> None:
        points = self._source_repair_points(outline)
        variant = outline.slide_index % 4
        if variant == 0:
            self._add_source_repair_split(slide, points, brand)
        elif variant == 1:
            self._add_source_repair_stage_row(slide, points, brand)
        elif variant == 2:
            self._add_source_repair_editorial_cards(slide, points, brand)
        else:
            self._add_source_repair_proof_strip(slide, points, brand)

    def _source_repair_points(self, outline: SlideOutline) -> list[str]:
        points = self._grid_items(outline)[:4]
        if not points:
            fallback = (
                outline.content_json.get("subheading")
                or outline.content_json.get("summary")
                or outline.label
            )
            points = [str(fallback)]
        cleaned: list[str] = []
        for point in points:
            text = self._truncate_phrase(point, 132)
            if text and not text.endswith((".", "?", "!")):
                text = f"{text}."
            if text:
                cleaned.append(text)
        return cleaned or ["Source evidence should be reviewed before scaling the decision."]

    def _add_source_repair_split(
        self,
        slide,
        points: list[str],
        brand: BrandDNA,
    ) -> None:
        panel = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.88),
            Inches(1.56),
            Inches(3.28),
            Inches(4.72),
        )
        panel.fill.solid()
        panel.fill.fore_color.rgb = self._rgb(brand.colors.primary)
        panel.line.color.rgb = self._rgb(brand.colors.primary)
        self._add_dark_text(
            slide,
            "SOURCE-BACKED SHIFT",
            1.16,
            1.94,
            2.3,
            0.22,
            brand,
            size=8,
            bold=True,
            color=brand.colors.accent,
        )
        self._add_dark_text(
            slide,
            f"{len(points[:3])} signals",
            1.16,
            2.42,
            2.1,
            0.46,
            brand,
            size=22,
            bold=True,
            color=brand.colors.text_light,
        )
        self._add_dark_text(
            slide,
            self._truncate_phrase(points[0], 118),
            1.18,
            3.26,
            2.42,
            1.24,
            brand,
            size=12,
            color=self._tint(brand.colors.primary, 0.78),
        )
        self._add_dark_text(
            slide,
            "Use the source evidence as the review gate, not as decorative slide filler.",
            1.18,
            5.3,
            2.34,
            0.48,
            brand,
            size=8,
            color=self._tint(brand.colors.primary, 0.64),
        )
        for idx, point in enumerate(points[:3]):
            y = 1.72 + idx * 1.42
            accent = self._icon_fill(brand, idx)
            self._add_card(slide, 4.72, y, 7.5, 1.02, "FFFFFF", brand.colors.background_light)
            self._add_badge(slide, f"{idx + 1:02d}", 5.0, y + 0.26, 0.46, brand, accent)
            self._add_body_text(
                slide,
                point,
                5.72,
                y + 0.22,
                5.94,
                0.44,
                brand,
                size=12,
            )

    def _add_source_repair_stage_row(
        self,
        slide,
        points: list[str],
        brand: BrandDNA,
    ) -> None:
        count = min(3, len(points))
        start_x = 1.0
        card_w = 3.55
        gap = 0.55
        y = 2.02
        for idx, point in enumerate(points[:count]):
            x = start_x + idx * (card_w + gap)
            accent = self._icon_fill(brand, idx)
            card_y = y + (0.28 if idx == 1 else 0)
            self._add_card(slide, x, card_y, card_w, 3.38, "FFFFFF", brand.colors.background_light)
            strip = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(x),
                Inches(card_y),
                Inches(card_w),
                Inches(0.14),
            )
            strip.fill.solid()
            strip.fill.fore_color.rgb = self._rgb(accent)
            strip.line.color.rgb = strip.fill.fore_color.rgb
            self._add_dark_text(
                slide,
                f"STAGE {idx + 1}",
                x + 0.34,
                card_y + 0.46,
                1.4,
                0.2,
                brand,
                size=8,
                bold=True,
                color=accent,
            )
            self._add_body_text(
                slide,
                point,
                x + 0.34,
                card_y + 1.08,
                card_w - 0.68,
                1.34,
                brand,
                size=13,
            )
            if idx < count - 1:
                arrow_y = card_y + 1.72
                self._add_arrow(
                    slide,
                    x + card_w + 0.12,
                    arrow_y,
                    x + card_w + gap - 0.14,
                    arrow_y,
                    brand.colors.accent,
                    width=1.5,
                )
                self._add_arrowhead(
                    slide,
                    x + card_w + gap - 0.14,
                    arrow_y,
                    0.2,
                    0,
                    brand.colors.accent,
                )

    def _add_source_repair_editorial_cards(
        self,
        slide,
        points: list[str],
        brand: BrandDNA,
    ) -> None:
        hero = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.92),
            Inches(1.6),
            Inches(11.35),
            Inches(1.36),
        )
        hero.fill.solid()
        hero.fill.fore_color.rgb = self._rgb(self._tint(brand.colors.primary, 0.92))
        hero.line.color.rgb = self._rgb(self._tint(brand.colors.primary, 0.78))
        self._add_body_text(
            slide,
            self._truncate_phrase(points[0], 142),
            1.28,
            2.02,
            10.36,
            0.42,
            brand,
            center=True,
            size=15,
        )
        support = points[1:4] or points[:1]
        for idx, point in enumerate(support):
            x = 1.02 + idx * 3.86
            accent = self._icon_fill(brand, idx)
            self._add_card(slide, x, 3.48, 3.36, 2.24, "FFFFFF", brand.colors.background_light)
            self._add_badge(slide, str(idx + 1), x + 0.3, 3.82, 0.42, brand, accent)
            self._add_body_text(
                slide,
                point,
                x + 0.9,
                3.8,
                2.0,
                0.92,
                brand,
                size=11,
            )

    def _add_source_repair_proof_strip(
        self,
        slide,
        points: list[str],
        brand: BrandDNA,
    ) -> None:
        for idx, point in enumerate(points[:4]):
            y = 1.5 + idx * 1.12
            accent = self._icon_fill(brand, idx)
            fill = "FFFFFF" if idx % 2 == 0 else "F4F6F8"
            self._add_card(slide, 0.92, y, 11.52, 0.86, fill, brand.colors.background_light)
            marker = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(0.92),
                Inches(y),
                Inches(0.1),
                Inches(0.86),
            )
            marker.fill.solid()
            marker.fill.fore_color.rgb = self._rgb(accent)
            marker.line.color.rgb = marker.fill.fore_color.rgb
            self._add_dark_text(
                slide,
                f"{idx + 1:02d}",
                1.26,
                y + 0.28,
                0.42,
                0.2,
                brand,
                size=8,
                bold=True,
                color=accent,
            )
            self._add_body_text(
                slide,
                point,
                1.92,
                y + 0.18,
                9.34,
                0.34,
                brand,
                size=12,
            )
