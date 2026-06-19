# ruff: noqa: F401
from pathlib import Path
import math
import re
import shutil
import subprocess
from typing import Any

from PIL import Image, ImageDraw
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt

from app.models.brand import BrandDNA
from app.models.outline import SlideOutline
from app.services.concept_diagram_renderer import ConceptDiagramRenderer, DiagramRenderError
from app.services.pptx_rendering.constants import ICON_SCALE, SLIDE_H, SLIDE_W


class TableLayoutRenderingMixin:
    def _add_comparison_table(self, slide, outline: SlideOutline, brand: BrandDNA) -> None:
        exhibit = self._exhibit(outline)
        columns, rows = self._comparison_rows(exhibit, outline)
        display_rows = rows[:5]
        if not display_rows:
            self._add_two_column(slide, outline, brand)
            return

        rail = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.9),
            Inches(1.5),
            Inches(2.72),
            Inches(4.82),
        )
        rail.fill.solid()
        rail.fill.fore_color.rgb = self._rgb(brand.colors.primary)
        rail.line.color.rgb = self._rgb(brand.colors.primary)
        self._add_dark_text(
            slide,
            "COMPARISON LENS",
            1.18,
            1.86,
            1.8,
            0.24,
            brand,
            size=10,
            color=brand.colors.accent,
        )
        self._add_dark_text(
            slide,
            f"{len(display_rows)} shifts",
            1.18,
            2.42,
            1.9,
            0.42,
            brand,
            size=22,
            bold=True,
        )
        self._add_dark_text(
            slide,
            self._comparison_panel_insight(columns, display_rows),
            1.2,
            3.18,
            2.0,
            1.14,
            brand,
            size=11,
            color=self._tint(brand.colors.primary, 0.76),
        )
        self._add_dark_text(
            slide,
            "Read left to right: what changes, what improves, and where the target model tightens control.",
            1.2,
            5.22,
            2.05,
            0.7,
            brand,
            size=9,
            color=self._tint(brand.colors.primary, 0.68),
        )

        current_header = columns[1] if len(columns) > 1 else "Current state"
        target_header = columns[2] if len(columns) > 2 else "Target state"
        self._add_label(slide, str(current_header).upper(), 4.9, 1.54, 2.6, brand, bold=True)
        self._add_label(slide, str(target_header).upper(), 8.78, 1.54, 2.6, brand, bold=True)

        row_count = len(display_rows)
        row_h = 0.72 if row_count <= 4 else 0.62
        gap = 0.22 if row_count <= 4 else 0.14
        start_y = 2.02
        for idx, row in enumerate(display_rows):
            cells = (row + ["", ""])[:3]
            dimension = self._truncate_at_word(cells[0], 28)
            current = self._truncate_at_word(cells[1], 72)
            target = self._truncate_at_word(cells[2], 72)
            y = start_y + idx * (row_h + gap)
            accent = self._icon_fill(brand, idx)

            self._add_card(
                slide,
                3.92,
                y,
                0.78,
                row_h,
                self._tint(accent, 0.84),
                self._tint(accent, 0.66),
            )
            self._add_body_text(slide, dimension, 4.02, y + 0.18, 0.56, 0.28, brand, center=True, size=8)

            self._add_card(slide, 4.92, y, 2.75, row_h, "FFFFFF", brand.colors.background_light)
            self._add_body_text(slide, current, 5.16, y + 0.17, 2.24, 0.3, brand, size=10)

            self._add_arrow(slide, 7.86, y + row_h / 2, 8.42, y + row_h / 2, brand.colors.accent, width=1.6)
            self._add_arrowhead(slide, 8.42, y + row_h / 2, 0.2, 0, brand.colors.accent)

            self._add_card(slide, 8.62, y, 3.62, row_h, "F4F6F8", brand.colors.background_light)
            self._add_body_text(slide, target, 8.88, y + 0.17, 3.02, 0.3, brand, size=10)

        self._add_card(slide, 4.92, 6.0, 7.32, 0.42, self._tint(brand.colors.accent, 0.88), self._tint(brand.colors.accent, 0.74))
        self._add_body_text(slide, "TARGET TEST", 5.18, 6.1, 1.3, 0.18, brand, size=8)
        self._add_body_text(
            slide,
            "Every row should make the behavior change explicit enough to inspect or assign.",
            6.48,
            6.08,
            4.96,
            0.2,
            brand,
            size=9,
        )

    def _comparison_panel_insight(
        self, columns: list[str], rows: list[list[str]]
    ) -> str:
        if not rows:
            return "Use the comparison to make the target-state behavior explicit."
        header_text = " ".join(str(column).lower() for column in columns)
        if "target" in header_text or "after" in header_text:
            return "Make the target model concrete by pairing each current behavior with the operating change."
        if "risk" in header_text or "impact" in header_text:
            return "Separate the issue from the consequence so leaders can judge priority."
        return "Turn the comparison into an operating choice, not a static list of differences."

    def _add_reference_table(self, slide, outline: SlideOutline, brand: BrandDNA) -> None:
        exhibit = self._exhibit(outline)
        columns, rows = self._reference_rows(exhibit, outline)
        display_rows = rows[:5]
        left = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.9),
            Inches(1.48),
            Inches(3.1),
            Inches(4.92),
        )
        left.fill.solid()
        left.fill.fore_color.rgb = self._rgb(brand.colors.primary)
        left.line.color.rgb = self._rgb(brand.colors.primary)
        self._add_dark_text(
            slide,
            "REFERENCE MAP",
            1.2,
            1.84,
            1.9,
            0.24,
            brand,
            size=10,
            color=brand.colors.accent,
        )
        self._add_dark_text(
            slide,
            f"{len(display_rows)} artifacts",
            1.2,
            2.42,
            2.05,
            0.42,
            brand,
            size=20,
            bold=True,
        )
        self._add_dark_text(
            slide,
            self._reference_table_insight(columns, display_rows),
            1.22,
            3.18,
            2.24,
            1.2,
            brand,
            size=11,
            color=self._tint(brand.colors.primary, 0.76),
        )
        self._add_dark_text(
            slide,
            "Use as the source of truth for owner and refresh triggers.",
            1.22,
            5.24,
            2.28,
            0.44,
            brand,
            size=8,
            color=self._tint(brand.colors.primary, 0.68),
        )

        x = 4.35
        y = 1.54
        row_h = 0.72
        gap = 0.16
        col_w = [2.2, 3.28, 2.46]
        headers = (columns + ["Purpose", "Update trigger"])[:3]
        for idx, header in enumerate(headers):
            self._add_label(
                slide,
                str(header).upper(),
                x + sum(col_w[:idx]) + idx * 0.16,
                y,
                col_w[idx],
                brand,
                bold=True,
            )
        for r_idx, row in enumerate(display_rows):
            row_y = y + 0.48 + r_idx * (row_h + gap)
            fill = "FFFFFF" if r_idx % 2 == 0 else "F4F6F8"
            self._add_card(slide, x, row_y, 8.26, row_h, fill, brand.colors.background_light)
            strip = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(x),
                Inches(row_y),
                Inches(0.08),
                Inches(row_h),
            )
            strip.fill.solid()
            strip.fill.fore_color.rgb = self._rgb(self._icon_fill(brand, r_idx))
            strip.line.color.rgb = strip.fill.fore_color.rgb
            cells = (row + ["", ""])[:3]
            self._add_body_text(
                slide,
                self._truncate_at_word(cells[0], 32),
                x + 0.28,
                row_y + 0.18,
                col_w[0] - 0.34,
                0.24,
                brand,
                size=10,
            )
            self._add_body_text(
                slide,
                self._truncate_at_word(cells[1], 52),
                x + col_w[0] + 0.36,
                row_y + 0.16,
                col_w[1] - 0.22,
                0.3,
                brand,
                size=9,
            )
            self._add_body_text(
                slide,
                self._truncate_at_word(cells[2], 42),
                x + col_w[0] + col_w[1] + 0.54,
                row_y + 0.16,
                col_w[2] - 0.28,
                0.3,
                brand,
                size=9,
            )
        self._add_card(slide, 4.35, 6.02, 8.26, 0.42, self._tint(brand.colors.accent, 0.88), self._tint(brand.colors.accent, 0.74))
        self._add_body_text(slide, "UPDATE CADENCE", 4.62, 6.12, 1.6, 0.18, brand, size=8)
        self._add_body_text(
            slide,
            "Refresh the artifact when source evidence, ownership, or risk changes.",
            6.22,
            6.1,
            5.72,
            0.2,
            brand,
            size=9,
        )

    def _reference_table_insight(
        self, columns: list[str], rows: list[list[str]]
    ) -> str:
        if not rows:
            return "Keep the reusable artifacts explicit so the operating model can be repeated."
        header_text = " ".join(str(column).lower() for column in columns)
        if "update" in header_text or "trigger" in header_text:
            return "Make refresh triggers visible so reference material stays current as conditions change."
        if "owner" in header_text or "timing" in header_text:
            return "Make ownership visible so each artifact has a clear accountable maintainer."
        return "Turn the table into a reusable operating reference, not a one-time appendix."

    def _set_table_cell(
        self,
        cell,
        text: str,
        brand: BrandDNA,
        size: int = 9,
        bold: bool = False,
        color: str | None = None,
        fill: str = "FFFFFF",
        center: bool = False,
    ) -> None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = self._rgb(fill)
        cell.margin_left = Inches(0.08)
        cell.margin_right = Inches(0.08)
        cell.margin_top = Inches(0.05)
        cell.margin_bottom = Inches(0.04)
        frame = cell.text_frame
        frame.word_wrap = True
        frame.clear()
        paragraph = frame.paragraphs[0]
        if center:
            paragraph.alignment = 1
        run = paragraph.add_run()
        run.text = text
        run.font.name = brand.fonts.body
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = self._rgb(color or brand.colors.text_dark)

    def _comparison_rows(
        self, exhibit: dict[str, Any], outline: SlideOutline
    ) -> tuple[list[str], list[list[str]]]:
        columns = exhibit.get("columns")
        rows = exhibit.get("rows")
        if isinstance(columns, list) and isinstance(rows, list) and rows:
            normalized_rows: list[list[str]] = []
            for row in rows:
                if isinstance(row, dict):
                    normalized_rows.append(
                        [
                            self._table_cell_text(row.get("label", "")),
                            *[self._table_cell_text(value) for value in row.get("values", [])],
                        ]
                    )
                elif isinstance(row, list):
                    normalized_rows.append([self._table_cell_text(value) for value in row])
            if normalized_rows:
                return [self._table_header_text(value) for value in columns], normalized_rows
        table_rows = self._table_rows(outline)
        if len(table_rows) >= 2:
            return [str(value) for value in table_rows[0]], [
                [str(value) for value in row] for row in table_rows[1:]
            ]
        bullets = self._bullets(outline)[:3]
        return ["Dimension", "Current state", "Target state"], [
            [f"Area {idx + 1}", bullet, "Managed behavior"]
            for idx, bullet in enumerate(bullets or ["Context", "Review", "Cadence"])
        ]

    def _reference_rows(
        self, exhibit: dict[str, Any], outline: SlideOutline
    ) -> tuple[list[str], list[list[str]]]:
        columns = exhibit.get("columns")
        rows = exhibit.get("rows")
        if isinstance(columns, list) and isinstance(rows, list) and rows:
            return [self._table_header_text(value) for value in columns], [
                [self._table_cell_text(value) for value in row]
                if isinstance(row, list)
                else [self._table_cell_text(row)]
                for row in rows
            ]
        table_rows = self._table_rows(outline)
        if len(table_rows) >= 2:
            return [str(value) for value in table_rows[0]], [
                [str(value) for value in row] for row in table_rows[1:]
            ]
        bullets = self._bullets(outline)[:5]
        return ["Item", "Implication"], [[self._truncate_at_word(item, 28), item] for item in bullets]

    def _table_header_text(self, value: Any) -> str:
        if isinstance(value, dict):
            for key in ("name", "label", "title", "header"):
                text = str(value.get(key) or "").strip()
                if text:
                    return text
            return ""
        return self._table_cell_text(value)

    def _table_cell_text(self, value: Any) -> str:
        if isinstance(value, dict):
            for key in ("text", "label", "name", "value"):
                text = str(value.get(key) or "").strip()
                if text:
                    return text
            return ""
        return str(value)

    def _column_widths(self, total_width: float, count: int) -> list[float]:
        if count <= 1:
            return [total_width]
        if count == 2:
            return [total_width * 0.34, total_width * 0.66]
        first = total_width * 0.26
        remaining = (total_width - first) / (count - 1)
        return [first] + [remaining] * (count - 1)

    def _table_rows(self, outline: SlideOutline) -> list[list[Any]]:
        blocks = outline.content_json.get("content_blocks") or []
        for block in blocks:
            if block.get("type") == "table" and isinstance(block.get("body"), list):
                rows = [row for row in block["body"] if isinstance(row, list)]
                if rows:
                    return rows
        bullets = self._bullets(outline)[:5]
        if bullets:
            return [["Step", "Action"]] + [
                [str(index + 1), bullet] for index, bullet in enumerate(bullets)
            ]
        return [["Action", "Owner", "Timing"], ["Align", "Sponsor", "Week 1"], ["Execute", "Team", "Week 2"]]
