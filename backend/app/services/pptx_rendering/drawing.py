# ruff: noqa: F401
from pathlib import Path
import ast
import json
import math
import re
import shutil
import subprocess
from typing import Any

from PIL import Image, ImageDraw
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from app.models.brand import BrandDNA
from app.models.outline import SlideOutline
from app.services.concept_diagram_renderer import ConceptDiagramRenderer, DiagramRenderError
from app.services.pptx_rendering.constants import ICON_SCALE, SLIDE_H, SLIDE_W


class DrawingMixin:
    def _add_card(self, slide, x: float, y: float, w: float, h: float, fill: str, line: str) -> None:
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = self._rgb(fill)
        shape.line.color.rgb = self._rgb(line)

    def _add_label(self, slide, text: str, x: float, y: float, w: float, brand: BrandDNA, bold: bool = False) -> None:
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.35))
        frame = box.text_frame
        frame.clear()
        run = frame.paragraphs[0].add_run()
        run.text = text
        run.font.name = brand.fonts.body
        run.font.size = Pt(12)
        run.font.bold = bold
        run.font.color.rgb = self._rgb(brand.colors.primary)

    def _add_dark_text(
        self,
        slide,
        text: str,
        x: float,
        y: float,
        w: float,
        h: float,
        brand: BrandDNA,
        size: int = 12,
        bold: bool = False,
        color: str | None = None,
    ) -> None:
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.clear()
        self._autofit(frame)
        para = frame.paragraphs[0]
        para.line_spacing = 1.12
        run = para.add_run()
        run.text = text[:360]
        run.font.name = brand.fonts.heading if bold else brand.fonts.body
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = self._rgb(color or brand.colors.text_light)

    def _add_code_text(
        self,
        slide,
        lines: list[str],
        x: float,
        y: float,
        w: float,
        h: float,
        brand: BrandDNA,
    ) -> None:
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.clear()
        self._autofit(frame)
        for idx, line in enumerate(lines[:8]):
            para = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
            para.line_spacing = 1.0
            run = para.add_run()
            run.text = line
            run.font.name = "Courier New"
            run.font.size = Pt(10)
            run.font.color.rgb = self._rgb(brand.colors.text_light)
            para.space_after = Pt(2)

    def _add_arrow(
        self,
        slide,
        x1: float,
        y1: float,
        x2: float,
        y2: float,
        color: str,
        width: float = 1.4,
    ) -> None:
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(x1),
            Inches(y1),
            Inches(x2),
            Inches(y2),
        )
        connector.line.color.rgb = self._rgb(color)
        connector.line.width = Pt(width)

    def _add_arrowhead(
        self,
        slide,
        x: float,
        y: float,
        dx: float,
        dy: float,
        color: str,
    ) -> None:
        arrow = slide.shapes.add_shape(
            MSO_SHAPE.ISOSCELES_TRIANGLE,
            Inches(x - 0.08),
            Inches(y - 0.08),
            Inches(0.16),
            Inches(0.16),
        )
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = self._rgb(color)
        arrow.line.color.rgb = self._rgb(color)
        arrow.rotation = math.degrees(math.atan2(dy, dx)) + 90

    def _add_badge(
        self,
        slide,
        text: str,
        x: float,
        y: float,
        size: float,
        brand: BrandDNA,
        fill: str,
    ) -> None:
        badge = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(x),
            Inches(y),
            Inches(size),
            Inches(size),
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = self._rgb(fill)
        badge.line.color.rgb = self._rgb(fill)
        frame = badge.text_frame
        frame.clear()
        frame.margin_left = Inches(0.02)
        frame.margin_right = Inches(0.02)
        frame.margin_top = Inches(0.02)
        frame.margin_bottom = Inches(0.02)
        para = frame.paragraphs[0]
        para.alignment = 1
        run = para.add_run()
        run.text = text[:2]
        run.font.name = brand.fonts.heading
        run.font.size = Pt(max(10, int(size * 22)))
        run.font.bold = True
        run.font.color.rgb = self._rgb(self._readable_text_color(fill, brand))

    def _add_body_text(
        self,
        slide,
        text: str,
        x: float,
        y: float,
        w: float,
        h: float,
        brand: BrandDNA,
        center: bool = False,
        size: int = 13,
    ) -> None:
        size = max(11, size)
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.clear()
        self._autofit(frame)
        items = self._split_multi_item_text(text)
        if len(items) > 1:
            for idx, item in enumerate(items[:4]):
                para = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
                self._format_body_paragraph(
                    para,
                    item,
                    brand,
                    size=max(11, size - 1 if len(items) >= 3 else size),
                    center=center,
                    bold_lead=True,
                )
                para.space_after = Pt(4)
            return
        self._format_body_paragraph(
            frame.paragraphs[0],
            text,
            brand,
            size=size,
            center=center,
            bold_lead=False,
        )

    def _format_body_paragraph(
        self,
        para,
        text: str,
        brand: BrandDNA,
        size: int,
        center: bool,
        bold_lead: bool,
    ) -> None:
        para.line_spacing = 1.15
        if center:
            para.alignment = 1
        cleaned = self._truncate_at_word(self._clean_display_text(text), 420)
        if bold_lead:
            lead, rest = self._split_lead(cleaned)
            if rest:
                lead_run = para.add_run()
                lead_run.text = lead
                lead_run.font.name = brand.fonts.body
                lead_run.font.size = Pt(size)
                lead_run.font.bold = True
                lead_run.font.color.rgb = self._rgb(brand.colors.text_dark)
                body_run = para.add_run()
                body_run.text = f": {rest}" if not lead.endswith(":") else f" {rest}"
                body_run.font.name = brand.fonts.body
                body_run.font.size = Pt(size)
                body_run.font.color.rgb = self._rgb(brand.colors.text_dark)
                return
        run = para.add_run()
        run.text = cleaned[:420]
        run.font.name = brand.fonts.body
        run.font.size = Pt(size)
        run.font.color.rgb = self._rgb(brand.colors.text_dark)

    def _add_bullets(self, slide, bullets: list[str], x: float, y: float, w: float, h: float, brand: BrandDNA) -> None:
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.clear()
        self._autofit(frame)
        for idx, bullet in enumerate(bullets[:4]):
            para = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
            para.level = 0
            self._apply_native_bullet(para)
            para.text = self._truncate_at_word(bullet, 175)
            para.font.name = brand.fonts.body
            para.font.size = Pt(13)
            para.font.color.rgb = self._rgb(brand.colors.text_dark)
            para.line_spacing = 1.12
            para.space_after = Pt(6)

    def _apply_native_bullet(self, para) -> None:
        p_pr = para._p.get_or_add_pPr()
        for child in list(p_pr):
            if child.tag.endswith(("}buChar", "}buAutoNum", "}buNone")):
                p_pr.remove(child)
        bullet = OxmlElement("a:buChar")
        bullet.set("char", "\u2022")
        p_pr.insert(0, bullet)

    def _add_big_number(self, slide, text: str, x: float, y: float, w: float, brand: BrandDNA) -> None:
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.8))
        frame = box.text_frame
        frame.clear()
        self._autofit(frame)
        para = frame.paragraphs[0]
        para.alignment = 1
        run = para.add_run()
        run.text = text[:12]
        run.font.name = brand.fonts.heading
        run.font.size = Pt(42)
        run.font.bold = True
        run.font.color.rgb = self._rgb(brand.colors.primary)

    def _add_text_in_shape(
        self,
        shape,
        text: str,
        brand: BrandDNA,
        size: int = 10,
        bold: bool = False,
        color: str | None = None,
        center: bool = False,
    ) -> None:
        frame = shape.text_frame
        frame.clear()
        self._autofit(frame)
        frame.margin_left = Inches(0.05)
        frame.margin_right = Inches(0.05)
        para = frame.paragraphs[0]
        if center:
            para.alignment = 1
        run = para.add_run()
        run.text = text[:90]
        run.font.name = brand.fonts.body
        run.font.size = Pt(size)
        run.font.bold = bold
        fill_hex = self._shape_fill_hex(shape)
        if fill_hex:
            if color and self._contrast_ratio(color, fill_hex) >= 4.5:
                chosen = color
            else:
                chosen = self._readable_text_color(fill_hex, brand)
        else:
            chosen = color or brand.colors.text_dark
        run.font.color.rgb = self._rgb(chosen)

    def _exhibit(self, outline: SlideOutline) -> dict[str, Any]:
        exhibit = outline.content_json.get("exhibit_spec")
        if isinstance(exhibit, dict):
            return exhibit
        return {}

    def _bullets(self, outline: SlideOutline) -> list[str]:
        bullets = outline.content_json.get("bullets")
        if isinstance(bullets, list) and bullets:
            return self._clean_item_texts(bullets)
        blocks = outline.content_json.get("content_blocks") or []
        collected: list[str] = []
        for block in blocks:
            for item in block.get("body", []):
                if isinstance(item, (str, dict)):
                    collected.extend(self._clean_item_texts([item]))
        return collected

    def _clean_item_texts(self, items: list[Any]) -> list[str]:
        cleaned_items: list[str] = []
        for item in items:
            cleaned = self._clean_display_text(self._coerce_item_text(item))
            if not cleaned:
                continue
            for part in self._split_multi_item_text(cleaned):
                item_text = self._complete_display_item(part)
                if item_text and not self._is_placeholder_bullet(item_text):
                    cleaned_items.append(item_text)
        return cleaned_items

    def _is_placeholder_bullet(self, text: str) -> bool:
        normalized = " ".join(str(text).lower().split()).strip(" .:-")
        return normalized in {
            "clarify the implication and required management action",
            "architecture overview",
            "closing remarks",
            "executive summary",
            "item",
            "introduction",
            "metric",
            "metrics",
            "prediction",
            "predictions",
            "sourced metric",
            "signal",
            "the model contract",
            "the problem",
            "the solution",
        }

    def _is_renderer_filler_copy(self, text: str) -> bool:
        normalized = re.sub(r"[^a-z0-9]+", " ", str(text).casefold())
        normalized = " ".join(normalized.split())
        if not normalized:
            return False
        phrases = {
            "tie the claim source backed evaluation artifact",
            "tie the claim to a source backed evaluation artifact",
            "connect the source evidence to the decision before scaling",
            "make the handoff inspectable before the decision moves",
            "each step should pair evidence ownership and timing",
            "make the next move visible enough to own inspect and revise",
            "name the review gate before expanding the benchmark",
            "update the benchmark when source evidence changes",
            "name the decision and success criteria",
            "define acceptance criteria",
            "assign evidence and review owners",
            "make the operating implication explicit",
            "assign the next review gate before scaling",
            "track what changes when the source evidence moves",
            "adopt the operating model through a named pilot and review gate",
            "confirm owner scope and timing",
            "confirm ownership and timing",
            "deterministic beat for evidence",
            "deterministic beat for implementation",
            "deterministic beat for reference",
            "deterministic beat for decision",
            "deterministic beat for closing",
            "evidence tension",
            "reliability risk",
            "operating move",
            "evidence gap",
            "validation risk",
            "operating choice",
        }
        if normalized in phrases:
            return True
        return bool(
            re.search(
                r"\bconnect\b.{0,90}\b(?:to\s+)?an explicit review gate\b|"
                r"\bmake\b.{0,90}\bvisible before execution starts\b|"
                r"\buse the model to decide what must be explicit\b|"
                r"\bdeterministic beat for\b",
                normalized,
            )
        )

    def _complete_display_item(self, text: str, limit: int = 150) -> str:
        raw = " ".join(str(text).split()).strip(" -:;")
        if self._is_incomplete_display_fragment(raw):
            return ""
        cleaned = self._clean_display_text(text)
        if not cleaned or self._is_renderer_filler_copy(cleaned):
            return ""
        sentences = re.findall(r"[^.!?]+[.!?]", cleaned)
        for sentence in sentences:
            candidate = self._repair_display_fragment(sentence.strip())
            if (
                len(candidate.split()) >= 4
                and not self._is_placeholder_bullet(candidate)
                and not self._is_renderer_filler_copy(candidate)
                and not self._is_incomplete_display_fragment(candidate)
            ):
                truncated = self._truncate_at_word(candidate, limit)
                if self._is_incomplete_display_fragment(truncated):
                    return ""
                return truncated
        if (
            self._is_placeholder_bullet(cleaned)
            or self._is_renderer_filler_copy(cleaned)
            or self._is_incomplete_display_fragment(cleaned)
        ):
            return ""
        truncated = self._truncate_at_word(cleaned, limit)
        if (
            self._is_placeholder_bullet(truncated)
            or self._is_renderer_filler_copy(truncated)
            or self._is_incomplete_display_fragment(truncated)
        ):
            return ""
        return truncated

    def _is_incomplete_display_fragment(self, text: str) -> bool:
        cleaned = " ".join(str(text).split()).strip(" .,:;-")
        if len(cleaned.split()) < 3:
            return False
        lowered = cleaned.casefold()
        if re.match(
            r"^(?:as|and|or|but|rather than|instead of|with|by|from|to|of|themselves)\b",
            lowered,
        ):
            return True
        if re.search(
            r"\b(?:has|have|had|has been|have been|was|were|will|would|should|"
            r"can|cannot|could|may|might|must|rather than|treating|requiring|including|"
            r"continued|consisted of|declares|outpaced the development|"
            r"outpaced the development of evaluation|"
            r"because|because it|desire to test these|enabling|managed|operating|"
            r"reproducible|scalable|valid|varying)$",
            lowered,
        ):
            return True
        if re.search(
            r"\b(?:are|presented|become|becomes|create|creates|define|defines|"
            r"evaluate|evaluated|make|makes|reflect|specified|specifies|"
            r"standardize|standardizes)$",
            lowered,
        ):
            return True
        if re.search(
            r"\b(?:a|an|and|as|by|for|from|in|into|of|or|that|the|their|through|to|which|with)$",
            lowered,
        ):
            return True
        if re.search(r"\b(?:because|while|when|where|that|which)\s+(?:it|they|the)$", lowered):
            return True
        if re.search(r"\b(?:cannot|can|make|makes|around|through|from|to)\s+is$", lowered):
            return True
        if re.search(
            r"\b(?:of|through|from|as|into)\s+"
            r"(?:implicit|agent-assisted|manual|systematic|operational|declarative)$",
            lowered,
        ):
            return True
        if re.search(r",\s+the\s+[a-z-]+$", lowered):
            return True
        if re.search(r"\b(?:the machinery|following phases)$", lowered):
            return True
        if re.search(
            r"\b(?:development of evaluation|reflect their actual use)$",
            lowered,
        ):
            return True
        if re.search(r"\b(?:a core|to help)$", lowered):
            return True
        if re.match(r"^[a-z][a-z -]+\)", cleaned):
            return True
        if cleaned[:1].islower() and len(cleaned.split()) >= 3:
            return True
        if re.search(r"\bacross enterprise$", lowered):
            return True
        return False

    def _coerce_item_text(self, item: Any) -> str:
        """Render a content item as display text.

        Bullets occasionally arrive as metric dicts (e.g.
        ``{"label": "Adoption", "value": 95, "unit": "%"}``); never let a raw
        dict reach a text frame as ``str(dict)``.
        """
        if isinstance(item, str):
            stripped = item.strip()
            # A stringified dict/JSON object must never reach a text frame as raw
            # ``{'label': ...}``. Parse it and coerce like a real dict.
            if stripped.startswith("{") and stripped.endswith("}") and ":" in stripped:
                parsed = self._try_parse_mapping(stripped)
                if parsed is not None:
                    return self._coerce_item_text(parsed)
            return item
        if isinstance(item, dict):
            label = str(
                item.get("label") or item.get("name") or item.get("title") or ""
            ).strip()
            if str(item.get("value", "")).strip() != "":
                value = self._format_metric_value(item)
                return f"{label}: {value}" if label else value
            for key in ("action", "text", "description"):
                text = str(item.get(key) or "").strip()
                if text:
                    return (
                        f"{label}: {text}"
                        if label and label.lower() not in text.lower()
                        else text
                    )
            return label
        return ""

    def _emphasis_panel_text(
        self,
        outline: SlideOutline,
        bullets: list[str],
        default_headline: str,
        default_support: str,
    ) -> tuple[str, str, str]:
        """A side-panel takeaway derived from the slide so two slides with the
        same archetype never show an identical hardcoded panel.

        Returns (kicker, headline, support). The headline prefers the slide's
        own subheading/summary (the "so what" of this exhibit); the support line
        prefers a bullet not already prominent, falling back to sane defaults.
        """
        title = self._clean_display_text(
            str(
                outline.content_json.get("action_title")
                or outline.content_json.get("title")
                or outline.label
            )
        )
        sub = str(
            outline.content_json.get("subheading")
            or outline.content_json.get("summary")
            or ""
        )
        if self._looks_like_meta_subheading(sub):
            sub = ""
        sub = self._clean_display_text(sub)
        headline = default_headline
        if sub and 3 <= len(sub.split()) <= 18 and sub.casefold() != title.casefold():
            headline = sub if sub.endswith(".") else f"{sub}."
        support = default_support
        for bullet in bullets:
            cleaned = self._clean_display_text(bullet)
            if (
                cleaned
                and cleaned.casefold() not in headline.casefold()
                and cleaned.casefold() != title.casefold()
                and len(cleaned.split()) >= 4
            ):
                support = self._truncate_phrase(cleaned, 120)
                break
        kicker = self._panel_kicker(title)
        return kicker, self._truncate_phrase(headline, 130), support

    def _split_lead(self, text: str) -> tuple[str, str]:
        """Split copy only at a real semantic boundary.

        The authored layouts use a bold lead plus body detail. Arbitrarily taking
        the first N words turns source sentences into nonsense ("Synthetic
        benchmarks can" / "Validation loops..."), so plain prose is split at the
        subject/verb boundary or left intact.
        """
        cleaned = self._clean_display_text(text)
        if not cleaned:
            return "", ""
        for delimiter in (": ", " — ", " - "):
            if delimiter in cleaned:
                lead, rest = cleaned.split(delimiter, 1)
                if (
                    1 <= len(lead.split()) <= 6
                    and rest.strip()
                    and not self._is_incomplete_display_fragment(lead)
                    and not self._looks_like_predicate_fragment(lead)
                ):
                    lead, rest = self._polish_split_pair(lead, rest)
                    return (
                        self._truncate_phrase(lead, 46),
                        self._sentence_case_fragment(self._truncate_phrase(rest, 130)),
                    )
        words = cleaned.split()
        split_at = self._subject_predicate_split_index(words)
        if split_at:
            lead = " ".join(words[:split_at]).rstrip(",.;:")
            rest = " ".join(words[split_at:]).rstrip(" ,;:")
            lead, rest = self._polish_split_pair(lead, rest)
            shortened_lead = self._truncate_phrase(lead, 54)
            shortened_rest = self._sentence_case_fragment(self._truncate_phrase(rest, 130))
            if (
                shortened_lead
                and shortened_rest
                and not self._is_incomplete_display_fragment(shortened_lead)
                and not self._looks_like_predicate_fragment(shortened_lead)
                and not self._is_incomplete_display_fragment(shortened_rest)
            ):
                return shortened_lead, shortened_rest
        return self._truncate_phrase(cleaned, 120), ""

    def _subject_predicate_split_index(self, words: list[str]) -> int | None:
        if len(words) < 6:
            return None
        predicate_markers = {
            "can",
            "cannot",
            "could",
            "should",
            "must",
            "will",
            "would",
            "adopts",
            "answer",
            "answering",
            "become",
            "becomes",
            "build",
            "builds",
            "capture",
            "captures",
            "commit",
            "commits",
            "connect",
            "connects",
            "create",
            "creates",
            "define",
            "defines",
            "determine",
            "determines",
            "discover",
            "discovers",
            "evaluate",
            "evaluates",
            "exist",
            "exists",
            "find",
            "finds",
            "give",
            "gives",
            "govern",
            "governs",
            "help",
            "helps",
            "improve",
            "improves",
            "make",
            "makes",
            "map",
            "maps",
            "matter",
            "matters",
            "moves",
            "need",
            "needs",
            "reflect",
            "reflects",
            "reuse",
            "reuses",
            "scale",
            "scales",
            "shifts",
            "specify",
            "specifies",
            "standardize",
            "standardizes",
            "turn",
            "turns",
            "uses",
            "is",
            "are",
            "was",
            "were",
        }
        for index, word in enumerate(words[1:7], start=1):
            normalized = re.sub(r"[^A-Za-z]", "", word).casefold()
            if normalized not in predicate_markers:
                continue
            if index > 5:
                continue
            if index < 2 and (index != 1 or len(words[0].strip(".,;:")) < 5):
                continue
            lead = " ".join(words[:index]).strip(" ,;:")
            rest = " ".join(words[index:]).strip(" ,;:")
            split_index = index
            lead_words = lead.split()
            if (
                index > 2
                and lead_words
                and lead_words[-1].casefold().strip(".,;:") in {"already", "also", "still", "now"}
            ):
                split_index = index - 1
                lead = " ".join(words[:split_index]).strip(" ,;:")
                rest = " ".join(words[split_index:]).strip(" ,;:")
            if len(lead.split()) < 2 and not (
                split_index == 1 and len(lead.strip(".,;:")) >= 5
            ):
                continue
            if len(rest.split()) < 3:
                continue
            return split_index
        return None

    def _looks_like_predicate_fragment(self, text: str) -> bool:
        cleaned = " ".join(str(text).split()).strip(" .,:;-")
        if len(cleaned.split()) < 2:
            return False
        return bool(
            re.search(
                r"\b(?:can|cannot|could|should|must|will|would|make|makes|"
                r"create|creates|turn|turns|give|gives|evaluate|evaluates|"
                r"specify|specifies|determine|determines|standardize|"
                r"standardizes)\.?$",
                cleaned,
                flags=re.IGNORECASE,
            )
        )

    def _sentence_case_fragment(self, text: str) -> str:
        cleaned = " ".join(str(text).split()).strip()
        if not cleaned:
            return ""
        return f"{cleaned[:1].upper()}{cleaned[1:]}"

    def _polish_split_pair(self, lead: str, rest: str) -> tuple[str, str]:
        polished_lead = " ".join(str(lead).split()).strip(" -:;,.")
        polished_rest = " ".join(str(rest).split()).strip(" -:;")
        lower_lead = polished_lead.casefold()
        lower_rest = polished_rest.casefold()
        if "shift" in lower_lead and lower_rest.startswith("is from "):
            polished_lead = re.sub(
                r"^the\s+",
                "",
                polished_lead,
                flags=re.IGNORECASE,
            )
            if polished_lead:
                polished_lead = f"{polished_lead[:1].upper()}{polished_lead[1:]}"
            polished_rest = "moves from " + polished_rest[8:]
        return polished_lead, polished_rest

    def _panel_kicker(self, title: str) -> str:
        lowered = title.lower()
        if any(token in lowered for token in ("review", "verify", "evidence", "quality")):
            return "WHY IT MATTERS"
        if any(token in lowered for token in ("memory", "context", "persist", "update")):
            return "KEEP IT CURRENT"
        if any(token in lowered for token in ("rule", "spec", "standard", "codify")):
            return "THE OPERATING TEST"
        if any(token in lowered for token in ("reset", "cycle", "loop", "cadence")):
            return "THE CADENCE"
        return "THE NEXT MOVE"

    def _try_parse_mapping(self, text: str) -> dict[str, Any] | None:
        for parser in (json.loads, ast.literal_eval):
            try:
                value = parser(text)
            except (ValueError, SyntaxError, TypeError):
                continue
            if isinstance(value, dict):
                return value
        return None

    def _clean_display_text(self, text: str) -> str:
        cleaned = re.sub(
            r"\[\s*diagram description\s*:[^\]]+\]",
            "",
            str(text),
            flags=re.IGNORECASE,
        )
        cleaned = re.sub(
            r"\bdiagram description\s*:[^.]+\.?",
            "",
            cleaned,
            flags=re.IGNORECASE,
        )
        cleaned = re.sub(r"\b(\d+)\.0(?=[A-Za-z]|\b)", r"\1", cleaned)
        cleaned = re.sub(
            r"(\d(?:[\d,]*\.?\d*)?)(tokens\b)",
            r"\1 \2",
            cleaned,
            flags=re.IGNORECASE,
        )
        cleaned = re.sub(
            r"\(\s*owner\s*/\s*next\s*\)|\bowner\s*/\s*next\b",
            "",
            cleaned,
            flags=re.IGNORECASE,
        )
        convert_match = re.match(
            r"\s*convert\s+(.+?)\s+into\s+an?(?:\s+owned)?(?:\s+action)?\.?\s*$",
            cleaned,
            flags=re.IGNORECASE,
        )
        if convert_match:
            cleaned = convert_match.group(1)
        if "|" in cleaned:
            cleaned = self._pipe_row_to_display_text(cleaned)
        cleaned = re.sub(
            r"^\s*(?:left|right)\s+column\s*:\s*",
            "",
            cleaned,
            flags=re.IGNORECASE,
        )
        cleaned = re.sub(r"\.{3,}", "", cleaned)
        return self._repair_display_fragment(" ".join(cleaned.split()).strip(" -:;"))

    def _repair_display_fragment(self, text: str) -> str:
        repaired = re.sub(
            r"\b(data catalogs make benchmark)\s+easier\b",
            r"\1 discovery easier",
            str(text),
            flags=re.IGNORECASE,
        )
        repaired = re.sub(
            r"\b(connect the harness-centric view)\s+an explicit review gate\b",
            r"\1 to an explicit review gate",
            repaired,
            flags=re.IGNORECASE,
        )
        repaired = re.sub(
            r"\b(tie the claim)\s+source-backed\b",
            r"\1 to a source-backed",
            repaired,
            flags=re.IGNORECASE,
        )
        repaired = re.sub(
            r"\bwhat must\.?$",
            "what must be tested.",
            repaired,
            flags=re.IGNORECASE,
        )
        repaired = re.sub(
            r"\bbuild valid\.?$",
            "build valid benchmark cases.",
            repaired,
            flags=re.IGNORECASE,
        )
        trailing = {
            "a",
            "an",
            "and",
            "as",
            "by",
            "for",
            "from",
            "has",
            "have",
            "in",
            "into",
            "need",
            "needs",
            "of",
            "or",
            "provide",
            "provides",
            "the",
            "their",
            "through",
            "to",
            "with",
        }
        words = repaired.split()
        while len(words) > 3 and words[-1].lower().strip(".") in trailing:
            words.pop()
        return " ".join(words).strip(" -:;")

    def _repair_truncated_display_text(self, text: str) -> str:
        repaired = self._repair_display_fragment(text)
        repaired = re.sub(
            r"\s+(?:that|which|because|while|where|when|whose|if|rather than|instead of)\.?$",
            "",
            repaired,
            flags=re.IGNORECASE,
        )
        if re.search(r"\b(?:and|or)\s+\w+\.?$", repaired, re.IGNORECASE):
            repaired = re.sub(
                r"\s+(?:and|or)\s+\w+\.?$",
                "",
                repaired,
                flags=re.IGNORECASE,
            )
        return repaired.strip(" -:;,.")

    def _pipe_row_to_display_text(self, text: str) -> str:
        cells = [
            " ".join(cell.strip(" -:;").split())
            for cell in str(text).split("|")
            if cell.strip(" -:;")
        ]
        header_cells = {
            "artifact",
            "purpose",
            "update trigger",
            "signal",
            "implication",
            "action",
            "owner",
            "timing",
        }
        meaningful = [
            cell
            for cell in cells
            if cell.casefold() not in header_cells
            and not re.fullmatch(r"-+", cell)
        ]
        if not meaningful:
            return ""
        return max(meaningful, key=lambda cell: (len(cell.split()), len(cell)))

    def _split_multi_item_text(self, text: str) -> list[str]:
        cleaned = self._clean_display_text(text)
        if not cleaned:
            return []
        marker_re = re.compile(
            r"(?<!\w)(?:Step|Phase|Stage)\s+\d+\s*[:.)-]?|\b\d{1,2}[.)](?=\s+[A-Z])",
            re.IGNORECASE,
        )
        matches = list(marker_re.finditer(cleaned))
        if len(matches) < 2:
            return [cleaned]
        parts: list[str] = []
        prefix = cleaned[: matches[0].start()].strip(" -:;")
        if prefix and len(prefix.split()) >= 4:
            parts.append(prefix)
        for idx, match in enumerate(matches):
            end = matches[idx + 1].start() if idx + 1 < len(matches) else len(cleaned)
            part = cleaned[match.start() : end].strip(" -:;")
            if part:
                parts.append(part)
        return parts or [cleaned]

    def _truncate_phrase(self, text: str, limit: int) -> str:
        cleaned = self._clean_display_text(text)
        if len(cleaned) <= limit:
            return cleaned
        window = cleaned[:limit]
        for delimiter in (". ", "; ", ": ", " - "):
            index = window.rfind(delimiter)
            if index >= max(24, int(limit * 0.45)):
                return self._repair_truncated_display_text(
                    window[: index + len(delimiter)]
                )
        truncated = self._repair_truncated_display_text(window.rsplit(" ", 1)[0])
        return truncated or self._truncate_at_word(cleaned, limit)

    def _nonduplicate_closing_recommendation(
        self,
        recommendation: str,
        title: str,
    ) -> str:
        cleaned_recommendation = self._clean_display_text(recommendation)
        cleaned_title = self._clean_display_text(title)
        if cleaned_recommendation.casefold() != cleaned_title.casefold():
            return cleaned_recommendation
        lowered = cleaned_title.lower()
        if "reset" in lowered:
            return "Make the reset habit the default operating rule before scaling."
        if "six-phase" in lowered or "cycle" in lowered or "loop" in lowered:
            return "Use the operating loop as the default delivery cadence."
        if "memory" in lowered or "context" in lowered:
            return "Keep persistent context current through named ownership."
        if any(token in lowered for token in ("benchmark", "harness", "ground truth", "evaluation")):
            return "Launch a governed benchmark pilot tied to validated source evidence."
        return "Approve the first governed pilot and make the evidence standard explicit."

    def _icons(self, outline: SlideOutline) -> list[str]:
        icons = outline.layout_json.get("icons") or []
        normalized = [str(icon) for icon in icons if str(icon).strip()]
        if not normalized:
            normalized = ["default"]
        while len(normalized) < 4:
            normalized.append(normalized[-1])
        return normalized[:4]

    def _line(
        self,
        draw: ImageDraw.ImageDraw,
        points: list[tuple[float, float]],
        color,
        width: int,
        rounded: bool = True,
    ) -> None:
        scaled = [self._point(point) for point in points]
        stroke_width = self._stroke(width)
        draw.line(scaled, fill=color, width=stroke_width, joint="curve")
        if rounded:
            radius = stroke_width / 2
            for x, y in scaled:
                draw.ellipse((x - radius, y - radius, x + radius, y + radius), fill=color)

    def _ellipse(
        self,
        draw: ImageDraw.ImageDraw,
        box: tuple[float, float, float, float],
        outline=None,
        fill=None,
        width: int = 1,
    ) -> None:
        draw.ellipse(self._box(box), outline=outline, fill=fill, width=self._stroke(width))

    def _arc(
        self,
        draw: ImageDraw.ImageDraw,
        box: tuple[float, float, float, float],
        start: float,
        end: float,
        color,
        width: int,
    ) -> None:
        draw.arc(self._box(box), start, end, fill=color, width=self._stroke(width))

    def _polygon(self, draw: ImageDraw.ImageDraw, points: list[tuple[float, float]], fill) -> None:
        draw.polygon([self._point(point) for point in points], fill=fill)

    def _rounded_rectangle(
        self,
        draw: ImageDraw.ImageDraw,
        box: tuple[float, float, float, float],
        outline=None,
        fill=None,
        width: int = 1,
        radius: int = 8,
    ) -> None:
        draw.rounded_rectangle(
            self._box(box),
            radius=self._scale(radius),
            outline=outline,
            fill=fill,
            width=self._stroke(width),
        )

    def _dot(self, draw: ImageDraw.ImageDraw, x: float, y: float, radius: float, color) -> None:
        scaled_x, scaled_y = self._point((x, y))
        scaled_radius = self._scale(radius)
        draw.ellipse(
            (
                scaled_x - scaled_radius,
                scaled_y - scaled_radius,
                scaled_x + scaled_radius,
                scaled_y + scaled_radius,
            ),
            fill=color,
        )

    def _point(self, point: tuple[float, float]) -> tuple[int, int]:
        return (self._scale(point[0]), self._scale(point[1]))

    def _box(self, box: tuple[float, float, float, float]) -> tuple[int, int, int, int]:
        return tuple(self._scale(value) for value in box)

    def _scale(self, value: float) -> int:
        return int(round(value * ICON_SCALE))

    def _stroke(self, width: int | float) -> int:
        return max(1, self._scale(width))

    def _tint(self, hex_color: str, white_mix: float) -> str:
        cleaned = self._clean_hex(hex_color)
        try:
            red = int(cleaned[0:2], 16)
            green = int(cleaned[2:4], 16)
            blue = int(cleaned[4:6], 16)
        except ValueError:
            red, green, blue = 0, 0, 0
        white_mix = min(1.0, max(0.0, white_mix))
        tinted = [
            round(channel * (1 - white_mix) + 255 * white_mix)
            for channel in (red, green, blue)
        ]
        return "".join(f"{channel:02X}" for channel in tinted)

    def _relative_luminance(self, hex_color: str) -> float:
        cleaned = self._clean_hex(hex_color)
        channels = []
        for offset in (0, 2, 4):
            value = int(cleaned[offset : offset + 2], 16) / 255.0
            value = (
                value / 12.92
                if value <= 0.03928
                else ((value + 0.055) / 1.055) ** 2.4
            )
            channels.append(value)
        red, green, blue = channels
        return 0.2126 * red + 0.7152 * green + 0.0722 * blue

    def _contrast_ratio(self, fg_hex: str, bg_hex: str) -> float:
        light = self._relative_luminance(fg_hex)
        dark = self._relative_luminance(bg_hex)
        lighter, darker = max(light, dark), min(light, dark)
        return (lighter + 0.05) / (darker + 0.05)

    def _readable_text_color(
        self,
        background_hex: str,
        brand: BrandDNA,
        light: str | None = None,
        dark: str | None = None,
        min_ratio: float = 4.5,
    ) -> str:
        """Pick the most legible text color for a given background.

        Prefer the brand's light/dark text color that best contrasts with the
        background; if neither clears AA (4.5:1), fall back to the higher-contrast
        of white/near-black so text on a light brand fill never goes illegible.
        """
        bg = self._clean_hex(background_hex)
        light_hex = self._clean_hex(light or brand.colors.text_light)
        dark_hex = self._clean_hex(dark or brand.colors.text_dark)
        brand_best = max(
            (light_hex, dark_hex), key=lambda candidate: self._contrast_ratio(candidate, bg)
        )
        if self._contrast_ratio(brand_best, bg) >= min_ratio:
            return brand_best
        return (
            "FFFFFF"
            if self._contrast_ratio("FFFFFF", bg) >= self._contrast_ratio("111111", bg)
            else "111111"
        )

    def _shape_fill_hex(self, shape) -> str | None:
        try:
            rgb = shape.fill.fore_color.rgb
        except Exception:
            return None
        if rgb is None:
            return None
        return str(rgb)

    def _fit_font_size(
        self,
        text: str,
        box_width_in: float,
        sizes: list[int],
        max_lines: int = 2,
        char_factor: float = 0.52,
    ) -> int:
        """Pick the largest size (from sizes, largest-first) whose wrapped text fits.

        Width-aware: estimates characters-per-line from the box width and point
        size instead of the raw character count, so a wide box keeps large type.
        Pair with a TEXT_TO_FIT_SHAPE autofit frame to absorb any residual overflow.
        """
        length = len(" ".join(str(text).split()))
        if length == 0 or not sizes:
            return sizes[0] if sizes else 12
        for size in sizes:
            chars_per_line = max(1, int((box_width_in * 72) / (char_factor * size)))
            if chars_per_line * max_lines >= length:
                return size
        return sizes[-1]

    def _autofit(self, frame) -> None:
        try:
            frame.word_wrap = True
            frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        except Exception:
            pass

    def _rgba(self, hex_color: str) -> tuple[int, int, int, int]:
        cleaned = self._clean_hex(hex_color)
        try:
            return (
                int(cleaned[0:2], 16),
                int(cleaned[2:4], 16),
                int(cleaned[4:6], 16),
                255,
            )
        except ValueError:
            return (0, 0, 0, 255)

    def _clean_hex(self, hex_color: str) -> str:
        cleaned = (hex_color or "000000").replace("#", "")[:6]
        if len(cleaned) != 6 or not re.fullmatch(r"[0-9A-Fa-f]{6}", cleaned):
            return "000000"
        return cleaned.upper()

    def _rgb(self, hex_color: str) -> RGBColor:
        return RGBColor.from_string(self._clean_hex(hex_color))

    def _truncate_at_word(self, text: str, limit: int) -> str:
        cleaned = " ".join(str(text).split())
        if len(cleaned) <= limit:
            return cleaned
        return self._repair_truncated_display_text(cleaned[:limit].rsplit(" ", 1)[0])
