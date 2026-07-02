"""Native primitive builders — one per slide kind, a faithful port of the HTML
``templates.py`` primitives using native editable shapes. Reuses the existing
content-extraction helpers (which return plain data, not HTML) and draws with the
``components`` polish toolkit + ``geometry`` auto-layout.
"""

from __future__ import annotations

import math

from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from app.services.slide_design import fit
from app.services.slide_design.content import (
    _as_text,
    _clean_sentence,
    _eyebrow_text,
    _normalize_items,
    _trim_dangling,
)

from . import components as C
from . import icons
from .geometry import Rect, column_split, grid, stack
from .theme import Theme, card_gap_in, pt

Para = C.Para

# Leading list marker for supporting ticks. Must NOT be one of the unicode
# bullet glyphs the rendered-slide audit rejects ([•◦▪●]);
# an en dash reads as a clean list marker and passes.
_TICK = "–  "


def _title_pt(title: str, theme: Theme) -> float:
    n = len(title or "")
    base = 46 if n <= 38 else (38 if n <= 64 else 31)
    return pt(base, theme.heading_scale)


def _text_box_height(
    text: str, size_pt: float, width_in: float, *, min_lines: int = 1, slack_in: float = 0.05
) -> float:
    """Box height (inches) that exactly holds ``text`` at ``size_pt`` across
    ``width_in`` per the shared fit model, so a header text box neither overflows
    (cut-off risk) nor over-reserves vertical space (which made the title box
    swallow the subhead box and read as an overlap to the audit)."""
    # Conservative wrap estimate: real serif-bold rendering is up to ~15% wider
    # than the generic 0.52x-pt model. A title sized to "exactly one line" by the
    # optimistic estimate wrapped to two in Office and overlapped the subhead
    # below it, so reserve for the wider real-world metrics.
    cpl = max(8, int(fit.chars_per_line(width_in * 72, size_pt) * 0.85))
    lines = max(min_lines, math.ceil(len(text) / max(1, cpl)))
    # capacity_lines = h_pt / (size_pt * 1.18); invert so capacity ~= lines.
    return (lines * size_pt * 1.18) / 72 + slack_in


def _short(value, max_chars: int) -> str:
    """Shorten text to ``max_chars`` at a word boundary, never mid-word, and trim
    any trailing dangling connective. Used for title-fallback display so a long
    untitled body never renders as ``"...experience goals, a"``."""
    t = " ".join(_as_text(value).split())
    if len(t) > max_chars:
        # word-boundary cut, then drop the trailing comma/dash the cut left behind
        t = t[:max_chars].rsplit(" ", 1)[0].rstrip(" ,;:—–-")
    return _trim_dangling(t)


# Icon-circle glyphs: a meaning symbol when the item's hint/lead signals one,
# else a monogram of the lead word. Never an empty dot (reads as placeholder)
# and never the audit-flagged bullet glyphs [•◦▪●].
_GLYPH_RULES: list[tuple[tuple[str, ...], str]] = [
    (("check", "valid", "success", "complete", "pass", "proof", "benefit", "win"), "✓"),
    (("risk", "fail", "problem", "warning", "gap", "break", "anti", "threat", "error"), "!"),
    (("shift", "transition", "migrat", "adopt", "handoff", "flow", "pipeline"), "→"),
    (("grow", "increase", "scale", "improve", "gain", "expand", "accelerat"), "↑"),
    (("reduc", "declin", "fewer", "lower", "shrink", "cut "), "↓"),
    (("question", "unknown", "uncertain", "decision", "decide"), "?"),
]


def _item_glyph(it: dict) -> str:
    signal = f"{it.get('icon') or ''} {it.get('title') or ''}".lower()
    for keys, glyph in _GLYPH_RULES:
        if any(k in signal for k in keys):
            return glyph
    lead = (it.get("title") or it.get("body") or "").strip()
    return lead[:1].upper() if lead[:1].isalpha() else (lead[:1] or "")


def _add_item_icon(slide, cx: float, cy: float, d: float, accent, it: dict, index: int) -> None:
    """Reference-deck icon chip: accent circle + white react-icons (Feather)
    icon resolved from the item's hint/lead. Glyph/monogram fallback when the
    Node icon worker is unavailable."""
    name = icons.resolve_icon(
        str(it.get("icon") or ""), str(it.get("title") or ""), str(it.get("body") or ""), index
    )
    C.add_icon_circle(slide, cx, cy, d, accent, _item_glyph(it),
                      icon_path=icons.icon_file(name, "FFFFFF"))


def _lead_body_paras(it: dict, theme: Theme, pal, *, lead_pt: float = 18.0,
                     body_pt: float = 15.5, lead_scale: float = 1.0,
                     lead_color=None, body_color=None,
                     serif_lead: bool = True, line_spacing: float = 1.35) -> list:
    """Paragraphs for an item card/row (sizes in CSS px, like every primitive).
    An untitled item renders its FULL body as the content, slightly larger —
    never a chopped pseudo-title with the rest of the sentence discarded."""
    title = str(it.get("title") or "")
    body_text = str(it.get("body") or "")
    lead_color = lead_color if lead_color is not None else pal.headline
    body_color = body_color if body_color is not None else pal.body
    paras: list = []
    if title:
        paras.append(Para(title, pt(lead_pt, lead_scale), lead_color, bold=True,
                          font=_serif(theme) if serif_lead else None,
                          line_spacing=1.1, space_after_pt=3))
        if body_text:
            paras.append(Para(body_text, pt(body_pt), body_color, line_spacing=line_spacing))
    elif body_text:
        paras.append(Para(body_text, pt(body_pt + 1.5), body_color, line_spacing=line_spacing))
    return paras


def _header(slide, content: dict, theme: Theme, pal, area: Rect, *, title_size: float | None = None) -> Rect:
    """Draw eyebrow + title + subhead at the top of ``area``; return the body rect.

    Each text box is sized to its own estimated height (and ``y`` advances by the
    same amount) so consecutive boxes stack without overlapping.
    """
    eyebrow = _eyebrow_text(content)
    title = _trim_dangling(_clean_sentence(content.get("action_title") or content.get("title") or ""))
    sub = _trim_dangling(_clean_sentence(content.get("subheading") or content.get("summary") or ""))
    y = area.y
    if eyebrow:
        eh = _text_box_height(eyebrow.upper(), pt(14), area.w, slack_in=0.02)
        C.add_paragraphs(slide, Rect(area.x, y, area.w, eh),
                         [Para(eyebrow.upper(), pt(14), pal.eyebrow, bold=True)])
        y += eh + 0.06
    if title:
        size = title_size or _title_pt(title, theme)
        th = _text_box_height(title, size, area.w)
        C.add_paragraphs(slide, Rect(area.x, y, area.w, th),
                         [Para(title, size, pal.headline, bold=True, font=_serif(theme), line_spacing=1.05)])
        y += th + 0.08
    if sub and len(sub.split()) > 6:
        sh = _text_box_height(sub, pt(18), area.w)
        C.add_paragraphs(slide, Rect(area.x, y, area.w, sh),
                         [Para(sub, pt(18), pal.body, line_spacing=1.3)])
        y += sh + 0.10
    y += 0.12
    return Rect(area.x, y, area.w, max(0.5, area.bottom - y))


def _serif(theme: Theme) -> str | None:
    # first family in the serif stack (the brand/display font)
    return (theme.serif.split(",")[0].strip().strip('"') or None)


def _footer(slide, content: dict, theme: Theme, pal, index: int, total: int) -> None:
    labels = content.get("source_labels") or content.get("sources") or content.get("source_refs") or []
    src = ""
    for label in labels:
        t = _clean_sentence(_as_text(label))
        if t and t.lower() not in ("uploaded source", "source needed", "[source needed]"):
            src = t
            break
    if src:
        C.add_paragraphs(slide, Rect(0.875, 7.04, 9.0, 0.3), [Para(f"Source: {src}", pt(11.5), pal.body)])
    C.add_paragraphs(slide, Rect(11.0, 7.04, 1.45, 0.3),
                     [Para(f"{index + 1:02d} / {total:02d}", pt(11.5), pal.body, align=PP_ALIGN.RIGHT)])


# --------------------------------------------------------------------------- #
# Item-card grid (shared by cards/steps/checklist/anti_pattern)
# --------------------------------------------------------------------------- #
def _card_content_height(items: list[dict], text_w_in: float, *, lead_pt: float = 18.0,
                         body_pt: float = 15.5) -> float:
    """Tallest card's content height (inches): title + body wrap estimate plus
    padding — so cards hug their content instead of floating text in a mostly
    empty box."""
    tallest = 0.0
    width_pt = max(1.0, text_w_in * 72)
    for it in items:
        h = 0.5  # top/bottom padding
        title = str(it.get("title") or "")
        body_text = str(it.get("body") or "")
        if title:
            h += fit.estimate_lines([title], width_pt, lead_pt) * (lead_pt * 1.15) / 72 + 0.05
        if body_text:
            size = body_pt if title else body_pt + 1.5
            h += fit.estimate_lines([body_text], width_pt, size) * (size * 1.4) / 72
        tallest = max(tallest, h)
    return tallest


def _card_grid(slide, content, theme, pal, body: Rect, *, numbered=False, max_n=6):
    items = _normalize_items(content)[:max_n]
    n = len(items) or 1
    cols = 1 if n == 1 else (2 if n in (2, 4) else 3)
    gap = card_gap_in(theme)
    n_rows = -(-n // cols)
    # Few items get BIGGER panels and type, not more whitespace — two tiny
    # cards floating in an empty canvas is the opposite of consultant-grade.
    duo = n <= 2
    lead_pt_v = 22.0 if duo else 18.0
    body_pt_v = 17.0 if duo else 15.5
    icon_d = 0.6 if duo else 0.5
    # content-sized card height (clamped), vertically centered block — matches
    # the CSS ``align-content: center`` without the dead space a fixed-height
    # card leaves under two lines of text.
    cap_h = 3.4 if duo else (2.7 if n_rows == 1 else 2.4)
    cell_w = (body.w - gap * (cols - 1)) / cols
    content_h = _card_content_height(items, cell_w - 0.36 - 0.62,
                                     lead_pt=lead_pt_v, body_pt=body_pt_v)
    # Duo panels get a height floor so two cards command the canvas instead of
    # floating as thin strips in empty space.
    card_h = min((body.h - gap * (n_rows - 1)) / n_rows, cap_h,
                 max(2.4 if duo else 1.15, content_h))
    used_h = card_h * n_rows + gap * (n_rows - 1)
    top = body.y + max(0.0, (body.h - used_h) / 2)
    cells = grid(n, cols, Rect(body.x, top, body.w, used_h), gap)
    for i, (it, cell) in enumerate(zip(items, cells)):
        C.add_card(slide, cell, theme, pal)
        inner = cell.inset(0.18, 0.2 if not duo else 0.3)
        cy = inner.y + 0.26 + (0.06 if duo else 0.0)
        accent = pal.teal if i % 2 == 0 else pal.gold
        if numbered:
            C.add_badge(slide, inner.x + 0.24, cy, 0.42, accent, number=str(i + 1))
        else:
            _add_item_icon(slide, inner.x + 0.26, cy, icon_d, accent, it, i)
        tx = inner.x + (0.72 if duo else 0.62)
        tw = inner.right - tx
        paras = _lead_body_paras(it, theme, pal, lead_pt=lead_pt_v, body_pt=body_pt_v,
                                 line_spacing=1.4)
        if paras:
            C.add_paragraphs(slide, Rect(tx, inner.y + 0.04, tw, inner.h - 0.04), paras)


def cards(slide, content, theme, pal, body):
    _card_grid(slide, content, theme, pal, body, numbered=False, max_n=6)


def steps(slide, content, theme, pal, body):
    _card_grid(slide, content, theme, pal, body, numbered=True, max_n=8)


def rows(slide, content, theme, pal, body):
    items = _normalize_items(content)[:4]
    n = len(items) or 1
    cells = stack(body, n, card_gap_in(theme))
    for i, (it, cell) in enumerate(zip(items, cells)):
        C.add_card(slide, cell, theme, pal)
        inner = cell.inset(0.12, 0.24)
        accent = pal.teal if i % 2 == 0 else pal.gold
        _add_item_icon(slide, inner.x + 0.26, inner.cy, 0.5, accent, it, i)
        tx = inner.x + 0.66
        if it.get("title"):
            C.add_paragraphs(slide, Rect(tx, cell.y, 3.0, cell.h),
                             [Para(it["title"], pt(19), pal.headline, bold=True, line_spacing=1.05)],
                             anchor=MSO_ANCHOR.MIDDLE)
            if it.get("body"):
                C.add_paragraphs(slide, Rect(inner.x + 3.8, cell.y, inner.right - (inner.x + 3.8), cell.h),
                                 [Para(it["body"], pt(15.5), pal.body, line_spacing=1.35)],
                                 anchor=MSO_ANCHOR.MIDDLE)
        else:
            # Untitled: the full sentence IS the row — render it whole across
            # the row instead of a 60-char chop with the remainder discarded.
            C.add_paragraphs(slide, Rect(tx, cell.y, inner.right - tx, cell.h),
                             [Para(it.get("body") or "", pt(16.5), pal.body, line_spacing=1.35)],
                             anchor=MSO_ANCHOR.MIDDLE)


def statement(slide, content, theme, pal, body):
    title = _clean_sentence(content.get("action_title") or content.get("title") or "")
    items = _normalize_items(content)
    eyebrow = _eyebrow_text(content) or "The Big Idea"
    if len(items) >= 2:
        left, right = column_split(body, [1.05, 0.95], 0.5)
        C.add_paragraphs(slide, Rect(left.x, left.y, left.w, 0.3), [Para(eyebrow.upper(), pt(14), pal.eyebrow, bold=True)])
        C.add_paragraphs(slide, Rect(left.x, left.y + 0.4, left.w, left.h - 0.4),
                         [Para(title, pt(40, theme.heading_scale), pal.headline, bold=True, font=_serif(theme), line_spacing=1.08)])
        ticks = []
        for it in items[:4]:
            label = (it.get("title") + ". " if it.get("title") else "") + (it.get("body") or "")
            ticks.append(Para(f"{_TICK}{label.strip()}", pt(16), pal.body, line_spacing=1.3, space_after_pt=10))
        C.add_paragraphs(slide, right, ticks, anchor=MSO_ANCHOR.MIDDLE)
    else:
        C.add_paragraphs(slide, Rect(body.x, body.y, body.w, 0.3), [Para(eyebrow.upper(), pt(14), pal.eyebrow, bold=True)])
        size = pt(52 if len(title) <= 50 else 38, theme.heading_scale)
        C.add_paragraphs(slide, Rect(body.x, body.y + 0.4, body.w, body.h - 0.4),
                         [Para(title, size, pal.headline, bold=True, font=_serif(theme), line_spacing=1.1)],
                         anchor=MSO_ANCHOR.MIDDLE)


def quote(slide, content, theme, pal, body):
    title = _clean_sentence(content.get("action_title") or "")
    ex = content.get("exhibit_spec") or {}
    q = ""
    # Prefer the exhibit's grounded source quote; otherwise use the first bullet
    # that is a COMPLETE sentence — a truncated fragment displayed as a pull
    # quote ("In professional software development, decisions") reads broken.
    candidates = [ex.get("quote"), ex.get("key_idea")] + list(content.get("bullets") or [])
    for b in candidates:
        t = _trim_dangling(_clean_sentence(_as_text(b)))
        if 24 <= len(t) <= 180 and (t[-1:] in ".!?\"”’" or t == title):
            q = t
            break
    q = q or title
    attr = _clean_sentence(content.get("subheading") or "")
    C.add_paragraphs(slide, Rect(body.x, body.y, 1.5, 1.0), [Para("“", pt(110), pal.gold, font=_serif(theme))])
    paras = [Para(q, pt(34, theme.heading_scale), pal.headline, italic=True, font=_serif(theme), line_spacing=1.2)]
    if attr:
        paras.append(Para(f"— {attr}", pt(16), pal.eyebrow, bold=True, space_after_pt=0))
    C.add_paragraphs(slide, Rect(body.x, body.y + 0.9, body.w, body.h - 0.9), paras, anchor=MSO_ANCHOR.MIDDLE)


def split(slide, content, theme, pal, body):
    items = _normalize_items(content)
    left_i = items[0] if items else {"title": "", "body": ""}
    right_i = items[1] if len(items) > 1 else left_i
    cols = column_split(body, [1.0, 1.0], 0.5)
    for idx, (it, col) in enumerate(zip((left_i, right_i), cols)):
        accent = idx == 1
        C.add_card(slide, col, theme, pal, fill=pal.teal if accent else None)
        inner = col.inset(0.3)
        kicker = "TODAY" if idx == 0 else "WITH THE SHIFT"
        kcolor = pal.on_accent if accent else pal.eyebrow
        tcolor = pal.on_accent if accent else pal.headline
        lead = it.get("title") or _as_text(it.get("body"))
        C.add_paragraphs(slide, inner, [
            Para(kicker, pt(14), kcolor, bold=True, space_after_pt=8),
            Para(lead, pt(24, theme.heading_scale), tcolor, italic=True, font=_serif(theme), line_spacing=1.2),
        ], anchor=MSO_ANCHOR.MIDDLE)


def callout_list(slide, content, theme, pal, body):
    items = _normalize_items(content)
    if not items:
        return cards(slide, content, theme, pal, body)
    feat = items[0]
    rest = items[1:5]
    cols = column_split(body, [0.82, 1.18], 0.3)
    C.add_card(slide, cols[0], theme, pal, fill=pal.teal)
    finner = cols[0].inset(0.3)
    feat_title = feat.get("title") or ""
    feat_body = feat.get("body") or ""
    feat_paras = []
    if feat_title:
        feat_paras.append(Para(feat_title, pt(25, theme.heading_scale), pal.on_accent, bold=True,
                               font=_serif(theme), space_after_pt=8))
        if feat_body:
            feat_paras.append(Para(feat_body, pt(15.5), pal.on_accent, line_spacing=1.4))
    else:
        # An untitled feature is a single statement: set it large so the panel
        # reads as a deliberate callout, not a lost caption in an empty box.
        size = C.fit_size(feat_body, finner, pt(24, theme.heading_scale), min_pt=15.5)
        feat_paras.append(Para(feat_body, size, pal.on_accent, bold=True,
                               font=_serif(theme), line_spacing=1.25))
    C.add_paragraphs(slide, finner, feat_paras, anchor=MSO_ANCHOR.MIDDLE)
    if rest:
        mini = stack(cols[1], len(rest), card_gap_in(theme))
        solo = len(rest) == 1  # one supporting card fills the column, larger type
        for i, (it, cell) in enumerate(zip(rest, mini)):
            C.add_card(slide, cell, theme, pal)
            inner = cell.inset(0.1, 0.2)
            accent = pal.teal if i % 2 == 0 else pal.gold
            _add_item_icon(slide, inner.x + 0.22, inner.cy, 0.5 if solo else 0.42, accent, it, i)
            paras = _lead_body_paras(
                it, theme, pal,
                lead_pt=20 if solo else 17,
                body_pt=16 if solo else 14.5,
                serif_lead=False, line_spacing=1.3,
            )
            C.add_paragraphs(slide, Rect(inner.x + 0.56, cell.y, inner.right - (inner.x + 0.56), cell.h),
                             paras, anchor=MSO_ANCHOR.MIDDLE)


def _metric_value_display(value: str) -> str:
    """Compact a large bare integer for the metric tile (200000 -> 200K, 3500000 ->
    3.5M). Leaves values with units/decimals/percent untouched."""
    raw = str(value or "").strip()
    if not raw.isdigit():
        return raw
    n = int(raw)
    if n >= 1_000_000:
        return f"{n / 1_000_000:.1f}".rstrip("0").rstrip(".") + "M"
    if n >= 10_000:
        return f"{round(n / 1000)}K"
    return raw


def _metric_tile_value(m: dict) -> str:
    """Value + unit for a metric tile — a bare '95' with no unit reads as noise;
    '95%' or '95 hrs' reads as a statistic."""
    value = _metric_value_display(m.get("value", ""))
    unit = str(m.get("unit") or "").strip()
    if not unit or not value or unit.lower() in value.lower():
        return value
    return f"{value}{unit}" if unit in ("%", "x", "×") else f"{value} {unit}"


def metrics(slide, content, theme, pal, body):
    data = []
    for m in content.get("metrics") or []:
        if isinstance(m, dict):
            data.append((_metric_tile_value(m), str(m.get("label") or m.get("name") or ""), str(m.get("description") or "")))
    if not data:
        chart = content.get("chart_spec") or {}
        for p in chart.get("data_points") or []:
            if isinstance(p, dict):
                data.append((_metric_tile_value(p), str(p.get("label", "")), ""))
    if not data:
        return cards(slide, content, theme, pal, body)
    data = data[:4]
    cells = grid(len(data), len(data), body, card_gap_in(theme))
    for (value, label, desc), cell in zip(data, cells):
        paras = [
            Para(value, pt(60, theme.heading_scale), pal.headline, bold=True, font=_serif(theme), space_after_pt=4),
            Para(label.upper(), pt(15), pal.eyebrow, bold=True, space_after_pt=4),
        ]
        if desc:
            paras.append(Para(desc, pt(14), pal.body, line_spacing=1.3))
        C.add_paragraphs(slide, cell, paras, anchor=MSO_ANCHOR.MIDDLE)


big_stat = metrics


def table(slide, content, theme, pal, body):
    ex = content.get("exhibit_spec") or {}
    cols = [str(c) for c in (ex.get("columns") or [])]
    rows_data = ex.get("rows") or []
    if not cols or not rows_data:
        return cards(slide, content, theme, pal, body)
    rows_data = rows_data[:6]
    tbl_shape = slide.shapes.add_table(len(rows_data) + 1, len(cols), Inches(body.x), Inches(body.y),
                                       Inches(body.w), Inches(min(body.h, 0.5 * (len(rows_data) + 1)))).table
    # Kill the default PowerPoint table style (banded blue) so the table carries
    # the deck theme: ink header, alternating theme-tinted body rows.
    tbl_shape.first_row = False
    tbl_shape.horz_banding = False
    header_fill = C._resolve(theme.ink, pal.bg)
    row_fill = C._resolve(theme.card_bg("light"), pal.bg)
    row_alt_fill = C._resolve(theme.card_light_alt, pal.bg)
    for c, col in enumerate(cols):
        cell = tbl_shape.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(*header_fill)
        cell.text = _clean_sentence(col)
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for r, row in enumerate(rows_data, start=1):
        cells = (row.get("label", ""), *row.get("values", [])) if isinstance(row, dict) else tuple(row)
        fill = row_fill if r % 2 == 1 else row_alt_fill
        for c in range(len(cols)):
            val = cells[c] if c < len(cells) else ""
            tcell = tbl_shape.cell(r, c)
            tcell.fill.solid()
            tcell.fill.fore_color.rgb = RGBColor(*fill)
            tcell.text = _clean_sentence(str(val))
            for p in tcell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(12)
                    run.font.bold = c == 0
                    run.font.color.rgb = RGBColor(*C._resolve(theme.text_dark, fill))


def matrix(slide, content, theme, pal, body):
    items = _normalize_items(content)[:4]
    while len(items) < 4:
        items.append({"title": "", "body": ""})
    cells = grid(4, 2, body, card_gap_in(theme))
    labels = ["Top-left", "Top-right", "Bottom-left", "Bottom-right"]
    ex = content.get("exhibit_spec") or {}
    qlabels = [str(x) for x in (ex.get("quadrant_labels") or [])]
    for i, (it, cell) in enumerate(zip(items, cells)):
        C.add_card(slide, cell, theme, pal)
        inner = cell.inset(0.22, 0.24)
        paras = [
            Para((qlabels[i] if i < len(qlabels) else labels[i]).upper(), pt(14), pal.eyebrow, bold=True, space_after_pt=4),
            *_lead_body_paras(it, theme, pal, lead_pt=18, body_pt=14.5, serif_lead=False, line_spacing=1.35),
        ]
        C.add_paragraphs(slide, inner, paras)


def timeline(slide, content, theme, pal, body):
    items = _normalize_items(content)[:6]
    n = len(items) or 1
    cells = stack(body, n, 0.14)
    for i, (it, cell) in enumerate(zip(items, cells)):
        C.add_badge(slide, cell.x + 0.3, cell.cy, 0.46, pal.teal if i % 2 == 0 else pal.gold, number=str(i + 1))
        C.add_paragraphs(slide, Rect(cell.x + 0.7, cell.y, cell.w - 0.7, cell.h),
                         _lead_body_paras(it, theme, pal, lead_pt=18, body_pt=15,
                                          serif_lead=False, line_spacing=1.35),
                         anchor=MSO_ANCHOR.MIDDLE)


def layers(slide, content, theme, pal, body):
    items = _normalize_items(content)[:5]
    n = len(items) or 1
    cells = stack(body, n, card_gap_in(theme))
    for i, (it, cell) in enumerate(zip(items, cells)):
        C.add_card(slide, cell, theme, pal)
        inner = cell.inset(0.1, 0.24)
        if it.get("title"):
            C.add_paragraphs(slide, Rect(inner.x, cell.y, 2.8, cell.h),
                             [Para(it["title"], pt(18), pal.headline, bold=True)],
                             anchor=MSO_ANCHOR.MIDDLE)
            if it.get("body"):
                C.add_paragraphs(slide, Rect(inner.x + 3.0, cell.y, inner.right - (inner.x + 3.0), cell.h),
                                 [Para(it["body"], pt(15), pal.body, line_spacing=1.35)],
                                 anchor=MSO_ANCHOR.MIDDLE)
        else:
            C.add_paragraphs(slide, Rect(inner.x, cell.y, inner.w, cell.h),
                             [Para(it.get("body") or "", pt(16), pal.body, line_spacing=1.35)],
                             anchor=MSO_ANCHOR.MIDDLE)


def columns(slide, content, theme, pal, body):
    items = _normalize_items(content)[:3]
    if not items:
        return cards(slide, content, theme, pal, body)
    cols = column_split(body, [1.0] * len(items), card_gap_in(theme))
    for it, col in zip(items, cols):
        C.add_card(slide, col, theme, pal)
        inner = col.inset(0.26)
        C.add_paragraphs(slide, inner,
                         _lead_body_paras(it, theme, pal, lead_pt=21,
                                          lead_scale=theme.heading_scale,
                                          body_pt=15, line_spacing=1.4))


def cover(slide, content, theme, pal, body):
    title = _clean_sentence(content.get("action_title") or content.get("deck_title") or content.get("title") or "")
    sub = _clean_sentence(content.get("subheading") or content.get("summary") or "")
    C.add_paragraphs(slide, Rect(body.x, 2.2, body.w, 0.4),
                     [Para("WHITEPAPER  ·  EXECUTIVE BRIEFING", pt(15), pal.gold, bold=True)])
    size = pt(74 if len(title) <= 46 else (56 if len(title) <= 90 else 44), theme.heading_scale)
    th = _text_box_height(title, size, body.w)
    C.add_paragraphs(slide, Rect(body.x, 2.7, body.w, th),
                     [Para(title, size, pal.headline, bold=True, font=_serif(theme), line_spacing=1.04)])
    if sub:
        # Subtitle sits below the measured title block — a fixed y overlapped
        # the fourth line of a long deck title.
        sub_y = min(6.4, max(5.1, 2.7 + th + 0.3))
        C.add_paragraphs(slide, Rect(body.x, sub_y, min(body.w, 8.5), 0.9),
                         [Para(sub, pt(21), pal.body, line_spacing=1.4)])


def closing(slide, content, theme, pal, body):
    ex = content.get("exhibit_spec") or {}
    steps_list = ex.get("next_steps") or [_as_text(b) for b in (content.get("bullets") or [])]
    steps_list = [_clean_sentence(_as_text(s)) for s in steps_list if _as_text(s).strip()][:4]
    ask = _clean_sentence(ex.get("decision_ask") or "")
    # The closing is a real slide, not a bare checklist: it opens with its own
    # eyebrow + action title (this primitive draws its own header).
    title = _trim_dangling(_clean_sentence(content.get("action_title") or content.get("title") or ""))
    y = body.y
    if title:
        eyebrow = _eyebrow_text(content) or "THE DECISION"
        C.add_paragraphs(slide, Rect(body.x, y, body.w, 0.3),
                         [Para(eyebrow.upper(), pt(14), pal.eyebrow, bold=True)])
        y += 0.42
        size = _title_pt(title, theme)
        th = _text_box_height(title, size, body.w)
        C.add_paragraphs(slide, Rect(body.x, y, body.w, th),
                         [Para(title, size, pal.headline, bold=True, font=_serif(theme), line_spacing=1.05)])
        y += th + 0.25
    body = Rect(body.x, y, body.w, max(0.8, body.bottom - y))
    band_w = min(body.w, 8.6)
    # Size the ask band to its text — a fixed-height band overflows on a
    # three-line decision ask and the last line spills past the band edge.
    band_h = 0.0
    if ask:
        band_h = max(0.62, _text_box_height(ask, pt(19), band_w - 0.6, slack_in=0.0) + 0.24)
    cells = stack(Rect(body.x, body.y, min(body.w, 8.0), body.h - (band_h + 0.25 if ask else 0)),
                  max(1, len(steps_list)), 0.18, row_h=0.6)
    for i, (step, cell) in enumerate(zip(steps_list, cells), start=1):
        C.add_badge(slide, cell.x + 0.22, cell.cy, 0.4, pal.gold, number=str(i))
        C.add_paragraphs(slide, Rect(cell.x + 0.6, cell.y, cell.w - 0.6, cell.h),
                         [Para(step, pt(19), pal.body, line_spacing=1.3)], anchor=MSO_ANCHOR.MIDDLE)
    if ask:
        band_y = body.bottom - band_h - 0.05
        C.add_card(slide, Rect(body.x, band_y, band_w, band_h), theme, pal, fill=pal.gold, shadow=False)
        C.add_paragraphs(slide, Rect(body.x + 0.3, band_y + 0.08, band_w - 0.6, band_h - 0.16),
                         [Para(ask, pt(19), RGBColor(0x19, 0x12, 0x0A), bold=True, line_spacing=1.18)],
                         anchor=MSO_ANCHOR.MIDDLE)


# slide_type primitive -> builder
BUILDERS = {
    "cover": cover,
    "closing": closing,
    "statement": statement,
    "quote": quote,
    "cards": cards,
    "rows": rows,
    "steps": steps,
    "split": split,
    "table": table,
    "callout_list": callout_list,
    "metrics": metrics,
    "matrix": matrix,
    "timeline": timeline,
    "layers": layers,
    "columns": columns,
    "big_stat": big_stat,
}
