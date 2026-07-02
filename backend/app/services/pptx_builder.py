import json
from pathlib import Path
from typing import Iterable

from app.models.outline import SlideOutline
from app.models.template import TemplateProfile
from app.services.authored_pptx_renderer import AuthoredPptxRenderer
from app.services.brand_layout_renderer import BrandLayoutInstantiationRenderer
from app.services.brand_template_renderer import BrandTemplateCloneRenderer
from app.services.generation_editing_contract import GenerationEditingContract
from app.services.pptx_native.renderer import NativePptxRenderer
from app.services.hybrid_assembler import HybridAssembler, SlideReplacement
from app.services.pptx_renderer import DeterministicPptxRenderer
from app.services.strict_injector import StrictSlideInjector
from app.workers.node_runner import NodePptxGenRunner


class PptxBuilder:
    def __init__(
        self,
        node_runner: NodePptxGenRunner,
        renderer_engine: str = "authored",
        brand_layout_instantiation: bool = False,
    ) -> None:
        self.node_runner = node_runner
        self.hybrid_assembler = HybridAssembler()
        self.strict_injector = StrictSlideInjector()
        self.legacy_renderer = DeterministicPptxRenderer()
        self.authored_renderer = AuthoredPptxRenderer(self.legacy_renderer)
        self.native_renderer = NativePptxRenderer()
        self.brand_template_renderer = BrandTemplateCloneRenderer()
        self.brand_layout_renderer = BrandLayoutInstantiationRenderer()
        self.editing_contract = GenerationEditingContract()
        self.renderer_engine = renderer_engine.strip().lower() or "authored"
        self.brand_layout_instantiation = brand_layout_instantiation

    def build_deck(
        self,
        template: TemplateProfile,
        outlines: list[SlideOutline],
        output_path: Path,
        working_dir: Path,
    ) -> list[dict[str, str | int]]:
        if template.type in {"freeform", "brand"}:
            prepared = self.prepare_outlines(template, outlines)
            if template.type == "brand" and self.renderer_engine != "legacy":
                pre_warnings: list[dict[str, str | int]] = []
                if self.brand_layout_instantiation:
                    instantiated, instantiate_warnings = self.brand_layout_renderer.render(
                        template,
                        prepared,
                        output_path,
                        working_dir,
                    )
                    if instantiated:
                        return instantiate_warnings
                    pre_warnings = instantiate_warnings
                rendered, clone_warnings = self.brand_template_renderer.render(
                    template,
                    prepared,
                    output_path,
                    working_dir,
                )
                if rendered:
                    return [*pre_warnings, *clone_warnings]
                fallback_outlines = self._strip_template_frames(prepared)
                render_warnings = self._render_generated(
                    fallback_outlines,
                    template.brand,
                    output_path,
                )
                return [*pre_warnings, *clone_warnings, *render_warnings]
            return self._render_generated(
                prepared,
                template.brand,
                output_path,
            )

        strict_output = working_dir / "strict.pptx"
        flex_output = working_dir / "flexible.pptx"
        strict_warnings = self.strict_injector.inject(
            Path(template.source_file), template, outlines, strict_output
        )

        flexible_outlines = [outline for outline in outlines if outline.mode == "flexible"]
        if not flexible_outlines:
            output_path.write_bytes(strict_output.read_bytes())
            return strict_warnings

        render_warnings = self.legacy_renderer.render(
            flexible_outlines, template.brand, flex_output, enable_diagrams=False
        )

        replacements = []
        flex_number = 1
        for slide_spec in template.slides:
            if slide_spec.mode == "flexible":
                replacements.append(
                    SlideReplacement(
                        template_slide_number=slide_spec.index + 1,
                        flexible_slide_number=flex_number,
                    )
                )
                flex_number += 1

        self.hybrid_assembler.assemble(
            strict_output, flex_output, replacements, output_path
        )
        return [*strict_warnings, *render_warnings]

    def _renderer(self):
        if self.renderer_engine == "legacy":
            return self.legacy_renderer
        if self.renderer_engine == "authored":
            return self.authored_renderer
        # "native" and any unrecognized value (e.g. the removed "html" image
        # engine) resolve to the fully editable native renderer, so a stale env
        # setting can never silently select an image-inserting path.
        return self.native_renderer

    def _render_generated(
        self,
        outlines: list[SlideOutline],
        brand,
        output_path: Path,
    ) -> list[dict[str, str | int]]:
        """Render generated freeform/brand decks with the selected native engine."""
        warnings = self._renderer().render(outlines, brand, output_path)
        warnings.extend(self._editability_audit(output_path))
        return warnings

    def _editability_audit(self, output_path: Path) -> list[dict[str, str | int]]:
        """Generated decks must download as editable PowerPoint, never as slides
        flattened to images. Flag any slide whose area is dominated by pictures
        so a rasterizing regression (or an image-heavy engine) is caught at
        build time instead of by a user opening the file."""
        try:
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE

            prs = Presentation(Path(output_path).as_posix())
        except Exception:
            return []
        slide_area = float((prs.slide_width or 0) * (prs.slide_height or 0))
        if not slide_area:
            return []
        warnings: list[dict[str, str | int]] = []
        for index, slide in enumerate(prs.slides):
            picture_area = 0.0
            for shape in slide.shapes:
                try:
                    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        picture_area += float((shape.width or 0) * (shape.height or 0))
                except Exception:
                    continue
            coverage = picture_area / slide_area
            if coverage >= 0.4:
                warnings.append(
                    {
                        "slide_index": index,
                        "field": "editability",
                        "severity": "warning",
                        "message": (
                            f"slide {index + 1} is {coverage:.0%} picture coverage; "
                            "generated slides must stay editable text/shapes"
                        ),
                    }
                )
        return warnings

    def prepare_outlines(
        self,
        template: TemplateProfile,
        outlines: list[SlideOutline],
    ) -> list[SlideOutline]:
        if template.type not in {"freeform", "brand"}:
            return outlines
        prepared = (
            outlines
            if self.renderer_engine == "legacy"
            else self.authored_renderer.author_outlines(outlines)
        )
        if template.type != "brand":
            return prepared
        return self._attach_template_frames(template, prepared)

    def _attach_template_frames(
        self,
        template: TemplateProfile,
        outlines: list[SlideOutline],
    ) -> list[SlideOutline]:
        if not template.slides or not template.source_file:
            return outlines
        framed: list[SlideOutline] = []
        used_counts: dict[int, int] = {}
        for outline in outlines:
            if outline.mode != "flexible":
                framed.append(outline)
                continue
            frame = self.editing_contract.template_slide_for(
                outline, template, outlines, used_counts
            )
            if not frame:
                framed.append(outline)
                continue
            try:
                index = int(frame.get("index"))
                used_counts[index] = used_counts.get(index, 0) + 1
            except (TypeError, ValueError):
                pass
            slot_plan = self.editing_contract.slot_plan_for(outline, frame, template)
            revised = outline.model_copy(deep=True)
            frame_payload = {
                **frame,
                "source_file": template.source_file,
                "reuse_mode": "duplicate-slide-edit",
                "slot_plan": slot_plan,
            }
            revised.layout_json["template_frame"] = frame_payload
            revised.content_json["template_frame"] = {
                key: value
                for key, value in frame_payload.items()
                if key != "source_file"
            }
            framed.append(revised)
        return framed

    def _strip_template_frames(self, outlines: list[SlideOutline]) -> list[SlideOutline]:
        stripped: list[SlideOutline] = []
        for outline in outlines:
            if "template_frame" not in outline.layout_json and "template_frame" not in outline.content_json:
                stripped.append(outline)
                continue
            revised = outline.model_copy(deep=True)
            revised.layout_json.pop("template_frame", None)
            revised.content_json.pop("template_frame", None)
            stripped.append(revised)
        return stripped

    def _write_deck_json(
        self,
        template: TemplateProfile,
        outlines: Iterable[SlideOutline],
        working_dir: Path,
    ) -> Path:
        slides = []
        for outline in outlines:
            slides.append(
                {
                    "title": outline.content_json.get("title", outline.label),
                    "summary": outline.content_json.get("summary", ""),
                    "bullets": outline.content_json.get("bullets", []),
                    "metrics": outline.content_json.get("metrics", []),
                    "layout": outline.layout_json.get("layout", "icon_rows"),
                    "icons": outline.layout_json.get("icons", []),
                }
            )
        payload = {
            "brand": template.brand.model_dump(),
            "slides": slides,
        }
        deck_json_path = working_dir / "deck.json"
        deck_json_path.write_text(json.dumps(payload), encoding="utf-8")
        return deck_json_path
