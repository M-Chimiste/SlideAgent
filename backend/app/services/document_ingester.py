import re
import uuid
import csv
from datetime import UTC, datetime
from pathlib import Path
from typing import Iterable

try:
    from markitdown import FileConversionException
    from markitdown import MarkItDown
except Exception:  # pragma: no cover - exercised in ARM Docker images without markitdown
    FileConversionException = Exception
    MarkItDown = None

from app.models.document import (
    DocumentBundle,
    DocumentMetadata,
    DocumentMetric,
    DocumentRecord,
    DocumentSection,
    DocumentTable,
)

MONTH_OR_SEASON_PATTERN = (
    r"January|February|March|April|May|June|July|August|September|October|"
    r"November|December|Winter|Spring|Summer|Fall|Autumn"
)


class DocumentIngester:
    def __init__(self) -> None:
        self.markitdown = self._build_markitdown()

    def _build_markitdown(self):
        if MarkItDown is None:
            return None
        try:
            return MarkItDown()
        except Exception:
            return None

    def ingest_documents(
        self, job_id: str, file_paths: Iterable[Path]
    ) -> tuple[list[DocumentRecord], DocumentBundle]:
        records: list[DocumentRecord] = []
        sections: list[DocumentSection] = []
        tables: list[DocumentTable] = []
        metrics: list[DocumentMetric] = []
        inventory: list[str] = []
        source_index: dict[str, dict[str, str]] = {}
        metadata = DocumentMetadata()

        for file_path in file_paths:
            record_id = str(uuid.uuid4())
            markdown = self._convert_to_markdown(file_path)
            record = DocumentRecord(
                id=record_id,
                job_id=job_id,
                filename=file_path.name,
                file_path=file_path.as_posix(),
                markdown_content=markdown,
                structured_json=None,
                created_at=self._timestamp(),
            )
            records.append(record)
            parsed_sections = self._parse_sections(record_id, markdown)
            self._assign_section_source_ids(
                parsed_sections, record, source_index
            )
            sections.extend(parsed_sections)
            parsed_tables = self._parse_tables(record_id, markdown)
            self._assign_table_source_ids(parsed_tables, record, source_index)
            tables.extend(parsed_tables)
            parsed_metrics = self._parse_metrics(record_id, markdown)
            self._assign_metric_source_ids(parsed_metrics, record, source_index)
            metrics.extend(parsed_metrics)
            inventory.extend(self._extract_inventory(markdown))
            if not metadata.title:
                metadata.title = file_path.stem

        bundle = DocumentBundle(
            job_id=job_id,
            sections=sections,
            tables=tables,
            metrics=metrics,
            metadata=metadata,
            content_inventory=inventory,
            source_index=source_index,
        )
        return records, bundle

    def _assign_section_source_ids(
        self,
        sections: list[DocumentSection],
        record: DocumentRecord,
        source_index: dict[str, dict[str, str]],
    ) -> None:
        for index, section in enumerate(sections, start=1):
            title = self._clean_source_title(section.title) or f"Section {index}"
            source_id = f"{record.id}:section:{index}:{self._slug(title)}"
            section.source_id = source_id
            source_index[source_id] = {
                "kind": "section",
                "source_doc_id": record.id,
                "filename": record.filename,
                "title": title,
                "label": f"{Path(record.filename).stem} > {title}",
            }

    def _assign_table_source_ids(
        self,
        tables: list[DocumentTable],
        record: DocumentRecord,
        source_index: dict[str, dict[str, str]],
    ) -> None:
        for index, table in enumerate(tables, start=1):
            title = self._clean_source_title(table.title or "") or f"Table {index}"
            source_id = f"{record.id}:table:{index}:{self._slug(title)}"
            table.source_id = source_id
            source_index[source_id] = {
                "kind": "table",
                "source_doc_id": record.id,
                "filename": record.filename,
                "title": title,
                "label": f"{Path(record.filename).stem} > {title}",
            }

    def _assign_metric_source_ids(
        self,
        metrics: list[DocumentMetric],
        record: DocumentRecord,
        source_index: dict[str, dict[str, str]],
    ) -> None:
        for index, metric in enumerate(metrics, start=1):
            title = self._clean_source_title(metric.label) or f"Metric {index}"
            source_id = f"{record.id}:metric:{index}:{self._slug(title)}"
            metric.source_id = source_id
            source_index[source_id] = {
                "kind": "metric",
                "source_doc_id": record.id,
                "filename": record.filename,
                "title": title,
                "label": f"{Path(record.filename).stem} > {title}",
            }

    def _clean_source_title(self, title: str) -> str:
        return " ".join(str(title or "").split()).strip(" #|")

    def _slug(self, text: str) -> str:
        cleaned = re.sub(r"[^a-zA-Z0-9]+", "-", text.lower()).strip("-")
        return cleaned[:48] or "source"

    def _convert_to_markdown(self, file_path: Path) -> str:
        conversion_error: Exception | None = None
        if self.markitdown is not None:
            try:
                result = self.markitdown.convert(file_path.as_posix())
                return result.text_content or ""
            except FileConversionException as exc:
                conversion_error = exc
            except Exception as exc:
                conversion_error = exc
        fallback = self._convert_with_native_fallback(file_path)
        if fallback is not None:
            return fallback
        if conversion_error is not None:
            raise conversion_error
        return ""

    def _convert_with_native_fallback(self, file_path: Path) -> str | None:
        suffix = file_path.suffix.lower()
        if suffix == ".docx":
            return self._convert_docx_fallback(file_path)
        if suffix == ".pptx":
            return self._convert_pptx_fallback(file_path)
        if suffix == ".xlsx":
            return self._convert_xlsx_fallback(file_path)
        if suffix == ".csv":
            return self._convert_csv_fallback(file_path)
        if suffix in {".md", ".markdown", ".txt"}:
            return file_path.read_text(encoding="utf-8", errors="ignore")
        return None

    def _convert_docx_fallback(self, file_path: Path) -> str:
        from docx import Document

        document = Document(file_path.as_posix())
        lines: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = paragraph.style.name.lower() if paragraph.style else ""
            if "heading 1" in style_name:
                lines.append(f"# {text}")
            elif "heading 2" in style_name:
                lines.append(f"## {text}")
            elif "heading 3" in style_name:
                lines.append(f"### {text}")
            else:
                lines.append(text)
        for table in document.tables:
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            header = rows[0]
            lines.append("| " + " | ".join(header) + " |")
            lines.append("| " + " | ".join(["---"] * len(header)) + " |")
            for row in rows[1:]:
                lines.append("| " + " | ".join(row) + " |")
        return "\n\n".join(lines)

    def _convert_pptx_fallback(self, file_path: Path) -> str:
        from pptx import Presentation

        presentation = Presentation(file_path.as_posix())
        lines: list[str] = []
        for slide_index, slide in enumerate(presentation.slides, start=1):
            slide_lines: list[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    text = "\n".join(
                        paragraph.text.strip()
                        for paragraph in shape.text_frame.paragraphs
                        if paragraph.text.strip()
                    )
                    if text:
                        slide_lines.append(text)
                if getattr(shape, "has_table", False):
                    table = shape.table
                    rows = [
                        [cell.text.strip() for cell in row.cells]
                        for row in table.rows
                    ]
                    slide_lines.extend(self._markdown_table(rows))
            if slide_lines:
                lines.append(f"# Slide {slide_index}\n\n" + "\n\n".join(slide_lines))
        return "\n\n".join(lines)

    def _convert_xlsx_fallback(self, file_path: Path) -> str:
        from openpyxl import load_workbook

        workbook = load_workbook(file_path.as_posix(), data_only=True, read_only=True)
        lines: list[str] = []
        for worksheet in workbook.worksheets:
            rows = [
                ["" if cell is None else str(cell) for cell in row]
                for row in worksheet.iter_rows(values_only=True)
            ]
            rows = [row for row in rows if any(cell.strip() for cell in row)]
            if rows:
                lines.append(f"# {worksheet.title}")
                lines.extend(self._markdown_table(rows))
        workbook.close()
        return "\n\n".join(lines)

    def _convert_csv_fallback(self, file_path: Path) -> str:
        with file_path.open(newline="", encoding="utf-8", errors="ignore") as handle:
            rows = [[cell.strip() for cell in row] for row in csv.reader(handle)]
        rows = [row for row in rows if any(cell for cell in row)]
        return "\n".join(self._markdown_table(rows))

    def _markdown_table(self, rows: list[list[str]]) -> list[str]:
        if not rows:
            return []
        width = max(len(row) for row in rows)
        normalized = [row + [""] * (width - len(row)) for row in rows]
        header = normalized[0]
        lines = ["| " + " | ".join(header) + " |"]
        lines.append("| " + " | ".join(["---"] * width) + " |")
        for row in normalized[1:]:
            lines.append("| " + " | ".join(row) + " |")
        return lines

    def _parse_sections(self, doc_id: str, markdown: str) -> list[DocumentSection]:
        sections: list[DocumentSection] = []
        current_title = "Overview"
        current_level = 1
        current_lines: list[str] = []
        for line in markdown.splitlines():
            heading_match = re.match(r"^(#+)\s+(.*)", line)
            if heading_match:
                if current_lines:
                    sections.append(
                        DocumentSection(
                            title=current_title,
                            level=current_level,
                            content="\n".join(current_lines).strip(),
                            source_doc_id=doc_id,
                        )
                    )
                    current_lines = []
                current_level = len(heading_match.group(1))
                current_title = heading_match.group(2).strip()
            else:
                if line.strip():
                    current_lines.append(line)
        if current_lines:
            sections.append(
                DocumentSection(
                    title=current_title,
                    level=current_level,
                    content="\n".join(current_lines).strip(),
                    source_doc_id=doc_id,
                )
            )
        return sections

    def _parse_tables(self, doc_id: str, markdown: str) -> list[DocumentTable]:
        tables: list[DocumentTable] = []
        lines = markdown.splitlines()
        idx = 0
        while idx < len(lines) - 1:
            if "|" in lines[idx] and "|" in lines[idx + 1]:
                header = [cell.strip() for cell in lines[idx].split("|") if cell.strip()]
                separator = lines[idx + 1]
                if set(separator.replace("|", "").strip()) <= {"-", ":"}:
                    idx += 2
                    rows = []
                    while idx < len(lines) and "|" in lines[idx]:
                        row = [
                            cell.strip()
                            for cell in lines[idx].split("|")
                            if cell.strip()
                        ]
                        rows.append(row)
                        idx += 1
                    tables.append(
                        DocumentTable(
                            title=None,
                            headers=header,
                            rows=rows,
                            source_doc_id=doc_id,
                        )
                    )
                    continue
            idx += 1
        return tables

    def _parse_metrics(self, doc_id: str, markdown: str) -> list[DocumentMetric]:
        metrics: list[DocumentMetric] = []
        seen: set[tuple[str, float, str]] = set()
        for sentence in self._metric_sentences(markdown):
            occupied_spans: list[tuple[int, int]] = []
            for match in re.finditer(
                r"(?<![\w.])(\d{1,3}(?:,\d{3})+|\d+(?:\.\d+)?)\s*"
                r"(?:to|-|–|—)\s*"
                r"(\d{1,3}(?:,\d{3})+|\d+(?:\.\d+)?)\s*(tokens?)\b",
                sentence,
                flags=re.IGNORECASE,
            ):
                label = self._metric_label_from_context(sentence, match.start(), match.end(), "tokens")
                self._append_metric(
                    metrics,
                    seen,
                    f"{label} minimum",
                    self._parse_metric_value(match.group(1)),
                    "tokens",
                    doc_id,
                )
                self._append_metric(
                    metrics,
                    seen,
                    f"{label} maximum",
                    self._parse_metric_value(match.group(2)),
                    "tokens",
                    doc_id,
                )
                occupied_spans.append(match.span())

            for match in re.finditer(
                r"(?<![\w.])(\d{1,3}(?:,\d{3})+|\d+(?:\.\d+)?)\s*"
                r"(million\s+tokens?|tokens?)\b",
                sentence,
                flags=re.IGNORECASE,
            ):
                if self._span_is_occupied(match.span(), occupied_spans):
                    continue
                self._append_metric(
                    metrics,
                    seen,
                    self._metric_label_from_context(sentence, match.start(), match.end(), "tokens"),
                    self._parse_metric_value(match.group(1), match.group(2)),
                    "tokens",
                    doc_id,
                )

            for match in re.finditer(r"(?<![\w.])(\d+(?:\.\d+)?)\s*%", sentence):
                self._append_metric(
                    metrics,
                    seen,
                    self._metric_label_from_context(sentence, match.start(), match.end(), "%"),
                    self._parse_metric_value(match.group(1)),
                    "%",
                    doc_id,
                )
        return metrics[:12]

    def _metric_sentences(self, markdown: str) -> list[str]:
        compact = re.sub(r"\s+", " ", markdown.replace("|", " "))
        return [
            sentence.strip()
            for sentence in re.split(r"(?<=[.!?])\s+|\n+", compact)
            if sentence.strip()
        ]

    def _append_metric(
        self,
        metrics: list[DocumentMetric],
        seen: set[tuple[str, float, str]],
        label: str,
        value: float,
        unit: str | None,
        doc_id: str,
    ) -> None:
        cleaned_label = self._clean_metric_label(label)
        key = (cleaned_label.casefold(), round(float(value), 4), unit or "")
        if key in seen:
            return
        seen.add(key)
        metrics.append(
            DocumentMetric(
                label=cleaned_label,
                value=value,
                unit=unit,
                source_doc_id=doc_id,
            )
        )

    def _parse_metric_value(self, number_text: str, unit_text: str = "") -> float:
        value = float(number_text.replace(",", ""))
        if unit_text.lower().startswith("million"):
            value *= 1_000_000
        return int(value) if value.is_integer() else value

    def _metric_label_from_context(
        self, sentence: str, number_start: int, number_end: int, unit: str
    ) -> str:
        before = sentence[:number_start].strip(" ,.;:()[]")
        after = sentence[number_end:].strip(" ,.;:()[]")
        if unit == "%":
            return self._percent_metric_label(before, after)
        if unit == "tokens":
            return self._token_metric_label(before)
        return self._noun_phrase_before_metric(before) or after or "Sourced metric"

    def _percent_metric_label(self, before: str, after: str) -> str:
        after_clean = re.sub(r"^(?:of|for)\s+", "", after, flags=re.IGNORECASE)
        after_clean = re.split(
            r"[,.;]|\b(?:by|during|as of)\b",
            after_clean,
            maxsplit=1,
            flags=re.IGNORECASE,
        )[0].strip()
        noun_phrase = self._noun_phrase_before_metric(before)
        if after.lower().startswith(("of ", "for ")) and after_clean:
            return after_clean
        if after_clean and len(after_clean.split()) <= 4 and noun_phrase:
            return f"{after_clean} {noun_phrase}"
        return noun_phrase or after_clean or "Sourced share"

    def _token_metric_label(self, before: str) -> str:
        noun_phrase = self._noun_phrase_before_metric(before)
        if "window" in noun_phrase.lower() and (
            noun_phrase.lower().startswith("this window")
            or "ranges" in noun_phrase.lower()
            or noun_phrase.lower().startswith("scaled this window")
        ):
            return "Context window"
        if "context" in noun_phrase.lower() and "window" not in noun_phrase.lower():
            return f"{noun_phrase} window"
        return noun_phrase or "Context window"

    def _noun_phrase_before_metric(self, before: str) -> str:
        cleaned = re.sub(
            rf"\b(?:{MONTH_OR_SEASON_PATTERN})\s+\d{{1,2}},?\s+\d{{4}}\b",
            "",
            before,
            flags=re.IGNORECASE,
        )
        cleaned = re.sub(
            rf"\b(?:{MONTH_OR_SEASON_PATTERN})\s+\d{{4}}\b",
            "",
            cleaned,
            flags=re.IGNORECASE,
        )
        cleaned = re.split(r"[,.;:]", cleaned)[-1]
        for marker in (" had ", " has ", " have ", " with "):
            if marker in cleaned.lower():
                cleaned = cleaned[cleaned.lower().rfind(marker) + len(marker) :]
        cleaned = re.sub(
            r"\b(?:that|which)\s+(?:were|was|are|is)\s*$",
            "",
            cleaned,
            flags=re.IGNORECASE,
        )
        cleaned = re.sub(
            r"\b(?:reached|hit|was|were|is|are|had|have|about|roughly|nearly|over|under)\s*$",
            "",
            cleaned,
            flags=re.IGNORECASE,
        )
        words = cleaned.split()
        return " ".join(words[-7:])

    def _clean_metric_label(self, label: str) -> str:
        cleaned = re.sub(r"#+", "", str(label))
        cleaned = re.sub(r"\[[^\]]+\]\([^)]+\)", "", cleaned)
        cleaned = re.sub(
            rf"\b(?:{MONTH_OR_SEASON_PATTERN})\s+\d{{1,2}},?\s+\d{{4}}\b",
            "",
            cleaned,
            flags=re.IGNORECASE,
        )
        cleaned = re.sub(
            rf"\b(?:{MONTH_OR_SEASON_PATTERN})\s+\d{{4}}\b",
            "",
            cleaned,
            flags=re.IGNORECASE,
        )
        cleaned = re.sub(
            r"\b(?:by the end of|as of|in|during)\s+(?:19|20)\d{2}\b",
            "",
            cleaned,
            flags=re.IGNORECASE,
        )
        cleaned = re.sub(
            r"\b(?:were|was|are|is|be|that were|that was)\b",
            "",
            cleaned,
            flags=re.IGNORECASE,
        )
        cleaned = " ".join(cleaned.split()).strip(" .,:;-")
        if not cleaned:
            cleaned = "Sourced metric"
        if len(cleaned) > 64:
            cleaned = cleaned[:64].rsplit(" ", 1)[0].rstrip(".,;:")
        return cleaned[:1].upper() + cleaned[1:]

    def _span_is_occupied(
        self, span: tuple[int, int], occupied_spans: list[tuple[int, int]]
    ) -> bool:
        start, end = span
        return any(start < occupied_end and end > occupied_start for occupied_start, occupied_end in occupied_spans)

    def _extract_inventory(self, markdown: str) -> list[str]:
        sentences = re.split(r"[.!?]\s+", markdown)
        return [sentence.strip() for sentence in sentences if sentence.strip()][:20]

    def _timestamp(self) -> str:
        return datetime.now(UTC).isoformat().replace("+00:00", "Z")


def fuse_document_bundles(bundles: Iterable[DocumentBundle]) -> DocumentBundle:
    bundles = list(bundles)
    if not bundles:
        return DocumentBundle(
            job_id="unknown",
            sections=[],
            tables=[],
            metrics=[],
            metadata=DocumentMetadata(),
            content_inventory=[],
        )
    job_id = bundles[0].job_id
    sections = []
    tables = []
    metrics = []
    inventory = []
    for bundle in bundles:
        sections.extend(bundle.sections)
        tables.extend(bundle.tables)
        metrics.extend(bundle.metrics)
        inventory.extend(bundle.content_inventory)
    metadata = bundles[0].metadata
    return DocumentBundle(
        job_id=job_id,
        sections=sections,
        tables=tables,
        metrics=metrics,
        metadata=metadata,
        content_inventory=inventory,
    )
