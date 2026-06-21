# ruff: noqa: F401
import json
import posixpath
import re
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Optional

from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.models.outline import SlideOutline
from app.models.qa import QAEnvelope, QAIssue, QAResult
from app.services.rendering import RenderingError, render_pptx_to_images, render_pptx_to_pdf
from app.services.visual_qa.constants import (
    CONTENT_TYPES_NS,
    EMU_PER_INCH,
    QA_SYSTEM_PROMPT,
    REL_NS,
    SAFE_XML_PARSER,
)


class PreviewFallbackMixin:
    def _render_pptx_preview_fallback(self, pptx_path: Path, output_dir: Path) -> list[Path]:
        if not pptx_path.exists():
            return []
        try:
            prs = Presentation(pptx_path.as_posix())
        except Exception:
            return []
        output_dir.mkdir(parents=True, exist_ok=True)
        slide_width = max(int(self._emu_to_inches(prs.slide_width) * 120), 1)
        slide_height = max(int(self._emu_to_inches(prs.slide_height) * 120), 1)
        font = ImageFont.load_default()
        images: list[Path] = []
        for idx, slide in enumerate(prs.slides):
            image = Image.new("RGB", (slide_width, slide_height), "white")
            draw = ImageDraw.Draw(image)
            for shape in slide.shapes:
                self._draw_shape_preview(draw, shape, prs, font)
            image_path = output_dir / f"slide-{idx + 1}.jpg"
            image.save(image_path, "JPEG", quality=90)
            images.append(image_path)
        return images

    def _draw_shape_preview(self, draw: ImageDraw.ImageDraw, shape, prs, font) -> None:
        scale_x = (self._emu_to_inches(prs.slide_width) * 120) / int(prs.slide_width)
        scale_y = (self._emu_to_inches(prs.slide_height) * 120) / int(prs.slide_height)
        x = int(shape.left * scale_x)
        y = int(shape.top * scale_y)
        w = max(int(shape.width * scale_x), 1)
        h = max(int(shape.height * scale_y), 1)
        fill = self._shape_color(shape, "fill")
        outline = self._shape_color(shape, "line")
        is_textbox = getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.TEXT_BOX
        if not is_textbox or fill is not None:
            draw.rectangle(
                [x, y, x + w, y + h],
                fill=fill,
                outline=outline or (210, 215, 222),
            )
        if getattr(shape, "has_text_frame", False):
            text = " ".join(shape.text.split())
            if text:
                color = self._first_text_color(shape) or (20, 24, 31)
                text_font = self._font_for_shape(shape)
                self._draw_wrapped_text(
                    draw, text, (x + 6, y + 5, w - 12, h - 10), text_font, color
                )

    def _draw_wrapped_text(
        self,
        draw: ImageDraw.ImageDraw,
        text: str,
        box: tuple[int, int, int, int],
        font,
        color: tuple[int, int, int],
    ) -> None:
        x, y, width, height = box
        if width <= 0 or height <= 0:
            return
        words = text.split()
        lines: list[str] = []
        current = ""
        for word in words:
            candidate = f"{current} {word}".strip()
            if draw.textlength(candidate, font=font) <= width or not current:
                current = candidate
            else:
                lines.append(current)
                current = word
        if current:
            lines.append(current)
        bbox = font.getbbox("Ag")
        line_height = max((bbox[3] - bbox[1]) + 4, 12)
        max_lines = max(height // line_height, 1)
        for line in lines[:max_lines]:
            draw.text((x, y), line[:120], fill=color, font=font)
            y += line_height

    def _font_for_shape(self, shape):
        point_size = self._first_text_point_size(shape) or 12
        pixel_size = max(int(point_size * 120 / 72), 10)
        try:
            return ImageFont.truetype("Arial.ttf", pixel_size)
        except Exception:
            return ImageFont.load_default(size=pixel_size)

    def _first_text_point_size(self, shape) -> float | None:
        try:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.size is not None:
                        return float(run.font.size.pt)
                if paragraph.font.size is not None:
                    return float(paragraph.font.size.pt)
        except Exception:
            return None
        return None

    def _shape_color(self, shape, color_type: str) -> tuple[int, int, int] | None:
        try:
            color = shape.fill.fore_color if color_type == "fill" else shape.line.color
            rgb = color.rgb
        except Exception:
            return None
        if rgb is None:
            return None
        return tuple(int(str(rgb)[idx : idx + 2], 16) for idx in (0, 2, 4))

    def _first_text_color(self, shape) -> tuple[int, int, int] | None:
        try:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    rgb = run.font.color.rgb
                    if rgb is not None:
                        return tuple(int(str(rgb)[idx : idx + 2], 16) for idx in (0, 2, 4))
        except Exception:
            return None
        return None

    def _emu_to_inches(self, value) -> float:
        return int(value) / EMU_PER_INCH

    def export_pdf(self, pptx_path: Path, output_dir: Path) -> Optional[Path]:
        try:
            return render_pptx_to_pdf(pptx_path, output_dir)
        except RenderingError:
            return None
