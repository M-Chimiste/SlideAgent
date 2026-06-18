# ruff: noqa: F401
import json
import posixpath
import re
import zipfile
from io import BytesIO
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.models.outline import SlideOutline
from app.models.qa import QAEnvelope, QAIssue, QAResult
from app.services.rendering import RenderingError, render_pptx_to_images, render_pptx_to_pdf
from app.services.visual_qa.constants import (
    CONTENT_TYPES_NS,
    EMU_PER_INCH,
    QA_SYSTEM_PROMPT,
    REL_NS,
    SAFE_XML_PARSER,
)


class RuleAndPackageChecksMixin:
    def _rule_based_checks(self, outlines: list[SlideOutline]) -> list[QAIssue]:
        issues = []
        for outline in outlines:
            if outline.mode != "flexible":
                continue
            visuals = outline.layout_json.get("visual_elements", [])
            if not visuals:
                issues.append(
                    QAIssue(
                        severity="CRITICAL",
                        message="Text wall detected: no visual elements.",
                        slide_index=outline.slide_index,
                        category="text_wall",
                    )
                )
            sources = outline.content_json.get("sources") or []
            if not sources:
                issues.append(
                    QAIssue(
                        severity="WARNING",
                        message="Slide is missing source coverage metadata.",
                        slide_index=outline.slide_index,
                        category="source_coverage",
                    )
                )
            elif any("source needed" in str(source).lower() for source in sources):
                issues.append(
                    QAIssue(
                        severity="WARNING",
                        message="Slide still contains a [source needed] placeholder.",
                        slide_index=outline.slide_index,
                        category="source_placeholder",
                    )
                )
        issues.extend(self._layout_variety_checks(outlines))
        issues.extend(self._narrative_rhythm_checks(outlines))
        issues.extend(self._exhibit_checks(outlines))
        return issues

    def _layout_variety_checks(self, outlines: list[SlideOutline]) -> list[QAIssue]:
        flexible = [outline for outline in outlines if outline.mode == "flexible"]
        if len(flexible) < 4:
            return []
        layouts = [outline.layout_json.get("layout", "") for outline in flexible]
        issues: list[QAIssue] = []
        if len(set(layouts)) < min(3, len(layouts)):
            issues.append(
                QAIssue(
                    severity="WARNING",
                    message="Deck uses too little layout variety for a consulting narrative.",
                    category="layout_variety",
                )
            )
        run_length = 1
        for idx in range(1, len(layouts)):
            run_length = run_length + 1 if layouts[idx] == layouts[idx - 1] else 1
            if run_length > 2:
                issues.append(
                    QAIssue(
                        severity="WARNING",
                        message="Repeated adjacent layouts may make the deck feel repetitive.",
                        slide_index=flexible[idx].slide_index,
                        category="layout_repetition",
                    )
                )
                break
        return issues

    def _narrative_rhythm_checks(self, outlines: list[SlideOutline]) -> list[QAIssue]:
        flexible = [outline for outline in outlines if outline.mode == "flexible"]
        if len(flexible) < 6:
            return []
        archetypes = [
            str(
                outline.layout_json.get("archetype")
                or outline.content_json.get("archetype")
                or outline.layout_json.get("layout", "")
            )
            for outline in flexible
        ]
        roles = [
            str(outline.content_json.get("narrative_role") or outline.layout_json.get("narrative_role") or "")
            for outline in flexible
        ]
        issues: list[QAIssue] = []
        if len(set(archetypes)) < min(6, len(flexible)):
            issues.append(
                QAIssue(
                    severity="WARNING",
                    message="Deck uses too few distinct archetypes for authored narrative rhythm.",
                    category="narrative_rhythm",
                )
            )
        for idx in range(1, len(archetypes)):
            if archetypes[idx] == archetypes[idx - 1]:
                issues.append(
                    QAIssue(
                        severity="WARNING",
                        message="Adjacent slides repeat the same archetype.",
                        slide_index=flexible[idx].slide_index,
                        category="archetype_repetition",
                    )
                )
                break
        if len(flexible) >= 10:
            required = {"reference", "implementation"}
            missing = sorted(role for role in required if role not in set(roles))
            if missing:
                issues.append(
                    QAIssue(
                        severity="WARNING",
                        message="Source-rich deck is missing narrative beats: " + ", ".join(missing),
                        category="narrative_beats",
                    )
                )
        bullet_layouts = {"two_column", "icon_rows", "icon_grid", "callouts"}
        bullet_count = sum(
            1 for outline in flexible if outline.layout_json.get("layout") in bullet_layouts
        )
        if bullet_count > len(flexible) * 0.55:
            issues.append(
                QAIssue(
                    severity="WARNING",
                    message="Deck relies too heavily on bullet-card layouts.",
                    category="bullet_card_usage",
                )
            )
        return issues

    def _exhibit_checks(self, outlines: list[SlideOutline]) -> list[QAIssue]:
        issues: list[QAIssue] = []
        for outline in outlines:
            if outline.mode != "flexible":
                continue
            role = str(outline.content_json.get("narrative_role") or "")
            archetype = str(outline.content_json.get("archetype") or "")
            if role == "cover" or archetype == "cover":
                continue
            exhibit = outline.content_json.get("exhibit_spec")
            if not isinstance(exhibit, dict) or not exhibit.get("type"):
                issues.append(
                    QAIssue(
                        severity="WARNING",
                        message="Non-cover slide is missing a primary exhibit spec.",
                        slide_index=outline.slide_index,
                        category="missing_exhibit",
                    )
                )
                continue
            exhibit_text = json.dumps(exhibit, sort_keys=True).lower()
            if "diagram description" in exhibit_text or "placeholder" in exhibit_text:
                issues.append(
                    QAIssue(
                        severity="CRITICAL",
                        message="Exhibit spec contains placeholder or meta text.",
                        slide_index=outline.slide_index,
                        category="exhibit_placeholder",
                    )
                )
            exhibit_type = str(exhibit.get("type") or "").lower()
            if exhibit_type == "comparison_table":
                if not exhibit.get("columns") or not exhibit.get("rows"):
                    issues.append(
                        QAIssue(
                            severity="WARNING",
                            message="Comparison exhibit needs clear columns and rows.",
                            slide_index=outline.slide_index,
                            category="comparison_axes",
                        )
                    )
            if exhibit_type == "dependency_map" and len(exhibit.get("middle_nodes", [])) < 2:
                issues.append(
                    QAIssue(
                        severity="WARNING",
                        message="Dependency map needs at least two structured middle nodes.",
                        slide_index=outline.slide_index,
                        category="exhibit_structure",
                    )
                )
            if exhibit_type == "cycle" and len(exhibit.get("steps", [])) < 3:
                issues.append(
                    QAIssue(
                        severity="WARNING",
                        message="Cycle exhibit needs at least three ordered steps.",
                        slide_index=outline.slide_index,
                        category="exhibit_structure",
                    )
                )
            if exhibit_type == "checklist" and len(exhibit.get("items", [])) < 3:
                issues.append(
                    QAIssue(
                        severity="WARNING",
                        message="Checklist exhibit needs at least three action items.",
                        slide_index=outline.slide_index,
                        category="exhibit_structure",
                    )
                )
        return issues

    def _pptx_structure_checks(
        self, pptx_path: Path, outlines: list[SlideOutline]
    ) -> list[QAIssue]:
        if not pptx_path.exists():
            return [
                QAIssue(
                    severity="CRITICAL",
                    message="PPTX file does not exist for QA.",
                    category="office_compatibility",
                )
            ]
        try:
            with zipfile.ZipFile(pptx_path, "r") as pptx_zip:
                bad_entry = pptx_zip.testzip()
                names = set(pptx_zip.namelist())
        except zipfile.BadZipFile:
            return [
                QAIssue(
                    severity="CRITICAL",
                    message="PPTX package is not a valid zip file.",
                    category="office_compatibility",
                )
            ]
        issues: list[QAIssue] = []
        if bad_entry:
            issues.append(
                QAIssue(
                    severity="CRITICAL",
                    message=f"PPTX package has a corrupt entry: {bad_entry}.",
                    category="office_compatibility",
                )
            )
        if "[Content_Types].xml" not in names:
            issues.append(
                QAIssue(
                    severity="CRITICAL",
                    message="PPTX package is missing [Content_Types].xml.",
                    category="office_compatibility",
                )
            )
        else:
            issues.extend(self._content_type_checks(pptx_path, names))
        issues.extend(self._relationship_checks(pptx_path, names))
        try:
            prs = Presentation(pptx_path.as_posix())
        except Exception as exc:
            issues.append(
                QAIssue(
                    severity="CRITICAL",
                    message=f"python-pptx could not open the package: {exc}",
                    category="office_compatibility",
                )
            )
            return issues
        if outlines and len(prs.slides) != len(outlines):
            issues.append(
                QAIssue(
                    severity="WARNING",
                    message="Rendered PPTX slide count does not match planned outlines.",
                    category="slide_count",
                )
            )
        slide_width = self._emu_to_inches(prs.slide_width)
        slide_height = self._emu_to_inches(prs.slide_height)
        for idx, slide in enumerate(prs.slides):
            outline = outlines[idx] if idx < len(outlines) else None
            issues.extend(self._slide_geometry_checks(slide, idx, slide_width, slide_height, outline))
        return issues

    def _content_type_checks(self, pptx_path: Path, names: set[str]) -> list[QAIssue]:
        issues: list[QAIssue] = []
        try:
            with zipfile.ZipFile(pptx_path, "r") as pptx_zip:
                root = etree.fromstring(
                    pptx_zip.read("[Content_Types].xml"),
                    parser=SAFE_XML_PARSER,
                )
        except Exception as exc:
            return [
                QAIssue(
                    severity="CRITICAL",
                    message=f"[Content_Types].xml could not be parsed: {exc}",
                    category="office_compatibility",
                )
            ]
        defaults = {
            node.get("Extension")
            for node in root.findall(f".//{{{CONTENT_TYPES_NS}}}Default")
            if node.get("Extension")
        }
        overrides = {
            node.get("PartName", "").lstrip("/")
            for node in root.findall(f".//{{{CONTENT_TYPES_NS}}}Override")
            if node.get("PartName")
        }
        for name in sorted(names):
            if name == "[Content_Types].xml" or name.endswith("/"):
                continue
            extension = name.rsplit(".", 1)[-1] if "." in name else ""
            if name not in overrides and extension not in defaults:
                issues.append(
                    QAIssue(
                        severity="CRITICAL",
                        message=f"PPTX part has no content type declaration: {name}",
                        category="office_compatibility",
                    )
                )
        return issues

    def _relationship_checks(self, pptx_path: Path, names: set[str]) -> list[QAIssue]:
        issues: list[QAIssue] = []
        try:
            with zipfile.ZipFile(pptx_path, "r") as pptx_zip:
                rel_names = sorted(name for name in names if name.endswith(".rels"))
                for rel_name in rel_names:
                    try:
                        root = etree.fromstring(
                            pptx_zip.read(rel_name),
                            parser=SAFE_XML_PARSER,
                        )
                    except Exception as exc:
                        issues.append(
                            QAIssue(
                                severity="CRITICAL",
                                message=f"Relationship part could not be parsed: {rel_name}: {exc}",
                                category="office_compatibility",
                            )
                        )
                        continue
                    for rel in root.findall(f".//{{{REL_NS}}}Relationship"):
                        target = rel.get("Target")
                        if not target or rel.get("TargetMode") == "External":
                            continue
                        if self._is_external_target(target):
                            continue
                        resolved = self._resolve_relationship_target(rel_name, target)
                        if resolved not in names:
                            issues.append(
                                QAIssue(
                                    severity="CRITICAL",
                                    message=(
                                        f"Relationship target is missing: {rel_name} -> {target}"
                                    ),
                                    category="office_compatibility",
                                )
                            )
        except zipfile.BadZipFile:
            return issues
        return issues

    def _resolve_relationship_target(self, rels_name: str, target: str) -> str:
        if target.startswith("/"):
            return posixpath.normpath(target.lstrip("/"))
        source_part = self._source_part_for_rels(rels_name)
        base_dir = posixpath.dirname(source_part)
        return posixpath.normpath(posixpath.join(base_dir, target))

    def _source_part_for_rels(self, rels_name: str) -> str:
        if rels_name == "_rels/.rels":
            return ""
        return rels_name.replace("/_rels/", "/").removesuffix(".rels")

    def _is_external_target(self, target: str) -> bool:
        return bool(re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", target))

    def _slide_geometry_checks(
        self,
        slide,
        slide_index: int,
        slide_width: float,
        slide_height: float,
        outline: SlideOutline | None = None,
    ) -> list[QAIssue]:
        issues: list[QAIssue] = []
        body_chars = 0
        scan_text_shape_count = 0
        archetype = ""
        if outline:
            archetype = str(
                outline.content_json.get("archetype")
                or outline.layout_json.get("archetype")
                or outline.layout_json.get("layout")
                or ""
            )
        for shape in slide.shapes:
            left = self._emu_to_inches(shape.left)
            top = self._emu_to_inches(shape.top)
            width = self._emu_to_inches(shape.width)
            height = self._emu_to_inches(shape.height)
            if left < -0.02 or top < -0.02:
                issues.append(
                    QAIssue(
                        severity="CRITICAL",
                        message="Slide shape starts outside the slide bounds.",
                        slide_index=slide_index,
                        category="shape_bounds",
                    )
                )
            if left + width > slide_width + 0.05 or top + height > slide_height + 0.05:
                issues.append(
                    QAIssue(
                        severity="CRITICAL",
                        message="Slide shape extends beyond the slide bounds.",
                        slide_index=slide_index,
                        category="shape_bounds",
                    )
                )
            if not getattr(shape, "has_text_frame", False):
                continue
            text = " ".join(shape.text.split())
            if not text:
                continue
            if top < 0.95 or top > slide_height - 0.7:
                continue
            if self._counts_toward_scanability(text, archetype):
                scan_text_shape_count += 1
            body_chars += len(text)
            area = max(width * height, 0.1)
            if len(text) / area > 180:
                issues.append(
                    QAIssue(
                        severity="WARNING",
                        message="Text density creates possible overflow risk.",
                        slide_index=slide_index,
                        category="overflow_risk",
                    )
                )
        if body_chars > 1500:
            issues.append(
                QAIssue(
                    severity="WARNING",
                    message="Slide has high text density; consider splitting or simplifying.",
                    slide_index=slide_index,
                    category="text_density",
                )
            )
        object_threshold = self._scanability_object_threshold(archetype)
        if scan_text_shape_count > object_threshold:
            issues.append(
                QAIssue(
                    severity="WARNING",
                    message="Slide has many text objects and may be difficult to scan.",
                    slide_index=slide_index,
                    category="scanability",
                )
            )
        return issues

    def _counts_toward_scanability(self, text: str, archetype: str) -> bool:
        normalized = " ".join(text.split()).strip()
        if len(normalized) < 24:
            return False
        if re.fullmatch(r"[\d\s.,/%$KMBkmb+-]+(?:tokens|pts|x)?", normalized):
            return False
        if normalized.isupper() and len(normalized) <= 48:
            return False
        structured_archetypes = {
            "checklist",
            "code_panel",
            "comparison_table",
            "executive_summary",
            "framework_cycle",
            "metric_chart",
            "quote_sidebar",
            "table_reference",
        }
        if archetype in structured_archetypes and len(normalized) <= 52:
            return False
        return True

    def _scanability_object_threshold(self, archetype: str) -> int:
        thresholds = {
            "checklist": 14,
            "code_panel": 14,
            "comparison_table": 14,
            "executive_summary": 12,
            "framework_cycle": 14,
            "metric_chart": 14,
            "quote_sidebar": 12,
            "table_reference": 16,
        }
        return thresholds.get(archetype, 10)

