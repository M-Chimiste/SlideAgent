"""PPTXPipeline: Orchestrates PPTX template unpack → XML injection → repack.

Handles ZIP operations, staging directories, and output storage.
"""

import io
import logging
import shutil
import tempfile
import zipfile
from pathlib import Path

from app.models.schemas import DeckOutline, InjectionTarget, PipelineResult
from app.models.templates import LayoutDefinition
from app.services.xml_injector import XMLInjector
from app.storage.local import LocalStorage

logger = logging.getLogger(__name__)


class PPTXPipeline:
    def __init__(
        self,
        storage: LocalStorage,
        injector: XMLInjector,
        staging_base: str,
    ):
        self._storage = storage
        self._injector = injector
        self._staging_base = Path(staging_base)

    async def execute(
        self,
        template_path: str,
        targets: list[InjectionTarget],
        job_id: str,
    ) -> PipelineResult:
        """Full pipeline: unpack template → inject fields → repack → store output."""
        staging_dir = self._staging_base / job_id
        result: PipelineResult | None = None

        try:
            # 1. Read template bytes
            template_bytes = await self._storage.read_bytes(template_path)

            # 2. Unpack ZIP to staging dir
            self._unpack(template_bytes, staging_dir)

            # 3. Inject all targets into slide XMLs
            warnings = self._injector.inject(staging_dir, targets)

            # 4. Repack to output bytes
            output_bytes = self._repack(staging_dir)

            # 5. Verify the output is a valid ZIP with parseable XML
            self._verify(output_bytes)

            # 6. Store output
            output_key = f"outputs/{job_id}/output.pptx"
            await self._storage.write_bytes(output_key, output_bytes)

            result = PipelineResult(
                success=True,
                output_path=output_key,
                injection_targets=targets,
                warnings=warnings,
            )
            return result

        except Exception as e:
            logger.exception("Pipeline failed for job %s", job_id)
            result = PipelineResult(
                success=False,
                output_path="",
                injection_targets=targets,
                warnings=[],
                error=str(e),
                staging_dir=str(staging_dir),
            )
            return result

        finally:
            # Cleanup staging dir on success; preserve on failure for debugging
            if result is not None and result.success and staging_dir.exists():
                shutil.rmtree(staging_dir)

    async def execute_mode2(
        self,
        template_path: str,
        outline: DeckOutline,
        layouts: dict[str, LayoutDefinition],
        targets: list[InjectionTarget],
        job_id: str,
    ) -> PipelineResult:
        """Mode 2 pipeline: create slides from layouts → inject content → store.

        Uses python-pptx for slide creation (handles relationships correctly),
        then lxml for content injection (preserves formatting).
        """
        staging_dir = self._staging_base / job_id
        result: PipelineResult | None = None

        try:
            from pptx import Presentation

            # 1. Read template bytes and create slides from layouts
            template_bytes = await self._storage.read_bytes(template_path)
            prs = Presentation(io.BytesIO(template_bytes))

            # Remove any existing slides (Mode 2 template should be empty)
            nsmap = {
                "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
                "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
            }
            sldIdLst = prs.part._element.find(".//p:sldIdLst", nsmap)
            if sldIdLst is not None:
                for sldId in list(sldIdLst):
                    rId = sldId.get(
                        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
                    )
                    if rId:
                        prs.part.drop_rel(rId)
                    sldIdLst.remove(sldId)

            # Add slides based on outline
            for slide_entry in outline.slides:
                layout_def = layouts.get(slide_entry.layout_name)
                if layout_def is None:
                    logger.warning(
                        "Unknown layout '%s' for slide %d, skipping",
                        slide_entry.layout_name, slide_entry.slide_number,
                    )
                    continue

                slide_layout = prs.slide_layouts[layout_def.slide_layout_index]
                prs.slides.add_slide(slide_layout)

            # Save the presentation with new slides to a temp file
            with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as tmp:
                prs.save(tmp.name)
                tmp_path = Path(tmp.name)

            # 2. Unpack the modified PPTX
            modified_bytes = tmp_path.read_bytes()
            tmp_path.unlink()

            self._unpack(modified_bytes, staging_dir)

            # 3. Inject content into slide XMLs
            warnings = self._injector.inject(staging_dir, targets)

            # 4. Repack to output bytes
            output_bytes = self._repack(staging_dir)

            # 5. Verify output
            self._verify(output_bytes)

            # 6. Store output
            output_key = f"outputs/{job_id}/output.pptx"
            await self._storage.write_bytes(output_key, output_bytes)

            result = PipelineResult(
                success=True,
                output_path=output_key,
                injection_targets=targets,
                warnings=warnings,
            )
            return result

        except Exception as e:
            logger.exception("Mode 2 pipeline failed for job %s", job_id)
            result = PipelineResult(
                success=False,
                output_path="",
                injection_targets=targets,
                warnings=[],
                error=str(e),
                staging_dir=str(staging_dir),
            )
            return result

        finally:
            if result is not None and result.success and staging_dir.exists():
                shutil.rmtree(staging_dir)

    def _unpack(self, pptx_bytes: bytes, staging_dir: Path) -> None:
        """Extract PPTX ZIP to staging directory."""
        staging_dir.mkdir(parents=True, exist_ok=True)

        with zipfile.ZipFile(io.BytesIO(pptx_bytes), "r") as zf:
            zf.extractall(staging_dir)
        logger.info("Unpacked PPTX to %s (%d files)", staging_dir, len(list(staging_dir.rglob("*"))))

    def _repack(self, staging_dir: Path) -> bytes:
        """Repack staging directory into PPTX ZIP bytes.

        Preserves original file structure and uses appropriate compression
        per file type.
        """
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            for file_path in sorted(staging_dir.rglob("*")):
                if file_path.is_file():
                    arcname = str(file_path.relative_to(staging_dir))
                    # Use STORED for media/binary, DEFLATED for XML/text
                    compress = zipfile.ZIP_DEFLATED
                    suffix = file_path.suffix.lower()
                    if suffix in (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".wmf", ".emf"):
                        compress = zipfile.ZIP_STORED
                    zf.write(file_path, arcname, compress_type=compress)

        return buf.getvalue()

    def _verify(self, pptx_bytes: bytes) -> None:
        """Verify the output is a valid ZIP with parseable slide XMLs."""
        from lxml import etree

        with zipfile.ZipFile(io.BytesIO(pptx_bytes), "r") as zf:
            # Check ZIP is valid
            bad = zf.testzip()
            if bad is not None:
                raise ValueError(f"Output PPTX has corrupt entry: {bad}")

            # Verify each slide XML is well-formed
            for name in zf.namelist():
                if name.startswith("ppt/slides/slide") and name.endswith(".xml"):
                    xml_data = zf.read(name)
                    try:
                        etree.fromstring(xml_data)
                    except etree.XMLSyntaxError as e:
                        raise ValueError(f"Malformed XML in {name}: {e}") from e
