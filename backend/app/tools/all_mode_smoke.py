import argparse
import json
import tempfile
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.util import Inches

from app.clients.openai_compatible_client import OpenAICompatibleClient
from app.config import Settings
from app.models.brand import BrandDNA
from app.models.qa import QAIssue, QAResult
from app.models.template import SlideField, SlideSchema, SlideSpec, TemplateProfile
from app.services.content_planner import ContentPlanner
from app.services.design_agent import DesignAgent
from app.services.document_ingester import DocumentIngester
from app.services.pptx_builder import PptxBuilder
from app.services.visual_qa_agent import VisualQAAgent


REPO_ROOT = Path(__file__).resolve().parents[3]
DEFAULT_DOC = REPO_ROOT / "data" / "Beyond Vibe Coding.docx"
DEFAULT_OUT_DIR = Path(tempfile.gettempdir()) / "slideagent-beyond-vibe"
DEFAULT_INSTRUCTIONS = (
    "Create a consulting-quality executive deck about moving beyond vibe coding. "
    "Use action titles, source-grounded claims, varied slide layouts, and concise evidence."
)


def run_smoke(
    doc_path: Path = DEFAULT_DOC,
    out_dir: Path = DEFAULT_OUT_DIR,
    instructions: str = DEFAULT_INSTRUCTIONS,
    use_vision: bool = False,
    settings: Settings | None = None,
    vision_settings: Settings | None = None,
    run_label: str | None = None,
) -> dict[str, Any]:
    settings = settings or Settings(BEDROCK_VALIDATE=False)
    out_dir.mkdir(parents=True, exist_ok=True)
    label = run_label or _model_slug(settings.openai_compatible_model)

    client = OpenAICompatibleClient(settings)
    planner = ContentPlanner(llm_client=client)
    designer = DesignAgent()
    builder = PptxBuilder(node_runner=object())
    qa_client = OpenAICompatibleClient(vision_settings or settings) if use_vision else None
    qa_agent = VisualQAAgent(openai_client=qa_client)

    _, bundle = DocumentIngester().ingest_documents("all-mode-qwen-smoke", [doc_path])
    results: dict[str, Any] = {}
    for mode, template in [
        ("freeform", _template("freeform")),
        ("brand", _template("brand", _brand())),
    ]:
        outlines, planning_warnings = planner.plan(
            template,
            bundle,
            instructions=instructions,
            generation_mode=mode,
        )
        outlines = designer.apply_design(outlines)
        output_path = out_dir / f"{mode}-beyond-vibe-{label}.pptx"
        build_warnings = builder.build_deck(
            template, outlines, output_path, out_dir / f"{mode}-{label}-work"
        )
        qa_result, preview_images = qa_agent.inspect_deck(
            output_path, out_dir / f"{mode}-{label}-preview", outlines
        )
        qa_rounds = 0
        qa_history = [summarize_qa(qa_result.issues)]
        seen_actionable_signatures: set[tuple[tuple[int | None, str, str], ...]] = set()
        while qa_rounds < settings.qa_max_rounds:
            actionable_signature = _actionable_issue_signature(designer, qa_result)
            if (
                not actionable_signature
                or actionable_signature in seen_actionable_signatures
            ):
                break
            seen_actionable_signatures.add(actionable_signature)
            qa_rounds += 1
            outlines = designer.revise_deck_for_qa(outlines, qa_result.issues)
            repair_warnings = builder.build_deck(
                template, outlines, output_path, out_dir / f"{mode}-{label}-work"
            )
            build_warnings.extend(repair_warnings)
            qa_result, preview_images = qa_agent.inspect_deck(
                output_path, out_dir / f"{mode}-{label}-preview", outlines
            )
            qa_history.append(summarize_qa(qa_result.issues))
        rendered = Presentation(output_path.as_posix())
        results[mode] = deck_report(
            mode=mode,
            output_path=output_path,
            slide_count=len(rendered.slides),
            titles=[outline.label for outline in outlines],
            layouts=[outline.layout_json.get("layout", "") for outline in outlines],
            planning_warnings=planning_warnings,
            build_warnings=build_warnings,
            qa_result=qa_result,
            preview_images=preview_images,
            qa_rounds=qa_rounds,
            qa_history=qa_history,
        )

    strict_template = _strict_template(out_dir, label)
    strict_outlines, strict_planning_warnings = ContentPlanner().plan(
        strict_template,
        bundle,
        instructions="Populate the strict executive summary fields from the uploaded source.",
        generation_mode="strict",
    )
    strict_output = out_dir / f"strict-beyond-vibe-{label}.pptx"
    strict_build_warnings = builder.build_deck(
        strict_template, strict_outlines, strict_output, out_dir / f"strict-{label}-work"
    )
    strict_qa_result, strict_images = qa_agent.inspect_deck(
        strict_output, out_dir / f"strict-{label}-preview", strict_outlines
    )
    strict_rendered = Presentation(strict_output.as_posix())
    results["strict"] = deck_report(
        mode="strict",
        output_path=strict_output,
        slide_count=len(strict_rendered.slides),
        titles=[outline.label for outline in strict_outlines],
        layouts=[outline.layout_json.get("layout", "") for outline in strict_outlines],
        planning_warnings=strict_planning_warnings,
        build_warnings=strict_build_warnings,
        qa_result=strict_qa_result,
        preview_images=strict_images,
        qa_rounds=0,
        qa_history=[summarize_qa(strict_qa_result.issues)],
    )

    report = {
        "doc_path": doc_path.as_posix(),
        "out_dir": out_dir.as_posix(),
        "model": settings.openai_compatible_model,
        "base_url": settings.openai_compatible_base_url,
        "run_label": label,
        "vision_enabled": use_vision,
        "vision_model": (vision_settings or settings).openai_compatible_model
        if use_vision
        else None,
        "vision_base_url": (vision_settings or settings).openai_compatible_base_url
        if use_vision
        else None,
        "generated_at": datetime.now(UTC).isoformat().replace("+00:00", "Z"),
        "modes": results,
    }
    report_path = out_dir / f"all-mode-{label}-smoke-report.json"
    report_path.write_text(json.dumps(report, indent=2), encoding="utf-8")
    report["report_path"] = report_path.as_posix()
    return report


def deck_report(
    mode: str,
    output_path: Path,
    slide_count: int,
    titles: list[str],
    layouts: list[str],
    planning_warnings: list[dict[str, Any]],
    build_warnings: list[dict[str, Any]],
    qa_result: QAResult,
    preview_images: list[Path],
    qa_rounds: int = 0,
    qa_history: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    return {
        "mode": mode,
        "output": output_path.as_posix(),
        "slides": slide_count,
        "titles": titles,
        "layouts": layouts,
        "planning_warnings": planning_warnings,
        "build_warnings": build_warnings,
        "qa_rounds": qa_rounds,
        "qa_passed": qa_result.passed,
        "qa": summarize_qa(qa_result.issues),
        "qa_issues": serialize_qa_issues(qa_result.issues),
        "qa_history": qa_history or [summarize_qa(qa_result.issues)],
        "preview_images": [path.as_posix() for path in preview_images],
    }


def summarize_qa(issues: list[QAIssue]) -> dict[str, Any]:
    categories = Counter(issue.category or "uncategorized" for issue in issues)
    return {
        "count": len(issues),
        "critical": sum(1 for issue in issues if issue.severity == "CRITICAL"),
        "warning": sum(1 for issue in issues if issue.severity == "WARNING"),
        "info": sum(1 for issue in issues if issue.severity == "INFO"),
        "categories": dict(categories),
    }


def serialize_qa_issues(issues: list[QAIssue]) -> list[dict[str, Any]]:
    return [
        {
            "severity": issue.severity,
            "category": issue.category,
            "message": issue.message,
            "slide_index": issue.slide_index,
        }
        for issue in issues
    ]


def _actionable_issue_signature(
    designer: DesignAgent, qa_result: QAResult
) -> tuple[tuple[int, str, str], ...]:
    return tuple(
        sorted(
            (
                -1 if issue.slide_index is None else issue.slide_index,
                issue.category or "",
                " ".join(issue.message.lower().split())[:160],
            )
            for issue in qa_result.issues
            if issue.severity == "CRITICAL" or designer.is_actionable_qa_issue(issue)
        )
    )


def _template(mode: str, brand: BrandDNA | None = None, source_file: str = "") -> TemplateProfile:
    timestamp = datetime.now(UTC).isoformat().replace("+00:00", "Z")
    return TemplateProfile(
        id=f"smoke-{mode}",
        name=f"Smoke {mode}",
        type=mode,
        brand=brand or BrandDNA(),
        slides=[],
        source_file=source_file,
        created_at=timestamp,
        updated_at=timestamp,
    )


def _brand() -> BrandDNA:
    return BrandDNA(
        primary_color="#111827",
        secondary_color="#2563eb",
        accent_color="#16a34a",
        font_headings="Aptos Display",
        font_body="Aptos",
    )


def _strict_template(out_dir: Path, label: str) -> TemplateProfile:
    strict_source = out_dir / f"strict-smoke-template-{label}.pptx"
    _write_strict_template(strict_source)
    template = _template("strict", source_file=strict_source.as_posix())
    template.slides = [
        SlideSpec(
            index=0,
            mode="strict",
            label="Strict Executive Summary",
            schema=SlideSchema(
                fields=[
                    SlideField(
                        id="project_title",
                        type="text",
                        location="shape:ProjectTitle",
                        required=True,
                        max_chars=90,
                    ),
                    SlideField(
                        id="executive_summary",
                        type="text",
                        location="shape:ExecutiveSummary",
                        required=True,
                        max_chars=260,
                    ),
                    SlideField(
                        id="key_implication",
                        type="text",
                        location="shape:KeyImplication",
                        required=True,
                        max_chars=260,
                    ),
                ]
            ),
        )
    ]
    return template


def _write_strict_template(path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for name, text, top in [
        ("ProjectTitle", "Old project title", 0.7),
        ("ExecutiveSummary", "Old summary", 1.6),
        ("KeyImplication", "Old implication", 4.2),
    ]:
        box = slide.shapes.add_textbox(Inches(0.8), Inches(top), Inches(11.5), Inches(0.8))
        box.name = name
        box.text = text
    prs.save(path.as_posix())


def _model_slug(model_name: str) -> str:
    slug = "".join(
        char.lower() if char.isalnum() else "-"
        for char in model_name.replace(".", "-")
    )
    return "-".join(part for part in slug.split("-") if part)[:48] or "model"


def main() -> None:
    parser = argparse.ArgumentParser(description="Run all-mode local model smoke generation.")
    parser.add_argument("--doc", type=Path, default=DEFAULT_DOC)
    parser.add_argument("--out-dir", type=Path, default=DEFAULT_OUT_DIR)
    parser.add_argument("--instructions", default=DEFAULT_INSTRUCTIONS)
    parser.add_argument("--base-url", default=None)
    parser.add_argument("--model", default=None)
    parser.add_argument("--timeout-seconds", type=int, default=None)
    parser.add_argument("--vision-base-url", default=None)
    parser.add_argument("--vision-model", default=None)
    parser.add_argument("--vision-timeout-seconds", type=int, default=None)
    parser.add_argument("--label", default=None)
    parser.add_argument(
        "--vision",
        action="store_true",
        help="Run local OpenAI-compatible vision QA over generated previews.",
    )
    args = parser.parse_args()
    settings_kwargs: dict[str, Any] = {"BEDROCK_VALIDATE": False}
    if args.base_url:
        settings_kwargs["OPENAI_COMPATIBLE_BASE_URL"] = args.base_url
    if args.model:
        settings_kwargs["OPENAI_COMPATIBLE_MODEL"] = args.model
    if args.timeout_seconds:
        settings_kwargs["OPENAI_COMPATIBLE_TIMEOUT_SECONDS"] = args.timeout_seconds
    settings = Settings(**settings_kwargs)
    vision_settings = None
    if args.vision and (
        args.vision_base_url or args.vision_model or args.vision_timeout_seconds
    ):
        vision_kwargs: dict[str, Any] = {"BEDROCK_VALIDATE": False}
        vision_kwargs["OPENAI_COMPATIBLE_BASE_URL"] = (
            args.vision_base_url or settings.openai_compatible_base_url
        )
        vision_kwargs["OPENAI_COMPATIBLE_MODEL"] = (
            args.vision_model or settings.openai_compatible_model
        )
        if args.vision_timeout_seconds:
            vision_kwargs["OPENAI_COMPATIBLE_TIMEOUT_SECONDS"] = (
                args.vision_timeout_seconds
            )
        vision_settings = Settings(**vision_kwargs)
    report = run_smoke(
        doc_path=args.doc,
        out_dir=args.out_dir,
        instructions=args.instructions,
        use_vision=args.vision,
        settings=settings,
        vision_settings=vision_settings,
        run_label=args.label,
    )
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
