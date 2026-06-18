import argparse
import json
from collections import Counter
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def analyze_pptx(path: Path) -> dict[str, Any]:
    prs = Presentation(path.as_posix())
    slide_reports = []
    archetypes: Counter[str] = Counter()
    dark_light: Counter[str] = Counter()
    for index, slide in enumerate(prs.slides):
        texts = [
            " ".join(shape.text.split())
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and shape.text.strip()
        ]
        text_chars = sum(len(text) for text in texts)
        table_count = sum(1 for shape in slide.shapes if getattr(shape, "has_table", False))
        chart_count = sum(1 for shape in slide.shapes if getattr(shape, "has_chart", False))
        connector_count = sum(
            1
            for shape in slide.shapes
            if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.LINE
        )
        object_count = len(slide.shapes)
        rhythm = _dark_or_light(slide)
        archetype = _guess_archetype(
            index=index,
            texts=texts,
            table_count=table_count,
            chart_count=chart_count,
            connector_count=connector_count,
            object_count=object_count,
        )
        dark_light[rhythm] += 1
        archetypes[archetype] += 1
        slide_reports.append(
            {
                "slide_number": index + 1,
                "archetype_guess": archetype,
                "dark_light": rhythm,
                "object_count": object_count,
                "text_chars": text_chars,
                "text_shape_count": len(texts),
                "table_count": table_count,
                "chart_count": chart_count,
                "connector_count": connector_count,
            }
        )
    return {
        "path": path.as_posix(),
        "slide_count": len(prs.slides),
        "archetype_distribution": dict(archetypes),
        "dark_light_rhythm": dict(dark_light),
        "section_divider_cadence": [
            slide["slide_number"]
            for slide in slide_reports
            if slide["archetype_guess"] == "section_or_cover"
        ],
        "has_table_slide": any(
            slide["table_count"]
            or slide["archetype_guess"] in {"comparison_table", "table_reference"}
            for slide in slide_reports
        ),
        "has_diagram_slide": any(slide["connector_count"] >= 2 for slide in slide_reports),
        "has_reference_slide": any(
            slide["archetype_guess"] in {"reference_or_code", "table_reference"}
            for slide in slide_reports
        ),
        "slides": slide_reports,
    }


def _guess_archetype(
    index: int,
    texts: list[str],
    table_count: int,
    chart_count: int,
    connector_count: int,
    object_count: int,
) -> str:
    joined = " ".join(texts).lower()
    if index == 0:
        return "section_or_cover"
    if chart_count:
        return "metric_chart"
    if "dimension" in joined and any(
        token in joined for token in ("before", "after", "current state", "target state")
    ):
        return "comparison_table"
    if "artifact" in joined and ("purpose" in joined or "refresh" in joined):
        return "table_reference"
    if table_count and any(token in joined for token in ("owner", "timing", "artifact")):
        return "table_reference"
    if table_count:
        return "comparison_table"
    if connector_count >= 3:
        return "diagram_or_cycle"
    if any(
        token in joined
        for token in (
            "rules.md",
            "agents.md",
            "agent-brief.md",
            "update-protocol",
            "agentic-cycle",
            "activecontext.md",
            "progress.md",
            "code",
            "github",
        )
    ):
        return "reference_or_code"
    if any(token in joined for token in ("checklist", "week", "owner")):
        return "checklist"
    if any(token in joined for token in ("failure", "risk", "anti", "pattern")):
        return "anti_patterns"
    if object_count <= 6 and len(joined) < 260:
        return "section_or_cover"
    return "content"


def _dark_or_light(slide) -> str:
    if not slide.shapes:
        return "light"
    shape = slide.shapes[0]
    try:
        rgb = shape.fill.fore_color.rgb
    except Exception:
        return "light"
    if rgb is None:
        return "light"
    red = int(str(rgb)[0:2], 16)
    green = int(str(rgb)[2:4], 16)
    blue = int(str(rgb)[4:6], 16)
    luminance = 0.2126 * red + 0.7152 * green + 0.0722 * blue
    return "dark" if luminance < 110 else "light"


def main() -> None:
    parser = argparse.ArgumentParser(description="Inspect high-level PPTX rhythm signals.")
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--out", type=Path, default=None)
    args = parser.parse_args()
    report = analyze_pptx(args.pptx)
    output = json.dumps(report, indent=2)
    if args.out:
        args.out.write_text(output, encoding="utf-8")
    else:
        print(output)


if __name__ == "__main__":
    main()
